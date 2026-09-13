from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path("outputs")
OUT_FILE = OUT_DIR / "Documento_Escopo_Davantti_Volt_Corp_Advogada.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(90, 98, 110)
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
BLACK = RGBColor(0, 0, 0)


def set_run_font(run, name="Calibri", size=11, color=BLACK, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph(p, before=0, after=6, line=1.10, align=None):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        p.alignment = align


def add_para(doc, text="", size=11, color=BLACK, bold=False, italic=False, after=6, before=0, align=None, style=None):
    p = doc.add_paragraph(style=style)
    set_paragraph(p, before=before, after=after, align=align)
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])
    set_run_font(run, size=9, color=MUTED)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("bottom", bottom), ("start", start), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D7DBE2"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    total = sum(widths_dxa)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))

    existing_grid = table._tbl.tblGrid
    if existing_grid is not None:
        table._tbl.remove(existing_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    table._tbl.insert(1, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_cell_text(cell, text, bold=False, color=BLACK, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph(p, before=0, after=0, line=1.10)
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def add_table(doc, headers, rows, widths_dxa):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    mark_header_row(table.rows[0])
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_shading(header_cells[idx], LIGHT_GRAY)
        set_cell_text(header_cells[idx], header, bold=True, color=INK, size=9.5)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_shading(cells[idx], WHITE)
            set_cell_text(cells[idx], value, size=9)
    add_para(doc, "", after=4)
    return table


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color="D7DBE2")
    mark_header_row(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    cell.text = ""
    p1 = cell.paragraphs[0]
    set_paragraph(p1, after=2)
    r1 = p1.add_run(title)
    set_run_font(r1, size=10.5, color=INK, bold=True)
    p2 = cell.add_paragraph()
    set_paragraph(p2, after=0)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10, color=BLACK)
    add_para(doc, "", after=4)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    set_paragraph(header, after=0)
    header_run = header.add_run("Davantti / Volt Corp | Escopo de produtos e funcionalidades")
    set_run_font(header_run, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    prefix = footer.add_run("Documento de apoio para revisão jurídica | Página ")
    set_run_font(prefix, size=9, color=MUTED)
    add_page_number(footer)


def add_cover(doc):
    add_para(doc, "Documento de Escopo", size=11, color=MUTED, bold=True, after=2)
    add_para(
        doc,
        "Davantti e Volt Corp",
        size=26,
        color=INK,
        bold=True,
        after=4,
    )
    add_para(
        doc,
        "Produtos, módulos, funcionalidades atuais e roadmap inicial",
        size=14,
        color=MUTED,
        after=18,
    )

    rows = [
        ("Destinatário", "Assessoria jurídica / advogada"),
        ("Objetivo", "Mapear produtos, módulos, marcas, funcionalidades e pontos de atenção para revisão legal."),
        ("Data", "22 de junho de 2026"),
        ("Natureza", "Documento comercial e técnico de apoio; não substitui parecer jurídico."),
        ("Escopo", "Davantti Seller, Volt Corp, Volt Core, Volt Chat e produtos futuros de estoque/visualização 3D."),
    ]
    add_table(doc, ["Campo", "Descrição"], rows, [2400, 6960])

    add_callout(
        doc,
        "Nota de uso",
        "Este documento organiza a visão de produto para apoiar contratos, termos de uso, políticas de privacidade, "
        "registro de marcas, domínios e validação de responsabilidades perante clientes e plataformas parceiras.",
    )

    doc.add_page_break()


def build_doc():
    OUT_DIR.mkdir(exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    add_heading(doc, "1. Visão executiva", 1)
    add_para(
        doc,
        "A Davantti está sendo organizada como um ecossistema de produtos SaaS e ferramentas operacionais. "
        "A linha Seller concentra soluções para marketplaces e e-commerce. A linha Business será posicionada "
        "com a marca Volt Corp, concentrando produtos de gestão empresarial, comunicação corporativa e controle operacional.",
    )
    add_para(
        doc,
        "A estratégia recomendada é manter a Davantti como marca/ecossistema principal e tratar Volt Corp como "
        "marca de linha empresarial, com produtos próprios como Volt Core, Volt Chat e eventuais módulos de estoque "
        "virtual 3D. Essa separação ajuda a evitar confusão comercial entre ferramentas de marketplace e sistemas "
        "de gestão interna para empresas.",
    )

    add_heading(doc, "2. Arquitetura de marcas e produtos", 1)
    add_table(
        doc,
        ["Marca / linha", "Papel no ecossistema", "Produtos e módulos relacionados", "Status"],
        [
            (
                "Davantti",
                "Marca guarda-chuva do ecossistema, centralizando a oferta institucional e a evolução dos produtos.",
                "Hub administrativo, autenticação, gestão de empresas, usuários, permissões, faturamento/acesso e páginas institucionais.",
                "Existente / em evolução",
            ),
            (
                "Davantti Seller",
                "Linha voltada a sellers, marketplaces e e-commerce.",
                "Mercado Livre, Shopee, MadeiraMadeira, Magalu, logística/rastreio, SKU líder e módulos auxiliares.",
                "Existente + módulos em expansão",
            ),
            (
                "Volt Corp",
                "Marca da linha Business/Enterprise para gestão empresarial e operação interna.",
                "Volt Core, Volt Chat e produto futuro de estoque virtual 3D/controle de estoque.",
                "Em planejamento/implementação",
            ),
        ],
        [1800, 2800, 3200, 1560],
    )
    add_para(
        doc,
        "Sugestão de arquitetura digital: manter uma landing geral para Davantti, uma landing Davantti Seller, "
        "uma landing Volt Corp e páginas específicas para cada produto: /core, /chat e, futuramente, /estoque-3d "
        "ou nomenclatura equivalente.",
    )

    add_heading(doc, "3. Davantti Seller - módulos atuais e previstos", 1)
    add_para(
        doc,
        "A linha Seller reúne ferramentas conectadas a marketplaces, com foco em operação, análise, automação, "
        "cadastro, anúncios, promoções, estoque, relatórios e controle administrativo por empresa/usuário.",
    )
    add_table(
        doc,
        ["Módulo", "Funções principais", "Dados envolvidos", "Status / observação jurídica"],
        [
            (
                "Mercado Livre",
                "Painel operacional, conexão OAuth/contas, anúncios, clonagem, modelo em massa, características em massa, promoções, remoção de promoções, ranking, reputação, publicidade, análise de mercado, análise por IA, prazo de produção, validação de dimensões, estoque alerta, Full, atacado, fiscal de vendas, financeiro/custos/margem e jobs administrativos.",
                "Contas de seller, tokens, anúncios, SKUs, preços, custos, reputação, vendas, métricas, arquivos CSV e logs.",
                "Existente. Requer atenção a termos da API Mercado Livre, uso de dados, limites de automação e responsabilidade por alterações em anúncios/preços.",
            ),
            (
                "Shopee",
                "Autenticação, pedidos, produtos, sincronização, anúncios, campanhas, descontos, flash sale, métricas, qualidade de catálogo, logística, clone de anúncios, geo vendas, SEO, webhooks, jobs e relatórios.",
                "Lojas, tokens, pedidos, produtos, campanhas, endereços, logística, métricas e webhooks.",
                "Existente. Requer revisão dos termos da API Shopee, tratamento de dados de compradores e regras de campanhas/anúncios.",
            ),
            (
                "MadeiraMadeira",
                "Módulo dedicado ao marketplace MadeiraMadeira, com workspace, autenticação, sincronização, catálogo, cache de API, painel, usuários, convites, patch notes e integração com Hub.",
                "Token/cadastro de seller, catálogo, produtos, atributos, imagens, usuários e logs operacionais.",
                "Incluso no escopo. Revisar contrato/termos MadeiraMadeira, SLA de integrações e responsabilidades por cadastro/sincronização.",
            ),
            (
                "Magalu",
                "Módulo de marketplace a ser incluído na linha Seller para operação e automação relacionada ao marketplace Magalu.",
                "A definir conforme API/integração: produtos, pedidos, anúncios, preços, estoque, campanhas, tokens e relatórios.",
                "Previsto/em inclusão. Precisa de validação jurídica dos termos Magalu antes de oferta comercial ampla.",
            ),
            (
                "Logística / Avantracking / Logisync",
                "Rastreamento de pedidos, conciliação de status, integração com transportadoras e ERPs, notificações, relatórios, recálculo de frete e apoio a operações com Tray, Anymarket, Magazord, Intelipost, Correios, SSW e outros conectores.",
                "Pedidos, notas/XML, rastreios, transportadoras, fretes, destinatários, status logísticos e notificações.",
                "Existente/em evolução. Exige atenção a LGPD, compartilhamento de dados com terceiros e responsabilidade por prazos/status.",
            ),
            (
                "SKU líder / LeaderSku",
                "Ferramenta auxiliar para análise e operação de SKUs, com login, serviço próprio, integração de acesso Hub e interface dedicada.",
                "SKUs, produtos, usuários, empresas e métricas operacionais.",
                "Existente. Avaliar se será produto autônomo, módulo Seller ou ferramenta interna.",
            ),
        ],
        [1450, 3600, 2150, 2160],
    )

    add_heading(doc, "4. Volt Corp - linha Business", 1)
    add_para(
        doc,
        "Volt Corp será a linha de produtos empresariais do ecossistema, separada da comunicação de marketplace. "
        "A proposta é atender empresas pequenas, micro e médias que precisam de controle de vendas, estoque, caixa, "
        "clientes, relatórios e módulos adaptáveis por segmento.",
    )

    add_heading(doc, "4.1 Volt Core", 2)
    add_para(
        doc,
        "Volt Core é o produto inicial da linha Volt Corp. Ele nasce como um motor de gestão empresarial com núcleo "
        "base e capacidade futura de extensão por setor. O primeiro segmento planejado é ótica, mas a arquitetura "
        "prevê que novos setores sejam ativados por configuração de plano, empresa e painel master.",
    )
    add_table(
        doc,
        ["Área", "Funcionalidades do núcleo", "Observações"],
        [
            (
                "Administração e acesso",
                "Login, sessão, usuário master, empresas, usuários por empresa, permissões, painel master, integração com Hub e auditoria.",
                "Base importante para contrato SaaS, trilha de auditoria e responsabilização de ações por usuário.",
            ),
            (
                "Cadastros",
                "Clientes, produtos, configurações da empresa, métodos de pagamento, módulos/telas por segmento e dados fiscais futuros.",
                "Deve prever LGPD para dados de clientes e política de retenção/exclusão.",
            ),
            (
                "Estoque",
                "Entrada, saída, posição de estoque, movimentações, histórico, ajuste manual, baixa automática na venda e relatórios.",
                "Produto deve deixar claro que o estoque depende de uso correto pelo cliente e integrações futuras.",
            ),
            (
                "Vendas / balcão",
                "Venda com cliente, itens, descontos, formas de pagamento, baixa de estoque, cancelamento, recibo imprimível e registro de eventos.",
                "Recibo deve ser tratado como não fiscal enquanto emissão fiscal não estiver homologada.",
            ),
            (
                "Pagamentos e recebíveis",
                "Dinheiro, Pix, cartão de débito, cartão de crédito com parcelas, nota promissória, cheque, datas de vencimento, baixa/recebimento e alertas.",
                "A Davantti/Volt Core não deve se posicionar como instituição financeira se apenas registra recebíveis.",
            ),
            (
                "Caixa",
                "Abertura, fechamento, entradas, saídas, sangria/suprimento, resumo do dia e conferência.",
                "Ajuda em prestação de contas interna, mas depende de parametrização e uso adequado.",
            ),
            (
                "Financeiro e relatórios",
                "Dashboard acionável, vendas, estoque, recebíveis, despesas, indicadores, relatórios gerais e exportações futuras.",
                "Definir limites de precisão, responsabilidade do usuário pelos dados inseridos e escopo de relatórios.",
            ),
            (
                "Setores e módulos específicos",
                "Templates por segmento, telas e campos específicos por setor, começando por ótica com prescrições, ordens de serviço e fluxos próprios.",
                "Cliente escolhe setor/plano; controle fino fica no painel master.",
            ),
        ],
        [1700, 5100, 2560],
    )

    add_heading(doc, "4.2 Volt Chat", 2)
    add_para(
        doc,
        "Volt Chat está em planejamento/implementação como produto de comunicação empresarial dentro da Volt Corp. "
        "A visão inicial é oferecer chat corporativo para equipes, atendimento e colaboração, com controle de empresa, "
        "usuários, permissões e integração futura com Volt Core.",
    )
    add_table(
        doc,
        ["Função prevista", "Descrição", "Ponto jurídico"],
        [
            (
                "Conversas internas",
                "Chats por equipe, canais, mensagens diretas, histórico e busca.",
                "Definir política de retenção, exclusão, auditoria e titularidade das mensagens.",
            ),
            (
                "Atendimento / relacionamento",
                "Possível uso para comunicação com clientes ou atendimento interno/externo.",
                "Se houver dados de clientes, aplicar LGPD e consentimento/canais oficiais.",
            ),
            (
                "Arquivos e notificações",
                "Envio de anexos, alertas e notificações operacionais.",
                "Definir limites de armazenamento, conteúdos proibidos e responsabilidade do cliente.",
            ),
            (
                "Integração com Volt Core",
                "Vincular conversas a clientes, vendas, ordens de serviço ou tarefas.",
                "Revisar compartilhamento de dados entre produtos e permissões por perfil.",
            ),
        ],
        [2200, 4200, 2960],
    )

    add_heading(doc, "4.3 Estoque virtual 3D e controle de estoque", 2)
    add_para(
        doc,
        "Há intenção de desenvolver um produto ou módulo de estoque virtual 3D, possivelmente dentro da Volt Corp "
        "e integrado ao Volt Core. A proposta é representar visualmente estoque, posições, prateleiras, boxes, áreas, "
        "endereços e movimentações, oferecendo uma camada mais intuitiva de operação para empresas que precisam "
        "localizar itens com rapidez.",
    )
    add_table(
        doc,
        ["Recurso futuro", "Descrição", "Decisão estratégica pendente"],
        [
            (
                "Mapa 3D de estoque",
                "Visualização de áreas, prateleiras, endereços, posições e ocupação.",
                "Definir se será produto independente, módulo premium do Volt Core ou extensão setorial.",
            ),
            (
                "Localização de produtos",
                "Busca por SKU/produto e indicação visual da posição física.",
                "Definir responsabilidade do cliente pela atualização das posições físicas.",
            ),
            (
                "Inventário guiado",
                "Contagens, divergências, conferência por localização e ajustes.",
                "Definir trilha de auditoria e limites de responsabilização por perdas/divergências.",
            ),
            (
                "Integração com estoque",
                "Sincronização com entradas, saídas, vendas, compras e ajustes do Volt Core.",
                "Evitar conflito de marca/escopo com o núcleo de estoque padrão.",
            ),
        ],
        [2200, 4200, 2960],
    )

    add_heading(doc, "5. Matriz consolidada de status", 1)
    add_table(
        doc,
        ["Produto / módulo", "Linha", "Status atual", "Observação"],
        [
            ("Mercado Livre", "Davantti Seller", "Existente/em evolução", "Módulo maduro com várias telas, serviços, jobs e relatórios."),
            ("Shopee", "Davantti Seller", "Existente/em evolução", "Módulo com pedidos, produtos, anúncios, campanhas, métricas, logística e webhooks."),
            ("MadeiraMadeira", "Davantti Seller", "Incluso / existente", "Módulo dedicado com autenticação, workspace, catálogo, usuários e sincronização."),
            ("Magalu", "Davantti Seller", "Previsto / em inclusão", "Escopo funcional e termos de integração ainda devem ser fechados."),
            ("Avantracking / Logisync", "Davantti Seller / logística", "Existente/em evolução", "Rastreamento, integrações logísticas, status, frete e relatórios."),
            ("LeaderSku", "Davantti Seller / ferramenta auxiliar", "Existente", "Pode permanecer como ferramenta auxiliar ou virar módulo comercial."),
            ("Volt Core", "Volt Corp", "Planejamento/implementação", "Gestão empresarial, vendas, estoque, caixa, recebíveis, relatórios e módulos por setor."),
            ("Volt Chat", "Volt Corp", "Planejamento/implementação", "Chat corporativo, comunicação interna/externa, anexos, histórico e integração futura."),
            ("Estoque virtual 3D", "Volt Corp", "Conceito/roadmap", "Produto ou módulo para visualização e controle físico/virtual de estoque."),
        ],
        [2300, 2100, 1900, 3060],
    )

    add_heading(doc, "6. Pontos de atenção para revisão jurídica", 1)
    legal_points = [
        ("Marcas e nomes", "Validar disponibilidade e estratégia de proteção para Davantti, Davantti Seller, Volt Corp, Volt Core, Volt Chat e nome futuro do estoque 3D."),
        ("Contratos SaaS", "Criar ou revisar termos de uso, contrato de assinatura, política de cancelamento, SLA quando aplicável e limites de responsabilidade."),
        ("LGPD e privacidade", "Mapear dados pessoais tratados: clientes, compradores, vendedores, usuários internos, mensagens, pedidos, notas, rastreios, logs e dados financeiros."),
        ("Integrações com marketplaces", "Validar termos de uso/API de Mercado Livre, Shopee, MadeiraMadeira, Magalu e demais plataformas integradas."),
        ("Fiscal e recibos", "Distinguir recibo não fiscal de nota/cupom fiscal; só prometer emissão fiscal quando houver implementação e homologação compatíveis."),
        ("Financeiro e recebíveis", "Deixar claro quando o sistema apenas registra recebíveis, cheques, promissórias e cartões, sem atuar como banco, instituição de pagamento ou garantidor de recebimento."),
        ("Comunicação e chat", "Definir responsabilidade sobre conteúdo de mensagens, anexos, retenção, auditoria, exclusão e acesso por administradores."),
        ("Propriedade intelectual", "Formalizar titularidade do código, layouts, marcas, textos, imagens, scripts, integrações, documentação e contribuições de terceiros."),
        ("Open source e dependências", "Revisar licenças de bibliotecas usadas no ecossistema e obrigações de distribuição/atribuição quando houver."),
        ("Domínios e subdomínios", "Definir estratégia de domínios para Davantti e Volt Corp, evitando conflito ou confusão comercial entre linhas."),
    ]
    add_table(doc, ["Tema", "O que revisar"], legal_points, [2500, 6860])

    add_heading(doc, "7. Recomendações críticas", 1)
    add_callout(
        doc,
        "Recomendação principal",
        "Não permitir que cada cliente monte livremente o sistema. O cliente escolhe setor e plano; o painel master "
        "controla módulos, telas, permissões e liberações finas. Isso reduz complexidade técnica, risco de suporte e "
        "promessas comerciais difíceis de sustentar.",
    )
    add_para(
        doc,
        "Para o jurídico, é importante separar claramente: (i) produtos já existentes, (ii) módulos em implantação, "
        "(iii) funcionalidades planejadas e (iv) integrações sujeitas a termos de terceiros. Essa separação evita "
        "promessa comercial antecipada e ajuda na elaboração de contratos e materiais públicos.",
    )
    add_para(
        doc,
        "Também é recomendável tratar Volt Corp como linha empresarial com identidade própria, mas dentro do ecossistema "
        "Davantti. Assim, a empresa pode vender produtos próprios sem diluir a Davantti Seller, que tem foco em marketplace.",
    )

    add_heading(doc, "8. Próximos passos sugeridos", 1)
    add_table(
        doc,
        ["Prioridade", "Ação", "Responsável sugerido"],
        [
            ("Alta", "Validar nomes/marcas e estratégia de registro para Davantti, Volt Corp, Volt Core e Volt Chat.", "Jurídico + direção"),
            ("Alta", "Criar termos de uso, política de privacidade/LGPD e contrato SaaS base.", "Jurídico"),
            ("Alta", "Definir linguagem comercial para módulos em planejamento, evitando promessas de função ainda não homologada.", "Jurídico + produto"),
            ("Média", "Revisar termos de API/marketplace para ML, Shopee, MadeiraMadeira, Magalu e integrações logísticas.", "Jurídico + técnico"),
            ("Média", "Definir se estoque virtual 3D será produto separado, módulo premium ou extensão do Volt Core.", "Produto + direção"),
            ("Média", "Definir cláusulas sobre dados, logs, auditoria, backups, retenção, exclusão e suporte.", "Jurídico + técnico"),
        ],
        [1400, 5700, 2260],
    )

    add_heading(doc, "Anexo A - Referências internas de implementação", 1)
    add_para(
        doc,
        "A estrutura atual observada no repositório indica módulos e pastas como ml, shopee, MadeiraMadeira, "
        "avantracking, logisync, LeaderSku e business/volt_core. O Volt Core já possui backend, frontend, migrations "
        "de banco, rotas de autenticação, serviço persistente, registries de módulos/telas/setores e documentação de handoff.",
    )
    add_para(
        doc,
        "Este anexo é apenas uma referência técnica de apoio. A advogada pode usar a matriz de módulos acima como base "
        "para perguntas jurídicas, contratos e políticas, sem depender da leitura do código-fonte.",
    )

    doc.core_properties.title = "Documento de Escopo - Davantti e Volt Corp"
    doc.core_properties.subject = "Produtos, módulos, funcionalidades e pontos de atenção jurídica"
    doc.core_properties.author = "Davantti"
    doc.core_properties.comments = "Documento gerado para revisão jurídica."
    doc.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    path = build_doc()
    print(path)
