from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Dict, Optional
from threading import Lock
from time import perf_counter
import asyncio
import base64
import csv
import difflib
import html
import hashlib
import hmac
import io
import json
import os
import secrets
import re
import tempfile
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
import uuid
import zipfile

import boto3
from botocore.config import Config as BotoConfig

from fastapi import Depends, FastAPI, File, HTTPException, Query, Request, UploadFile, WebSocket, WebSocketDisconnect
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse, Response, StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from jose import JWTError, jwt
from passlib.context import CryptContext
from sqlalchemy import JSON, BigInteger, Boolean, Column, DateTime, ForeignKey, Integer, LargeBinary, String, Text, UniqueConstraint, and_, create_engine, event, func, or_, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.declarative import declarative_base
from sqlalchemy.orm import defer, relationship, sessionmaker
from sqlalchemy.pool import NullPool
import uvicorn


SECRET_KEY = os.getenv("SECRET_KEY", "voltcorp_secret_key_super_secure_2024")
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = int(os.getenv("ACCESS_TOKEN_EXPIRE_MINUTES", "1440"))
REMEMBERED_ACCESS_TOKEN_EXPIRE_DAYS = int(os.getenv("REMEMBERED_ACCESS_TOKEN_EXPIRE_DAYS", "30"))
IS_PRODUCTION = os.getenv("ENVIRONMENT", "development").lower() == "production"
APP_VERSION = os.getenv("APP_VERSION") or os.getenv("RENDER_GIT_COMMIT") or "local"
APP_BUILD_TIME = os.getenv("APP_BUILD_TIME")
DESKTOP_RELEASE_EXTERNAL_URL = os.getenv("VOLT_CHAT_DESKTOP_DOWNLOAD_URL", "").strip()
R2_ACCOUNT_ID = (os.getenv("R2_ACCOUNT_ID") or os.getenv("CLOUDFLARE_R2_ACCOUNT_ID") or "").strip()
R2_ACCESS_KEY_ID = (os.getenv("R2_ACCESS_KEY_ID") or os.getenv("CLOUDFLARE_R2_ACCESS_KEY_ID") or "").strip()
R2_SECRET_ACCESS_KEY = (os.getenv("R2_SECRET_ACCESS_KEY") or os.getenv("CLOUDFLARE_R2_SECRET_ACCESS_KEY") or "").strip()
R2_BUCKET_NAME = (os.getenv("R2_BUCKET_NAME") or os.getenv("CLOUDFLARE_R2_BUCKET_NAME") or "").strip()
R2_ENDPOINT_URL = (
    os.getenv("R2_ENDPOINT_URL")
    or os.getenv("CLOUDFLARE_R2_ENDPOINT_URL")
    or (f"https://{R2_ACCOUNT_ID}.r2.cloudflarestorage.com" if R2_ACCOUNT_ID else "")
).strip().rstrip("/")
R2_RELEASE_PREFIX = (os.getenv("R2_RELEASE_PREFIX") or "desktop/releases").strip().strip("/")
R2_RELEASE_PLATFORM = "windows"
R2_RELEASE_MANIFEST_KEY = "/".join(
    part for part in (R2_RELEASE_PREFIX, R2_RELEASE_PLATFORM, "latest.json") if part
)
R2_CHAT_FILE_PREFIX = (os.getenv("R2_CHAT_FILE_PREFIX") or "chat/files").strip().strip("/")
CHAT_FILE_STORAGE_MODE = (os.getenv("VOLT_CHAT_FILE_STORAGE_MODE") or "database").strip().lower()
R2_CHAT_FILE_URI_PREFIX = "r2://"
try:
    R2_PRESIGN_EXPIRES_SECONDS = max(300, min(int(os.getenv("R2_PRESIGN_EXPIRES_SECONDS", "14400")), 604800))
except ValueError:
    R2_PRESIGN_EXPIRES_SECONDS = 14400
DESKTOP_PACKAGE_EXTERNAL_URL = (
    os.getenv("VOLT_CHAT_DESKTOP_PACKAGE_URL", "").strip()
    or DESKTOP_RELEASE_EXTERNAL_URL
)
# Accepts dates typed by people in Brazil: 24/09/03, 24-09-2003 and 24092003.
# The database keeps the normalized DD-MM-YY or DD-MM-YYYY form so a four-digit
# year remains available for accurate age calculations.
BREVO_API_KEY = os.getenv("BREVO_API_KEY", "").strip()
BREVO_SENDER_EMAIL = os.getenv("BREVO_SENDER_EMAIL", "").strip()
BREVO_SENDER_NAME = os.getenv("BREVO_SENDER_NAME", "VoltChat").strip() or "VoltChat"
PASSWORD_RESET_TOKEN_EXPIRE_MINUTES = max(10, min(int(os.getenv("PASSWORD_RESET_TOKEN_EXPIRE_MINUTES", "30")), 1440))
PASSWORD_RESET_URL = (
    os.getenv("VOLT_CHAT_PASSWORD_RESET_URL", "").strip()
    or "https://www.voltcorporation.com.br/chat/reset-password"
)

BIRTHDAY_PATTERN = re.compile(r"^(0[1-9]|[12][0-9]|3[01])[-/](0[1-9]|1[0-2])[-/](\d{2}|\d{4})$")


def normalize_birthday(value) -> Optional[str]:
    birthday = str(value or "").strip()
    if not birthday:
        return None
    digits = re.sub(r"\D", "", birthday)
    if birthday.isdigit() and len(digits) in {6, 8}:
        year_length = 2 if len(digits) == 6 else 4
        birthday = f"{digits[:2]}-{digits[2:4]}-{digits[4:4 + year_length]}"
    match = BIRTHDAY_PATTERN.fullmatch(birthday)
    if not match:
        raise ValueError("Data de nascimento invalida. Use DD-MM-YY ou DD-MM-YYYY.")
    day, month, year_text = match.groups()
    day, month = int(day), int(month)
    current_year = datetime.now().year
    if len(year_text) == 2:
        short_year = int(year_text)
        current_short_year = current_year % 100
        full_year = (2000 if short_year <= current_short_year else 1900) + short_year
    else:
        full_year = int(year_text)
    try:
        birth_date = datetime(full_year, month, day).date()
    except ValueError as exc:
        raise ValueError("Data de nascimento invalida. Use DD-MM-YY ou DD-MM-YYYY.") from exc
    today = datetime.now().date()
    age = today.year - birth_date.year - ((today.month, today.day) < (birth_date.month, birth_date.day))
    if birth_date > today or age > 120:
        raise ValueError("Data de nascimento fora do intervalo permitido.")
    normalized_year = f"{short_year:02d}" if len(year_text) == 2 else f"{full_year:04d}"
    return f"{day:02d}-{month:02d}-{normalized_year}"
def birthday_or_http_error(value) -> Optional[str]:
    try:
        return normalize_birthday(value)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc



DEFAULT_FRONTEND_ORIGINS = [
    "https://www.voltcorporation.com.br",
    "https://voltcorporation.com.br",
]
DEFAULT_DEPARTMENTS = ["TI", "Suporte", "Comercial", "Financeiro", "Operacao", "Produto"]
DEFAULT_COMPANY_ID = "00000000-0000-0000-0000-000000000001"
DEFAULT_COMPANY_NAME = "Empresa Padrao"
USER_LEVELS = {"usuario", "coordenador", "master", "padrao"}
COORDINATOR_LEVELS = {"coordenador", "master"}
PLATFORM_ROLES = {"master_admin", "company_admin", "coordinator", "user"}
TENANT_ADMIN_ROLES = {"company_admin", "master_admin"}
ACTIVE_STATUS = "active"
INACTIVE_STATUS = "inactive"
INLINE_CHAT_MESSAGE_MAX_LENGTH = 1000
EMAIL_PATTERN = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
ALLOWED_UPLOAD_EXTENSIONS = {
    ".csv",
    ".doc",
    ".docx",
    ".gif",
    ".m4a",
    ".mp3",
    ".ogg",
    ".jpeg",
    ".jpg",
    ".pdf",
    ".png",
    ".txt",
    ".webm",
    ".webp",
    ".wav",
    ".xls",
    ".xlsx",
    ".zip",
}
DEFAULT_STICKER_KEYS = {
    "bolt-celebrate", "bolt-thumbs-up", "bolt-thinking", "bolt-urgent",
    "bolt-working", "bolt-confused", "bolt-applause", "bolt-sleepy",
    "mitty-hello", "mitty-calendar", "mitty-coffee", "mitty-done",
    "mitty-meeting", "mitty-focused", "mitty-time",
    "bolt-mitty-celebrate", "bolt-mitty-teamwork", "bolt-mitty-hello",
    "bolt-mitty-planning", "bolt-mitty-done", "bolt-mitty-coffee",
    "volt-hello-v2", "volt-celebrate-v2", "volt-love-v2", "volt-surprised-v2", "volt-working-v2",
    "mitty-hello-v2", "mitty-celebrate-v2", "mitty-love-v2", "mitty-surprised-v2", "mitty-working-v2",
    "muricoca-hello-v2", "muricoca-celebrate-v2", "muricoca-love-v2", "muricoca-surprised-v2", "muricoca-working-v2",
    "zoe-hello-v2", "zoe-celebrate-v2", "zoe-love-v2", "zoe-surprised-v2", "zoe-working-v2",
    "zara-hello-v2", "zara-celebrate-v2", "zara-love-v2", "zara-surprised-v2", "zara-working-v2",
    "team-hello-v2", "team-celebrate-v2", "team-love-v2", "team-surprised-v2", "team-working-v2",
    "volt-ok-v3", "volt-no-v3", "mitty-ok-v3", "mitty-no-v3",
    "muricoca-ok-v3", "muricoca-no-v3", "zoe-ok-v3", "zoe-no-v3",
    "zara-ok-v3", "zara-no-v3", "team-ok-v3", "team-no-v3",
    "macaco-beleza", "macaco-cafe", "macaco-banana", "macaco-bravo", "macaco-risada",
}
AUDIT_RETENTION_DAYS = 7
CHAT_ARCHIVE_INTERVAL = timedelta(days=30)
DATABASE_MAINTENANCE_INTERVAL_SECONDS = 24 * 60 * 60
DATABASE_MAINTENANCE_HOUR_UTC = 6
DATABASE_MAINTENANCE_MINUTE_UTC = 15
FILE_CLEANUP_BATCH_SIZE = 100
BOLT_USERNAME = "bolt"
BOLT_NAME = "Volt"
MITTY_USERNAME = "mitty"
MITTY_NAME = "Mitty"
BOLT_DAILY_SUMMARY_HOUR_UTC = max(0, min(23, int(os.getenv("BOLT_DAILY_SUMMARY_HOUR_UTC", "12"))))
BOLT_DAILY_SUMMARY_TIMEZONE = timezone(timedelta(hours=-3))
TICKET_URGENCY_LEVELS = ["Leve", "Moderado", "Alto", "Urgente", "Extrema Urgência"]


def normalize_database_url(database_url: str) -> str:
    if database_url.startswith("postgres://"):
        database_url = database_url.replace("postgres://", "postgresql://", 1)
    if database_url.startswith("postgresql://"):
        database_url = database_url.replace("postgresql://", "postgresql+psycopg://", 1)
    return database_url


def normalize_db_pool_mode(pool_mode: str) -> str:
    normalized_mode = (pool_mode or "").strip().lower()
    if normalized_mode in {"nopooling", "no_pooling", "no-pooling", "nullpool"}:
        return "null"
    return normalized_mode or "small"


DATABASE_URL = normalize_database_url(os.getenv("DATABASE_URL", "sqlite:///./voltcorp.db"))
DB_POOL_MODE = (
    "sqlite"
    if DATABASE_URL.startswith("sqlite")
    else normalize_db_pool_mode(os.getenv("VOLT_CHAT_DB_POOL_MODE", "null" if "-pooler." in DATABASE_URL else "small"))
)
if DB_POOL_MODE == "sqlite":
    engine_options = {"connect_args": {"check_same_thread": False}}
elif DB_POOL_MODE == "null":
    # O endpoint Neon ja usa PgBouncer. Nao mantenha um segundo pool cliente
    # aberto, pois conexoes ociosas podem impedir o compute de escalar para zero.
    engine_options = {
        "poolclass": NullPool,
        "connect_args": {"application_name": "voltchat-api", "connect_timeout": 10},
    }
else:
    engine_options = {
        "pool_pre_ping": True,
        "pool_size": max(1, int(os.getenv("VOLT_CHAT_DB_POOL_SIZE", "2"))),
        "max_overflow": max(0, int(os.getenv("VOLT_CHAT_DB_MAX_OVERFLOW", "1"))),
        "pool_timeout": 10,
        "pool_recycle": 300,
        "connect_args": {"application_name": "voltchat-api", "connect_timeout": 10},
    }

DB_QUERY_METRICS_ENABLED = os.getenv("VOLT_CHAT_DB_QUERY_METRICS", "false").strip().lower() in {"1", "true", "yes"}
try:
    DB_QUERY_SLOW_MS = max(10, min(int(os.getenv("VOLT_CHAT_DB_QUERY_SLOW_MS", "250")), 60000))
except ValueError:
    DB_QUERY_SLOW_MS = 250
DB_QUERY_METRICS_LOCK = Lock()
DB_QUERY_METRICS = {"total_queries": 0, "slow_queries": 0, "max_duration_ms": 0.0}

engine = create_engine(DATABASE_URL, **engine_options)
if DB_QUERY_METRICS_ENABLED:
    @event.listens_for(engine, "before_cursor_execute")
    def track_database_query_start(conn, cursor, statement, parameters, context, executemany):
        context._voltchat_query_started_at = perf_counter()

    @event.listens_for(engine, "after_cursor_execute")
    def track_database_query_end(conn, cursor, statement, parameters, context, executemany):
        started_at = getattr(context, "_voltchat_query_started_at", None)
        if started_at is None:
            return
        duration_ms = (perf_counter() - started_at) * 1000
        with DB_QUERY_METRICS_LOCK:
            DB_QUERY_METRICS["total_queries"] += 1
            DB_QUERY_METRICS["max_duration_ms"] = max(DB_QUERY_METRICS["max_duration_ms"], duration_ms)
            if duration_ms >= DB_QUERY_SLOW_MS:
                DB_QUERY_METRICS["slow_queries"] += 1


def database_query_metrics() -> dict:
    with DB_QUERY_METRICS_LOCK:
        return {
            "enabled": DB_QUERY_METRICS_ENABLED,
            "slow_threshold_ms": DB_QUERY_SLOW_MS,
            "total_queries": DB_QUERY_METRICS["total_queries"],
            "slow_queries": DB_QUERY_METRICS["slow_queries"],
            "max_duration_ms": round(DB_QUERY_METRICS["max_duration_ms"], 2),
        }


SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()
pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")
security = HTTPBearer()


class User(Base):
    __tablename__ = "users"

    id = Column(Integer, primary_key=True, index=True)
    uuid = Column(String(36), default=lambda: str(uuid.uuid4()), unique=True, nullable=True, index=True)
    username = Column(String(80), unique=True, index=True, nullable=False)
    email = Column(String(255), unique=True, index=True, nullable=False)
    full_name = Column(String(255), nullable=False)
    nickname = Column(String(80), nullable=True, index=True)
    assistant_address_name = Column(String(80), nullable=True)
    hashed_password = Column(String(255), nullable=False)
    phone = Column(String(40), nullable=True)
    is_platform_admin = Column(Boolean, default=False, nullable=False)
    must_change_password = Column(Boolean, default=False, nullable=False)
    auth_version = Column(Integer, default=0, nullable=False)
    status = Column(String(40), default=ACTIVE_STATUS, nullable=False)
    access_level = Column(String(40), default="usuario", nullable=False)
    department = Column(String(100), nullable=True)
    phone_extension = Column(String(40), nullable=True)
    birthday = Column(String(10), nullable=True)
    profile_photo = Column(Text, nullable=True)
    role_title = Column(String(100), nullable=True)
    is_active = Column(Boolean, default=True, nullable=False)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)
    updated_at = Column(DateTime, nullable=True)

    sent_messages = relationship("Message", foreign_keys="Message.sender_id", back_populates="sender")
    received_messages = relationship("Message", foreign_keys="Message.receiver_id", back_populates="receiver")
    uploaded_files = relationship("FileUpload", back_populates="uploader")


class PasswordResetToken(Base):
    __tablename__ = "password_reset_tokens"

    id = Column(Integer, primary_key=True, autoincrement=True)
    user_id = Column(Integer, ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    token_hash = Column(String(64), unique=True, nullable=False, index=True)
    expires_at = Column(DateTime(timezone=True), nullable=False, index=True)
    used_at = Column(DateTime(timezone=True), nullable=True, index=True)
    created_at = Column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), nullable=False)


class Company(Base):
    __tablename__ = "companies"

    id = Column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    name = Column(String(255), nullable=False, index=True)
    cnpj = Column(String(32), nullable=True, unique=True, index=True)
    responsible_name = Column(String(255), nullable=True)
    phone_primary = Column(String(40), nullable=True)
    phone_secondary = Column(String(40), nullable=True)
    status = Column(String(40), default=ACTIVE_STATUS, nullable=False, index=True)
    tenant_global_id = Column(String(64), nullable=True, unique=True, index=True)
    hub_enabled = Column(Boolean, default=False, nullable=False)
    hub_synced_at = Column(DateTime, nullable=True)
    general_chat_name = Column(String(120), nullable=True)
    allow_user_sticker_creation = Column(Boolean, default=True, nullable=False)
    allow_user_message_editing = Column(Boolean, default=False, nullable=False)
    history_cleared_at = Column(DateTime, nullable=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)
    updated_at = Column(DateTime, nullable=True)


class Department(Base):
    __tablename__ = "departments"

    id = Column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    company_id = Column(String(36), ForeignKey("companies.id"), nullable=False, index=True)
    name = Column(String(120), nullable=False, index=True)
    description = Column(Text, nullable=True)
    status = Column(String(40), default=ACTIVE_STATUS, nullable=False, index=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)


class DepartmentLink(Base):
    __tablename__ = "department_links"

    id = Column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    company_id = Column(String(36), ForeignKey("companies.id"), nullable=False, index=True)
    source_department_id = Column(String(36), ForeignKey("departments.id"), nullable=False, index=True)
    target_department_id = Column(String(36), ForeignKey("departments.id"), nullable=False, index=True)
    label = Column(String(120), nullable=True)
    created_by_id = Column(Integer, ForeignKey("users.id"), nullable=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)


class CompanyUser(Base):
    __tablename__ = "company_users"

    id = Column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    company_id = Column(String(36), ForeignKey("companies.id"), nullable=False, index=True)
    user_id = Column(Integer, ForeignKey("users.id"), nullable=False, index=True)
    department_id = Column(String(36), ForeignKey("departments.id"), nullable=True, index=True)
    role = Column(String(40), default="user", nullable=False, index=True)
    status = Column(String(40), default=ACTIVE_STATUS, nullable=False, index=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)


class AuditLog(Base):
    __tablename__ = "audit_logs"

    id = Column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    company_id = Column(String(36), ForeignKey("companies.id"), nullable=True, index=True)
    actor_user_id = Column(Integer, ForeignKey("users.id"), nullable=True, index=True)
    action = Column(String(100), nullable=False, index=True)
    entity_type = Column(String(100), nullable=False, index=True)
    entity_id = Column(String(80), nullable=True, index=True)
    metadata_json = Column("metadata", JSON, nullable=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False, index=True)


class Ticket(Base):
    __tablename__ = "tickets"

    id = Column(Integer, primary_key=True, index=True)
    company_id = Column(String(36), ForeignKey("companies.id"), default=DEFAULT_COMPANY_ID, nullable=True, index=True)
    title = Column(String(200), nullable=False, index=True)
    description = Column(Text, nullable=False)
    priority = Column(String(40), default="Media", nullable=False)
    delivery_due_at = Column(DateTime, nullable=True, index=True)
    status = Column(String(40), default="Aberto", nullable=False)
    department = Column(String(100), nullable=True, index=True)
    image_data = Column(Text, nullable=True)
    channel = Column(String(80), default="Web", nullable=False)
    created_by_id = Column(Integer, ForeignKey("users.id"), nullable=False)
    assigned_to_id = Column(Integer, ForeignKey("users.id"), nullable=True)
    attachment_file_id = Column(Integer, ForeignKey("file_uploads.id"), nullable=True)
    rating_score = Column(Integer, nullable=True)
    rating_comment = Column(Text, nullable=True)
    rated_at = Column(DateTime, nullable=True)
    first_response_at = Column(DateTime, nullable=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)
    updated_at = Column(DateTime, nullable=True)
    closed_at = Column(DateTime, nullable=True)
    closed_by_id = Column(Integer, ForeignKey("users.id"), nullable=True, index=True)
    close_reason = Column(Text, nullable=True)


class TicketAssignee(Base):
    __tablename__ = "ticket_assignees"

    id = Column(Integer, primary_key=True, index=True)
    ticket_id = Column(Integer, ForeignKey("tickets.id"), nullable=False, index=True)
    user_id = Column(Integer, ForeignKey("users.id"), nullable=False, index=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)


class TicketDepartment(Base):
    __tablename__ = "ticket_departments"

    id = Column(Integer, primary_key=True, index=True)
    ticket_id = Column(Integer, ForeignKey("tickets.id"), nullable=False, index=True)
    department_id = Column(String(36), ForeignKey("departments.id"), nullable=False, index=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)


class ChatGroup(Base):
    __tablename__ = "chat_groups"

    id = Column(Integer, primary_key=True, index=True)
    company_id = Column(String(36), ForeignKey("companies.id"), default=DEFAULT_COMPANY_ID, nullable=True, index=True)
    name = Column(String(120), nullable=False, index=True)
    description = Column(Text, nullable=True)
    department = Column(String(100), nullable=True, index=True)
    image_data = Column(Text, nullable=True)
    created_by_id = Column(Integer, ForeignKey("users.id"), nullable=False)
    is_active = Column(Boolean, default=True, nullable=False)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)


class ChatGroupMember(Base):
    __tablename__ = "chat_group_members"

    id = Column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    group_id = Column(Integer, ForeignKey("chat_groups.id", ondelete="CASCADE"), nullable=False, index=True)
    user_id = Column(Integer, ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)


class TaskItem(Base):
    __tablename__ = "tasks"

    id = Column(Integer, primary_key=True, index=True)
    company_id = Column(String(36), ForeignKey("companies.id"), default=DEFAULT_COMPANY_ID, nullable=True, index=True)
    title = Column(String(200), nullable=False, index=True)
    description = Column(Text, nullable=True)
    priority = Column(String(40), default="medium", nullable=False)
    category = Column(String(100), default="Operacao", nullable=False)
    status = Column(String(40), default="backlog", nullable=False, index=True)
    due_date = Column(String(10), nullable=True)
    created_by_id = Column(Integer, ForeignKey("users.id"), nullable=False)
    assigned_to_id = Column(Integer, ForeignKey("users.id"), nullable=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)
    updated_at = Column(DateTime, nullable=True)


class Message(Base):
    __tablename__ = "messages"

    id = Column(Integer, primary_key=True, index=True)
    company_id = Column(String(36), ForeignKey("companies.id"), default=DEFAULT_COMPANY_ID, nullable=True, index=True)
    content = Column(Text, nullable=False)
    sender_id = Column(Integer, ForeignKey("users.id"), nullable=False)
    receiver_id = Column(Integer, ForeignKey("users.id"), nullable=True)
    group_id = Column(Integer, ForeignKey("chat_groups.id", ondelete="CASCADE"), nullable=True, index=True)
    reply_to_id = Column(Integer, ForeignKey("messages.id"), nullable=True, index=True)
    message_type = Column(String(40), default="text", nullable=False)
    file_path = Column(String(500), nullable=True)
    timestamp = Column(DateTime, default=datetime.utcnow, nullable=False)
    is_read = Column(Boolean, default=False, nullable=False)
    archive_uid = Column(String(64), nullable=True, unique=True, index=True, default=lambda: uuid.uuid4().hex)

    sender = relationship("User", foreign_keys=[sender_id], back_populates="sent_messages")
    receiver = relationship("User", foreign_keys=[receiver_id], back_populates="received_messages")


class MessageReadReceipt(Base):
    __tablename__ = "message_read_receipts"
    __table_args__ = (
        UniqueConstraint("message_id", "user_id", name="uq_message_read_receipt_user"),
    )

    id = Column(Integer, primary_key=True, index=True)
    message_id = Column(Integer, ForeignKey("messages.id", ondelete="CASCADE"), nullable=False, index=True)
    user_id = Column(Integer, ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    read_at = Column(DateTime, default=datetime.utcnow, nullable=False)


class MessageEdit(Base):
    __tablename__ = "message_edits"

    id = Column(Integer, primary_key=True, index=True)
    message_id = Column(Integer, ForeignKey("messages.id", ondelete="CASCADE"), nullable=False, index=True)
    editor_user_id = Column(Integer, ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    previous_content = Column(Text, nullable=False)
    edited_at = Column(DateTime, default=datetime.utcnow, nullable=False, index=True)


class MessageMention(Base):
    __tablename__ = "message_mentions"
    __table_args__ = (
        UniqueConstraint("message_id", "user_id", name="uq_message_mention_user"),
    )

    id = Column(Integer, primary_key=True, index=True)
    message_id = Column(Integer, ForeignKey("messages.id", ondelete="CASCADE"), nullable=False, index=True)
    user_id = Column(Integer, ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)
    read_at = Column(DateTime, nullable=True, index=True)


class MessageReaction(Base):
    __tablename__ = "message_reactions"
    __table_args__ = (
        UniqueConstraint("message_id", "user_id", "emoji", name="uq_message_reaction_user_emoji"),
    )

    id = Column(Integer, primary_key=True, index=True)
    message_id = Column(Integer, ForeignKey("messages.id", ondelete="CASCADE"), nullable=False, index=True)
    user_id = Column(Integer, ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    emoji = Column(String(32), nullable=False)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)


class ChatSticker(Base):
    __tablename__ = "chat_stickers"

    id = Column(Integer, primary_key=True, index=True)
    company_id = Column(String(36), ForeignKey("companies.id", ondelete="CASCADE"), nullable=False, index=True)
    owner_user_id = Column(Integer, ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    name = Column(String(80), nullable=False)
    image_data = Column(LargeBinary, nullable=False)
    content_type = Column(String(80), default="image/webp", nullable=False)
    file_size = Column(Integer, nullable=False)
    source_reference = Column(String(160), nullable=True, index=True)
    is_active = Column(Boolean, default=True, nullable=False, index=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)
    updated_at = Column(DateTime, nullable=True)

class ChatAttentionLimit(Base):
    __tablename__ = "chat_attention_limits"
    __table_args__ = (
        UniqueConstraint(
            "company_id",
            "sender_id",
            "receiver_id",
            name="uq_chat_attention_sender_receiver",
        ),
    )

    id = Column(Integer, primary_key=True, index=True)
    company_id = Column(String(36), ForeignKey("companies.id", ondelete="CASCADE"), nullable=False, index=True)
    sender_id = Column(Integer, ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    receiver_id = Column(Integer, ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    usage_date = Column(String(10), nullable=False, index=True)
    daily_count = Column(Integer, default=0, nullable=False)
    last_sent_at = Column(DateTime, nullable=True)
    blocked_until = Column(DateTime, nullable=True, index=True)
    updated_at = Column(DateTime, default=datetime.utcnow, nullable=False)


class BoltDailyTicketSummary(Base):
    __tablename__ = "bolt_daily_ticket_summaries"
    __table_args__ = (
        UniqueConstraint("company_id", "user_id", "summary_date", name="uq_bolt_daily_ticket_summary"),
    )

    id = Column(Integer, primary_key=True, index=True)
    company_id = Column(String(36), ForeignKey("companies.id", ondelete="CASCADE"), nullable=False, index=True)
    user_id = Column(Integer, ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    summary_date = Column(String(10), nullable=False, index=True)
    open_ticket_count = Column(Integer, default=0, nullable=False)
    awaiting_reply_count = Column(Integer, default=0, nullable=False)
    message_id = Column(Integer, ForeignKey("messages.id", ondelete="SET NULL"), nullable=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)


class ChatHistoryBackup(Base):
    __tablename__ = "chat_history_backups"

    id = Column(Integer, primary_key=True, index=True)
    company_id = Column(String(36), ForeignKey("companies.id"), nullable=False, index=True)
    period_start = Column(DateTime, nullable=False)
    period_end = Column(DateTime, nullable=False)
    message_count = Column(Integer, nullable=False, default=0)
    original_size = Column(BigInteger, nullable=False, default=0)
    compressed_size = Column(BigInteger, nullable=False, default=0)
    compressed_data = Column(LargeBinary, nullable=False)
    format_version = Column(Integer, nullable=False, default=2)
    checksum_sha256 = Column(String(64), nullable=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)


class TicketMessage(Base):
    __tablename__ = "ticket_messages"

    id = Column(Integer, primary_key=True, index=True)
    company_id = Column(String(36), ForeignKey("companies.id"), default=DEFAULT_COMPANY_ID, nullable=True, index=True)
    ticket_id = Column(Integer, ForeignKey("tickets.id"), nullable=False, index=True)
    sender_id = Column(Integer, ForeignKey("users.id"), nullable=False, index=True)
    content = Column(Text, default="", nullable=False)
    file_id = Column(Integer, ForeignKey("file_uploads.id"), nullable=True)
    reply_to_id = Column(Integer, ForeignKey("ticket_messages.id", ondelete="SET NULL"), nullable=True, index=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)


class TicketMessageReaction(Base):
    __tablename__ = "ticket_message_reactions"
    __table_args__ = (
        UniqueConstraint("ticket_message_id", "user_id", "emoji", name="uq_ticket_message_reaction_user_emoji"),
    )

    id = Column(Integer, primary_key=True, index=True)
    ticket_message_id = Column(Integer, ForeignKey("ticket_messages.id", ondelete="CASCADE"), nullable=False, index=True)
    user_id = Column(Integer, ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    emoji = Column(String(32), nullable=False)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)


class FileUpload(Base):
    __tablename__ = "file_uploads"

    id = Column(Integer, primary_key=True, index=True)
    company_id = Column(String(36), ForeignKey("companies.id"), default=DEFAULT_COMPANY_ID, nullable=True, index=True)
    filename = Column(String(255), nullable=False)
    file_path = Column(String(500), nullable=False)
    file_size = Column(Integer, nullable=False)
    content_type = Column(String(120), nullable=False)
    binary_data = Column(LargeBinary, nullable=True)
    uploaded_by = Column(Integer, ForeignKey("users.id"), nullable=False)
    upload_date = Column(DateTime, default=datetime.utcnow, nullable=False)
    expires_at = Column(DateTime, nullable=True, index=True)

    uploader = relationship("User", back_populates="uploaded_files")



def sqlite_table_columns(conn, table_name: str) -> set[str]:
    rows = conn.exec_driver_sql(f"PRAGMA table_info({table_name})").fetchall()
    return {row[1] for row in rows}


def sqlite_add_column_if_missing(conn, table_name: str, column_name: str, ddl: str):
    if column_name not in sqlite_table_columns(conn, table_name):
        conn.exec_driver_sql(f"ALTER TABLE {table_name} ADD COLUMN {ddl}")


def ensure_sqlite_multi_tenant_schema():
    with engine.begin() as conn:
        user_columns = sqlite_table_columns(conn, "users")
        if "uuid" not in user_columns:
            conn.exec_driver_sql("ALTER TABLE users ADD COLUMN uuid VARCHAR(36)")
            rows = conn.exec_driver_sql("SELECT id FROM users WHERE uuid IS NULL").fetchall()
            for row in rows:
                conn.execute(text("UPDATE users SET uuid = :uuid WHERE id = :id"), {"uuid": str(uuid.uuid4()), "id": row[0]})
        if "phone" not in user_columns:
            conn.exec_driver_sql("ALTER TABLE users ADD COLUMN phone VARCHAR(40)")
        if "is_platform_admin" not in user_columns:
            conn.exec_driver_sql("ALTER TABLE users ADD COLUMN is_platform_admin BOOLEAN DEFAULT 0 NOT NULL")
        if "must_change_password" not in user_columns:
            conn.exec_driver_sql("ALTER TABLE users ADD COLUMN must_change_password BOOLEAN DEFAULT 0 NOT NULL")
        if "auth_version" not in user_columns:
            conn.exec_driver_sql("ALTER TABLE users ADD COLUMN auth_version INTEGER DEFAULT 0 NOT NULL")
        if "status" not in user_columns:
            conn.exec_driver_sql("ALTER TABLE users ADD COLUMN status VARCHAR(40) DEFAULT 'active' NOT NULL")
        if "profile_photo" not in user_columns:
            conn.exec_driver_sql("ALTER TABLE users ADD COLUMN profile_photo TEXT")
        if "nickname" not in user_columns:
            conn.exec_driver_sql("ALTER TABLE users ADD COLUMN nickname VARCHAR(80)")

        if "assistant_address_name" not in user_columns:
            conn.exec_driver_sql("ALTER TABLE users ADD COLUMN assistant_address_name VARCHAR(80)")

        sqlite_add_column_if_missing(conn, "companies", "tenant_global_id", "tenant_global_id VARCHAR(64)")
        sqlite_add_column_if_missing(conn, "companies", "general_chat_name", "general_chat_name VARCHAR(120)")
        sqlite_add_column_if_missing(conn, "companies", "allow_user_sticker_creation", "allow_user_sticker_creation BOOLEAN DEFAULT 1 NOT NULL")
        sqlite_add_column_if_missing(conn, "companies", "allow_user_message_editing", "allow_user_message_editing BOOLEAN DEFAULT 0 NOT NULL")
        sqlite_add_column_if_missing(conn, "companies", "history_cleared_at", "history_cleared_at DATETIME")
        sqlite_add_column_if_missing(conn, "chat_groups", "image_data", "image_data TEXT")
        sqlite_add_column_if_missing(conn, "tickets", "image_data", "image_data TEXT")
        sqlite_add_column_if_missing(conn, "tickets", "delivery_due_at", "delivery_due_at DATETIME")
        sqlite_add_column_if_missing(conn, "companies", "hub_enabled", "hub_enabled BOOLEAN DEFAULT 0 NOT NULL")
        sqlite_add_column_if_missing(conn, "companies", "hub_synced_at", "hub_synced_at DATETIME")
        conn.exec_driver_sql(
            "UPDATE companies SET tenant_global_id = id WHERE tenant_global_id IS NULL OR tenant_global_id = ''"
        )

        for table_name in ["tickets", "chat_groups", "tasks", "messages", "ticket_messages", "file_uploads"]:
            if table_name in {"ticket_messages", "tasks", "chat_groups", "tickets", "messages", "file_uploads"}:
                sqlite_add_column_if_missing(conn, table_name, "company_id", "company_id VARCHAR(36)")
                conn.execute(
                    text(f"UPDATE {table_name} SET company_id = :company_id WHERE company_id IS NULL"),
                    {"company_id": DEFAULT_COMPANY_ID},
                )
        sqlite_add_column_if_missing(conn, "messages", "archive_uid", "archive_uid VARCHAR(64)")
        sqlite_add_column_if_missing(conn, "chat_history_backups", "format_version", "format_version INTEGER DEFAULT 1 NOT NULL")
        sqlite_add_column_if_missing(conn, "chat_history_backups", "checksum_sha256", "checksum_sha256 VARCHAR(64)")
        conn.exec_driver_sql("UPDATE messages SET archive_uid = lower(hex(randomblob(16))) WHERE archive_uid IS NULL")
        conn.exec_driver_sql("CREATE UNIQUE INDEX IF NOT EXISTS ux_messages_archive_uid ON messages (archive_uid) WHERE archive_uid IS NOT NULL")

        sqlite_add_column_if_missing(conn, "file_uploads", "expires_at", "expires_at DATETIME")
        conn.exec_driver_sql(
            "UPDATE file_uploads SET expires_at = datetime(COALESCE(upload_date, CURRENT_TIMESTAMP), '+15 days') WHERE expires_at IS NULL"
        )
        conn.exec_driver_sql("""
            CREATE TABLE IF NOT EXISTS message_edits (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                message_id INTEGER NOT NULL,
                editor_user_id INTEGER NOT NULL,
                previous_content TEXT NOT NULL,
                edited_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(message_id) REFERENCES messages(id) ON DELETE CASCADE,
                FOREIGN KEY(editor_user_id) REFERENCES users(id) ON DELETE CASCADE
            )
        """)
        conn.exec_driver_sql("CREATE INDEX IF NOT EXISTS ix_message_edits_message ON message_edits(message_id, edited_at)")
        conn.exec_driver_sql("""
            CREATE TABLE IF NOT EXISTS message_mentions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                message_id INTEGER NOT NULL,
                user_id INTEGER NOT NULL,
                created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                read_at DATETIME NULL,
                FOREIGN KEY(message_id) REFERENCES messages(id) ON DELETE CASCADE,
                FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE,
                UNIQUE(message_id, user_id)
            )
        """)
        conn.exec_driver_sql("CREATE INDEX IF NOT EXISTS ix_message_mentions_user_unread ON message_mentions(user_id, read_at)")
        conn.exec_driver_sql("""
            CREATE TABLE IF NOT EXISTS chat_stickers (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              company_id VARCHAR(36) NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
              owner_user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
              name VARCHAR(80) NOT NULL,
              image_data BLOB NOT NULL,
              content_type VARCHAR(80) NOT NULL DEFAULT 'image/webp',
              file_size INTEGER NOT NULL,
              source_reference VARCHAR(160),
              is_active BOOLEAN NOT NULL DEFAULT 1,
              created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
              updated_at DATETIME
            )
        """)
        sqlite_add_column_if_missing(conn, "chat_stickers", "source_reference", "source_reference VARCHAR(160)")
        conn.exec_driver_sql("CREATE INDEX IF NOT EXISTS ix_chat_stickers_owner ON chat_stickers(company_id, owner_user_id, is_active)")
        conn.exec_driver_sql(
            "CREATE UNIQUE INDEX IF NOT EXISTS ux_chat_stickers_favorite_source "
            "ON chat_stickers(company_id, owner_user_id, source_reference) "
            "WHERE source_reference IS NOT NULL AND is_active = 1"
        )
        conn.exec_driver_sql("""
            CREATE TABLE IF NOT EXISTS meetings (
              id VARCHAR(36) PRIMARY KEY,
              company_id VARCHAR(36) NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
              creator_user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
              title VARCHAR(180) NOT NULL,
              description TEXT,
              meeting_type VARCHAR(30) NOT NULL DEFAULT 'reminder',
              starts_at DATETIME NOT NULL,
              ends_at DATETIME NOT NULL,
              color VARCHAR(20) NOT NULL DEFAULT '#2563eb',
              link_url TEXT,
              status VARCHAR(30) NOT NULL DEFAULT 'scheduled',
              reminder_day_sent BOOLEAN NOT NULL DEFAULT 0,
              reminder_30m_sent BOOLEAN NOT NULL DEFAULT 0,
              start_alert_sent BOOLEAN NOT NULL DEFAULT 0,
              created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
              updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
              CHECK (meeting_type IN ('reminder', 'task', 'video')),
              CHECK (ends_at > starts_at)
            )
        """)
        conn.exec_driver_sql("""
            CREATE TABLE IF NOT EXISTS meeting_participants (
              meeting_id VARCHAR(36) NOT NULL REFERENCES meetings(id) ON DELETE CASCADE,
              user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
              response_status VARCHAR(30) NOT NULL DEFAULT 'pending',
              joined_at DATETIME,
              left_at DATETIME,
              PRIMARY KEY (meeting_id, user_id)
            )
        """)
        conn.exec_driver_sql("CREATE INDEX IF NOT EXISTS ix_meetings_company_starts ON meetings(company_id, starts_at)")
        conn.exec_driver_sql("CREATE INDEX IF NOT EXISTS ix_meeting_participants_user ON meeting_participants(user_id, meeting_id)")


if DATABASE_URL.startswith("sqlite"):
    Base.metadata.create_all(bind=engine)
    ensure_sqlite_multi_tenant_schema()


def get_cors_origins():
    configured = os.getenv("FRONTEND_ORIGINS", "")
    origins = list(DEFAULT_FRONTEND_ORIGINS)
    if configured:
        origins.extend(origin.strip() for origin in configured.split(",") if origin.strip())
    if not IS_PRODUCTION:
        origins.extend(["http://127.0.0.1:3000", "http://localhost:3000"])
    return sorted(set(origins))


app = FastAPI(title="Volt Corp API", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=get_cors_origins(),
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.middleware("http")
async def enforce_required_password_change(request, call_next):
    allowed_paths = {
        "/",
        "/health",
        "/version",
        "/auth/login",
        "/auth/logout",
        "/auth/me",
        "/auth/change-password",
        "/auth/forgot-password",
        "/auth/reset-password",
    }
    path = request.url.path
    if path in allowed_paths or path.startswith("/downloads/"):
        return await call_next(request)

    authorization = request.headers.get("authorization") or ""
    if authorization.lower().startswith("bearer "):
        token = authorization.split(" ", 1)[1].strip()
        try:
            payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
            user_id = int(payload.get("sub"))
            with SessionLocal() as db:
                user = db.query(User).filter(User.id == user_id).first()
                if user:
                    db.expunge(user)
                    request.state.voltchat_current_user = user
                    if user.must_change_password:
                        return JSONResponse(
                            status_code=403,
                            content={
                                "detail": "Troca de senha obrigatoria no primeiro acesso.",
                                "must_change_password": True,
                            },
                        )
        except (JWTError, TypeError, ValueError):
            pass

    return await call_next(request)


def get_password_hash(password: str) -> str:
    return pwd_context.hash(password)


def verify_password(plain_password: str, hashed_password: str) -> bool:
    return pwd_context.verify(plain_password, hashed_password)


def password_reset_token_hash(raw_token: str) -> str:
    return hashlib.sha256(raw_token.encode("utf-8")).hexdigest()


def send_password_reset_email(user: User, raw_token: str) -> None:
    if not BREVO_API_KEY or not BREVO_SENDER_EMAIL:
        raise RuntimeError("Recuperacao de senha nao configurada no Brevo.")

    separator = "&" if "?" in PASSWORD_RESET_URL else "?"
    reset_link = f"{PASSWORD_RESET_URL}{separator}{urllib.parse.urlencode({'token': raw_token})}"
    safe_name = html.escape(display_user_name(user) or "usuario")
    safe_link = html.escape(reset_link, quote=True)
    payload = {
        "sender": {"name": BREVO_SENDER_NAME, "email": BREVO_SENDER_EMAIL},
        "to": [{"email": user.email, "name": display_user_name(user) or user.email}],
        "subject": "Redefinicao de senha - VoltChat",
        "htmlContent": (
            f"<p>Ola, {safe_name}.</p>"
            "<p>Recebemos uma solicitacao para redefinir sua senha do VoltChat.</p>"
            f"<p><a href=\"{safe_link}\">Redefinir minha senha</a></p>"
            f"<p>Este link expira em {PASSWORD_RESET_TOKEN_EXPIRE_MINUTES} minutos e pode ser usado uma unica vez.</p>"
            "<p>Se voce nao solicitou a redefinicao, ignore este e-mail.</p>"
        ),
    }
    request = urllib.request.Request(
        "https://api.brevo.com/v3/smtp/email",
        data=json.dumps(payload).encode("utf-8"),
        headers={"accept": "application/json", "content-type": "application/json", "api-key": BREVO_API_KEY},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=12):
            pass
    except (urllib.error.HTTPError, urllib.error.URLError, TimeoutError) as exc:
        raise RuntimeError(f"Falha ao enviar redefinicao de senha pelo Brevo: {exc}") from exc



def create_access_token(data: dict, expires_delta: Optional[timedelta] = None) -> str:
    to_encode = data.copy()
    expire = datetime.utcnow() + (expires_delta or timedelta(minutes=15))
    to_encode.update({"exp": expire})
    return jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)


def get_current_user(
    request: Request,
    credentials: HTTPAuthorizationCredentials = Depends(security),
):
    try:
        payload = jwt.decode(credentials.credentials, SECRET_KEY, algorithms=[ALGORITHM])
        user_id = int(payload.get("sub"))
        token_auth_version = int(payload.get("auth_version", 0) or 0)
    except (JWTError, TypeError, ValueError):
        raise HTTPException(status_code=401, detail="Token invalido")

    cached_user = getattr(request.state, "voltchat_current_user", None)
    if (
        cached_user
        and cached_user.id == user_id
        and int(getattr(cached_user, "auth_version", 0) or 0) == token_auth_version
    ):
        return cached_user

    with SessionLocal() as db:
        user = db.query(User).filter(User.id == user_id).first()
        if user is None:
            raise HTTPException(status_code=401, detail="Usuario nao encontrado")
        if int(getattr(user, "auth_version", 0) or 0) != token_auth_version:
            raise HTTPException(status_code=401, detail="Sessao encerrada. Faca login novamente.")
        db.expunge(user)
        return user


def normalize_access_level(access_level: Optional[str]) -> str:
    value = (access_level or "usuario").strip().lower()
    if value == "padrao":
        return "usuario"
    return value if value in USER_LEVELS else "usuario"


def normalize_status(status: Optional[str]) -> str:
    value = (status or ACTIVE_STATUS).strip().lower()
    if value in {"ativo", "active", "enabled", "habilitado"}:
        return ACTIVE_STATUS
    if value in {"inativo", "inactive", "disabled", "desabilitado"}:
        return INACTIVE_STATUS
    return value or ACTIVE_STATUS


def normalize_username_piece(value: Optional[str], fallback: str = "usuario") -> str:
    normalized = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii").lower()
    normalized = re.sub(r"[^a-z0-9._-]+", "", normalized).strip("._-")
    return normalized or fallback


def generate_unique_username(
    db,
    email: str,
    full_name: str,
    requested_username: Optional[str] = None,
    exclude_user_id: Optional[int] = None,
) -> str:
    """Gera um login global legivel sem misturar contas de empresas diferentes."""
    email_local = str(email or "").split("@", 1)[0]
    base = normalize_username_piece(requested_username or email_local)[:80]
    first_name = str(full_name or "").strip().split(" ", 1)[0]
    name_fragment = normalize_username_piece(first_name, "usr")[:3]

    def is_available(candidate: str) -> bool:
        query = db.query(User.id).filter(User.username == candidate)
        if exclude_user_id is not None:
            query = query.filter(User.id != exclude_user_id)
        return query.first() is None

    if is_available(base):
        return base

    decorated_base = f"{base[:76]}.{name_fragment}"[:80]
    if is_available(decorated_base):
        return decorated_base

    sequence = 2
    while sequence < 10000:
        suffix = str(sequence)
        candidate = f"{decorated_base[:80 - len(suffix)]}{suffix}"
        if is_available(candidate):
            return candidate
        sequence += 1
    raise HTTPException(status_code=409, detail="Nao foi possivel gerar um usuario unico.")

def normalize_platform_role(role: Optional[str]) -> str:
    value = (role or "user").strip().lower()
    aliases = {
        "master": "master_admin",
        "admin_master": "master_admin",
        "admin": "company_admin",
        "administrador": "company_admin",
        "coordenador": "coordinator",
        "usuario": "user",
        "usuário": "user",
    }
    value = aliases.get(value, value)
    return value if value in PLATFORM_ROLES else "user"


def legacy_access_for_role(role: str) -> str:
    role = normalize_platform_role(role)
    if role == "coordinator":
        return "coordenador"
    return "usuario"


def payload_flag(value) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"1", "true", "yes", "sim", "on", "enabled", "habilitado"}


def hub_config() -> tuple[str, str]:
    return (
        str(os.getenv("HUB_BASE_URL") or "").strip().rstrip("/"),
        str(os.getenv("HUB_INTERNAL_TOKEN") or "").strip(),
    )


def require_hub_usage_token(request: Request) -> None:
    configured_token = str(os.getenv("VOLT_CHAT_INTERNAL_USAGE_TOKEN") or "").strip()
    if not configured_token:
        raise HTTPException(status_code=503, detail="Consulta interna do Hub nao configurada.")

    authorization = str(request.headers.get("authorization") or "").strip()
    bearer_token = authorization[7:].strip() if authorization.lower().startswith("bearer ") else ""
    provided_token = str(request.headers.get("x-hub-internal-token") or bearer_token).strip()
    if not provided_token or not hmac.compare_digest(provided_token, configured_token):
        raise HTTPException(status_code=401, detail="Token interno invalido.")


def sync_hub_identity(company: Company, user: User, role: str) -> dict:
    hub_base_url, hub_token = hub_config()
    if not hub_base_url or not hub_token:
        raise HTTPException(
            status_code=503,
            detail="Hub global nao configurado. Defina HUB_BASE_URL e HUB_INTERNAL_TOKEN.",
        )

    tenant_global_id = str(company.tenant_global_id or company.id).strip()
    user_global_id = str(user.uuid or uuid.uuid4()).strip()
    if not user.uuid:
        user.uuid = user_global_id

    hub_role = "owner" if role == "company_admin" else ("admin" if role == "coordinator" else "member")
    body = json.dumps({
        "tenant_id": tenant_global_id,
        "company_name": company.name,
        "document_type": "cnpj" if company.cnpj else None,
        "document_number": company.cnpj,
        "user_id": user_global_id,
        "full_name": user.full_name,
        "email": user.email,
        "role": hub_role,
    }).encode("utf-8")
    request = urllib.request.Request(
        f"{hub_base_url}/v1/internal/identity/sync",
        data=body,
        method="POST",
        headers={"content-type": "application/json", "authorization": f"Bearer {hub_token}"},
    )

    try:
        with urllib.request.urlopen(request, timeout=15) as response:
            response_body = response.read().decode("utf-8")
            result = json.loads(response_body) if response_body else {}
    except urllib.error.HTTPError as exc:
        detail = {}
        try:
            detail = json.loads(exc.read().decode("utf-8"))
        except (json.JSONDecodeError, UnicodeDecodeError):
            pass
        raise HTTPException(
            status_code=502,
            detail=f"Hub recusou a criacao do login: {detail.get('error') or detail.get('reason') or exc.code}.",
        )
    except (urllib.error.URLError, TimeoutError) as exc:
        reason = getattr(exc, "reason", None) or str(exc)
        raise HTTPException(status_code=502, detail=f"Nao foi possivel comunicar com o Hub global: {reason}.")

    company.tenant_global_id = tenant_global_id
    company.hub_enabled = True
    company.hub_synced_at = datetime.utcnow()
    return result


def is_admin(user: User) -> bool:
    return bool(getattr(user, "is_platform_admin", False)) or normalize_access_level(user.access_level) == "master"


def is_coordinator(user: User) -> bool:
    return normalize_access_level(user.access_level) in COORDINATOR_LEVELS


def ensure_admin(user: User):
    if not is_admin(user):
        raise HTTPException(status_code=403, detail="Acesso restrito ao administrador.")


def ensure_utc_datetime(value: datetime) -> datetime:
    """Normaliza datas do SQLite/PostgreSQL para UTC com fuso."""
    if value.tzinfo is None:
        return value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc)


def parse_ticket_delivery_due_at(value) -> Optional[datetime]:
    """Recebe o prazo informado na tela e o guarda como instante UTC."""
    raw_value = str(value or "").strip()
    if not raw_value:
        return None
    try:
        parsed = datetime.fromisoformat(raw_value.replace("Z", "+00:00"))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Prazo de entrega invalido.") from exc
    # datetime-local nao inclui fuso; a interface do VoltChat usa horario de Brasilia.
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone(timedelta(hours=-3)))
    return parsed.astimezone(timezone.utc)


def ticket_is_overdue(ticket: Ticket) -> bool:
    if ticket.status == "Resolvido" or not ticket.delivery_due_at:
        return False
    return datetime.now(timezone.utc) > ensure_utc_datetime(ticket.delivery_due_at)

def archive_due_chat_history():
    """Cria no máximo um backup leve por empresa a cada 30 dias."""
    now = datetime.now(timezone.utc)
    next_run_seconds = CHAT_ARCHIVE_INTERVAL.total_seconds()
    with SessionLocal() as db:
        for company in db.query(Company).all():
            last_backup = (
                db.query(ChatHistoryBackup)
                .filter(ChatHistoryBackup.company_id == company.id)
                .order_by(ChatHistoryBackup.created_at.desc(), ChatHistoryBackup.id.desc())
                .first()
            )
            if last_backup:
                last_created_at = ensure_utc_datetime(last_backup.created_at)
                next_due_at = last_created_at + CHAT_ARCHIVE_INTERVAL
                if now < next_due_at:
                    next_run_seconds = min(next_run_seconds, (next_due_at - now).total_seconds())
                    continue
                period_start = ensure_utc_datetime(last_backup.period_end)
            else:
                earliest = (
                    db.query(Message.timestamp)
                    .filter(Message.company_id == company.id)
                    .order_by(Message.timestamp.asc())
                    .first()
                )
                if not earliest:
                    continue
                period_start = ensure_utc_datetime(earliest[0])

            period_end = period_start + CHAT_ARCHIVE_INTERVAL
            if now < period_end:
                next_run_seconds = min(next_run_seconds, (period_end - now).total_seconds())
                continue

            users = {
                row.id: {"name": row.full_name, "department": row.department}
                for row in db.query(User.id, User.full_name, User.department).all()
            }
            files_by_id = {}
            files_by_path = {}
            file_rows = db.query(
                FileUpload.id,
                FileUpload.file_path,
                FileUpload.filename,
                FileUpload.content_type,
                FileUpload.file_size,
            ).filter(FileUpload.company_id == company.id).all()
            for row in file_rows:
                metadata = {
                    "file_id": row.id,
                    "attachment_file_id": row.id,
                    "attachment_filename": row.filename,
                    "attachment_content_type": row.content_type,
                    "attachment_file_size": row.file_size,
                }
                files_by_id[row.id] = metadata
                files_by_path[str(row.file_path)] = metadata

            group_members = {}
            for group_id, user_id in (
                db.query(ChatGroupMember.group_id, ChatGroupMember.user_id)
                .join(ChatGroup, ChatGroup.id == ChatGroupMember.group_id)
                .filter(ChatGroup.company_id == company.id)
                .all()
            ):
                group_members.setdefault(group_id, []).append(user_id)

            message_count = 0
            original_size = 2
            with tempfile.SpooledTemporaryFile(max_size=2 * 1024 * 1024) as buffer:
                with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
                    with archive.open("messages.json", "w") as entry:
                        entry.write(b"[")
                        messages = (
                            db.query(Message)
                            .filter(
                                Message.company_id == company.id,
                                Message.timestamp >= period_start,
                                Message.timestamp < period_end,
                            )
                            .order_by(Message.timestamp.asc())
                            .yield_per(200)
                        )
                        for message in messages:
                            sender = users.get(message.sender_id, {})
                            receiver = users.get(message.receiver_id, {})
                            attachment = {}
                            if message.file_path:
                                try:
                                    attachment = files_by_id.get(int(message.file_path), {})
                                except (TypeError, ValueError):
                                    attachment = {}
                                attachment = attachment or files_by_path.get(str(message.file_path), {})
                            timestamp_text = message.timestamp.isoformat() if message.timestamp else ""
                            archive_uid = message.archive_uid or hashlib.sha256(
                                f"{company.id}:{message.id}:{timestamp_text}".encode("utf-8")
                            ).hexdigest()
                            audience_kind = "group" if message.group_id else (
                                "direct" if message.receiver_id else "general"
                            )
                            encoded = json.dumps({
                                "id": message.id,
                                "archive_uid": archive_uid,
                                "company_id": message.company_id,
                                "content": message.content,
                                "sender_id": message.sender_id,
                                "sender_name": sender.get("name", "Usuario"),
                                "sender_department": sender.get("department"),
                                "receiver_id": message.receiver_id,
                                "group_id": message.group_id,
                                "reply_to_id": message.reply_to_id,
                                "audience_kind": audience_kind,
                                "participant_user_ids": group_members.get(message.group_id, []),
                                "receiver_name": receiver.get("name"),
                                "receiver_department": receiver.get("department"),
                                "message_type": message.message_type,
                                "timestamp": timestamp_text or None,
                                "file_path": message.file_path,
                                "is_read": bool(message.is_read),
                                **attachment,
                            }, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
                            if message_count:
                                entry.write(b",")
                                original_size += 1
                            entry.write(encoded)
                            original_size += len(encoded)
                            message_count += 1
                        entry.write(b"]")
                buffer.seek(0)
                compressed_data = buffer.read()

            db.add(ChatHistoryBackup(
                company_id=company.id,
                period_start=period_start,
                period_end=period_end,
                message_count=message_count,
                original_size=original_size,
                compressed_size=len(compressed_data),
                compressed_data=compressed_data,
                format_version=2,
                checksum_sha256=hashlib.sha256(compressed_data).hexdigest(),
            ))
            db.query(Message).filter(
                Message.company_id == company.id,
                Message.timestamp >= period_start,
                Message.timestamp < period_end,
            ).delete(synchronize_session=False)
            db.commit()
    return max(60, int(next_run_seconds))



def ensure_coordinator(user: User):
    if not is_coordinator(user):
        raise HTTPException(status_code=403, detail="Acesso restrito a coordenadores.")


def serialize_company(company: Company) -> dict:
    return {
        "id": company.id,
        "name": company.name,
        "cnpj": company.cnpj,
        "responsible_name": company.responsible_name,
        "phone_primary": company.phone_primary,
        "phone_secondary": company.phone_secondary,
        "status": company.status,
        "tenant_global_id": company.tenant_global_id or company.id,
        "hub_enabled": bool(company.hub_enabled),
        "hub_synced_at": company.hub_synced_at.isoformat() if company.hub_synced_at else None,
        "allow_user_sticker_creation": bool(company.allow_user_sticker_creation),
        "allow_user_message_editing": bool(company.allow_user_message_editing),
        "created_at": company.created_at.isoformat() if company.created_at else None,
        "updated_at": company.updated_at.isoformat() if company.updated_at else None,
    }


def serialize_department(department: Department) -> dict:
    return {
        "id": department.id,
        "company_id": department.company_id,
        "name": department.name,
        "description": department.description,
        "status": department.status,
        "created_at": department.created_at.isoformat() if department.created_at else None,
    }


def serialize_department_link(link: DepartmentLink, db) -> dict:
    source = db.query(Department).filter(Department.id == link.source_department_id).first()
    target = db.query(Department).filter(Department.id == link.target_department_id).first()
    return {
        "id": link.id,
        "company_id": link.company_id,
        "source_department_id": link.source_department_id,
        "source_department_name": source.name if source else "Setor removido",
        "target_department_id": link.target_department_id,
        "target_department_name": target.name if target else "Setor removido",
        "label": link.label,
        "created_at": link.created_at.isoformat() if link.created_at else None,
    }


def serialize_company_user(company_user: CompanyUser, db) -> dict:
    company = db.query(Company).filter(Company.id == company_user.company_id).first()
    user = db.query(User).filter(User.id == company_user.user_id).first()
    department = db.query(Department).filter(Department.id == company_user.department_id).first() if company_user.department_id else None
    return {
        "id": company_user.id,
        "company_id": company_user.company_id,
        "company_name": company.name if company else None,
        "user_id": company_user.user_id,
        "user_name": user.full_name if user else None,
        "user_email": user.email if user else None,
        "department_id": company_user.department_id,
        "department_name": department.name if department else None,
        "role": company_user.role,
        "status": company_user.status,
        "created_at": company_user.created_at.isoformat() if company_user.created_at else None,
    }


def serialize_audit_log(log: AuditLog, db) -> dict:
    actor = db.query(User).filter(User.id == log.actor_user_id).first() if log.actor_user_id else None
    actor_label = actor.full_name if actor else ("Usuario removido" if log.actor_user_id else "Sistema")
    metadata = log.metadata_json or {}
    if isinstance(metadata, str):
        try:
            metadata = json.loads(metadata)
        except (json.JSONDecodeError, TypeError):
            metadata = {"raw": metadata}
    return {
        "id": log.id,
        "company_id": log.company_id,
        "actor_user_id": log.actor_user_id,
        "action": log.action,
        "entity_type": log.entity_type,
        "actor_name": actor_label,
        "actor_username": actor.username if actor else None,
        "actor_email": actor.email if actor else None,
        "entity_id": log.entity_id,
        "metadata": metadata,
        "created_at": log.created_at.isoformat() if log.created_at else None,
    }


def display_user_name(user: Optional[User]) -> str:
    if not user:
        return "Usuario"
    return (str(getattr(user, "nickname", "" ) or "" ).strip() or user.full_name or user.username or "Usuario")


def serialize_user(user: User, db=None) -> dict:
    memberships = []
    if db is not None:
        memberships = [
            serialize_company_user(item, db)
            for item in db.query(CompanyUser)
            .filter(CompanyUser.user_id == user.id, CompanyUser.status == ACTIVE_STATUS)
            .order_by(CompanyUser.created_at.asc())
            .all()
        ]

    return {
        "id": user.id,
        "uuid": user.uuid,
        "name": display_user_name(user),

        "nickname": user.nickname,
        "assistant_address_name": getattr(user, "assistant_address_name", None),

        "display_name": display_user_name(user),

        "username": user.username,
        "email": user.email,
        "full_name": user.full_name,
        "phone": getattr(user, "phone", None) or user.phone_extension,
        "is_platform_admin": bool(getattr(user, "is_platform_admin", False)),
        "must_change_password": bool(getattr(user, "must_change_password", False)),
        "status": getattr(user, "status", ACTIVE_STATUS) or (ACTIVE_STATUS if user.is_active else INACTIVE_STATUS),
        "access_level": normalize_access_level(user.access_level),
        "department": user.department,
        "phone_extension": user.phone_extension,
        "birthday": user.birthday,
        "profile_photo": user.profile_photo,
        "role_title": user.role_title,
        "is_active": user.is_active,
        "companies": memberships,
        "company_id": memberships[0]["company_id"] if memberships else None,
        "company_role": memberships[0]["role"] if memberships else ("master_admin" if is_admin(user) else None),
        "company_name": memberships[0]["company_name"] if memberships else None,
        "created_at": user.created_at.isoformat() if user.created_at else None,
        "updated_at": user.updated_at.isoformat() if user.updated_at else None,
    }


def log_audit(db, actor_user_id: Optional[int], action: str, entity_type: str, entity_id: Optional[str], company_id: Optional[str] = None, metadata: Optional[dict] = None):
    db.add(
        AuditLog(
            company_id=company_id,
            actor_user_id=actor_user_id,
            action=action,
            entity_type=entity_type,
            entity_id=str(entity_id) if entity_id is not None else None,
            metadata_json=metadata or {},
        )
    )
def purge_expired_audit_logs() -> int:
    cutoff = datetime.utcnow() - timedelta(days=AUDIT_RETENTION_DAYS)
    with SessionLocal() as db:
        deleted = (
            db.query(AuditLog)
            .filter(AuditLog.created_at < cutoff)
            .delete(synchronize_session=False)
        )
        db.commit()
        return deleted





def ensure_bolt_user(db, company_id: str) -> User:
    """Mantem um usuario de sistema real para mensagens e notificacoes do assistente."""
    bolt = db.query(User).filter(User.username == BOLT_USERNAME).first()
    if not bolt:
        bolt = User(
            username=BOLT_USERNAME,
            email="bolt@voltcorporation.com.br",
            full_name=BOLT_NAME,
            hashed_password=get_password_hash(uuid.uuid4().hex),
            is_platform_admin=False,
            must_change_password=False,
            status=ACTIVE_STATUS,
            access_level="usuario",
            department="Automacao",
            role_title="Assistente virtual",
            is_active=True,
        )
        db.add(bolt)
        db.flush()
    elif bolt.full_name != BOLT_NAME:
        bolt.full_name = BOLT_NAME

    department = get_or_create_department(db, company_id, "Automacao", "Assistente virtual Volt")
    membership = (
        db.query(CompanyUser)
        .filter(CompanyUser.company_id == company_id, CompanyUser.user_id == bolt.id)
        .first()
    )
    if not membership:
        db.add(CompanyUser(
            id=str(uuid.uuid4()),
            company_id=company_id,
            user_id=bolt.id,
            department_id=department.id,
            role="user",
            status=ACTIVE_STATUS,
        ))
    elif membership.status != ACTIVE_STATUS or membership.department_id != department.id:
        membership.status = ACTIVE_STATUS
        membership.department_id = department.id
    return bolt


def ensure_mitty_user(db, company_id: str) -> User:
    """Mantem a secretaria virtual como usuario real em cada empresa."""
    mitty = db.query(User).filter(User.username == MITTY_USERNAME).first()
    if not mitty:
        mitty = User(
            username=MITTY_USERNAME,
            email="mitty@voltcorporation.com.br",
            full_name=MITTY_NAME,
            hashed_password=get_password_hash(uuid.uuid4().hex),
            is_platform_admin=False,
            must_change_password=False,
            status=ACTIVE_STATUS,
            access_level="usuario",
            department="Automacao",
            role_title="Secretaria virtual",
            is_active=True,
        )
        db.add(mitty)
        db.flush()

    department = get_or_create_department(
        db,
        company_id,
        "Automacao",
        "Assistentes virtuais Volt e Mitty",
    )
    membership = (
        db.query(CompanyUser)
        .filter(CompanyUser.company_id == company_id, CompanyUser.user_id == mitty.id)
        .first()
    )
    if not membership:
        db.add(CompanyUser(
            id=str(uuid.uuid4()),
            company_id=company_id,
            user_id=mitty.id,
            department_id=department.id,
            role="user",
            status=ACTIVE_STATUS,
        ))
    elif membership.status != ACTIVE_STATUS or membership.department_id != department.id:
        membership.status = ACTIVE_STATUS
        membership.department_id = department.id
    return mitty


def ensure_virtual_assistant_accounts():
    with SessionLocal() as db:
        company_ids = [
            row[0]
            for row in db.query(Company.id)
            .filter(Company.status == ACTIVE_STATUS)
            .all()
        ]
        for company_id in company_ids:
            ensure_bolt_user(db, company_id)
            ensure_mitty_user(db, company_id)
        db.commit()


def assistant_address_name(user: User) -> str:
    configured = str(getattr(user, "assistant_address_name", "") or "").strip()
    return configured or display_user_name(user) or user.full_name or user.username


def assistant_greeting(user: User, message: str) -> str:
    name = assistant_address_name(user)
    greetings = [
        f"Olá, {name}!",
        f"Certo, {name}.",
        f"Entendi, {name} —",
        f"Perfeito, {name}.",
        f"Vamos lá, {name}:",
    ]
    greeting = greetings[sum(ord(char) for char in message) % len(greetings)]
    return f"{greeting} {message}"

def get_or_create_department(db, company_id: str, name: str, description: Optional[str] = None) -> Department:
    clean_name = str(name or "").strip() or "Geral"
    department = (
        db.query(Department)
        .filter(Department.company_id == company_id, Department.name.ilike(clean_name))
        .first()
    )
    if department:
        return department

    department = Department(
        id=str(uuid.uuid4()),
        company_id=company_id,
        name=clean_name,
        description=description,
        status=ACTIVE_STATUS,
    )
    db.add(department)
    db.flush()
    return department


def ensure_default_tenant_records():
    with SessionLocal() as db:
        company = db.query(Company).filter(Company.id == DEFAULT_COMPANY_ID).first()
        if not company:
            company = Company(
                id=DEFAULT_COMPANY_ID,
                tenant_global_id=DEFAULT_COMPANY_ID,
                name=DEFAULT_COMPANY_NAME,
                responsible_name="Admin Master",
                status=ACTIVE_STATUS,
            )
            db.add(company)
            db.flush()

        for legacy_membership in db.query(CompanyUser).filter(CompanyUser.company_id == DEFAULT_COMPANY_ID).all():
            has_other_company = (
                db.query(CompanyUser.id)
                .filter(
                    CompanyUser.user_id == legacy_membership.user_id,
                    CompanyUser.company_id != DEFAULT_COMPANY_ID,
                    CompanyUser.status == ACTIVE_STATUS,
                )
                .first()
            )
            if has_other_company:
                db.delete(legacy_membership)

        department_by_name = {}
        for name in sorted(set(DEFAULT_DEPARTMENTS + ["Administracao"])):
            department = get_or_create_department(db, DEFAULT_COMPANY_ID, name)
            department_by_name[name.lower()] = department

        for user in db.query(User).all():
            if not user.is_platform_admin and db.query(CompanyUser.id).filter(CompanyUser.user_id == user.id, CompanyUser.company_id != DEFAULT_COMPANY_ID, CompanyUser.status == ACTIVE_STATUS).first():
                continue
            if not user.uuid:
                user.uuid = str(uuid.uuid4())
            if normalize_access_level(user.access_level) == "master" and not user.is_platform_admin:
                user.is_platform_admin = True
            if not getattr(user, "status", None):
                user.status = ACTIVE_STATUS if user.is_active else INACTIVE_STATUS
            department = department_by_name.get((user.department or "Operacao").lower()) or get_or_create_department(
                db, DEFAULT_COMPANY_ID, user.department or "Operacao"
            )
            membership = (
                db.query(CompanyUser)
                .filter(CompanyUser.company_id == DEFAULT_COMPANY_ID, CompanyUser.user_id == user.id)
                .first()
            )
            if not membership:
                role = "master_admin" if user.is_platform_admin else ("coordinator" if normalize_access_level(user.access_level) == "coordenador" else "user")
                db.add(
                    CompanyUser(
                        id=str(uuid.uuid4()),
                        company_id=DEFAULT_COMPANY_ID,
                        user_id=user.id,
                        department_id=department.id,
                        role=role,
                        status=ACTIVE_STATUS if user.is_active else INACTIVE_STATUS,
                    )
                )
        db.commit()


def active_company_memberships(db, user: User) -> list[CompanyUser]:
    return (
        db.query(CompanyUser)
        .join(Company, Company.id == CompanyUser.company_id)
        .filter(
            CompanyUser.user_id == user.id,
            CompanyUser.status == ACTIVE_STATUS,
            Company.status == ACTIVE_STATUS,
        )
        .order_by(CompanyUser.created_at.asc())
        .all()
    )


def ensure_company_access(db, user: User, company_id: Optional[str] = None) -> str:
    if is_admin(user):
        target_company_id = str(company_id or DEFAULT_COMPANY_ID)
        company = db.query(Company).filter(Company.id == target_company_id).first()
        if not company:
            raise HTTPException(status_code=404, detail="Empresa nao encontrada.")
        return company.id

    memberships = active_company_memberships(db, user)
    if company_id:
        for membership in memberships:
            if membership.company_id == company_id:
                return membership.company_id
        raise HTTPException(status_code=403, detail="Usuario sem vinculo ativo com esta empresa.")

    if not memberships:
        raise HTTPException(status_code=403, detail="Usuario sem vinculo ativo com empresa.")
    return memberships[0].company_id


def get_company_membership(db, user: User, company_id: str) -> Optional[CompanyUser]:
    return (
        db.query(CompanyUser)
        .filter(
            CompanyUser.company_id == company_id,
            CompanyUser.user_id == user.id,
            CompanyUser.status == ACTIVE_STATUS,
        )
        .first()
    )


def is_company_admin(db, user: User, company_id: str) -> bool:
    if is_admin(user):
        return True
    membership = get_company_membership(db, user, company_id)
    return bool(membership and membership.role in TENANT_ADMIN_ROLES)


def ensure_company_admin(db, user: User, company_id: str):
    if not is_company_admin(db, user, company_id):
        raise HTTPException(status_code=403, detail="Acesso restrito ao administrador da empresa.")


def can_create_company_stickers(db, user: User, company_id: str) -> bool:
    company = db.query(Company).filter(Company.id == company_id).first()
    return bool(company and (company.allow_user_sticker_creation or is_company_admin(db, user, company_id)))


def ensure_company_sticker_creation(db, user: User, company_id: str):
    if not can_create_company_stickers(db, user, company_id):
        raise HTTPException(
            status_code=403,
            detail="A criacao de figurinhas foi desabilitada pelo administrador da empresa.",
        )


def user_department_name(db, user: User, company_id: str) -> Optional[str]:
    membership = get_company_membership(db, user, company_id)
    if not membership or not membership.department_id:
        return user.department
    department = db.query(Department).filter(Department.id == membership.department_id).first()
    return department.name if department else user.department


def user_can_access_department(db, user: User, company_id: str, department_name: Optional[str]) -> bool:
    if is_company_admin(db, user, company_id):
        return True
    membership = get_company_membership(db, user, company_id)
    if not membership:
        return False
    if membership.role == "coordinator":
        return bool(department_name and user_department_name(db, user, company_id) == department_name)
    return False


def ensure_creation_scope(db, user: User, company_id: str, department_name: Optional[str], intent: str):
    """Valida o escopo de tickets; tarefas sao sempre pessoais e independentes de setor."""
    if is_company_admin(db, user, company_id):
        return
    membership = get_company_membership(db, user, company_id)
    if not membership:
        raise HTTPException(status_code=403, detail="Usuario sem vinculo ativo com esta empresa.")
    if intent != "ticket":
        return
    own_department = user_department_name(db, user, company_id)
    if membership.role == "coordinator" and department_name != own_department:
        raise HTTPException(status_code=403, detail="Coordenador so pode criar itens do proprio setor.")


def ensure_user_in_company(db, user_id: int, company_id: str) -> User:
    user = db.query(User).filter(User.id == user_id, User.is_active == True).first()
    if not user:
        raise HTTPException(status_code=400, detail="Usuario nao encontrado.")
    if is_admin(user):
        return user
    membership = (
        db.query(CompanyUser)
        .filter(CompanyUser.company_id == company_id, CompanyUser.user_id == user.id, CompanyUser.status == ACTIVE_STATUS)
        .first()
    )
    if not membership:
        raise HTTPException(status_code=403, detail="Usuario nao pertence a esta empresa.")
    return user


def serialize_ticket(ticket: Ticket, db) -> dict:
    created_by = db.query(User).filter(User.id == ticket.created_by_id).first()
    assigned_to = db.query(User).filter(User.id == ticket.assigned_to_id).first() if ticket.assigned_to_id else None
    closed_by = db.query(User).filter(User.id == ticket.closed_by_id).first() if ticket.closed_by_id else None
    attachment = db.query(FileUpload).filter(FileUpload.id == ticket.attachment_file_id).first() if ticket.attachment_file_id else None
    first_response_minutes = None
    if ticket.created_at and ticket.first_response_at:
        first_response_minutes = max(0, int((ticket.first_response_at - ticket.created_at).total_seconds() // 60))
    assigned_users = (
        db.query(User)
        .join(TicketAssignee, TicketAssignee.user_id == User.id)
        .filter(TicketAssignee.ticket_id == ticket.id)
        .order_by(User.full_name.asc())
        .all()
    )
    if not assigned_users and assigned_to:
        assigned_users = [assigned_to]
    assigned_departments = (
        db.query(Department)
        .join(TicketDepartment, TicketDepartment.department_id == Department.id)
        .filter(TicketDepartment.ticket_id == ticket.id)
        .order_by(Department.name.asc())
        .all()
    )
    assigned_department_payload = [{"id": department.id, "name": department.name} for department in assigned_departments]
    if not assigned_department_payload and ticket.department:
        assigned_department_payload = [{"id": None, "name": ticket.department}]

    return {
        "id": ticket.id,
        "company_id": ticket.company_id,
        "title": ticket.title,
        "description": ticket.description,
        "priority": ticket.priority,
        "delivery_due_at": ticket.delivery_due_at.isoformat() if ticket.delivery_due_at else None,
        "is_overdue": ticket_is_overdue(ticket),
        "status": ticket.status,
        "department": ticket.department,
        "channel": ticket.channel,
        "created_by_id": ticket.created_by_id,
        "created_by_name": created_by.full_name if created_by else None,
        "assigned_to_id": ticket.assigned_to_id,
        "assigned_to_name": assigned_to.full_name if assigned_to else None,
        "assigned_users": [{"id": user.id, "name": user.full_name, "email": user.email} for user in assigned_users],
        "assigned_departments": assigned_department_payload,
        "is_group_ticket": len(assigned_users) > 1 or len(assigned_department_payload) > 1,
        "attachment_file_id": ticket.attachment_file_id,
        "attachment_filename": attachment.filename if attachment else None,
        "attachment_content_type": attachment.content_type if attachment else None,
        "attachment_file_size": attachment.file_size if attachment else None,
        "rating_score": ticket.rating_score,
        "rating_comment": ticket.rating_comment,
        "rated_at": ticket.rated_at.isoformat() if ticket.rated_at else None,
        "first_response_at": ticket.first_response_at.isoformat() if ticket.first_response_at else None,
        "first_response_minutes": first_response_minutes,
        "created_at": ticket.created_at.isoformat() if ticket.created_at else None,
        "updated_at": ticket.updated_at.isoformat() if ticket.updated_at else None,
        "closed_at": ticket.closed_at.isoformat() if ticket.closed_at else None,
        "closed_by_id": ticket.closed_by_id,
        "closed_by_name": closed_by.full_name if closed_by else None,
        "close_reason": ticket.close_reason,
    }


def serialize_ticket_message_reactions(message_id: int, db) -> list[dict]:
    rows = (
        db.query(TicketMessageReaction, User)
        .join(User, User.id == TicketMessageReaction.user_id)
        .filter(TicketMessageReaction.ticket_message_id == message_id)
        .order_by(TicketMessageReaction.created_at.asc())
        .all()
    )
    grouped = {}
    for reaction, user in rows:
        item = grouped.setdefault(reaction.emoji, {
            "emoji": reaction.emoji, "count": 0, "users": [], "user_ids": [],
        })
        item["count"] += 1
        item["users"].append(user.full_name)
        item["user_ids"].append(user.id)
    return list(grouped.values())


def serialize_ticket_message(message: TicketMessage, db) -> dict:
    sender = db.query(User).filter(User.id == message.sender_id).first()
    attachment = db.query(FileUpload).filter(FileUpload.id == message.file_id).first() if message.file_id else None
    reply_to = db.query(TicketMessage).filter(TicketMessage.id == message.reply_to_id).first() if message.reply_to_id else None
    reply_sender = db.query(User).filter(User.id == reply_to.sender_id).first() if reply_to else None
    reply_attachment = db.query(FileUpload).filter(FileUpload.id == reply_to.file_id).first() if reply_to and reply_to.file_id else None
    return {
        "id": message.id,
        "company_id": message.company_id,
        "ticket_id": message.ticket_id,
        "sender_id": message.sender_id,
        "sender_name": display_user_name(sender),
        "sender_full_name": sender.full_name if sender else None,
        "sender_nickname": sender.nickname if sender else None,
        "content": message.content,
        "file_id": message.file_id,
        "reply_to_id": message.reply_to_id,
        "reply_to": {
            "id": reply_to.id,
            "content": reply_to.content,
            "sender_id": reply_to.sender_id,
            "sender_name": display_user_name(reply_sender),
            "sender_full_name": reply_sender.full_name if reply_sender else None,
            "sender_nickname": reply_sender.nickname if reply_sender else None,
            "attachment_filename": reply_attachment.filename if reply_attachment else None,
        } if reply_to else None,
        "reactions": serialize_ticket_message_reactions(message.id, db),
        "attachment_filename": attachment.filename if attachment else None,
        "attachment_content_type": attachment.content_type if attachment else None,
        "attachment_file_size": attachment.file_size if attachment else None,
        "created_at": message.created_at.isoformat() if message.created_at else None,
    }


def serialize_message_attachment(file_ref, db) -> dict:
    if not file_ref:
        return {}

    attachment = None
    parsed_file_id = None

    try:
        parsed_file_id = int(file_ref)
    except (TypeError, ValueError):
        parsed_file_id = None

    if parsed_file_id:
        attachment = db.query(FileUpload).filter(FileUpload.id == parsed_file_id).first()

    if not attachment:
        attachment = db.query(FileUpload).filter(FileUpload.file_path == str(file_ref)).first()

    if attachment:
        return {
            "file_id": attachment.id,
            "attachment_file_id": attachment.id,
            "attachment_filename": attachment.filename,
            "attachment_content_type": attachment.content_type,
            "attachment_file_size": attachment.file_size,
            "attachment_expires_at": attachment.expires_at.isoformat() if attachment.expires_at else None,
        }

    if parsed_file_id:
        # O registro pode ter sido removido pela retenção obrigatória de 15 dias.
        # Preserve a mensagem, mas informe ao cliente para não repetir downloads 404.
        return {
            "file_id": parsed_file_id,
            "attachment_file_id": parsed_file_id,
            "attachment_unavailable": True,
        }

    return {}


def serialize_message_reactions(message_id: int, db) -> list[dict]:
    rows = (
        db.query(MessageReaction, User)
        .join(User, User.id == MessageReaction.user_id)
        .filter(MessageReaction.message_id == message_id)
        .order_by(MessageReaction.created_at.asc())
        .all()
    )
    grouped = {}
    for reaction, user in rows:
        item = grouped.setdefault(reaction.emoji, {
            "emoji": reaction.emoji, "count": 0, "users": [], "user_ids": [],
        })
        item["count"] += 1
        item["users"].append(user.full_name)
        item["user_ids"].append(user.id)
    return list(grouped.values())


def serialize_message_readers(message_id: int, db) -> list[dict]:
    rows = (
        db.query(MessageReadReceipt, User)
        .join(User, User.id == MessageReadReceipt.user_id)
        .filter(MessageReadReceipt.message_id == message_id)
        .order_by(MessageReadReceipt.read_at.asc())
        .all()
    )
    return [
        {
            "id": user.id,
            "name": display_user_name(user),
            "full_name": user.full_name,
            "profile_photo": user.profile_photo,
            "read_at": receipt.read_at.isoformat() if receipt.read_at else None,
        }
        for receipt, user in rows
    ]


def ensure_chat_group_member(db, group_id: int, user_id: int, company_id: str) -> ChatGroup:
    """Aplica a regra unica de acesso aos grupos do chat.

    Administradores da empresa podem acessar todos os grupos. Coordenadores e
    usuarios comuns somente podem acessar grupos em que participam explicitamente.
    O departamento do grupo nao concede acesso por si so.
    """
    group = db.query(ChatGroup).filter(
        ChatGroup.id == group_id,
        ChatGroup.company_id == company_id,
        ChatGroup.is_active == True,
    ).first()
    if not group:
        raise HTTPException(status_code=404, detail="Grupo nao encontrado.")

    membership = db.query(ChatGroupMember).filter(
        ChatGroupMember.group_id == group.id,
        ChatGroupMember.user_id == user_id,
    ).first()
    if membership:
        return group

    user = db.query(User).filter(User.id == user_id, User.is_active == True).first()
    if user and is_company_admin(db, user, company_id):
        return group

    raise HTTPException(status_code=403, detail="Voce nao participa deste grupo.")


def accessible_messages_query(db, user_id: int, company_id: str):
    user = db.query(User).filter(User.id == user_id, User.is_active == True).first()
    if user and is_company_admin(db, user, company_id):
        group_ids = [
            row[0]
            for row in db.query(ChatGroup.id).filter(
                ChatGroup.company_id == company_id,
                ChatGroup.is_active == True,
            ).all()
        ]
    else:
        group_ids = [
            row[0]
            for row in (
                db.query(ChatGroupMember.group_id)
                .join(ChatGroup, ChatGroup.id == ChatGroupMember.group_id)
                .filter(
                    ChatGroupMember.user_id == user_id,
                    ChatGroup.company_id == company_id,
                    ChatGroup.is_active == True,
                )
                .all()
            )
        ]
    conditions = [
        and_(
            Message.group_id.is_(None),
            Message.receiver_id.isnot(None),
            or_(Message.sender_id == user_id, Message.receiver_id == user_id),
        ),
    ]
    if group_ids:
        conditions.append(Message.group_id.in_(group_ids))
    return db.query(Message).filter(Message.company_id == company_id, or_(*conditions))

def serialize_message_edit_history(message_id: int, db) -> list[dict]:
    rows = (
        db.query(MessageEdit, User)
        .join(User, User.id == MessageEdit.editor_user_id)
        .filter(MessageEdit.message_id == message_id)
        .order_by(MessageEdit.edited_at.asc(), MessageEdit.id.asc())
        .all()
    )
    return [
        {
            "id": edit.id,
            "previous_content": edit.previous_content,
            "edited_at": edit.edited_at.isoformat() if edit.edited_at else None,
            "editor_user_id": edit.editor_user_id,
            "editor_name": display_user_name(editor),
        }
        for edit, editor in rows
    ]


def message_mention_user_ids(message_id: int, db) -> list[int]:
    return [
        row[0]
        for row in db.query(MessageMention.user_id)
        .filter(MessageMention.message_id == message_id)
        .order_by(MessageMention.id.asc())
        .all()
    ]


def validated_group_mention_ids(
    db,
    group: ChatGroup,
    mention_ids,
    company_id: str,
    sender_id: int,
    content: str = "",
) -> set[int]:
    raw_ids = mention_ids if isinstance(mention_ids, list) else []
    requested = {int(item) for item in raw_ids if str(item).isdigit()}

    member_users = (
        db.query(User)
        .join(ChatGroupMember, ChatGroupMember.user_id == User.id)
        .filter(ChatGroupMember.group_id == group.id)
        .all()
    )
    member_ids = {user.id for user in member_users}

    # Compatibilidade com clientes 0.1.35: eles escreviam @usuario no texto,
    # mas ainda nao enviavam mention_user_ids no payload WebSocket.
    if not requested and content:
        tokens = {token.lower() for token in re.findall(r"(?:^|\s)@([A-Za-z0-9._-]+)", str(content))}
        for member in member_users:
            candidates = [member.username, member.nickname, member.full_name]
            for candidate in candidates:
                normalized = re.sub(r"[^A-Za-z0-9._-]", "", re.sub(r"\s+", ".", str(candidate or "").strip())).lower()
                if normalized and normalized in tokens:
                    requested.add(member.id)
                    break

    requested.discard(int(sender_id))
    if not requested:
        return set()
    valid_ids = requested & member_ids
    if len(valid_ids) != len(requested):
        raise HTTPException(status_code=400, detail="Mencoes devem apontar apenas para participantes do grupo.")
    for user_id in valid_ids:
        ensure_user_in_company(db, user_id, company_id)
    return valid_ids


def sync_message_mentions(db, message: Message, mention_ids: set[int]) -> tuple[set[int], set[int]]:
    existing_rows = db.query(MessageMention).filter(MessageMention.message_id == message.id).all()
    existing_ids = {row.user_id for row in existing_rows}
    removed = existing_ids - mention_ids
    added = mention_ids - existing_ids
    for row in existing_rows:
        if row.user_id in removed:
            db.delete(row)
    for user_id in added:
        db.add(MessageMention(message_id=message.id, user_id=user_id))
    return added, removed


def serialize_message_sticker(message: Message, db) -> dict:
    if message.message_type != "sticker":
        return {}
    reference = str(message.file_path or "")
    if reference.startswith("default:"):
        key = reference.split(":", 1)[1]
        return {
            "sticker_key": key,
            "sticker_name": message.content or key,
            "sticker_is_default": True,
        }
    if reference.startswith("sticker:"):
        try:
            sticker_id = int(reference.split(":", 1)[1])
        except (TypeError, ValueError):
            return {}
        sticker = db.query(ChatSticker).filter(ChatSticker.id == sticker_id).first()
        return {
            "sticker_id": sticker_id,
            "sticker_name": sticker.name if sticker else (message.content or "Figurinha"),
            "sticker_is_default": False,
        }
    return {}

def paginate_message_history(query, db, before_message_id: Optional[int] = None, limit: int = 50) -> dict:
    """Return a stable keyset page, ordered for insertion above the current view."""
    limit = max(1, min(int(limit), 100))
    if before_message_id:
        cursor = query.filter(Message.id == int(before_message_id)).first()
        if not cursor:
            raise HTTPException(status_code=400, detail="Cursor de historico invalido.")
        query = query.filter(or_(
            Message.timestamp < cursor.timestamp,
            and_(Message.timestamp == cursor.timestamp, Message.id < cursor.id),
        ))
    rows = query.order_by(Message.timestamp.desc(), Message.id.desc()).limit(limit + 1).all()
    has_more = len(rows) > limit
    page = rows[:limit]
    return {
        "messages": serialize_messages(reversed(page), db),
        "next_cursor": page[-1].id if has_more and page else None,
        "has_more": has_more,
    }


def serialize_messages(messages, db) -> list[dict]:
    """Serialize a history with shared related-data queries."""
    messages = list(messages or [])
    if not messages:
        return []
    reply_ids = {item.reply_to_id for item in messages if item.reply_to_id}
    replies = {item.id: item for item in db.query(Message).filter(Message.id.in_(reply_ids)).all()} if reply_ids else {}
    user_ids = {user_id for item in messages for user_id in (item.sender_id, item.receiver_id) if user_id}
    user_ids.update(item.sender_id for item in replies.values() if item.sender_id)
    users = {item.id: item for item in db.query(User).filter(User.id.in_(user_ids)).all()} if user_ids else {}
    group_ids = {item.group_id for item in messages if item.group_id}
    groups = {item.id: item for item in db.query(ChatGroup).filter(ChatGroup.id.in_(group_ids)).all()} if group_ids else {}
    related = {"users": users, "groups": groups, "replies": replies}
    return [serialize_message(item, db, related=related) for item in messages]


def serialize_message(message: Message, db, related=None) -> dict:
    related = related or {}
    users = related.get("users", {})
    groups = related.get("groups", {})
    replies = related.get("replies", {})
    sender = users.get(message.sender_id) if users else db.query(User).filter(User.id == message.sender_id).first()
    receiver = users.get(message.receiver_id) if message.receiver_id and users else (db.query(User).filter(User.id == message.receiver_id).first() if message.receiver_id else None)
    group = groups.get(message.group_id) if message.group_id and groups else (db.query(ChatGroup).filter(ChatGroup.id == message.group_id).first() if message.group_id else None)
    reply_to = replies.get(message.reply_to_id) if message.reply_to_id and replies else (db.query(Message).filter(Message.id == message.reply_to_id).first() if message.reply_to_id else None)
    reply_sender = users.get(reply_to.sender_id) if reply_to and users else (db.query(User).filter(User.id == reply_to.sender_id).first() if reply_to else None)
    edit_history = serialize_message_edit_history(message.id, db)
    mention_user_ids = message_mention_user_ids(message.id, db)
    return {
        "id": message.id,
        "archive_uid": message.archive_uid,
        "company_id": message.company_id,
        "content": message.content,
        "sender_id": message.sender_id,
        "sender_name": display_user_name(sender),
        "sender_full_name": sender.full_name if sender else None,
        "sender_nickname": sender.nickname if sender else None,
        "sender_profile_photo": sender.profile_photo if sender else None,
        "sender_department": sender.department if sender else None,
        "receiver_id": message.receiver_id,
        "group_id": message.group_id,
        "group_name": group.name if group else None,
        "receiver_name": display_user_name(receiver) if receiver else None,
        "receiver_full_name": receiver.full_name if receiver else None,
        "receiver_nickname": receiver.nickname if receiver else None,
        "receiver_department": receiver.department if receiver else None,
        "reply_to_id": message.reply_to_id,
        "reply_to": {
            "id": reply_to.id,
            "content": reply_to.content,
            "sender_id": reply_to.sender_id,
            "sender_name": display_user_name(reply_sender),
            "sender_full_name": reply_sender.full_name if reply_sender else None,
            "sender_nickname": reply_sender.nickname if reply_sender else None,
            "message_type": reply_to.message_type,
        } if reply_to else None,
        "reactions": serialize_message_reactions(message.id, db),
        "readers": serialize_message_readers(message.id, db),
        "message_type": message.message_type,
        "timestamp": message.timestamp.isoformat() if message.timestamp else None,
        "file_path": message.file_path,
        "is_read": bool(message.is_read),
        "is_edited": bool(edit_history),
        "edit_count": len(edit_history),
        "edited_at": edit_history[-1]["edited_at"] if edit_history else None,
        "edit_history": edit_history,
        "mention_user_ids": mention_user_ids,
        **serialize_message_attachment(message.file_path, db),
        **serialize_message_sticker(message, db),
    }


def ticket_department_names(ticket: Ticket, db) -> set[str]:
    names = set()
    if ticket.department:
        names.add(str(ticket.department).strip())
    rows = (
        db.query(Department.name)
        .join(TicketDepartment, TicketDepartment.department_id == Department.id)
        .filter(TicketDepartment.ticket_id == ticket.id)
        .all()
    )
    names.update(str(row[0]).strip() for row in rows if row and row[0])
    return {name for name in names if name}


def can_manage_ticket(ticket: Ticket, user: User, db) -> bool:
    """Acesso gerencial sem transformar o gestor em participante do ticket.

    Administradores da empresa controlam qualquer ticket da empresa. Coordenadores
    controlam tickets classificados no proprio setor. Esse acesso permite abrir,
    responder e executar acoes operacionais, mas nao cria TicketAssignee.
    """
    company_id = ticket.company_id or DEFAULT_COMPANY_ID
    try:
        ensure_company_access(db, user, company_id)
    except HTTPException:
        return False

    if is_company_admin(db, user, company_id):
        return True
    membership = get_company_membership(db, user, company_id)
    if not membership or membership.role != "coordinator":
        return False
    own_department = user_department_name(db, user, company_id)
    if not own_department:
        return False
    return own_department in ticket_department_names(ticket, db)


def can_access_ticket(ticket: Ticket, user: User, db) -> bool:
    """Visibilidade normal: participantes explicitos, mais administradores da empresa.

    Coordenador nao passa a aparecer como participante apenas por gerenciar o setor.
    A visao gerencial e tratada separadamente por ``can_manage_ticket``.
    """
    company_id = ticket.company_id or DEFAULT_COMPANY_ID
    try:
        ensure_company_access(db, user, company_id)
    except HTTPException:
        return False

    if is_company_admin(db, user, company_id):
        return True
    if ticket.created_by_id == user.id or ticket.assigned_to_id == user.id:
        return True
    return bool(
        db.query(TicketAssignee)
        .filter(TicketAssignee.ticket_id == ticket.id, TicketAssignee.user_id == user.id)
        .first()
    )


def sync_ticket_group_assignments(db, ticket: Ticket, company_id: str, user_ids=None, department_names=None):
    if user_ids is not None:
        try:
            normalized_user_ids = sorted({int(user_id) for user_id in user_ids if user_id not in (None, "")})
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Responsaveis do ticket invalidos.")
        assigned_users = [ensure_user_in_company(db, user_id, company_id) for user_id in normalized_user_ids]
        db.query(TicketAssignee).filter(TicketAssignee.ticket_id == ticket.id).delete(synchronize_session=False)
        for assigned_user in assigned_users:
            db.add(TicketAssignee(ticket_id=ticket.id, user_id=assigned_user.id))
        ticket.assigned_to_id = assigned_users[0].id if assigned_users else None

    if department_names is not None:
        normalized_names = sorted({str(name).strip() for name in department_names if str(name).strip()})
        assigned_departments = []
        for name in normalized_names:
            department = db.query(Department).filter(Department.company_id == company_id, Department.name == name, Department.status == ACTIVE_STATUS).first()
            if not department:
                raise HTTPException(status_code=400, detail=f"Setor nao encontrado ou inativo: {name}.")
            assigned_departments.append(department)
        db.query(TicketDepartment).filter(TicketDepartment.ticket_id == ticket.id).delete(synchronize_session=False)
        for department in assigned_departments:
            db.add(TicketDepartment(ticket_id=ticket.id, department_id=department.id))


def ensure_ticket_access(ticket: Ticket, user: User, db):
    if not (can_access_ticket(ticket, user, db) or can_manage_ticket(ticket, user, db)):
        raise HTTPException(status_code=403, detail="Sem permissao para acessar este ticket.")


async def notify_ticket_user(user_id: Optional[int], title: str, description: str, ticket_id: int):
    if not user_id:
        return
    payload = {
        "type": "notification",
        "notification": {
            "id": f"ticket-{ticket_id}-{int(datetime.utcnow().timestamp() * 1000)}",
            "type": "Ticket",
            "title": title,
            "description": description,
            "time": "Agora",
            "unread": True,
            "ticket_id": ticket_id,
            "link": "/tickets",
        },
    }
    await manager.send_personal_message(json.dumps(payload), user_id)


def get_ticket_participant_ids(db, ticket: Ticket) -> set[int]:
    """Retorna apenas pessoas explicitamente ligadas ao ticket.

    O criador participa automaticamente. Setor/departamento nao expande a audiencia.
    """
    participant_ids = set()
    if ticket.created_by_id:
        participant_ids.add(ticket.created_by_id)
    if ticket.assigned_to_id:
        participant_ids.add(ticket.assigned_to_id)

    participant_ids.update(
        row[0]
        for row in db.query(TicketAssignee.user_id)
        .filter(TicketAssignee.ticket_id == ticket.id)
        .all()
        if row and row[0]
    )
    return {user_id for user_id in participant_ids if user_id}


async def notify_ticket_participants(
    db,
    ticket: Ticket,
    current_user: User,
    title: str,
    description: str,
):
    recipient_ids = get_ticket_participant_ids(db, ticket)
    if current_user and current_user.id in recipient_ids:
        recipient_ids.discard(current_user.id)
    for recipient_id in sorted(recipient_ids):
        await notify_ticket_user(recipient_id, title, description, ticket.id)


def create_bolt_ticket_messages(db, ticket: Ticket, plan: dict, requester: User) -> list[tuple[int, dict]]:
    """Persiste a notificacao do Volt somente para participantes explicitos do ticket."""
    bolt = ensure_bolt_user(db, ticket.company_id)
    db.add(TicketMessage(
        company_id=ticket.company_id,
        ticket_id=ticket.id,
        sender_id=bolt.id,
        content=(f"Volt abriu o ticket #{ticket.id}. Tarefa a realizar: {plan['subject']}. "
                 f"Urgência: {plan['ticket_priority']}. Registre aqui o andamento e a resolução."),
    ))
    recipient_ids = {requester.id, *plan.get("assigned_to_ids", [])}
    department_names = plan.get("departments", [])

    people_text = ", ".join(plan.get("assigned_to_names", [])) or requester.full_name
    departments_text = ", ".join(department_names) or ticket.department or "sem setor"
    content = (
        f"Volt notificou você sobre o ticket #{ticket.id}. Assunto: {plan['subject']}. "
        f"Urgência: {plan['ticket_priority']}. Participantes: {people_text}. "
        f"Setor de referência: {departments_text}. Aberto por {requester.full_name}."
    )
    notifications = []
    for recipient_id in sorted(recipient_ids):
        if recipient_id == bolt.id:
            continue
        message = Message(
            company_id=ticket.company_id,
            content=content,
            sender_id=bolt.id,
            receiver_id=recipient_id,
            message_type="ticket_notification",
        )
        db.add(message)
        db.flush()
        notifications.append((recipient_id, serialize_message(message, db)))
    return notifications


async def dispatch_bolt_ticket_notifications(ticket: Ticket, notifications: list[tuple[int, dict]]):
    for recipient_id, message_payload in notifications:
        await manager.send_personal_message(
            json.dumps({"type": "new_message", "message": message_payload}),
            recipient_id,
        )
        await notify_ticket_user(
            recipient_id,
            f"Volt • Ticket #{ticket.id}",
            message_payload["content"],
            ticket.id,
        )


async def notify_task_user(user_id: Optional[int], title: str, description: str, task_id: int):
    if not user_id:
        return
    payload = {
        "type": "notification",
        "notification": {
            "id": f"task-{task_id}-{int(datetime.utcnow().timestamp() * 1000)}",
            "type": "Tarefa",
            "title": title,
            "description": description,
            "time": "Agora",
            "unread": True,
            "task_id": task_id,
            "link": "/tasks",
        },
    }
    await manager.send_personal_message(json.dumps(payload), user_id)


def serialize_groups(groups, db) -> list[dict]:
    """Serialize group lists with shared creator and member queries."""
    groups = list(groups or [])
    if not groups:
        return []
    group_ids = [group.id for group in groups]
    creator_ids = {group.created_by_id for group in groups if group.created_by_id}
    creators = {
        user.id: user
        for user in db.query(User).filter(User.id.in_(creator_ids)).all()
    } if creator_ids else {}
    members_by_group = {group_id: [] for group_id in group_ids}
    rows = (
        db.query(ChatGroupMember.group_id, User)
        .join(User, User.id == ChatGroupMember.user_id)
        .filter(ChatGroupMember.group_id.in_(group_ids))
        .order_by(User.full_name.asc())
        .all()
    )
    for group_id, member in rows:
        members_by_group.setdefault(group_id, []).append(member)
    payloads = []
    for group in groups:
        members = members_by_group.get(group.id, [])
        creator = creators.get(group.created_by_id)
        payloads.append({
            "id": group.id,
            "company_id": group.company_id,
            "name": group.name,
            "description": group.description,
            "department": group.department,
            "image_data": group.image_data,
            "created_by_id": group.created_by_id,
            "created_by_name": display_user_name(creator),
            "member_ids": [member.id for member in members],
            "members": [serialize_user(member, db) for member in members],
            "is_active": group.is_active,
            "created_at": group.created_at.isoformat() if group.created_at else None,
        })
    return payloads


def serialize_group(group: ChatGroup, db) -> dict:
    creator = db.query(User).filter(User.id == group.created_by_id).first()
    members = (
        db.query(User)
        .join(ChatGroupMember, ChatGroupMember.user_id == User.id)
        .filter(ChatGroupMember.group_id == group.id)
        .order_by(User.full_name.asc())
        .all()
    )
    return {
        "id": group.id,
        "company_id": group.company_id,
        "name": group.name,
        "description": group.description,
        "department": group.department,
        "image_data": group.image_data,
        "created_by_id": group.created_by_id,
        "created_by_name": display_user_name(creator),
        "member_ids": [member.id for member in members],
        "members": [serialize_user(member, db) for member in members],
        "is_active": group.is_active,
        "created_at": group.created_at.isoformat() if group.created_at else None,
    }

def serialize_task(task: TaskItem, db) -> dict:
    created_by = db.query(User).filter(User.id == task.created_by_id).first()
    assigned_to = db.query(User).filter(User.id == task.assigned_to_id).first() if task.assigned_to_id else None
    return {
        "id": task.id,
        "company_id": task.company_id,
        "title": task.title,
        "description": task.description,
        "priority": task.priority,
        "category": task.category,
        "status": task.status,
        "due_date": task.due_date,
        "created_by_id": task.created_by_id,
        "created_by_name": created_by.full_name if created_by else None,
        "assigned_to_id": task.assigned_to_id,
        "assigned_to_name": assigned_to.full_name if assigned_to else None,
        "created_at": task.created_at.isoformat() if task.created_at else None,
        "updated_at": task.updated_at.isoformat() if task.updated_at else None,
    }


class ConnectionManager:
    def __init__(self):
        self.active_connections: Dict[int, WebSocket] = {}
        self.user_connections: Dict[int, set[int]] = {}
        self.connection_users: Dict[int, int] = {}
        self.connection_company_scope: Dict[int, str] = {}
        self.connection_client_keys: Dict[int, str] = {}
        self.client_connections: Dict[str, int] = {}
        self.presence_status: Dict[int, str] = {}

    @staticmethod
    def _client_key(user_id: int, client_id: Optional[str]) -> str:
        normalized = re.sub(r"[^A-Za-z0-9._:-]", "", str(client_id or "").strip())[:80]
        return f"{user_id}:{normalized}" if normalized else ""

    async def _safe_send(self, websocket: WebSocket, message: str, timeout_seconds: float = 8.0) -> bool:
        try:
            await asyncio.wait_for(websocket.send_text(message), timeout=timeout_seconds)
            return True
        except Exception:
            return False

    async def connect(self, websocket: WebSocket, user_id: int, company_id: str, client_id: Optional[str] = None):
        was_online = bool(self.user_connections.get(user_id))
        await websocket.accept()
        connection_id = id(websocket)
        client_key = self._client_key(user_id, client_id)
        previous_connection_id = self.client_connections.get(client_key) if client_key else None

        self.active_connections[connection_id] = websocket
        self.user_connections.setdefault(user_id, set()).add(connection_id)
        self.connection_users[connection_id] = user_id
        self.connection_company_scope[connection_id] = company_id
        if client_key:
            self.connection_client_keys[connection_id] = client_key
            self.client_connections[client_key] = connection_id

        if previous_connection_id and previous_connection_id != connection_id:
            previous_socket = self.active_connections.get(previous_connection_id)
            if previous_socket:
                try:
                    await asyncio.wait_for(
                        previous_socket.close(code=4002, reason="Conexao substituida pelo mesmo cliente"),
                        timeout=2.0,
                    )
                except Exception:
                    pass
            self.disconnect(user_id, previous_connection_id)

        if not was_online:
            await self.broadcast_user_status(user_id, True, company_id)
        await self.send_online_users(websocket, company_id)

    def disconnect(self, user_id: int, connection_id: Optional[int] = None) -> bool:
        connection_ids = self.user_connections.get(user_id, set())
        targets = set(connection_ids) if connection_id is None else {connection_id}
        for current_id in targets:
            if current_id not in connection_ids:
                continue
            connection_ids.discard(current_id)
            self.connection_users.pop(current_id, None)
            self.connection_company_scope.pop(current_id, None)
            client_key = self.connection_client_keys.pop(current_id, "")
            if client_key and self.client_connections.get(client_key) == current_id:
                self.client_connections.pop(client_key, None)
            self.active_connections.pop(current_id, None)
        if connection_ids:
            return True
        self.user_connections.pop(user_id, None)
        return False

    async def send_personal_message(self, message: str, user_id: int):
        disconnected = []
        for connection_id in list(self.user_connections.get(user_id, set())):
            websocket = self.active_connections.get(connection_id)
            if not websocket:
                disconnected.append(connection_id)
                continue
            if not await self._safe_send(websocket, message):
                disconnected.append(connection_id)
        for connection_id in disconnected:
            self.disconnect(user_id, connection_id)

    async def broadcast(self, message: str, exclude_user: Optional[int] = None, company_id: Optional[str] = None):
        disconnected = []
        for connection_id, websocket in list(self.active_connections.items()):
            user_id = self.connection_users.get(connection_id)
            if exclude_user and user_id == exclude_user:
                continue
            if company_id and self.connection_company_scope.get(connection_id) != company_id:
                continue
            if websocket and not await self._safe_send(websocket, message):
                disconnected.append((user_id, connection_id))

        for user_id, connection_id in disconnected:
            if user_id is not None:
                self.disconnect(user_id, connection_id)

    async def broadcast_user_status(self, user_id: int, is_online: bool, company_id: Optional[str] = None):
        with SessionLocal() as db:
            user = db.query(User).filter(User.id == user_id).first()
            if not user:
                return
            message = {
                "type": "user_status",
                "user_id": user_id,
                "username": user.username,
                "full_name": user.full_name,
                "nickname": user.nickname,
                "is_online": is_online,
                "presence_status": self.presence_status.get(user_id, "online"),
            }
            if not company_id:
                connection_ids = self.user_connections.get(user_id, set())
                connection_id = next(iter(connection_ids), None)
                company_id = self.connection_company_scope.get(connection_id)
        await self.broadcast(json.dumps(message), exclude_user=user_id, company_id=company_id)

    async def send_online_users(self, websocket: WebSocket, company_id: str):
        online_users = []
        seen_user_ids = set()
        with SessionLocal() as db:
            for user_id in self.user_connections.keys():
                connection_ids = self.user_connections.get(user_id, set())
                if not any(self.connection_company_scope.get(item) == company_id for item in connection_ids):
                    continue
                user = db.query(User).filter(User.id == user_id).first()
                if user:
                    seen_user_ids.add(user.id)
                    online_users.append({"id": user.id, "username": user.username, "full_name": user.full_name, "nickname": user.nickname, "presence_status": self.presence_status.get(user.id, "online")})

            assistants = (
                db.query(User)
                .join(CompanyUser, CompanyUser.user_id == User.id)
                .filter(
                    CompanyUser.company_id == company_id,
                    CompanyUser.status == ACTIVE_STATUS,
                    User.is_active == True,
                    User.username.in_([BOLT_USERNAME, MITTY_USERNAME]),
                )
                .all()
            )
            for assistant in assistants:
                if assistant.id in seen_user_ids:
                    continue
                online_users.append({"id": assistant.id, "username": assistant.username, "full_name": assistant.full_name, "nickname": assistant.nickname, "presence_status": "online", "is_virtual_assistant": True})

        await self._safe_send(websocket, json.dumps({"type": "online_users", "users": online_users}))


manager = ConnectionManager()


DEFAULT_ADMIN_USERNAME = "admin"
DEFAULT_ADMIN_EMAIL = "admin@voltcorp.com"
DEFAULT_ADMIN_NAME = "Volt Admin"
DEFAULT_ADMIN_PASSWORD = os.getenv("VOLT_CHAT_ADMIN_PASSWORD", "VoltAdmin@172839")


def ensure_default_admin_account():
    with SessionLocal() as db:
        admin = db.query(User).filter(
            or_(User.username == DEFAULT_ADMIN_USERNAME, User.email == DEFAULT_ADMIN_EMAIL)
        ).first()
        if not admin:
            admin = User(
                username=DEFAULT_ADMIN_USERNAME,
                uuid=str(uuid.uuid4()),
                email=DEFAULT_ADMIN_EMAIL,
                full_name=DEFAULT_ADMIN_NAME,
                hashed_password=get_password_hash(DEFAULT_ADMIN_PASSWORD),
                phone="1000",
                is_platform_admin=True,
                must_change_password=False,
                status=ACTIVE_STATUS,
                is_active=True,
                access_level="master",
                department="Administracao",
                role_title="Administrador do sistema",
                phone_extension="1000",
                birthday="01-01-90",
            )
            db.add(admin)
        else:
            admin.username = DEFAULT_ADMIN_USERNAME
            admin.email = DEFAULT_ADMIN_EMAIL
            admin.full_name = DEFAULT_ADMIN_NAME
            admin.is_platform_admin = True
            admin.must_change_password = False
            admin.status = ACTIVE_STATUS
            admin.is_active = True
            admin.access_level = "master"
            if not verify_password(DEFAULT_ADMIN_PASSWORD, admin.hashed_password):
                admin.hashed_password = get_password_hash(DEFAULT_ADMIN_PASSWORD)
            admin.updated_at = datetime.utcnow()
        db.commit()


def create_default_users():
    default_users = [
        {
            "username": DEFAULT_ADMIN_USERNAME,
            "email": DEFAULT_ADMIN_EMAIL,
            "full_name": DEFAULT_ADMIN_NAME,
            "password": DEFAULT_ADMIN_PASSWORD,
            "access_level": "master",
            "department": "Administracao",
            "role_title": "Administrador do sistema",
            "phone_extension": "1000",
            "phone": "1000",
            "birthday": "01-01-90",
        },
        {
            "username": "coordenador",
            "email": "coord@voltcorp.com",
            "full_name": "Coordenador Sistema",
            "password": "coord123",
            "access_level": "coordenador",
            "department": "TI",
            "role_title": "Coordenador de TI",
            "phone_extension": "2000",
            "phone": "2000",
            "birthday": "02-02-90",
        },
        {
            "username": "usuario",
            "email": "user@voltcorp.com",
            "full_name": "Usuario Padrao",
            "password": "user123",
            "access_level": "usuario",
            "department": "TI",
            "role_title": "Analista",
            "phone_extension": "2001",
            "phone": "2001",
            "birthday": "03-03-90",
        },
    ]

    with SessionLocal() as db:
        for user_data in default_users:
            existing = db.query(User).filter(User.username == user_data["username"]).first()
            if existing:
                changed = False
                for field in ["department", "role_title", "phone_extension", "birthday"]:
                    if not getattr(existing, field, None):
                        setattr(existing, field, user_data[field])
                        changed = True
                if existing.access_level == "padrao":
                    existing.access_level = "usuario"
                    changed = True
                if existing.access_level == "master" and not existing.is_platform_admin:
                    existing.is_platform_admin = True
                    changed = True
                if not existing.status:
                    existing.status = ACTIVE_STATUS if existing.is_active else INACTIVE_STATUS
                    changed = True
                if not existing.phone:
                    existing.phone = user_data["phone"]
                    changed = True
                if changed:
                    existing.updated_at = datetime.utcnow()
                continue

            db.add(
                User(
                    username=user_data["username"],
                    uuid=str(uuid.uuid4()),
                    email=user_data["email"],
                    full_name=user_data["full_name"],
                    hashed_password=get_password_hash(user_data["password"]),
                    phone=user_data["phone"],
                    is_platform_admin=user_data["access_level"] == "master",
                    must_change_password=False,
                    status=ACTIVE_STATUS,
                    access_level=user_data["access_level"],
                    department=user_data["department"],
                    role_title=user_data["role_title"],
                    phone_extension=user_data["phone_extension"],
                    birthday=user_data["birthday"],
                )
            )
        db.commit()


def split_sql_statements(sql: str) -> list[str]:
    statements = []
    current = []
    in_single_quote = False
    in_double_quote = False

    for char in sql:
        if char == "'" and not in_double_quote:
            in_single_quote = not in_single_quote
        elif char == '"' and not in_single_quote:
            in_double_quote = not in_double_quote

        if char == ";" and not in_single_quote and not in_double_quote:
            statement = "".join(current).strip()
            if statement:
                statements.append(statement)
            current = []
            continue

        current.append(char)

    statement = "".join(current).strip()
    if statement:
        statements.append(statement)
    return statements


def normalize_text(value: str) -> str:
    replacements = {
        "á": "a",
        "à": "a",
        "ã": "a",
        "â": "a",
        "é": "e",
        "ê": "e",
        "í": "i",
        "ó": "o",
        "ô": "o",
        "õ": "o",
        "ú": "u",
        "ç": "c",
    }
    normalized = str(value or "").lower()
    for source, target in replacements.items():
        normalized = normalized.replace(source, target)
    normalized = unicodedata.normalize("NFKD", normalized)
    normalized = "".join(character for character in normalized if not unicodedata.combining(character))
    return " ".join(normalized.split())


def contains_phrase(text: str, phrase: str) -> bool:
    normalized_text = normalize_text(text)
    normalized_phrase = normalize_text(phrase)
    if not normalized_phrase:
        return False
    return bool(re.search(rf"(?<!\w){re.escape(normalized_phrase)}(?!\w)", normalized_text))


def contains_fuzzy_word(text: str, expected_words: list[str], cutoff: float = 0.72) -> bool:
    """Aceita erros curtos de digitacao e transposicao sem corrigir o texto do usuario."""
    tokens = re.findall(r"[a-z0-9]+", normalize_text(text))
    for token in tokens:
        if len(token) < 4:
            continue
        for expected in expected_words:
            normalized_expected = normalize_text(expected)
            if token == normalized_expected:
                return True
            if abs(len(token) - len(normalized_expected)) <= 2 and difflib.SequenceMatcher(
                None, token, normalized_expected
            ).ratio() >= cutoff:
                return True
    return False


def normalize_ticket_urgency(value: Optional[str]) -> Optional[str]:
    normalized = normalize_text(value)
    aliases = {
        "leve": "Leve",
        "baixa": "Leve",
        "baixo": "Leve",
        "moderado": "Moderado",
        "moderada": "Moderado",
        "medio": "Moderado",
        "media": "Moderado",
        "alto": "Alto",
        "alta": "Alto",
        "urgente": "Urgente",
        "extrema urgencia": "Extrema Urgência",
        "urgencia extrema": "Extrema Urgência",
        "critico": "Extrema Urgência",
        "critica": "Extrema Urgência",
    }
    if not normalized:
        return None
    urgency = aliases.get(normalized)
    if not urgency:
        raise HTTPException(
            status_code=400,
            detail=f"Urgencia invalida. Escolha: {', '.join(TICKET_URGENCY_LEVELS)}.",
        )
    return urgency


def infer_ticket_urgency(text: str) -> Optional[str]:
    normalized = normalize_text(text)
    ordered_aliases = [
        (["extrema urgencia", "urgencia extrema", "prioridade maxima", "critico", "critica"], "Extrema Urgência"),
        (["urgente"], "Urgente"),
        (["alto", "alta"], "Alto"),
        (["moderado", "moderada", "medio", "media"], "Moderado"),
        (["leve", "baixo", "baixa", "sem pressa", "quando der"], "Leve"),
    ]
    for aliases, urgency in ordered_aliases:
        if any(contains_phrase(normalized, alias) for alias in aliases):
            return urgency
    return None


def infer_priority(text: str, task_format: bool = False) -> str:
    normalized = normalize_text(text)
    if any(contains_phrase(normalized, word) for word in ["urgente", "critico", "critica", "alta", "prioridade maxima", "importante"]):
        return "high" if task_format else "Alta"
    if any(contains_phrase(normalized, word) for word in ["baixa", "sem pressa", "quando der", "normal baixa"]):
        return "low" if task_format else "Baixa"
    return "medium" if task_format else "Media"


def infer_due_date(text: str) -> Optional[str]:
    normalized = normalize_text(text)
    today = datetime.utcnow().date()
    if "depois de amanha" in normalized:
        return (today + timedelta(days=2)).isoformat()
    if "amanha" in normalized:
        return (today + timedelta(days=1)).isoformat()
    if "hoje" in normalized:
        return today.isoformat()
    if "proxima semana" in normalized or "semana que vem" in normalized:
        return (today + timedelta(days=7)).isoformat()

    tokens = normalized.replace(",", " ").replace(".", " ").split()
    for token in tokens:
        parts = token.split("/")
        if len(parts) not in {2, 3}:
            continue
        try:
            day = int(parts[0])
            month = int(parts[1])
            year = int(parts[2]) if len(parts) == 3 else today.year
            return datetime(year, month, day).date().isoformat()
        except ValueError:
            continue
    return None


def infer_intent(text: str) -> str:
    normalized = normalize_text(text)
    task_words = ["tarefa", "task", "kanban", "atividade", "afazer", "a fazer", "pendencia"]
    ticket_words = ["ticket", "chamado", "solicitacao", "incidente", "atendimento"]
    meeting_words = ["reuniao", "meeting", "videochamada", "chamada de video", "videoconferencia"]
    reminder_words = ["compromisso", "lembrete", "agenda", "agendamento", "evento"]

    task_positions = [normalized.find(normalize_text(word)) for word in task_words if contains_phrase(normalized, word)]
    ticket_positions = [normalized.find(normalize_text(word)) for word in ticket_words if contains_phrase(normalized, word)]
    if not task_positions and contains_fuzzy_word(normalized, ["tarefa", "task", "atividade", "pendencia"]):
        task_positions = [0]
    if not ticket_positions and contains_fuzzy_word(normalized, ["ticket", "chamado", "solicitacao", "incidente"]):
        ticket_positions = [0]
    if not task_positions and not ticket_positions and any(contains_phrase(normalized, word) for word in ["suporte", "problema"]):
        ticket_positions = [0]

    unsupported_actions = ["feche", "fechar", "exclua", "excluir", "apague", "apagar", "remova", "remover", "cancele", "cancelar"]
    create_actions = [
        "crie", "criar", "abra", "abrir", "registre", "registrar", "adicione",
        "adicionar", "novo", "nova", "agende", "agendar", "marque", "marcar",
        "lembre", "lembrar",
    ]
    if any(contains_phrase(normalized, action) for action in unsupported_actions) and not any(
        contains_phrase(normalized, action) for action in create_actions
    ):
        raise HTTPException(
            status_code=400,
            detail="Esse comando altera ou remove um item existente e ainda nao e executado automaticamente.",
        )

    if ticket_positions and (not task_positions or min(ticket_positions) <= min(task_positions)):
        return "ticket"
    if any(contains_phrase(normalized, word) for word in meeting_words):
        return "meeting"
    if task_positions:
        scheduled_markers = ["agende", "agendar", "marque", "marcar", "na agenda", "calendario"]
        return "agenda_task" if any(contains_phrase(normalized, marker) for marker in scheduled_markers) else "task"
    if any(contains_phrase(normalized, word) for word in reminder_words):
        return "reminder"

    raise HTTPException(
        status_code=400,
        detail=(
            "Nao identifiquei ticket, tarefa, reuniao ou compromisso. "
            "Explique o que deseja criar e, para agenda, informe dia e horario."
        ),
    )


def extract_assistant_time_matches(value: str) -> list[tuple[int, int, int]]:
    """Aceita 14h, 14h30, 14:30, 1430, 930, 14 30 e formas com 'horas'."""
    normalized = normalize_text(value)
    matches: list[tuple[int, int, int, int]] = []

    def add(start: int, end: int, hour: int, minute: int = 0):
        if not (0 <= hour <= 23 and 0 <= minute <= 59):
            return
        if any(not (end <= current[0] or start >= current[1]) for current in matches):
            return
        matches.append((start, end, hour, minute))

    detailed_patterns = [
        r"(?<!\d)([01]?\d|2[0-3])\s*(?:h|:)\s*([0-5]\d)(?:\s*min(?:utos?)?)?\b",
        r"(?<!\d)([01]?\d|2[0-3])\s+horas?\s*(?:e\s*)?([0-5]\d)(?:\s*min(?:utos?)?)?\b",
        r"(?<!\d)([01]?\d|2[0-3])\s+([0-5]\d)(?!\d)",
    ]
    for pattern in detailed_patterns:
        for match in re.finditer(pattern, normalized):
            add(match.start(), match.end(), int(match.group(1)), int(match.group(2)))

    compact_patterns = [
        r"(?:\b(?:as|das|ate as|até as|para as|a partir das|horario|horário)\s+|complemento do usuario:\s*|^)(\d{3,4})\b",
        r"(?<![\d/])(\d{3,4})\s*$",
    ]
    for pattern in compact_patterns:
        for match in re.finditer(pattern, normalized):
            digits = match.group(1)
            hour, minute = int(digits[:-2]), int(digits[-2:])
            add(match.start(), match.end(), hour, minute)

    hour_patterns = [
        r"(?<!\d)([01]?\d|2[0-3])\s*(?:h|hr|hrs)\b",
        r"(?<!\d)([01]?\d|2[0-3])\s+horas?\b",
        r"\b(?:as|das|ate as|até as|para as|a partir das)\s+([01]?\d|2[0-3])\b",
        r"(?<![\d/:])([01]?\d|2[0-3])\s*$",
    ]
    for pattern in hour_patterns:
        for match in re.finditer(pattern, normalized):
            add(match.start(), match.end(), int(match.group(1)), 0)

    matches.sort(key=lambda item: item[0])
    return [(start, hour, minute) for start, _end, hour, minute in matches]


def assistant_schedule_requirements(value: str) -> list[str]:
    normalized = normalize_text(value)
    date_markers = [
        "hoje", "amanha", "depois de amanha", "segunda", "terca", "quarta",
        "quinta", "sexta", "sabado", "domingo",
    ]
    has_date = bool(
        re.search(r"(?<!\d)\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?(?!\d)", normalized)
        or any(contains_phrase(normalized, marker) for marker in date_markers)
    )
    missing = []
    if not has_date:
        missing.append("date")
    if not extract_assistant_time_matches(normalized):
        missing.append("time")
    return missing


def infer_assistant_schedule(value: str) -> tuple[datetime, datetime]:
    normalized = normalize_text(value)
    local_now = datetime.now(BOLT_DAILY_SUMMARY_TIMEZONE)
    target_date = local_now.date()

    explicit_date = re.search(r"(?<!\d)(\d{1,2})[/-](\d{1,2})(?:[/-](\d{2,4}))?(?!\d)", normalized)
    if explicit_date:
        day, month = int(explicit_date.group(1)), int(explicit_date.group(2))
        year_text = explicit_date.group(3)
        year = local_now.year if not year_text else int(year_text)
        if year < 100:
            year += 2000
        try:
            target_date = datetime(year, month, day).date()
        except ValueError as exc:
            raise HTTPException(status_code=400, detail="A data informada para a agenda é inválida.") from exc
    elif "depois de amanha" in normalized:
        target_date += timedelta(days=2)
    elif "amanha" in normalized:
        target_date += timedelta(days=1)
    elif "hoje" not in normalized:
        weekdays = {
            "segunda": 0, "terca": 1, "quarta": 2, "quinta": 3,
            "sexta": 4, "sabado": 5, "domingo": 6,
        }
        for label, weekday in weekdays.items():
            if contains_phrase(normalized, label):
                days_ahead = (weekday - target_date.weekday()) % 7
                target_date += days_ahead or 7
                break

    time_matches = extract_assistant_time_matches(normalized)
    if not time_matches:
        raise HTTPException(
            status_code=400,
            detail="A Mitty precisa do horário. Você pode informar 14h, 14h30, 14:30 ou 1430.",
        )

    _, start_hour, start_minute = time_matches[0]
    starts_local = datetime(target_date.year, target_date.month, target_date.day, start_hour, start_minute, tzinfo=BOLT_DAILY_SUMMARY_TIMEZONE)
    if len(time_matches) > 1:
        _, end_hour, end_minute = time_matches[1]
        ends_local = datetime(target_date.year, target_date.month, target_date.day, end_hour, end_minute, tzinfo=BOLT_DAILY_SUMMARY_TIMEZONE)
        if ends_local <= starts_local:
            ends_local += timedelta(days=1)
    else:
        ends_local = starts_local + timedelta(hours=1)

    if ends_local <= local_now:
        raise HTTPException(status_code=400, detail="O horário informado já passou. Escolha um horário futuro.")
    return starts_local.astimezone(timezone.utc), ends_local.astimezone(timezone.utc)

def infer_assistant_link(value: str) -> Optional[str]:
    match = re.search(r"https?://[^\s<>\"]+", value, flags=re.IGNORECASE)
    return normalize_meeting_link(match.group(0).rstrip(".,;)")) if match else None

def infer_departments(text: str, department_names: list[str], fallback: Optional[str]) -> list[str]:
    normalized_text = normalize_text(text)
    available = {normalize_text(name): name for name in department_names if name}
    aliases = {
        "rh": "RH",
        "recursos humanos": "RH",
        "financeiro": "Financeiro",
        "financas": "Financeiro",
        "ti": "TI",
        "tecnologia": "TI",
        "suporte": "Suporte",
        "comercial": "Comercial",
        "operacao": "Operacao",
        "operacoes": "Operacao",
        "produto": "Produto",
    }
    candidates = {}
    for normalized_name, name in available.items():
        candidates[normalized_name] = name
    for alias, default_name in aliases.items():
        resolved_name = available.get(normalize_text(default_name), default_name)
        candidates.setdefault(normalize_text(alias), resolved_name)

    occurrences = []
    for phrase, resolved_name in candidates.items():
        for match in re.finditer(rf"(?<!\w){re.escape(phrase)}(?!\w)", normalized_text):
            occurrences.append((match.start(), match.end(), resolved_name))
    occurrences.sort(key=lambda item: (item[0], -(item[1] - item[0])))

    matches = []
    previous_end = None
    for start, end, resolved_name in occurrences:
        prefix = normalized_text[max(0, start - 70):start]
        direct_target = bool(
            re.search(
                r"(?:para|em|no|na|nos|nas|setor|setores|departamento|departamentos|area|areas)\s+(?:o|a|os|as|de|do|da)?\s*$",
                prefix,
            )
        )
        intent_target = bool(re.search(r"(?:ticket|chamado|tarefa|atividade)\s+(?:de|do|da)?\s*$", prefix))
        chained_target = previous_end is not None and bool(re.fullmatch(r"\s*(?:,|e|/|\+|\s)+\s*", normalized_text[previous_end:start]))
        if direct_target or intent_target or chained_target:
            if resolved_name not in [item[2] for item in matches]:
                matches.append((start, end, resolved_name))
            previous_end = end

    ordered = [item[2] for item in matches]
    return ordered or [fallback or "Operacao"]


def infer_assignees(text: str, users: list[User]) -> list[User]:
    normalized = normalize_text(text)
    short_names = {}
    short_prefixes = {}
    for user in users:
        first_name = normalize_text(user.full_name).split(" ", 1)[0]
        if len(first_name) >= 3:
            short_names.setdefault(first_name, []).append(user)
            # Accept a short name only when it identifies one user in the company.
            for length in range(3, len(first_name)):
                short_prefixes.setdefault(first_name[:length], []).append(user)

    matches = {}
    for user in users:
        candidates = [user.email, user.username, user.full_name, user.nickname or ""]
        first_name = normalize_text(user.full_name).split(" ", 1)[0]
        if len(short_names.get(first_name, [])) == 1:
            candidates.append(first_name)
        positions = []
        for candidate in candidates:
            normalized_candidate = normalize_text(candidate)
            if len(normalized_candidate) < 3:
                continue
            match = re.search(rf"(?<!\w){re.escape(normalized_candidate)}(?!\w)", normalized)
            if match:
                positions.append(match.start())
        for prefix, prefix_users in short_prefixes.items():
            if len(prefix_users) != 1 or prefix_users[0].id != user.id:
                continue
            match = re.search(rf"(?<!\w){re.escape(prefix)}(?!\w)", normalized)
            if match:
                positions.append(match.start())
        if positions:
            matches[user.id] = (min(positions), user)
    return [item[1] for item in sorted(matches.values(), key=lambda item: item[0])]
def extract_external_meeting_counterpart(text: str, assignees: list[User]) -> Optional[str]:
    """Extrai um contato externo citado em "com ..." sem transforma-lo em usuario interno."""
    clean = re.sub(r"https?://[^\s<>\"]+", " ", str(text or ""), flags=re.IGNORECASE)
    matches = list(re.finditer(
        r"\bcom\s+(.+?)(?=\s+(?:no\s+dia|dia|em\s+\d|para\s+o\s+dia|às|as\s+\d|hoje|amanh[ãa]|depois\s+de\s+amanh[ãa])\b|[.,;]|$)",
        clean,
        flags=re.IGNORECASE,
    ))
    if not matches:
        return None
    candidate = matches[-1].group(1).strip(" .,:;-")
    candidate = re.sub(r"^(?:o|a|os|as)\s+", "", candidate, count=1, flags=re.IGNORECASE).strip()
    if not candidate:
        return None

    normalized_candidate = normalize_text(candidate)
    for user in assignees:
        internal_names = {
            normalize_text(user.full_name),
            normalize_text(user.full_name).split(" ", 1)[0],
            normalize_text(user.nickname or ""),
            normalize_text(user.username),
        }
        if any(name and contains_phrase(normalized_candidate, name) for name in internal_names):
            return None
        first_name = normalize_text(user.full_name).split(" ", 1)[0]
        candidate_words = normalized_candidate.split()
        if any(len(word) >= 3 and first_name.startswith(word) for word in candidate_words):
            return None
    return candidate[:160]


def clean_meeting_subject(subject: str, external_counterpart: Optional[str]) -> str:
    cleaned = str(subject or "").strip(" .,:;-")
    if external_counterpart:
        counterpart_pattern = re.escape(external_counterpart)
        cleaned = re.sub(
            rf"\s+com\s+(?:o|a|os|as)?\s*{counterpart_pattern}\s*$",
            "",
            cleaned,
            count=1,
            flags=re.IGNORECASE,
        ).strip(" .,:;-")
        cleaned = re.sub(
            rf"^com\s+(?:o|a|os|as)?\s*{counterpart_pattern}$",
            "",
            cleaned,
            count=1,
            flags=re.IGNORECASE,
        ).strip(" .,:;-")

    cleaned = re.sub(r"^(?:de|sobre)\s+", "", cleaned, count=1, flags=re.IGNORECASE).strip()
    if normalize_text(cleaned) in {"", "para", "em", "com", "reuniao", "reuniao para"}:
        cleaned = ""
    return cleaned[:160]


def extract_assistant_subject(text: str, intent: str) -> str:
    clean = " ".join(text.strip().split())
    # URLs are meeting metadata. Remove them before interpreting a colon as a title separator.
    clean = re.sub(r"https?://[^\s<>\"]+", " ", clean, flags=re.IGNORECASE)
    clean = re.sub(r"^(?:mitty|volt)\s*[,;:-]?\s*", "", clean, count=1, flags=re.IGNORECASE)
    clean = re.sub(
        r"^(?:crie|criar|abra|abrir|envie|enviar|registre|registrar|adicione|adicionar|agende|agendar|marque|marcar)\s+"
        r"(?:(?:um|uma|o|a)\s+)?(?:ticket|tiket|ticekt|tickte|chamado|tarefa|tareaf|task|reuniao|meeting|compromisso|lembrete|evento)\b",
        "",
        clean,
        count=1,
        flags=re.IGNORECASE,
    ).strip(" :-") or clean
    prefixes = [
        "crie um ticket",
        "crie o ticket",
        "criar ticket",
        "abra um ticket",
        "abra o ticket",
        "abrir ticket",
        "envie um ticket",
        "enviar ticket",
        "registrar chamado",
        "registre chamado",
        "crie uma tarefa",
        "crie a tarefa",
        "criar tarefa",
        "adicione uma tarefa",
        "adicionar tarefa",
        "agende uma reuniao",
        "agendar reuniao",
        "marque uma reuniao",
        "marcar reuniao",
        "agende um compromisso",
        "agendar compromisso",
        "crie um lembrete",
        "criar lembrete",
        "registrar",
        "registre",
    ]
    normalized = normalize_text(clean)
    subject = clean
    for prefix in prefixes:
        if normalized.startswith(prefix):
            subject = clean[len(prefix):].strip(" :-")
            break

    if intent in {"meeting", "reminder", "agenda_task"}:
        subject = re.sub(
            r"^(?:(?:um|uma|o|a)\s+)?(?:reuni(?:\u00e3o|ao)|meeting|compromisso|lembrete|evento)\b\s*",
            "",
            subject,
            count=1,
            flags=re.IGNORECASE,
        ).strip(" :-")
        subject = re.sub(
            r"^(?:para\s+)?(?:hoje|amanh(?:\u00e3|a)|depois\s+de\s+amanh(?:\u00e3|a)|segunda|terca|quarta|quinta|sexta|sabado|domingo)\b\s*",
            "",
            subject,
            count=1,
            flags=re.IGNORECASE,
        ).strip(" :-")

    lowered = normalize_text(subject)
    if " sobre " in lowered:
        index = lowered.find(" sobre ")
        subject = subject[index + len(" sobre "):].strip(" :-")
    elif ":" in subject:
        subject = subject.split(":", 1)[1].strip()

    cut_markers = [" e atribua", " atribua", " ate amanha", " para amanha", " amanha as ", " hoje as ", " amanha ", " hoje ", " às ", " as "]
    lowered = normalize_text(subject)
    for marker in cut_markers:
        index = lowered.find(normalize_text(marker))
        if index > 8:
            subject = subject[:index].strip(" .,;-")
            lowered = normalize_text(subject)

    # Em itens de agenda, data, hora e participantes são metadados — nunca parte do título.
    # Ex.: "Agende um compromisso ECF Shopee dia 31/07 às 11h e inclua Ana".
    if intent in {"meeting", "reminder", "agenda_task"}:
        schedule_suffix_patterns = [
            r"\s+(?:no\s+)?dia\s+\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?\b.*$",
            r"\s+\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?\b.*$",
            r"\s+(?:hoje|amanha|depois\s+de\s+amanha|segunda|terca|quarta|quinta|sexta|sabado|domingo)\b.*$",
            r"\s+(?:às|as)\s+\d{1,2}(?:(?:h|:)\d{0,2})?\b.*$",
            r"\s+(?:e\s+)?(?:inclua|incluindo|participantes?|convide|convidar)\b.*$",
        ]
        for pattern in schedule_suffix_patterns:
            subject = re.sub(pattern, "", subject, count=1, flags=re.IGNORECASE).strip(" .,;-")
    filler = ["urgente", "alta prioridade", "prioridade alta", "por favor"]
    words = [word for word in subject.split() if normalize_text(word) not in filler]
    subject = " ".join(words).strip(" .,;-")
    if not subject:
        subject = "atendimento interno" if intent == "ticket" else "atividade operacional"
    return subject[:160]


def build_assistant_title(subject: str, intent: str) -> str:
    words = subject.split()
    short_subject = " ".join(words[:9]).strip()
    if len(words) > 9:
        short_subject = f"{short_subject}..."
    prefixes = {
        "ticket": "Ticket",
        "task": "Tarefa",
        "agenda_task": "Tarefa agendada",
        "meeting": "Reunião",
        "reminder": "Compromisso",
    }
    return f"{prefixes.get(intent, 'Item')}: {short_subject}".strip()[:180]

def remove_leading_assignees(subject: str, assignees: list[User]) -> str:
    if not assignees or not contains_phrase(subject, "para") or not normalize_text(subject).startswith("para "):
        return subject
    cleaned = re.sub(r"^para\s+", "", subject, count=1, flags=re.IGNORECASE).strip()
    for user in assignees:
        candidates = [user.full_name, user.full_name.split(" ", 1)[0], user.nickname or "", user.username]
        matched = False
        for candidate in sorted({item for item in candidates if item}, key=len, reverse=True):
            updated = re.sub(rf"^{re.escape(candidate)}(?=\s|,|$)", "", cleaned, count=1, flags=re.IGNORECASE)
            if updated != cleaned:
                cleaned = re.sub(r"^(?:\s*,\s*|\s+e\s+|\s+)", "", updated, count=1, flags=re.IGNORECASE).strip()
                matched = True
                break
        if not matched:
            continue
    return cleaned or subject


def remove_meeting_assignees(subject: str, assignees: list[User]) -> str:
    """Remove do titulo participantes internos, inclusive apelidos curtos unicos."""
    cleaned = str(subject or "")
    for user in assignees:
        full_name = str(user.full_name or "").strip()
        first_name = full_name.split(" ", 1)[0]
        variants = {full_name, first_name, str(user.nickname or "").strip(), str(user.username or "").strip()}
        if len(first_name) >= 3:
            variants.update(first_name[:length] for length in range(3, len(first_name) + 1))
        for candidate in sorted((item for item in variants if len(item) >= 3), key=len, reverse=True):
            cleaned = re.sub(
                rf"\s+com\s+(?:o|a)\s*{re.escape(candidate)}(?=\s|,|$)",
                "",
                cleaned,
                flags=re.IGNORECASE,
            )
    return cleaned.strip(" .,:;-") or subject


def build_assistant_scope(subject: str, intent: str, departments: list[str], assignees: list[User], due_date: Optional[str]) -> str:
    owner = ", ".join(user.full_name for user in assignees) if assignees else "o proprio usuario"
    department_text = ", ".join(departments)
    due_text = f" Prazo identificado: {due_date}." if due_date else ""
    if intent == "ticket":
        return (
            f"Escopo: analisar {subject}, identificar causa, registrar evidencias e conduzir o atendimento ate a "
            f"resolucao ou proximo encaminhamento. Setores: {department_text}. Responsaveis: {owner}.{due_text}"
        )
    if intent in {"meeting", "reminder", "agenda_task"}:
        return (
            f"Agenda organizada pela Mitty: {subject}. Participantes: {owner}. "
            "Todos receberao os lembretes configurados."
        )
    return (
        f"Escopo: executar {subject}, organizar as etapas necessarias, registrar andamento no Kanban e concluir com "
        f"validacao do responsavel. Responsavel: {owner}.{due_text}"
    )

def plan_assistant_action(text: str, current_user: User, db, company_id: str, selected_urgency: Optional[str] = None) -> dict:
    users = (
        db.query(User)
        .join(CompanyUser, CompanyUser.user_id == User.id)
        .filter(
            User.is_active == True,
            ~User.username.in_([BOLT_USERNAME, MITTY_USERNAME]),
            CompanyUser.company_id == company_id,
            CompanyUser.status == ACTIVE_STATUS,
        )
        .all()
    )
    department_names = [
        department.name
        for department in db.query(Department)
        .filter(
            Department.company_id == company_id,
            Department.status == ACTIVE_STATUS,
            Department.name != "Automacao",
        )
        .all()
    ]
    intent = infer_intent(text)
    assignees = infer_assignees(text, users)
    if intent == "task" and any(int(user.id) != int(current_user.id) for user in assignees):
        raise HTTPException(
            status_code=400,
            detail="Tarefas sao pessoais. Posso criar esta tarefa apenas para voce.",
        )
    explicit_people_markers = [
        "atribua", "atribuir", "responsavel", "responsaveis", "participante",
        "participantes", "convide", "convidar",
    ]
    if any(contains_phrase(text, marker) for marker in explicit_people_markers) and not assignees:
        raise HTTPException(
            status_code=400,
            detail="Nao encontrei na empresa a pessoa citada. Informe o nome completo, usuario ou e-mail.",
        )

    departments = (
        infer_departments(text, department_names, user_department_name(db, current_user, company_id))
        if intent == "ticket"
        else []
    )
    department = departments[0] if departments else None
    due_date = infer_due_date(text)
    external_counterpart = (
        extract_external_meeting_counterpart(text, assignees)
        if intent in {"meeting", "reminder", "agenda_task"}
        else None
    )
    subject = extract_assistant_subject(text, intent)
    subject = remove_leading_assignees(subject, assignees)
    if intent in {"meeting", "reminder", "agenda_task"}:
        subject = remove_meeting_assignees(subject, assignees)
        subject = clean_meeting_subject(subject, external_counterpart)
        if not subject and external_counterpart:
            subject = external_counterpart
        if not subject:
            subject = "reunião" if intent == "meeting" else "compromisso"
        subject = subject[:1].upper() + subject[1:]
    title = build_assistant_title(subject, intent)
    # Um compromisso com convidados é uma reunião de agenda, mesmo sem chamada de vídeo.
    if intent == "reminder" and assignees:
        title = f"Reunião: {subject}".strip()[:180]
    ticket_urgency = (
        normalize_ticket_urgency(selected_urgency)
        if selected_urgency
        else infer_ticket_urgency(text)
    )

    starts_at = None
    ends_at = None
    meeting_type = None
    if intent in {"meeting", "reminder", "agenda_task"}:
        starts_at, ends_at = infer_assistant_schedule(text)
        meeting_type = {
            "meeting": "video",
            "reminder": "reminder",
            "agenda_task": "task",
        }[intent]

    description = build_assistant_scope(subject, intent, departments, assignees, due_date)
    if external_counterpart:
        description = f"{description} Contato externo mencionado: {external_counterpart}."

    return {
        "intent": intent,
        "assistant_name": BOLT_NAME if intent == "ticket" else MITTY_NAME,
        "title": title,
        "description": description,
        "subject": subject,
        "external_counterpart": external_counterpart,
        "department": department,
        "departments": departments,
        "assigned_to_id": assignees[0].id if assignees else None,
        "assigned_to_name": assignees[0].full_name if assignees else None,
        "assigned_to_ids": [user.id for user in assignees],
        "assigned_to_names": [user.full_name for user in assignees],
        "ticket_priority": ticket_urgency if intent == "ticket" else None,
        "urgency_options": TICKET_URGENCY_LEVELS,
        "task_priority": infer_priority(text, task_format=True),
        "due_date": due_date,
        "meeting_type": meeting_type,
        "starts_at": starts_at.isoformat() if starts_at else None,
        "ends_at": ends_at.isoformat() if ends_at else None,
        "link_url": infer_assistant_link(text),
        "color": "#7c3aed" if intent in {"meeting", "reminder", "agenda_task"} else None,
    }

def is_ticket_summary_request(value: str) -> bool:
    """Reconhece consultas pessoais de tickets sem confundir com pedidos de criação."""
    normalized = normalize_text(value)
    if not contains_fuzzy_word(value, ["ticket", "tickets", "tiket", "ticekt"], cutoff=0.66):
        return False
    creation_markers = [
        "abra", "abrir", "crie", "criar", "registre", "registrar", "novo ticket", "nova tarefa",
    ]
    if any(contains_phrase(value, marker) for marker in creation_markers):
        return False
    summary_markers = [
        "quantos", "quais", "meus ticket", "meu ticket", "em aberto", "aguardando",
        "pendente", "resumo", "situacao", "status", "minha resposta", "minhas respostas",
    ]
    return any(marker in normalized for marker in summary_markers)


def build_user_ticket_summary(db, user: User, company_id: str) -> dict:
    """Resume apenas tickets criados pelo usuario ou atribuidos explicitamente a ele."""
    assigned_ticket_ids = {
        row[0]
        for row in db.query(TicketAssignee.ticket_id)
        .filter(TicketAssignee.user_id == user.id)
        .all()
    }

    ownership_conditions = [
        Ticket.created_by_id == user.id,
        Ticket.assigned_to_id == user.id,
    ]
    if assigned_ticket_ids:
        ownership_conditions.append(Ticket.id.in_(assigned_ticket_ids))

    tickets = (
        db.query(Ticket)
        .filter(
            Ticket.company_id == company_id,
            func.lower(func.coalesce(Ticket.status, "")) != "resolvido",
            or_(*ownership_conditions),
        )
        .order_by(Ticket.created_at.desc(), Ticket.id.desc())
        .all()
    )

    latest_by_ticket = {}
    ticket_ids = [ticket.id for ticket in tickets]
    if ticket_ids:
        latest_message_ids = (
            db.query(
                TicketMessage.ticket_id.label("ticket_id"),
                func.max(TicketMessage.id).label("message_id"),
            )
            .filter(TicketMessage.ticket_id.in_(ticket_ids))
            .group_by(TicketMessage.ticket_id)
            .subquery()
        )
        latest_messages = (
            db.query(TicketMessage)
            .join(latest_message_ids, TicketMessage.id == latest_message_ids.c.message_id)
            .all()
        )
        latest_by_ticket = {message.ticket_id: message for message in latest_messages}

    items = []
    for ticket in tickets:
        latest_message = latest_by_ticket.get(ticket.id)
        is_responsible = (
            ticket.assigned_to_id == user.id
            or ticket.id in assigned_ticket_ids
        )
        awaiting_user_response = (
            latest_message.sender_id != user.id
            if latest_message
            else is_responsible
        )
        items.append({
            "id": ticket.id,
            "title": ticket.title,
            "priority": ticket.priority,
            "status": ticket.status,
            "awaiting_user_response": awaiting_user_response,
            "waiting_label": (
                "Aguardando sua resposta"
                if awaiting_user_response
                else "Aguardando resposta de outro participante"
            ),
            "last_message_at": (
                latest_message.created_at.isoformat()
                if latest_message and latest_message.created_at
                else None
            ),
        })

    awaiting_count = sum(1 for item in items if item["awaiting_user_response"])
    return {
        "company_id": company_id,
        "open_ticket_count": len(items),
        "awaiting_user_response_count": awaiting_count,
        "awaiting_other_response_count": len(items) - awaiting_count,
        "tickets": items,
    }


def build_ticket_summary_reply(summary: dict, daily: bool = False) -> str:
    open_count = int(summary.get("open_ticket_count") or 0)
    awaiting_count = int(summary.get("awaiting_user_response_count") or 0)
    other_count = int(summary.get("awaiting_other_response_count") or 0)
    prefix = "Resumo diário do Volt: " if daily else ""
    if open_count == 0:
        return f"{prefix}você não possui tickets em aberto no momento."

    ticket_word = "ticket" if open_count == 1 else "tickets"
    waiting_word = "ticket aguarda" if awaiting_count == 1 else "tickets aguardam"
    other_word = "ticket aguarda" if other_count == 1 else "tickets aguardam"
    reply = (
        f"{prefix}você tem {open_count} {ticket_word} em aberto. "
        f"{awaiting_count} {waiting_word} sua resposta e {other_count} {other_word} "
        "resposta de outro participante."
    )
    awaiting_items = [
        item for item in summary.get("tickets", []) if item.get("awaiting_user_response")
    ]
    if awaiting_items:
        preview = "; ".join(
            f"#{item['id']} - {item['title']}" for item in awaiting_items[:8]
        )
        remaining = len(awaiting_items) - 8
        if remaining > 0:
            preview += f"; e mais {remaining}"
        reply += f" Aguardando sua resposta: {preview}."
    return reply


def is_mitty_summary_request(value: str) -> bool:
    normalized = normalize_text(value)
    query_markers = ("qual", "quais", "quant", "minha", "meus", "tenho", "listar", "mostre", "semana", "aberto")
    subject_markers = ("agenda", "tarefa", "task", "compromisso", "reuniao")
    creation_markers = ("crie", "criar", "agende", "agendar", "marque", "marcar", "adicione", "adicionar")
    # "hoje" sozinho e "na agenda" podem fazer parte de um convite. Uma consulta
    # precisa ter um marcador explicito, ou a forma natural "agenda de hoje".
    is_explicit_query = (
        any(marker in normalized for marker in query_markers)
        or contains_phrase(normalized, "agenda de hoje")
    )
    return (
        is_explicit_query
        and any(marker in normalized for marker in subject_markers)
        and not any(marker in normalized for marker in creation_markers)
    )
def build_mitty_schedule_summary(db, user: User, company_id: str, query: str) -> dict:
    normalized = normalize_text(query)
    local_now = datetime.now(BOLT_DAILY_SUMMARY_TIMEZONE)
    start_local = local_now.replace(hour=0, minute=0, second=0, microsecond=0)
    end_local = start_local + timedelta(days=1)
    period_label = "hoje"
    if "semana" in normalized:
        start_local -= timedelta(days=start_local.weekday())
        end_local = start_local + timedelta(days=7)
        period_label = "esta semana"

    include_tasks = any(marker in normalized for marker in ("tarefa", "task"))
    include_agenda = any(marker in normalized for marker in ("agenda", "compromisso", "reuniao"))
    if not include_tasks and not include_agenda:
        include_tasks = include_agenda = True

    tasks = []
    if include_tasks:
        tasks = (
            db.query(TaskItem)
            .filter(
                TaskItem.company_id == company_id,
                TaskItem.assigned_to_id == user.id,
                TaskItem.status != "done",
            )
            .order_by(TaskItem.created_at.desc())
            .limit(12)
            .all()
        )

    meetings = []
    if include_agenda:
        meetings = db.execute(text("""
            SELECT DISTINCT m.id, m.title, m.meeting_type, m.starts_at, m.ends_at
            FROM meetings m
            LEFT JOIN meeting_participants mp ON mp.meeting_id = m.id
            WHERE m.company_id = :company_id
              AND m.status = 'scheduled'
              AND (m.creator_user_id = :user_id OR mp.user_id = :user_id)
              AND m.starts_at < :end_at
              AND m.ends_at > :start_at
            ORDER BY m.starts_at
            LIMIT 20
        """), {
            "company_id": company_id,
            "user_id": user.id,
            "start_at": start_local.astimezone(timezone.utc),
            "end_at": end_local.astimezone(timezone.utc),
        }).all()

    lines = []
    if include_tasks:
        if tasks:
            lines.append(f"Voce tem {len(tasks)} tarefa(s) em aberto:")
            lines.extend(f"- {task.title} ({task.status})" for task in tasks)
        else:
            lines.append("Voce nao tem tarefas em aberto.")
    if include_agenda:
        if meetings:
            lines.append(f"Na sua agenda de {period_label}:")
            for meeting in meetings:
                starts_at = ensure_utc_datetime(meeting.starts_at).astimezone(BOLT_DAILY_SUMMARY_TIMEZONE)
                kind = "reuniao" if meeting.meeting_type == "video" else (
                    "tarefa" if meeting.meeting_type == "task" else "compromisso"
                )
                lines.append(f"- {starts_at.strftime('%d/%m as %H:%M')} · {kind}: {meeting.title}")
        else:
            lines.append(f"Voce nao tem compromissos na agenda de {period_label}.")
    return {
        "task_count": len(tasks),
        "meeting_count": len(meetings),
        "period": period_label,
        "reply": assistant_greeting(user, "\n".join(lines)),
    }

def prepare_bolt_daily_ticket_summaries(
    summary_date: Optional[str] = None,
    company_id: Optional[str] = None,
) -> list[tuple[int, dict, str]]:
    """Persiste no máximo um resumo por usuário, empresa e dia."""
    summary_day = summary_date or datetime.now(BOLT_DAILY_SUMMARY_TIMEZONE).date().isoformat()
    prepared = []
    with SessionLocal() as db:
        company_query = db.query(Company).filter(Company.status == ACTIVE_STATUS)
        if company_id:
            company_query = company_query.filter(Company.id == company_id)

        for company in company_query.order_by(Company.id.asc()).all():
            bolt = ensure_bolt_user(db, company.id)
            db.commit()
            bolt_id = bolt.id
            members = (
                db.query(User)
                .join(CompanyUser, CompanyUser.user_id == User.id)
                .filter(
                    CompanyUser.company_id == company.id,
                    CompanyUser.status == ACTIVE_STATUS,
                    User.is_active == True,
                    ~User.username.in_([BOLT_USERNAME, MITTY_USERNAME]),
                )
                .order_by(User.id.asc())
                .all()
            )
            seen_user_ids = set()
            for user in members:
                if user.id in seen_user_ids:
                    continue
                seen_user_ids.add(user.id)
                already_sent = db.query(BoltDailyTicketSummary.id).filter(
                    BoltDailyTicketSummary.company_id == company.id,
                    BoltDailyTicketSummary.user_id == user.id,
                    BoltDailyTicketSummary.summary_date == summary_day,
                ).first()
                if already_sent:
                    continue

                summary = build_user_ticket_summary(db, user, company.id)
                message = Message(
                    company_id=company.id,
                    content=build_ticket_summary_reply(summary, daily=True),
                    sender_id=bolt_id,
                    receiver_id=user.id,
                    message_type="ticket_daily_summary",
                )
                db.add(message)
                db.flush()
                db.add(BoltDailyTicketSummary(
                    company_id=company.id,
                    user_id=user.id,
                    summary_date=summary_day,
                    open_ticket_count=summary["open_ticket_count"],
                    awaiting_reply_count=summary["awaiting_user_response_count"],
                    message_id=message.id,
                ))
                try:
                    db.commit()
                except IntegrityError:
                    db.rollback()
                    continue
                db.refresh(message)
                prepared.append((user.id, serialize_message(message, db), summary_day))
    return prepared


async def dispatch_bolt_daily_ticket_summaries(prepared: list[tuple[int, dict, str]]):
    for user_id, message_payload, summary_day in prepared:
        await manager.send_personal_message(
            json.dumps({"type": "new_message", "message": message_payload}),
            user_id,
        )
        await manager.send_personal_message(
            json.dumps({
                "type": "notification",
                "notification": {
                    "id": f"bolt-daily-{summary_day}-{user_id}",
                    "type": "Volt",
                    "title": "Volt - Resumo diário de tickets",
                    "description": message_payload["content"],
                    "time": "Agora",
                    "unread": True,
                    "link": "/tickets",
                },
            }),
            user_id,
        )

def run_startup_migrations():
    if DATABASE_URL.startswith("sqlite"):
        return
    if os.getenv("AUTO_MIGRATE_DB", "true").lower() != "true":
        return

    migrations_dir = Path(__file__).resolve().parent / "db" / "migrations"
    if not migrations_dir.exists():
        return

    with engine.begin() as conn:
        conn.exec_driver_sql("SELECT pg_advisory_xact_lock(hashtext('voltchat_startup_migrations'))")
        conn.exec_driver_sql(
            """
            CREATE TABLE IF NOT EXISTS schema_migrations (
              version TEXT PRIMARY KEY,
              applied_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
            """
        )

        applied_versions = {
            row[0]
            for row in conn.exec_driver_sql("SELECT version FROM schema_migrations")
        }

        for migration in sorted(migrations_dir.glob("*.sql")):
            if migration.stem in applied_versions:
                continue
            sql = migration.read_text(encoding="utf-8")
            for statement in split_sql_statements(sql):
                conn.exec_driver_sql(statement)
            conn.exec_driver_sql(
                "INSERT INTO schema_migrations (version) VALUES (%s) ON CONFLICT (version) DO NOTHING",
                (migration.stem,),
            )


@app.on_event("startup")
async def startup_event():
    run_startup_migrations()
    if DATABASE_URL.startswith("sqlite") or os.getenv("AUTO_SEED_USERS", "false").lower() == "true":
        create_default_users()
    ensure_default_admin_account()
    ensure_default_tenant_records()
    ensure_virtual_assistant_accounts()
    asyncio.create_task(periodic_database_maintenance())
    if not DATABASE_URL.startswith("sqlite"):
        asyncio.create_task(periodic_meeting_reminders())
    if (
        not DATABASE_URL.startswith("sqlite")
        and os.getenv("ENABLE_BOLT_DAILY_SUMMARY", "true").lower() == "true"
    ):
        asyncio.create_task(periodic_bolt_daily_ticket_summaries())


@app.get("/")
async def root():
    return {"message": "Volt Corp API", "version": "1.0.0", "status": "online"}


@app.get("/health")
async def health_check():
    return {
        "status": "healthy",
        "database": "postgres" if "postgresql" in DATABASE_URL else "sqlite",
        "database_pool": DB_POOL_MODE,
        "chat_file_storage": "r2" if r2_chat_file_storage_enabled() else "database",
    }


@app.get("/admin/database-metrics")
async def database_metrics(current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    return database_query_metrics()


@app.get("/internal/hub/tenants/{tenant_global_id}/usage")
async def hub_tenant_usage(tenant_global_id: str, request: Request):
    """Returns aggregate commercial usage for one Hub tenant only.

    This endpoint deliberately exposes no user names, emails, messages, or files.
    The Hub uses the shared tenant_global_id to reconcile the Chat plan capacity.
    """
    require_hub_usage_token(request)
    normalized_tenant_id = str(tenant_global_id or "").strip()
    if not normalized_tenant_id or len(normalized_tenant_id) > 64:
        raise HTTPException(status_code=400, detail="Tenant global invalido.")

    with SessionLocal() as db:
        company = (
            db.query(Company)
            .filter(Company.tenant_global_id == normalized_tenant_id)
            .first()
        )
        if not company:
            raise HTTPException(status_code=404, detail="Empresa do tenant nao encontrada.")

        total_users = (
            db.query(func.count(CompanyUser.id))
            .filter(CompanyUser.company_id == company.id)
            .scalar()
            or 0
        )
        active_users = (
            db.query(func.count(CompanyUser.id))
            .join(User, User.id == CompanyUser.user_id)
            .filter(
                CompanyUser.company_id == company.id,
                CompanyUser.status == ACTIVE_STATUS,
                User.is_active.is_(True),
            )
            .scalar()
            or 0
        )
        stored_file_bytes = (
            db.query(func.coalesce(func.sum(FileUpload.file_size), 0))
            .filter(
                FileUpload.company_id == company.id,
                or_(FileUpload.expires_at.is_(None), FileUpload.expires_at >= datetime.utcnow()),
            )
            .scalar()
            or 0
        )
        backup_bytes = (
            db.query(func.coalesce(func.sum(ChatHistoryBackup.compressed_size), 0))
            .filter(ChatHistoryBackup.company_id == company.id)
            .scalar()
            or 0
        )
        inline_ticket_image_bytes = (
            db.query(func.coalesce(func.sum(func.length(Ticket.image_data)), 0))
            .filter(
                Ticket.company_id == company.id,
                Ticket.image_data.isnot(None),
            )
            .scalar()
            or 0
        )

    storage_breakdown = {
        "uploaded_files_bytes": int(stored_file_bytes),
        "chat_history_backups_bytes": int(backup_bytes),
        "inline_ticket_images_bytes": int(inline_ticket_image_bytes),
    }
    storage_bytes = sum(storage_breakdown.values())
    return {
        "ok": True,
        "source": "volt_chat",
        "tenant_global_id": normalized_tenant_id,
        "company": {
            "id": company.id,
            "name": company.name,
            "status": company.status,
            "hub_enabled": bool(company.hub_enabled),
        },
        "usage": {
            "active_users": int(active_users),
            "total_users": int(total_users),
            "storage_bytes": int(storage_bytes),
            "storage_gb": round(storage_bytes / (1024 ** 3), 4),
            "storage_breakdown": storage_breakdown,
        },
        "measured_at": datetime.now(timezone.utc).isoformat(),
    }


@app.get("/version")
async def version_check():
    return {
        "service": "api",
        "app": "Volt Corp",
        "version": APP_VERSION,
        "commit": os.getenv("RENDER_GIT_COMMIT") or APP_VERSION,
        "build_time": APP_BUILD_TIME,
    }


def r2_is_configured() -> bool:
    return all((R2_ENDPOINT_URL, R2_ACCESS_KEY_ID, R2_SECRET_ACCESS_KEY, R2_BUCKET_NAME))


def r2_client():
    if not r2_is_configured():
        raise HTTPException(
            status_code=503,
            detail="Cloudflare R2 nao configurado no servidor.",
        )
    return boto3.client(
        service_name="s3",
        endpoint_url=R2_ENDPOINT_URL,
        aws_access_key_id=R2_ACCESS_KEY_ID,
        aws_secret_access_key=R2_SECRET_ACCESS_KEY,
        region_name="auto",
        config=BotoConfig(signature_version="s3v4", retries={"max_attempts": 3, "mode": "standard"}),
    )


def r2_chat_file_storage_enabled() -> bool:
    return CHAT_FILE_STORAGE_MODE == "r2" and r2_is_configured()


def r2_chat_file_storage_key(company_id: str, filename: str) -> str:
    safe_company_id = re.sub(r"[^a-zA-Z0-9_-]", "", str(company_id or "")) or "unknown"
    extension = Path(str(filename or "")).suffix.lower()
    return "/".join(part for part in (R2_CHAT_FILE_PREFIX, safe_company_id, f"{uuid.uuid4().hex}{extension}") if part)


def is_r2_chat_file_path(file_path: str) -> bool:
    return str(file_path or "").startswith(R2_CHAT_FILE_URI_PREFIX)


def r2_chat_file_key(file_path: str) -> str:
    if not is_r2_chat_file_path(file_path):
        raise ValueError("Caminho de arquivo R2 invalido.")
    return str(file_path)[len(R2_CHAT_FILE_URI_PREFIX):].strip().lstrip("/")


def upload_r2_chat_file(company_id: str, filename: str, content: bytes, content_type: str) -> str:
    if CHAT_FILE_STORAGE_MODE == "r2" and not r2_is_configured():
        raise HTTPException(status_code=503, detail="Armazenamento privado R2 nao configurado no servidor.")
    key = r2_chat_file_storage_key(company_id, filename)
    try:
        r2_client().put_object(
            Bucket=R2_BUCKET_NAME,
            Key=key,
            Body=content,
            ContentType=content_type,
            Metadata={"company_id": str(company_id)},
        )
    except HTTPException:
        raise
    except Exception as exc:
        print(f"[chat-files] Falha ao enviar arquivo ao R2: {exc}")
        raise HTTPException(status_code=503, detail="Nao foi possivel salvar o arquivo no armazenamento privado.") from exc
    return f"{R2_CHAT_FILE_URI_PREFIX}{key}"


def presign_r2_chat_file(file_path: str) -> str:
    try:
        return r2_client().generate_presigned_url(
            "get_object",
            Params={"Bucket": R2_BUCKET_NAME, "Key": r2_chat_file_key(file_path)},
            ExpiresIn=R2_PRESIGN_EXPIRES_SECONDS,
        )
    except HTTPException:
        raise
    except Exception as exc:
        print(f"[chat-files] Falha ao assinar download R2: {exc}")
        raise HTTPException(status_code=503, detail="Nao foi possivel preparar o download do arquivo.") from exc


def delete_r2_chat_file(file_path: str) -> bool:
    try:
        r2_client().delete_object(Bucket=R2_BUCKET_NAME, Key=r2_chat_file_key(file_path))
        return True
    except Exception as exc:
        print(f"[chat-files] Falha ao remover arquivo R2 expirado: {exc}")
        return False


def load_r2_release_manifest() -> dict:
    """Le o manifesto latest.json diretamente do R2. O banco nao participa do updater."""
    try:
        response = r2_client().get_object(Bucket=R2_BUCKET_NAME, Key=R2_RELEASE_MANIFEST_KEY)
        body = response.get("Body")
        raw = body.read(128 * 1024) if body is not None else b""
        manifest = json.loads(raw.decode("utf-8"))
    except HTTPException:
        raise
    except Exception as exc:
        print(f"[desktop-update] Falha ao ler manifesto do R2: {exc}")
        raise HTTPException(
            status_code=503,
            detail="Nao foi possivel consultar a versao no Cloudflare R2.",
        ) from exc

    required = ("version", "filename", "file_size", "sha256", "storage_key")
    if not isinstance(manifest, dict) or any(not manifest.get(field) for field in required):
        raise HTTPException(status_code=503, detail="Manifesto de atualizacao do R2 invalido.")
    if not re.fullmatch(r"[a-fA-F0-9]{64}", str(manifest.get("sha256") or "")):
        raise HTTPException(status_code=503, detail="SHA-256 invalido no manifesto do R2.")
    return manifest


def presign_r2_release(storage_key: str) -> str:
    if not storage_key:
        raise HTTPException(status_code=503, detail="Release R2 sem chave de armazenamento.")
    try:
        return r2_client().generate_presigned_url(
            "get_object",
            Params={"Bucket": R2_BUCKET_NAME, "Key": storage_key},
            ExpiresIn=R2_PRESIGN_EXPIRES_SECONDS,
        )
    except HTTPException:
        raise
    except Exception as exc:
        print(f"[desktop-update] Falha ao assinar URL do R2: {exc}")
        raise HTTPException(status_code=503, detail="Nao foi possivel preparar o download no Cloudflare R2.") from exc


def r2_release_metadata(manifest: dict) -> dict:
    # Ponte para clientes 0.1.36 e anteriores. A fonte de verdade e o latest.json do R2.
    # O download continua direto R2 -> cliente, sem bytes no Render e sem consulta ao Neon.
    return {
        "version": str(manifest["version"]),
        "platform": str(manifest.get("platform") or R2_RELEASE_PLATFORM),
        "filename": str(manifest["filename"]),
        "content_type": str(manifest.get("content_type") or "application/octet-stream"),
        "file_size": int(manifest["file_size"]),
        "sha256": str(manifest["sha256"]).lower(),
        "storage_mode": "r2",
        "storage_key": str(manifest["storage_key"]),
        "created_at": manifest.get("published_at"),
        "download_url": presign_r2_release(str(manifest["storage_key"])),
        "external_url": False,
        "direct_download": True,
        "source": "cloudflare-r2",
    }


@app.get("/downloads/desktop/latest/meta")
async def latest_desktop_release_meta(client_version: str = Query(default="", max_length=80)):
    # client_version e mantido apenas para compatibilidade com clientes antigos.
    _ = client_version
    return r2_release_metadata(load_r2_release_manifest())


@app.get("/admin/websocket-health")
async def websocket_health(current_user: User = Depends(get_current_user)):
    """Resumo seguro das conexoes em tempo real para diagnostico de churn/duplicidade."""
    ensure_admin(current_user)

    connection_ids = list(manager.active_connections.keys())
    user_ids = [
        manager.connection_users.get(connection_id)
        for connection_id in connection_ids
        if manager.connection_users.get(connection_id) is not None
    ]
    by_company = {}
    for connection_id in connection_ids:
        company_id = manager.connection_company_scope.get(connection_id) or "sem_empresa"
        user_id = manager.connection_users.get(connection_id)
        stats = by_company.setdefault(company_id, {"connections": 0, "user_ids": set()})
        stats["connections"] += 1
        if user_id is not None:
            stats["user_ids"].add(user_id)

    client_key_counts = {}
    legacy_connections = 0
    for connection_id in connection_ids:
        client_key = manager.connection_client_keys.get(connection_id, "")
        if not client_key:
            legacy_connections += 1
            continue
        client_key_counts[client_key] = client_key_counts.get(client_key, 0) + 1

    duplicate_client_instances = sum(1 for count in client_key_counts.values() if count > 1)
    return {
        "connections": len(connection_ids),
        "online_users": len(set(user_ids)),
        "client_instances": len(client_key_counts),
        "legacy_connections": legacy_connections,
        "duplicate_client_instances": duplicate_client_instances,
        "companies": [
            {
                "company_id": company_id,
                "connections": stats["connections"],
                "online_users": len(stats["user_ids"]),
            }
            for company_id, stats in sorted(by_company.items(), key=lambda item: str(item[0]))
        ],
    }


@app.post("/admin/desktop-updates/force-check")
async def force_desktop_update_check(current_user: User = Depends(get_current_user)):
    """Dispara uma checagem imediata de update em TODOS os desktops online, em todas as empresas."""
    # Esta operacao e global e, por seguranca, continua restrita ao master/platform admin.
    ensure_admin(current_user)

    connection_ids = list(manager.active_connections.keys())
    online_user_ids = {
        manager.connection_users.get(connection_id)
        for connection_id in connection_ids
        if manager.connection_users.get(connection_id) is not None
    }

    companies = {}
    for connection_id in connection_ids:
        company_id = manager.connection_company_scope.get(connection_id) or "sem_empresa"
        user_id = manager.connection_users.get(connection_id)
        company_stats = companies.setdefault(company_id, {"connections": 0, "user_ids": set()})
        company_stats["connections"] += 1
        if user_id is not None:
            company_stats["user_ids"].add(user_id)

    payload = {
        "type": "force_update_check",
        "requested_by": current_user.id,
        "requested_at": datetime.now(timezone.utc).isoformat(),
        "source": "master_admin_global",
        "scope": "all_companies",
    }
    # company_id=None faz broadcast para todas as conexoes WebSocket ativas,
    # independentemente do tenant/empresa.
    await manager.broadcast(json.dumps(payload), company_id=None)

    company_summary = [
        {
            "company_id": company_id,
            "online_users": len(stats["user_ids"]),
            "connections": stats["connections"],
        }
        for company_id, stats in sorted(companies.items(), key=lambda item: str(item[0]))
    ]

    return {
        "message": "Verificacao de atualizacao disparada globalmente para todas as empresas com usuarios online.",
        "scope": "all_companies",
        "companies_online": len(company_summary),
        "online_users": len(online_user_ids),
        "connections": len(connection_ids),
        "companies": company_summary,
    }


def desktop_package_external_url() -> str:
    return DESKTOP_PACKAGE_EXTERNAL_URL or DESKTOP_RELEASE_EXTERNAL_URL


@app.head("/downloads/desktop/package")
async def head_desktop_package():
    package_url = desktop_package_external_url()
    if package_url.startswith("https://"):
        return RedirectResponse(package_url, status_code=307)
    raise HTTPException(status_code=404, detail="Pacote ZIP do VoltChat ainda nao publicado")


@app.get("/downloads/desktop/package")
async def download_desktop_package():
    package_url = desktop_package_external_url()
    if package_url.startswith("https://"):
        return RedirectResponse(package_url, status_code=307)
    raise HTTPException(status_code=404, detail="Pacote ZIP do VoltChat ainda nao publicado")


@app.head("/downloads/desktop/latest")
async def head_latest_desktop_release():
    manifest = load_r2_release_manifest()
    return RedirectResponse(presign_r2_release(str(manifest["storage_key"])), status_code=307)


@app.get("/downloads/desktop/latest")
async def download_latest_desktop_release():
    manifest = load_r2_release_manifest()
    return RedirectResponse(presign_r2_release(str(manifest["storage_key"])), status_code=307)

@app.post("/auth/login")
async def login(credentials: dict):
    username = credentials.get("username")
    password = credentials.get("password")
    remember_me = bool(credentials.get("remember_me", False))

    if not username or not password:
        raise HTTPException(status_code=400, detail="Username e password sao obrigatorios")

    with SessionLocal() as db:
        user = db.query(User).filter(or_(User.username == username, User.email == username)).first()
        if not user or not verify_password(password, user.hashed_password):
            raise HTTPException(status_code=401, detail="Credenciais invalidas")
        if not user.is_active or normalize_status(getattr(user, "status", ACTIVE_STATUS)) != ACTIVE_STATUS:
            raise HTTPException(status_code=401, detail="Usuario inativo")

        expires_delta = timedelta(days=REMEMBERED_ACCESS_TOKEN_EXPIRE_DAYS) if remember_me else timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
        access_token = create_access_token(
            data={"sub": str(user.id), "auth_version": int(getattr(user, "auth_version", 0) or 0)},
            expires_delta=expires_delta,
        )
        return {
            "access_token": access_token,
            "token_type": "bearer",
            "user": serialize_user(user, db),
        }


@app.post("/auth/forgot-password")
async def request_password_reset(payload: dict):
    email = str(payload.get("email") or "").strip().lower()
    generic_response = {
        "message": "Se existir uma conta vinculada a este e-mail, enviaremos um link para redefinir sua senha.",
    }
    if not EMAIL_PATTERN.match(email):
        return generic_response

    if not BREVO_API_KEY or not BREVO_SENDER_EMAIL:
        raise HTTPException(status_code=503, detail="Recuperacao de senha nao esta configurada. Contate o administrador.")

    user = None
    raw_token = None
    with SessionLocal() as db:
        user = (
            db.query(User)
            .filter(User.email == email, User.is_active.is_(True))
            .first()
        )
        if user and normalize_status(getattr(user, "status", ACTIVE_STATUS)) == ACTIVE_STATUS:
            now = datetime.now(timezone.utc)
            db.query(PasswordResetToken).filter(
                PasswordResetToken.user_id == user.id,
                PasswordResetToken.used_at.is_(None),
            ).update({PasswordResetToken.used_at: now}, synchronize_session=False)
            raw_token = secrets.token_urlsafe(32)
            db.add(
                PasswordResetToken(
                    user_id=user.id,
                    token_hash=password_reset_token_hash(raw_token),
                    expires_at=now + timedelta(minutes=PASSWORD_RESET_TOKEN_EXPIRE_MINUTES),
                )
            )
            db.commit()
            db.expunge(user)

    if user and raw_token:
        try:
            send_password_reset_email(user, raw_token)
        except RuntimeError as exc:
            print(f"[voltchat] {exc}")

    return generic_response


@app.post("/auth/reset-password")
async def reset_password_from_email(payload: dict):
    raw_token = str(payload.get("token") or "").strip()
    new_password = str(payload.get("new_password") or "").strip()
    if len(new_password) < 6:
        raise HTTPException(status_code=400, detail="A nova senha deve ter pelo menos 6 caracteres.")
    if len(raw_token) < 32:
        raise HTTPException(status_code=400, detail="Link de redefinicao invalido ou expirado.")

    now = datetime.now(timezone.utc)
    with SessionLocal() as db:
        reset_token = (
            db.query(PasswordResetToken)
            .filter(
                PasswordResetToken.token_hash == password_reset_token_hash(raw_token),
                PasswordResetToken.used_at.is_(None),
                PasswordResetToken.expires_at > now,
            )
            .first()
        )
        if not reset_token:
            raise HTTPException(status_code=400, detail="Link de redefinicao invalido ou expirado.")

        user = db.query(User).filter(User.id == reset_token.user_id).first()
        if not user or not user.is_active:
            raise HTTPException(status_code=400, detail="Link de redefinicao invalido ou expirado.")

        reset_token.used_at = now
        db.query(PasswordResetToken).filter(
            PasswordResetToken.user_id == user.id,
            PasswordResetToken.used_at.is_(None),
        ).update({PasswordResetToken.used_at: now}, synchronize_session=False)
        user.hashed_password = get_password_hash(new_password)
        user.auth_version = int(getattr(user, "auth_version", 0) or 0) + 1
        user.must_change_password = False
        user.updated_at = datetime.utcnow()
        log_audit(
            db, user.id, "senha_redefinida_por_email", "user", user.id,
            ensure_company_access(db, user) if not is_admin(user) else None,
            {"self_service": True},
        )
        db.commit()

    return {"message": "Senha redefinida com sucesso. Entre novamente no VoltChat."}


@app.post("/auth/logout")
async def logout(current_user: User = Depends(get_current_user)):
    return {"message": "Logout realizado com sucesso"}


@app.get("/auth/me")
async def get_current_user_info(current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        user = db.query(User).filter(User.id == current_user.id).first()
        if not user:
            raise HTTPException(status_code=404, detail="Usuario nao encontrado.")
        return serialize_user(user, db)


@app.put("/auth/profile")
async def update_own_profile(payload: dict, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        user = db.query(User).filter(User.id == current_user.id).first()
        if not user:
            raise HTTPException(status_code=404, detail="Usuario nao encontrado.")

        if "email" in payload:
            email = str(payload.get("email") or "").strip().lower()
            if not EMAIL_PATTERN.match(email):
                raise HTTPException(status_code=400, detail="Email invalido.")
            duplicate = db.query(User).filter(User.email == email, User.id != user.id).first()
            if duplicate:
                raise HTTPException(status_code=409, detail="Email ja utilizado por outro usuario.")
            user.email = email

        if "phone_extension" in payload:
            user.phone_extension = str(payload.get("phone_extension") or "").strip()[:40] or None
        if "nickname" in payload:
            nickname = " ".join(str(payload.get("nickname") or "").split())
            if len(nickname) > 80:
                raise HTTPException(status_code=400, detail="Apelido deve ter no maximo 80 caracteres.")
            user.nickname = nickname or None

        if "birthday" in payload:
            user.birthday = birthday_or_http_error(payload.get("birthday"))
        if "profile_photo" in payload:
            profile_photo = str(payload.get("profile_photo") or "").strip()
            if profile_photo:
                if not re.match(r"^data:image/(?:webp|jpeg|png);base64,", profile_photo, re.IGNORECASE):
                    raise HTTPException(status_code=400, detail="Foto de perfil deve ser uma imagem compactada.")
                if len(profile_photo.encode("utf-8")) > 220_000:
                    raise HTTPException(status_code=413, detail="Foto de perfil compactada excede 220 KB.")
            user.profile_photo = profile_photo or None

        user.updated_at = datetime.utcnow()
        company_id = ensure_company_access(db, user)
        log_audit(db, user.id, "perfil_atualizado", "user", user.id, company_id, {"self_service": True})
        db.commit()
        db.refresh(user)
        return {"message": "Perfil atualizado.", "user": serialize_user(user, db)}


@app.post("/auth/change-password")
async def change_password(payload: dict, current_user: User = Depends(get_current_user)):
    new_password = str(payload.get("new_password") or "").strip()
    current_password = str(payload.get("current_password") or "").strip()
    if len(new_password) < 6:
        raise HTTPException(status_code=400, detail="A nova senha deve ter pelo menos 6 caracteres.")

    with SessionLocal() as db:
        user = db.query(User).filter(User.id == current_user.id).first()
        if not user:
            raise HTTPException(status_code=404, detail="Usuario nao encontrado.")
        if current_password and not verify_password(current_password, user.hashed_password):
            raise HTTPException(status_code=401, detail="Senha atual invalida.")

        user.hashed_password = get_password_hash(new_password)
        user.auth_version = int(getattr(user, "auth_version", 0) or 0) + 1
        user.must_change_password = False
        user.updated_at = datetime.utcnow()
        log_audit(
            db,
            user.id,
            "senha_alterada",
            "user",
            user.id,
            ensure_company_access(db, user) if not is_admin(user) else None,
            {"self_service": True},
        )
        db.commit()
        db.refresh(user)
        return {"message": "Senha alterada com sucesso.", "user": serialize_user(user, db)}


@app.get("/users/")
async def list_users(company_id: Optional[str] = Query(default=None), current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        query = db.query(User)
        if is_admin(current_user) and not company_id:
            users = query.order_by(User.full_name.asc()).all()
            return [serialize_user(item, db) for item in users]

        scoped_company_id = ensure_company_access(db, current_user, company_id)
        query = query.join(CompanyUser, CompanyUser.user_id == User.id).filter(
            CompanyUser.company_id == scoped_company_id,
            CompanyUser.status == ACTIVE_STATUS,
        )
        if not is_company_admin(db, current_user, scoped_company_id):
            membership = get_company_membership(db, current_user, scoped_company_id)
            if membership and membership.role == "coordinator" and membership.department_id:
                query = query.filter(CompanyUser.department_id == membership.department_id)
            else:
                query = query.filter(User.id == current_user.id)
        users = query.order_by(User.full_name.asc()).all()
        return [serialize_user(item, db) for item in users]


@app.get("/birthdays/")
async def list_birthdays(current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user)
        users = (
            db.query(User)
            .join(CompanyUser, CompanyUser.user_id == User.id)
            .filter(User.is_active == True, User.birthday.isnot(None), User.birthday != "")
            .filter(CompanyUser.company_id == company_id, CompanyUser.status == ACTIVE_STATUS)
            .order_by(User.full_name.asc())
            .all()
        )
        return [serialize_user(item, db) for item in users]


@app.get("/chat/contacts")
async def list_chat_contacts(
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, company_id)
        users = (
            db.query(User)
            .join(CompanyUser, CompanyUser.user_id == User.id)
            .filter(
                CompanyUser.company_id == company_id,
                CompanyUser.status == ACTIVE_STATUS,
                User.is_active == True,
                User.id != current_user.id,
            )
            .order_by(User.full_name.asc())
            .all()
        )
        return [serialize_user(item, db) for item in users]


@app.post("/users/")
async def create_user(payload: dict, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        membership = get_company_membership(db, current_user, company_id)
        can_create = is_admin(current_user) or (membership and membership.role in {"company_admin", "master_admin", "coordinator"})
    if not can_create:
        raise HTTPException(status_code=403, detail="Sem permissao para criar usuarios.")

    required_fields = ["email", "password"]
    missing = [field for field in required_fields if not str(payload.get(field) or "").strip()]
    if missing:
        raise HTTPException(status_code=400, detail=f"Campos obrigatorios: {', '.join(missing)}")

    email = str(payload["email"]).strip().lower()
    if not EMAIL_PATTERN.match(email):
        raise HTTPException(status_code=400, detail="Email invalido.")
    full_name = str(payload.get("full_name") or payload.get("name") or email).strip()

    role = normalize_platform_role(payload.get("role") or payload.get("access_level") or "user")
    if role == "master_admin" and not is_admin(current_user):
        raise HTTPException(status_code=403, detail="Somente Admin Master cria administradores globais.")

    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        ensure_company_admin(db, current_user, company_id) if role in {"company_admin", "coordinator"} else None
        department_id = payload.get("department_id")
        if department_id:
            department = db.query(Department).filter(Department.id == str(department_id), Department.company_id == company_id).first()
            if not department:
                raise HTTPException(status_code=400, detail="Setor nao encontrado.")
        else:
            department = get_or_create_department(db, company_id, str(payload.get("department") or current_user.department or "Operacao"))

        if not is_company_admin(db, current_user, company_id):
            current_membership = get_company_membership(db, current_user, company_id)
            if not current_membership or current_membership.role != "coordinator" or current_membership.department_id != department.id:
                raise HTTPException(status_code=403, detail="Coordenador so pode criar usuarios do proprio setor.")
            role = "user"

        user = db.query(User).filter(User.email == email).first()
        if not user:
            username = generate_unique_username(db, email, full_name, payload.get("username"))
            user = User(
                uuid=str(uuid.uuid4()),
                username=username,
                email=email,
                full_name=full_name,
                hashed_password=get_password_hash(str(payload["password"])),
                phone=str(payload.get("phone") or payload.get("phone_extension") or "").strip() or None,
                must_change_password=True,
                status=normalize_status(payload.get("status")),
                is_active=normalize_status(payload.get("status")) == ACTIVE_STATUS,
                is_platform_admin=role == "master_admin",
                access_level="master" if role == "master_admin" else legacy_access_for_role(role),
                department=department.name,
                phone_extension=str(payload.get("phone_extension") or payload.get("phone") or "").strip() or None,
                birthday=birthday_or_http_error(payload.get("birthday")),
                role_title=str(payload.get("role_title") or "").strip() or None,
            )
            db.add(user)
            db.flush()
            log_audit(db, current_user.id, "usuario_criado", "user", user.id, company_id, {"email": user.email, "role": role})

        existing_link = (
            db.query(CompanyUser)
            .filter(CompanyUser.company_id == company_id, CompanyUser.user_id == user.id)
            .first()
        )
        if existing_link:
            existing_link.department_id = department.id
            existing_link.role = role if role != "master_admin" else "company_admin"
            existing_link.status = normalize_status(payload.get("status"))
        else:
            db.add(
                CompanyUser(
                    id=str(uuid.uuid4()),
                    company_id=company_id,
                    user_id=user.id,
                    department_id=department.id,
                    role=role if role != "master_admin" else "company_admin",
                    status=normalize_status(payload.get("status")),
                )
            )
            log_audit(db, current_user.id, "vinculo_criado", "company_user", user.id, company_id, {"role": role})
        db.commit()
        db.refresh(user)
        return serialize_user(user, db)


@app.put("/users/{user_id}")
async def update_user(user_id: int, payload: dict, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        target = db.query(User).filter(User.id == user_id).first()
        if not target:
            raise HTTPException(status_code=404, detail="Usuario nao encontrado.")

        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        target_membership = (
            db.query(CompanyUser)
            .filter(CompanyUser.company_id == company_id, CompanyUser.user_id == target.id)
            .first()
        )
        if not target_membership:
            raise HTTPException(status_code=404, detail="Usuario nao pertence a esta empresa.")

        actor_is_platform_admin = is_admin(current_user)
        actor_is_company_admin = is_company_admin(db, current_user, company_id)
        if target.is_platform_admin and not actor_is_platform_admin:
            raise HTTPException(status_code=403, detail="Administrador da empresa nao pode editar Admin Master.")

        if not actor_is_company_admin:
            current_membership = get_company_membership(db, current_user, company_id)
            if (
                not current_membership
                or current_membership.role != "coordinator"
                or target_membership.department_id != current_membership.department_id
            ):
                raise HTTPException(status_code=403, detail="Sem permissao para editar este usuario.")

        if "full_name" in payload:
            full_name = str(payload.get("full_name") or "").strip()
            if not full_name:
                raise HTTPException(status_code=400, detail="Nome completo e obrigatorio.")
            target.full_name = full_name

        if "email" in payload:
            email = str(payload.get("email") or "").strip().lower()
            if not EMAIL_PATTERN.match(email):
                raise HTTPException(status_code=400, detail="Email invalido.")
            duplicate_email = db.query(User.id).filter(User.email == email, User.id != target.id).first()
            if duplicate_email:
                raise HTTPException(status_code=409, detail="Email ja cadastrado para outro usuario.")
            target.email = email

        if "username" in payload:
            username = normalize_username_piece(payload.get("username"))
            duplicate_username = db.query(User.id).filter(
                User.username == username,
                User.id != target.id,
            ).first()
            if duplicate_username:
                raise HTTPException(status_code=409, detail="Nome de usuario ja cadastrado.")
            target.username = username

        for field in ["nickname", "phone", "phone_extension", "role_title"]:
            if field in payload:
                setattr(target, field, str(payload.get(field) or "").strip() or None)
        if "birthday" in payload:
            target.birthday = birthday_or_http_error(payload.get("birthday"))

        if "password" in payload and str(payload.get("password") or "").strip():
            password = str(payload.get("password") or "").strip()
            if len(password) < 6:
                raise HTTPException(status_code=400, detail="Nova senha deve ter pelo menos 6 caracteres.")
            target.hashed_password = get_password_hash(password)
            target.must_change_password = True

        if "department_id" in payload:
            department = db.query(Department).filter(
                Department.id == str(payload.get("department_id")),
                Department.company_id == company_id,
                Department.status == ACTIVE_STATUS,
            ).first()
            if not department:
                raise HTTPException(status_code=400, detail="Setor ativo nao encontrado.")
            if not actor_is_company_admin:
                current_membership = get_company_membership(db, current_user, company_id)
                if not current_membership or current_membership.department_id != department.id:
                    raise HTTPException(status_code=403, detail="Coordenador nao pode mover usuario para outro setor.")
            target_membership.department_id = department.id
            target.department = department.name

        if "role" in payload:
            if target.id == current_user.id:
                raise HTTPException(status_code=400, detail="Nao e permitido alterar o proprio nivel de acesso.")
            role = normalize_platform_role(payload.get("role"))
            if role == "master_admin" and not actor_is_platform_admin:
                raise HTTPException(status_code=403, detail="Somente Admin Master altera administrador global.")
            if not actor_is_company_admin and role != "user":
                raise HTTPException(status_code=403, detail="Coordenador nao pode promover usuarios.")
            target_membership.role = "company_admin" if role == "master_admin" else role
            if actor_is_platform_admin:
                target.is_platform_admin = role == "master_admin"
            target.access_level = "master" if role == "master_admin" else legacy_access_for_role(role)
            log_audit(
                db,
                current_user.id,
                "nivel_alterado",
                "company_user",
                target_membership.id,
                company_id,
                {"role": target_membership.role},
            )

        requested_status = None
        if "status" in payload:
            requested_status = normalize_status(payload.get("status"))
        elif "is_active" in payload:
            requested_status = ACTIVE_STATUS if bool(payload.get("is_active")) else INACTIVE_STATUS
        if requested_status:
            if requested_status not in {ACTIVE_STATUS, INACTIVE_STATUS}:
                raise HTTPException(status_code=400, detail="Status de usuario invalido.")
            if target.id == current_user.id and requested_status == INACTIVE_STATUS:
                raise HTTPException(status_code=400, detail="Nao e permitido desativar a propria conta.")
            target_membership.status = requested_status
            if requested_status == ACTIVE_STATUS:
                target.status = ACTIVE_STATUS
                target.is_active = True
            else:
                other_active_membership = db.query(CompanyUser.id).filter(
                    CompanyUser.user_id == target.id,
                    CompanyUser.company_id != company_id,
                    CompanyUser.status == ACTIVE_STATUS,
                ).first()
                target.is_active = bool(other_active_membership or target.is_platform_admin)
                target.status = ACTIVE_STATUS if target.is_active else INACTIVE_STATUS

        target.updated_at = datetime.utcnow()
        log_audit(
            db,
            current_user.id,
            "usuario_editado",
            "user",
            target.id,
            company_id,
            {
                "email": target.email,
                "username": target.username,
                "membership_status": target_membership.status,
                "role": target_membership.role,
            },
        )
        db.commit()
        db.refresh(target)
        return serialize_user(target, db)


@app.delete("/users/{user_id}")
async def delete_user(
    user_id: int,
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        ensure_company_admin(db, current_user, scoped_company_id)
        target = db.query(User).filter(User.id == user_id).first()
        if not target:
            raise HTTPException(status_code=404, detail="Usuario nao encontrado.")
        if target.id == current_user.id:
            raise HTTPException(status_code=400, detail="Nao e permitido excluir a propria conta.")
        if target.username in {BOLT_USERNAME, DEFAULT_ADMIN_USERNAME}:
            raise HTTPException(status_code=400, detail="Usuario de sistema protegido e nao pode ser excluido.")
        if target.is_platform_admin and not is_admin(current_user):
            raise HTTPException(status_code=403, detail="Administrador da empresa nao pode excluir Admin Master.")

        membership = db.query(CompanyUser).filter(
            CompanyUser.company_id == scoped_company_id,
            CompanyUser.user_id == target.id,
        ).first()
        if not membership:
            raise HTTPException(status_code=404, detail="Usuario nao pertence a esta empresa.")
        if membership.status == INACTIVE_STATUS:
            raise HTTPException(status_code=409, detail="Usuario ja foi excluido desta empresa.")

        company_ticket_ids = db.query(Ticket.id).filter(Ticket.company_id == scoped_company_id)
        db.query(TicketAssignee).filter(
            TicketAssignee.user_id == target.id,
            TicketAssignee.ticket_id.in_(company_ticket_ids),
        ).delete(synchronize_session=False)
        db.query(Ticket).filter(
            Ticket.company_id == scoped_company_id,
            Ticket.assigned_to_id == target.id,
        ).update({Ticket.assigned_to_id: None}, synchronize_session=False)
        db.query(TaskItem).filter(
            TaskItem.company_id == scoped_company_id,
            TaskItem.assigned_to_id == target.id,
        ).update({TaskItem.assigned_to_id: None}, synchronize_session=False)
        company_group_ids = db.query(ChatGroup.id).filter(ChatGroup.company_id == scoped_company_id)
        db.query(ChatGroupMember).filter(
            ChatGroupMember.user_id == target.id,
            ChatGroupMember.group_id.in_(company_group_ids),
        ).delete(synchronize_session=False)

        membership.status = INACTIVE_STATUS
        other_active_membership = db.query(CompanyUser.id).filter(
            CompanyUser.user_id == target.id,
            CompanyUser.company_id != scoped_company_id,
            CompanyUser.status == ACTIVE_STATUS,
        ).first()
        if not other_active_membership and not target.is_platform_admin:
            target.is_active = False
            target.status = INACTIVE_STATUS
        target.updated_at = datetime.utcnow()

        log_audit(
            db,
            current_user.id,
            "usuario_excluido",
            "user",
            target.id,
            scoped_company_id,
            {
                "email": target.email,
                "username": target.username,
                "membership_id": membership.id,
                "account_deactivated": not target.is_active,
                "history_preserved": True,
            },
        )
        db.commit()
        return {
            "message": "Usuario excluido da empresa. Historico preservado.",
            "user_id": target.id,
            "company_id": scoped_company_id,
            "account_deactivated": not target.is_active,
        }

@app.get("/departments/")
async def list_departments(current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user)
        query = db.query(Department).filter(Department.company_id == company_id, Department.status == ACTIVE_STATUS)
        if not is_company_admin(db, current_user, company_id):
            membership = get_company_membership(db, current_user, company_id)
            if membership and membership.department_id:
                query = query.filter(Department.id == membership.department_id)
        return [department.name for department in query.order_by(Department.name.asc()).all()]


@app.get("/tickets/")
async def list_tickets(
    status: Optional[str] = Query(default=None),
    department: Optional[str] = Query(default=None),
    company_id: Optional[str] = Query(default=None),
    management: bool = Query(default=False, description="Visao gerencial para administradores e coordenadores"),
    limit: int = Query(default=100, ge=1, le=200),
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        query = db.query(Ticket).filter(Ticket.company_id == scoped_company_id)
        if department:
            query = query.filter(Ticket.department == department)

        if status:
            query = query.filter(Ticket.status == status)

        tickets = query.order_by(Ticket.created_at.desc()).all()
        if not is_company_admin(db, current_user, scoped_company_id):
            membership = get_company_membership(db, current_user, scoped_company_id)
            if management and membership and membership.role == "coordinator":
                tickets = [ticket for ticket in tickets if can_manage_ticket(ticket, current_user, db)]
            else:
                tickets = [ticket for ticket in tickets if can_access_ticket(ticket, current_user, db)]
        tickets = tickets[:limit]
        return [serialize_ticket(ticket, db) for ticket in tickets]


def minutes_between(start_value: Optional[datetime], end_value: Optional[datetime]) -> Optional[float]:
    if not start_value or not end_value:
        return None
    seconds = (ensure_utc_datetime(end_value) - ensure_utc_datetime(start_value)).total_seconds()
    return max(0, round(seconds / 60, 1))


def build_ticket_report_section(contexts: list[dict], user_lookup: dict[int, User], department_name: Optional[str] = None) -> dict:
    scoped_contexts = contexts if not department_name else [
        context for context in contexts if department_name in context["departments"]
    ]
    user_metrics = {}

    def ensure_metric(user_id: Optional[int]):
        if not user_id:
            return None
        user = user_lookup.get(user_id)
        return user_metrics.setdefault(user_id, {
            "user_id": user_id,
            "user_name": user.full_name if user else "Usuario removido",
            "opened_tickets": 0,
            "assigned_tickets": 0,
            "closed_tickets": 0,
            "messages_sent": 0,
            "response_times": [],
            "tickets_responded": set(),
        })

    first_response_times = []
    resolution_times = []
    details = []
    for context in scoped_contexts:
        ticket = context["ticket"]
        first_response = minutes_between(ticket.created_at, ticket.first_response_at)
        resolution = minutes_between(ticket.created_at, ticket.closed_at)
        if first_response is not None:
            first_response_times.append(first_response)
        if resolution is not None:
            resolution_times.append(resolution)

        opener_metric = ensure_metric(ticket.created_by_id)
        if opener_metric:
            opener_metric["opened_tickets"] += 1
        for assigned_user_id in context["assigned_user_ids"]:
            assigned_metric = ensure_metric(assigned_user_id)
            if assigned_metric:
                assigned_metric["assigned_tickets"] += 1
        closer_metric = ensure_metric(context["closed_by_id"])
        if closer_metric and ticket.status == "Resolvido":
            closer_metric["closed_tickets"] += 1

        previous_sender_id = ticket.created_by_id
        previous_activity_at = ticket.created_at
        for message in context["messages"]:
            sender_metric = ensure_metric(message.sender_id)
            if sender_metric:
                sender_metric["messages_sent"] += 1
                sender_metric["tickets_responded"].add(ticket.id)
            if previous_sender_id != message.sender_id:
                response_minutes = minutes_between(previous_activity_at, message.created_at)
                if sender_metric and response_minutes is not None:
                    sender_metric["response_times"].append(response_minutes)
            previous_sender_id = message.sender_id
            previous_activity_at = message.created_at

        details.append({
            "id": ticket.id,
            "subject": ticket.title,
            "description": ticket.description,
            "priority": ticket.priority,
            "status": ticket.status,
            "departments": context["departments"],
            "opened_by_id": ticket.created_by_id,
            "opened_by_name": context["opened_by_name"],
            "opened_at": ticket.created_at.isoformat() if ticket.created_at else None,
            "assigned_users": context["assigned_user_names"],
            "first_response_minutes": first_response,
            "closed_by_id": context["closed_by_id"],
            "closed_by_name": context["closed_by_name"],
            "closed_at": ticket.closed_at.isoformat() if ticket.closed_at else None,
            "close_reason": context["close_reason"],
            "total_progress_minutes": resolution,
            "rating_score": ticket.rating_score,
            "rating_comment": ticket.rating_comment,
        })

    users = []
    for metric in user_metrics.values():
        response_times = metric.pop("response_times")
        tickets_responded = metric.pop("tickets_responded")
        users.append({
            **metric,
            "tickets_responded": len(tickets_responded),
            "responses_measured": len(response_times),
            "average_response_minutes": round(sum(response_times) / len(response_times), 1) if response_times else None,
        })
    users.sort(key=lambda item: (-item["closed_tickets"], item["user_name"].lower()))

    open_count = sum(1 for context in scoped_contexts if context["ticket"].status == "Aberto")
    closed_count = sum(1 for context in scoped_contexts if context["ticket"].status == "Resolvido")
    in_progress_count = sum(1 for context in scoped_contexts if context["ticket"].status == "Em andamento")
    return {
        "department": department_name,
        "summary": {
            "total_tickets": len(scoped_contexts),
            "open_tickets": open_count,
            "in_progress_tickets": in_progress_count,
            "closed_tickets": closed_count,
            "average_first_response_minutes": round(sum(first_response_times) / len(first_response_times), 1) if first_response_times else None,
            "average_total_progress_minutes": round(sum(resolution_times) / len(resolution_times), 1) if resolution_times else None,
        },
        "users": users,
        "tickets": details,
    }


@app.get("/tickets/reports/control")
async def ticket_control_report(
    date_from: Optional[str] = Query(default=None),
    date_to: Optional[str] = Query(default=None),
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        membership = get_company_membership(db, current_user, scoped_company_id)
        company_admin = is_company_admin(db, current_user, scoped_company_id)
        coordinator_department = user_department_name(db, current_user, scoped_company_id)
        coordinator_allowed = bool(membership and membership.role == "coordinator" and coordinator_department)
        if not company_admin and not coordinator_allowed:
            raise HTTPException(status_code=403, detail="Relatorio disponivel apenas para coordenadores e administradores da empresa.")

        query = db.query(Ticket).filter(Ticket.company_id == scoped_company_id)
        try:
            if date_from:
                start_at = datetime.strptime(date_from, "%Y-%m-%d").replace(tzinfo=timezone.utc)
                query = query.filter(Ticket.created_at >= start_at)
            if date_to:
                end_at = datetime.strptime(date_to, "%Y-%m-%d").replace(tzinfo=timezone.utc) + timedelta(days=1)
                query = query.filter(Ticket.created_at < end_at)
        except ValueError:
            raise HTTPException(status_code=400, detail="Periodo invalido. Use datas no formato AAAA-MM-DD.")

        tickets = query.order_by(Ticket.created_at.desc()).all()
        if not company_admin:
            # No relatorio, coordenador exerce visao gerencial do proprio setor sem
            # ser adicionado como participante dos tickets.
            tickets = [ticket for ticket in tickets if can_manage_ticket(ticket, current_user, db)]
        ticket_ids = [ticket.id for ticket in tickets]
        assignments_by_ticket = {ticket_id: set() for ticket_id in ticket_ids}
        departments_by_ticket = {ticket_id: set() for ticket_id in ticket_ids}
        messages_by_ticket = {ticket_id: [] for ticket_id in ticket_ids}
        user_ids = {
            user_id for ticket in tickets
            for user_id in [ticket.created_by_id, ticket.assigned_to_id, ticket.closed_by_id]
            if user_id
        }

        if ticket_ids:
            for assignment in db.query(TicketAssignee).filter(TicketAssignee.ticket_id.in_(ticket_ids)).all():
                assignments_by_ticket[assignment.ticket_id].add(assignment.user_id)
                user_ids.add(assignment.user_id)
            department_rows = (
                db.query(TicketDepartment, Department)
                .join(Department, Department.id == TicketDepartment.department_id)
                .filter(TicketDepartment.ticket_id.in_(ticket_ids))
                .all()
            )
            for ticket_department, department in department_rows:
                departments_by_ticket[ticket_department.ticket_id].add(department.name)
            message_rows = (
                db.query(TicketMessage)
                .filter(TicketMessage.ticket_id.in_(ticket_ids))
                .order_by(TicketMessage.ticket_id.asc(), TicketMessage.created_at.asc(), TicketMessage.id.asc())
                .all()
            )
            for message in message_rows:
                messages_by_ticket[message.ticket_id].append(message)
                user_ids.add(message.sender_id)

        user_lookup = {
            user.id: user for user in db.query(User).filter(User.id.in_(user_ids)).all()
        } if user_ids else {}
        contexts = []
        observed_departments = set()
        for ticket in tickets:
            assigned_user_ids = set(assignments_by_ticket[ticket.id])
            if ticket.assigned_to_id:
                assigned_user_ids.add(ticket.assigned_to_id)
            assigned_user_names = {
                user_lookup[user_id].full_name for user_id in assigned_user_ids if user_id in user_lookup
            }
            departments = set(departments_by_ticket[ticket.id])
            if ticket.department:
                departments.add(ticket.department)
            observed_departments.update(departments)
            messages = messages_by_ticket[ticket.id]
            closed_by_id = ticket.closed_by_id
            close_reason = ticket.close_reason
            if ticket.status == "Resolvido" and (not closed_by_id or not close_reason) and messages:
                closing_message = messages[-1]
                close_gap = minutes_between(closing_message.created_at, ticket.closed_at)
                if close_gap is not None and close_gap <= 10:
                    closed_by_id = closed_by_id or closing_message.sender_id
                    close_reason = close_reason or closing_message.content
            closed_by = user_lookup.get(closed_by_id)
            opened_by = user_lookup.get(ticket.created_by_id)
            contexts.append({
                "ticket": ticket,
                "departments": sorted(departments) or ["Sem setor"],
                "opened_by_name": opened_by.full_name if opened_by else "Usuario removido",
                "assigned_user_ids": sorted(assigned_user_ids),
                "assigned_user_names": sorted(assigned_user_names),
                "closed_by_id": closed_by_id,
                "closed_by_name": closed_by.full_name if closed_by else None,
                "close_reason": close_reason or ("Motivo nao informado" if ticket.status == "Resolvido" else None),
                "messages": messages,
            })

        if company_admin:
            active_departments = {
                row[0] for row in db.query(Department.name).filter(
                    Department.company_id == scoped_company_id,
                    Department.status == ACTIVE_STATUS,
                ).all()
            }
            department_names = sorted(active_departments | observed_departments)
            sections = [build_ticket_report_section(contexts, user_lookup, None)] + [
                build_ticket_report_section(contexts, user_lookup, department_name)
                for department_name in department_names
            ]
            scope = {"type": "company", "company_id": scoped_company_id, "department": None}
        else:
            sections = [build_ticket_report_section(contexts, user_lookup, coordinator_department)]
            scope = {"type": "department", "company_id": scoped_company_id, "department": coordinator_department}

        return {
            "scope": scope,
            "period": {"date_from": date_from, "date_to": date_to},
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "sections": sections,
        }

@app.get("/tickets/assignment-options")
async def ticket_assignment_options(company_id: Optional[str] = Query(default=None), current_user: User = Depends(get_current_user)):
    """Lista o time da empresa para que o dono monte um ticket colaborativo."""
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        users = (
            db.query(User)
            .join(CompanyUser, CompanyUser.user_id == User.id)
            .filter(
                CompanyUser.company_id == scoped_company_id,
                CompanyUser.status == ACTIVE_STATUS,
                User.is_active == True,
                ~User.username.in_([BOLT_USERNAME, MITTY_USERNAME]),
            )
            .order_by(User.full_name.asc())
            .all()
        )
        departments = (
            db.query(Department)
            .filter(
                Department.company_id == scoped_company_id,
                Department.status == ACTIVE_STATUS,
                Department.name != "Automacao",
            )
            .order_by(Department.name.asc())
            .all()
        )
        return {
            "users": [{"id": user.id, "full_name": user.full_name, "username": user.username} for user in users],
            "departments": [department.name for department in departments],
        }


@app.post("/tickets/")
async def create_ticket(payload: dict, current_user: User = Depends(get_current_user)):
    title = str(payload.get("title") or "").strip()
    description = str(payload.get("description") or "").strip()
    if not title or not description:
        raise HTTPException(status_code=400, detail="Titulo e descricao sao obrigatorios.")

    assigned_to_id = payload.get("assigned_to_id")
    assigned_to_ids = payload.get("assigned_to_ids")
    assigned_departments = payload.get("assigned_departments")
    attachment_file_id = payload.get("attachment_file_id")
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        department = str(payload.get("department") or user_department_name(db, current_user, company_id) or "Operacao").strip()
        get_or_create_department(db, company_id, department)
        ensure_creation_scope(db, current_user, company_id, department, "ticket")
        if assigned_to_ids is None:
            assigned_to_ids = [assigned_to_id] if assigned_to_id else []
        if not isinstance(assigned_to_ids, list):
            raise HTTPException(status_code=400, detail="Responsaveis do ticket devem ser uma lista.")
        if assigned_departments is None:
            assigned_departments = [department]
        if not isinstance(assigned_departments, list):
            raise HTTPException(status_code=400, detail="Setores do ticket devem ser uma lista.")
        if attachment_file_id:
            attachment = db.query(FileUpload).filter(FileUpload.id == int(attachment_file_id), FileUpload.company_id == company_id).first()
            if not attachment:
                raise HTTPException(status_code=400, detail="Anexo nao encontrado.")

        ticket = Ticket(
            company_id=company_id,
            title=title,
            description=description,
            priority=normalize_ticket_urgency(payload.get("priority")) or "Moderado",
            delivery_due_at=parse_ticket_delivery_due_at(payload.get("delivery_due_at")),
            status=str(payload.get("status") or "Aberto"),
            department=department,
            image_data=str(payload.get("image_data") or "").strip() or None,
            channel=str(payload.get("channel") or "Web"),
            created_by_id=current_user.id,
            assigned_to_id=None,
            attachment_file_id=int(attachment_file_id) if attachment_file_id else None,
        )
        db.add(ticket)
        db.flush()
        sync_ticket_group_assignments(db, ticket, company_id, assigned_to_ids, assigned_departments)
        db.commit()
        db.refresh(ticket)
        ticket_response = serialize_ticket(ticket, db)

        await notify_ticket_participants(
            db,
            ticket,
            current_user,
            "🎫 Novo ticket criado",
            f"{current_user.full_name} abriu: {ticket.title}",
        )

        return ticket_response


@app.get("/tickets/{ticket_id}")
async def get_ticket(ticket_id: int, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
        if not ticket:
            raise HTTPException(status_code=404, detail="Ticket nao encontrado.")
        ensure_ticket_access(ticket, current_user, db)
        return serialize_ticket(ticket, db)


@app.get("/tickets/{ticket_id}/messages")
async def list_ticket_messages(ticket_id: int, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
        if not ticket:
            raise HTTPException(status_code=404, detail="Ticket nao encontrado.")
        ensure_ticket_access(ticket, current_user, db)

        messages = (
            db.query(TicketMessage)
            .filter(TicketMessage.ticket_id == ticket_id)
            .order_by(TicketMessage.created_at.asc())
            .all()
        )
        return [serialize_ticket_message(message, db) for message in messages]


@app.post("/tickets/{ticket_id}/messages")
async def create_ticket_message(ticket_id: int, payload: dict, current_user: User = Depends(get_current_user)):
    content = str(payload.get("content") or "").strip()
    file_id = payload.get("file_id")
    reply_to_id = payload.get("reply_to_id")
    if not content and not file_id:
        raise HTTPException(status_code=400, detail="Informe uma mensagem ou anexo.")

    with SessionLocal() as db:
        ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
        if not ticket:
            raise HTTPException(status_code=404, detail="Ticket nao encontrado.")
        ensure_ticket_access(ticket, current_user, db)
        if ticket.status == "Resolvido":
            raise HTTPException(status_code=400, detail="Ticket resolvido nao recebe novas mensagens.")

        if file_id:
            attachment = db.query(FileUpload).filter(FileUpload.id == int(file_id), FileUpload.company_id == ticket.company_id).first()
            if not attachment:
                raise HTTPException(status_code=400, detail="Anexo nao encontrado.")

        reply_to = None
        if reply_to_id:
            try:
                normalized_reply_to_id = int(reply_to_id)
            except (TypeError, ValueError):
                raise HTTPException(status_code=400, detail="Mensagem respondida invalida.")
            reply_to = db.query(TicketMessage).filter(
                TicketMessage.id == normalized_reply_to_id,
                TicketMessage.ticket_id == ticket.id,
            ).first()
            if not reply_to:
                raise HTTPException(status_code=400, detail="Mensagem respondida nao pertence a este ticket.")

        message = TicketMessage(
            company_id=ticket.company_id,
            ticket_id=ticket.id,
            sender_id=current_user.id,
            content=content,
            file_id=int(file_id) if file_id else None,
            reply_to_id=reply_to.id if reply_to else None,
        )
        db.add(message)

        if not ticket.first_response_at and current_user.id != ticket.created_by_id:
            ticket.first_response_at = datetime.utcnow()
        if ticket.status == "Aberto" and current_user.id != ticket.created_by_id:
            ticket.status = "Em andamento"
        ticket.updated_at = datetime.utcnow()

        db.commit()
        db.refresh(message)
        db.refresh(ticket)
        message_response = serialize_ticket_message(message, db)

        await notify_ticket_participants(
            db,
            ticket,
            current_user,
            "💬 Nova mensagem no ticket",
            f"{current_user.full_name} comentou em: {ticket.title}",
        )

        return message_response


@app.post("/tickets/{ticket_id}/messages/{message_id}/reactions")
async def toggle_ticket_message_reaction(
    ticket_id: int,
    message_id: int,
    payload: dict,
    current_user: User = Depends(get_current_user),
):
    emoji = str(payload.get("emoji") or "").strip()
    if not emoji or len(emoji) > 32:
        raise HTTPException(status_code=400, detail="Informe um emoji valido.")

    with SessionLocal() as db:
        ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
        if not ticket:
            raise HTTPException(status_code=404, detail="Ticket nao encontrado.")
        ensure_ticket_access(ticket, current_user, db)
        message = db.query(TicketMessage).filter(
            TicketMessage.id == message_id,
            TicketMessage.ticket_id == ticket.id,
        ).first()
        if not message:
            raise HTTPException(status_code=404, detail="Mensagem do ticket nao encontrada.")

        existing = db.query(TicketMessageReaction).filter(
            TicketMessageReaction.ticket_message_id == message.id,
            TicketMessageReaction.user_id == current_user.id,
            TicketMessageReaction.emoji == emoji,
        ).first()
        if existing:
            db.delete(existing)
            action = "removed"
        else:
            db.add(TicketMessageReaction(
                ticket_message_id=message.id,
                user_id=current_user.id,
                emoji=emoji,
            ))
            action = "added"
        db.commit()
        reactions = serialize_ticket_message_reactions(message.id, db)

    return {"action": action, "reactions": reactions}

@app.patch("/tickets/{ticket_id}/transfer")
async def transfer_ticket(ticket_id: int, payload: dict, current_user: User = Depends(get_current_user)):
    assigned_to_id = payload.get("assigned_to_id")
    if not assigned_to_id:
        raise HTTPException(status_code=400, detail="Selecione um novo responsavel.")

    note = str(payload.get("message") or "").strip()
    with SessionLocal() as db:
        ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
        if not ticket:
            raise HTTPException(status_code=404, detail="Ticket nao encontrado.")
        ensure_ticket_access(ticket, current_user, db)

        assigned_user = ensure_user_in_company(db, int(assigned_to_id), ticket.company_id)
        if not is_company_admin(db, current_user, ticket.company_id) and assigned_user.department != ticket.department:
            raise HTTPException(status_code=403, detail="Responsavel precisa estar no mesmo setor.")

        ticket.assigned_to_id = assigned_user.id
        sync_ticket_group_assignments(db, ticket, ticket.company_id, [assigned_user.id], None)
        ticket.status = "Em andamento" if ticket.status != "Resolvido" else ticket.status
        ticket.updated_at = datetime.utcnow()
        db.add(TicketMessage(
            company_id=ticket.company_id,
            ticket_id=ticket.id,
            sender_id=current_user.id,
            content=note or f"Ticket repassado para {assigned_user.full_name}.",
        ))
        db.commit()
        db.refresh(ticket)
        ticket_response = serialize_ticket(ticket, db)

        await notify_ticket_participants(
            db,
            ticket,
            current_user,
            "🔁 Ticket repassado",
            f"{current_user.full_name} repassou: {ticket.title}",
        )

        return ticket_response


@app.post("/tickets/{ticket_id}/close")
async def close_ticket(ticket_id: int, payload: dict, current_user: User = Depends(get_current_user)):
    note = str(payload.get("message") or "").strip()
    if not note:
        raise HTTPException(status_code=400, detail="Informe o motivo do fechamento do ticket.")
    with SessionLocal() as db:
        ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
        if not ticket:
            raise HTTPException(status_code=404, detail="Ticket nao encontrado.")
        ensure_ticket_access(ticket, current_user, db)

        ticket.status = "Resolvido"
        ticket.closed_at = ticket.closed_at or datetime.utcnow()
        ticket.closed_by_id = current_user.id
        ticket.close_reason = note or "Ticket fechado."
        ticket.updated_at = datetime.utcnow()
        db.add(TicketMessage(
            company_id=ticket.company_id,
            ticket_id=ticket.id,
            sender_id=current_user.id,
            content=note or "Ticket fechado.",
        ))
        db.commit()
        db.refresh(ticket)
        ticket_response = serialize_ticket(ticket, db)

        await notify_ticket_participants(
            db,
            ticket,
            current_user,
            "✅ Ticket fechado",
            f"{current_user.full_name} fechou: {ticket.title}. Avalie o atendimento.",
        )

        return ticket_response


@app.post("/tickets/{ticket_id}/rating")
async def rate_ticket(ticket_id: int, payload: dict, current_user: User = Depends(get_current_user)):
    try:
        rating_score = int(payload.get("rating_score"))
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="Informe uma avaliacao de 1 a 5.")
    if rating_score < 1 or rating_score > 5:
        raise HTTPException(status_code=400, detail="A avaliacao deve ficar entre 1 e 5.")

    with SessionLocal() as db:
        ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
        if not ticket:
            raise HTTPException(status_code=404, detail="Ticket nao encontrado.")
        ensure_ticket_access(ticket, current_user, db)
        if ticket.created_by_id != current_user.id:
            raise HTTPException(status_code=403, detail="Somente o dono do ticket pode avaliar.")
        if ticket.status != "Resolvido":
            raise HTTPException(status_code=400, detail="Avalie apenas depois do fechamento.")

        ticket.rating_score = rating_score
        ticket.rating_comment = str(payload.get("rating_comment") or "").strip() or None
        ticket.rated_at = datetime.utcnow()
        ticket.updated_at = datetime.utcnow()
        db.commit()
        db.refresh(ticket)
        return serialize_ticket(ticket, db)


@app.patch("/tickets/{ticket_id}")
async def update_ticket(ticket_id: int, payload: dict, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        ticket = db.query(Ticket).filter(Ticket.id == ticket_id).first()
        if not ticket:
            raise HTTPException(status_code=404, detail="Ticket nao encontrado.")

        ensure_company_access(db, current_user, ticket.company_id)
        can_edit = can_access_ticket(ticket, current_user, db) or can_manage_ticket(ticket, current_user, db)
        if not can_edit:
            raise HTTPException(status_code=403, detail="Sem permissao para editar este ticket.")

        previous_assigned_to_id = ticket.assigned_to_id

        if "delivery_due_at" in payload:
            ticket.delivery_due_at = parse_ticket_delivery_due_at(payload.get("delivery_due_at"))

        for field in ["title", "description", "priority", "status", "department", "channel"]:
            if field in payload:
                value = str(payload.get(field) or "").strip() or None
                if field == "department" and value:
                    get_or_create_department(db, ticket.company_id, value)
                    if not is_company_admin(db, current_user, ticket.company_id) and not user_can_access_department(db, current_user, ticket.company_id, value):
                        raise HTTPException(status_code=403, detail="Sem permissao para mover ticket para outro setor.")
                setattr(ticket, field, value)

        if "assigned_to_id" in payload:
            assigned_to_id = payload.get("assigned_to_id")
            if assigned_to_id:
                assigned_user = ensure_user_in_company(db, int(assigned_to_id), ticket.company_id)
                if not is_company_admin(db, current_user, ticket.company_id) and assigned_user.department != ticket.department:
                    raise HTTPException(status_code=403, detail="Responsavel precisa estar no mesmo setor.")
                ticket.assigned_to_id = assigned_user.id
            else:
                ticket.assigned_to_id = None

        if "assigned_to_ids" in payload or "assigned_departments" in payload:
            if not (is_company_admin(db, current_user, ticket.company_id) or ticket.created_by_id == current_user.id):
                raise HTTPException(status_code=403, detail="Somente o dono do ticket pode alterar o time do ticket grupal.")
            sync_ticket_group_assignments(
                db,
                ticket,
                ticket.company_id,
                payload.get("assigned_to_ids") if "assigned_to_ids" in payload else None,
                payload.get("assigned_departments") if "assigned_departments" in payload else None,
            )

        if "attachment_file_id" in payload:
            attachment_file_id = payload.get("attachment_file_id")
            if attachment_file_id:
                attachment = db.query(FileUpload).filter(FileUpload.id == int(attachment_file_id), FileUpload.company_id == ticket.company_id).first()
                if not attachment:
                    raise HTTPException(status_code=400, detail="Anexo nao encontrado.")
                ticket.attachment_file_id = attachment.id
            else:
                ticket.attachment_file_id = None

        if ticket.status == "Resolvido":
            ticket.closed_at = ticket.closed_at or datetime.utcnow()
            ticket.closed_by_id = ticket.closed_by_id or current_user.id
            ticket.close_reason = str(payload.get("close_reason") or ticket.close_reason or "Status alterado para resolvido.").strip()
        else:
            ticket.closed_at = None
            ticket.closed_by_id = None
            ticket.close_reason = None
        ticket.updated_at = datetime.utcnow()
        db.commit()
        db.refresh(ticket)
        ticket_response = serialize_ticket(ticket, db)

        await notify_ticket_participants(
            db,
            ticket,
            current_user,
            "✏️ Ticket atualizado",
            f"{current_user.full_name} atualizou: {ticket.title}",
        )

        return ticket_response


@app.get("/tasks/")
async def list_tasks(
    status: Optional[str] = Query(default=None),
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        query = db.query(TaskItem).filter(TaskItem.company_id == scoped_company_id)
        if status:
            query = query.filter(TaskItem.status == status)
        if not is_company_admin(db, current_user, scoped_company_id):
            query = query.filter(or_(TaskItem.created_by_id == current_user.id, TaskItem.assigned_to_id == current_user.id))
        tasks = query.order_by(TaskItem.created_at.desc()).all()
        return [serialize_task(task, db) for task in tasks]


@app.post("/tasks/")
async def create_task(payload: dict, current_user: User = Depends(get_current_user)):
    title = str(payload.get("title") or "").strip()
    if not title:
        raise HTTPException(status_code=400, detail="Titulo da tarefa e obrigatorio.")

    assigned_to_id = payload.get("assigned_to_id")
    if assigned_to_id and int(assigned_to_id) != int(current_user.id):
        raise HTTPException(
            status_code=403,
            detail="Tarefas sao pessoais e nao podem ser criadas para outros usuarios.",
        )

    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        category = str(payload.get("category") or "Pessoal").strip() or "Pessoal"
        ensure_creation_scope(db, current_user, company_id, None, "task")
        assigned_user = current_user

        task = TaskItem(
            company_id=company_id,
            title=title,
            description=str(payload.get("description") or "").strip() or None,
            priority=str(payload.get("priority") or "medium").strip(),
            category=category,
            status=str(payload.get("status") or "backlog").strip(),
            due_date=str(payload.get("due_date") or "").strip() or None,
            created_by_id=current_user.id,
            assigned_to_id=assigned_user.id if assigned_user else None,
        )
        db.add(task)
        db.commit()
        db.refresh(task)
        return serialize_task(task, db)


@app.patch("/tasks/{task_id}")
async def update_task(task_id: int, payload: dict, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        task = db.query(TaskItem).filter(TaskItem.id == task_id).first()
        if not task:
            raise HTTPException(status_code=404, detail="Tarefa nao encontrada.")

        ensure_company_access(db, current_user, task.company_id)
        can_edit = is_company_admin(db, current_user, task.company_id) or task.created_by_id == current_user.id or task.assigned_to_id == current_user.id
        if not can_edit:
            raise HTTPException(status_code=403, detail="Sem permissao para editar esta tarefa.")

        for field in ["title", "description", "priority", "category", "status", "due_date"]:
            if field in payload:
                value = str(payload.get(field) or "").strip() or None
                if field in ["title", "priority", "category", "status"]:
                    value = value or getattr(task, field)
                if field == "category":
                    value = value or "Pessoal"
                setattr(task, field, value)

        if "assigned_to_id" in payload:
            assigned_to_id = payload.get("assigned_to_id")
            if assigned_to_id and int(assigned_to_id) != int(current_user.id):
                raise HTTPException(
                    status_code=403,
                    detail="Tarefas pessoais nao podem ser transferidas para outros usuarios.",
                )
            task.assigned_to_id = current_user.id

        task.updated_at = datetime.utcnow()
        db.commit()
        db.refresh(task)
        return serialize_task(task, db)


@app.delete("/tasks/{task_id}")
async def delete_task(task_id: int, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        task = db.query(TaskItem).filter(TaskItem.id == task_id).first()
        if not task:
            raise HTTPException(status_code=404, detail="Tarefa nao encontrada.")
        ensure_company_access(db, current_user, task.company_id)
        can_delete = (
            is_company_admin(db, current_user, task.company_id)
            or task.created_by_id == current_user.id
            or task.assigned_to_id == current_user.id
        )
        if not can_delete:
            raise HTTPException(status_code=403, detail="Sem permissao para remover esta tarefa.")
        if task.status != "done":
            raise HTTPException(status_code=400, detail="Conclua a tarefa antes de apaga-la.")
        task_title = task.title
        company_id = task.company_id
        db.delete(task)
        log_audit(db, current_user.id, "tarefa_removida", "task", task_id, company_id, {
            "title": task_title,
            "status": "done",
        })
        db.commit()
        return {"message": "Tarefa concluida removida."}


@app.get("/assistant/preferences")
async def get_assistant_preferences(current_user: User = Depends(get_current_user)):
    return {
        "address_name": assistant_address_name(current_user),
        "configured": bool(str(getattr(current_user, "assistant_address_name", "") or "").strip()),
        "default_name": display_user_name(current_user),
    }


@app.put("/assistant/preferences")
async def update_assistant_preferences(payload: dict, current_user: User = Depends(get_current_user)):
    raw_name = str(payload.get("address_name") or "").strip()
    if len(raw_name) > 80:
        raise HTTPException(status_code=400, detail="O nome de tratamento deve ter no maximo 80 caracteres.")
    if raw_name and (re.search(r"[\x00-\x1f<>]", raw_name) or not re.search(r"\w", raw_name)):
        raise HTTPException(status_code=400, detail="Informe um nome de tratamento valido.")
    with SessionLocal() as db:
        user = db.query(User).filter(User.id == current_user.id).first()
        if not user:
            raise HTTPException(status_code=404, detail="Usuario nao encontrado.")
        user.assistant_address_name = raw_name or None
        user.updated_at = datetime.utcnow()
        db.commit()
        db.refresh(user)
        return {
            "address_name": assistant_address_name(user),
            "configured": bool(user.assistant_address_name),
            "default_name": display_user_name(user),
        }

def is_assistant_greeting_only(value: str) -> bool:
    normalized = normalize_text(value).strip(" .,!?-")
    greetings = {
        "oi", "ola", "olá", "ei", "hey", "e ai", "e aí", "opa", "salve",
        "bom dia", "boa tarde", "boa noite", "tudo bem", "como vai",
        "oi tudo bem", "ola tudo bem", "olá tudo bem",
    }
    return normalized in {normalize_text(item) for item in greetings}


def assistant_social_message(value: str) -> Optional[str]:
    """Responde interacoes sociais sem confundi-las com uma solicitacao operacional."""
    normalized = normalize_text(value).strip(" .,!?-")
    normalized = re.sub(r"^(?:mitty|volt)\s*[,;:-]?\s*", "", normalized).strip()
    normalized = re.sub(r"\s*[,;:-]?\s*(?:mitty|volt)$", "", normalized).strip()
    if is_assistant_greeting_only(normalized):
        return "como posso ajudar hoje? Posso consultar ou organizar tickets, tarefas e sua agenda."
    thanks = {
        "obrigado", "obrigada", "obrigdo", "obrigda", "brigado", "obg", "obgd", "obrigadissimo", "obrigadissima", "valeu",
        "muito obrigado", "muito obrigada", "agradeco", "grato", "grata",
        "perfeito obrigado", "perfeito obrigada", "show obrigado", "show obrigada",
        "obrigado pela ajuda", "obrigada pela ajuda", "valeu pela ajuda",
    }
    if normalized in {normalize_text(item) for item in thanks}:
        return "por nada! Quando precisar, estou por aqui para ajudar."
    farewells = {"tchau", "ate mais", "ate logo", "falou", "bom trabalho"}
    if normalized in {normalize_text(item) for item in farewells}:
        return "ate mais! Se surgir algo, e so me chamar."
    return None


def assistant_conversation_payload(current_user: User, command: str, assistant_name: str, reply: str) -> dict:
    return {
        "assistant_name": assistant_name,
        "address_name": assistant_address_name(current_user),
        "command": command,
        "intent": "conversation",
        "plan": {},
        "executed": True,
        "needs_input": False,
        "missing_fields": [],
        "continue_context": False,
        "requires_confirmation": False,
        "requires_urgency": False,
        "confirmation_summary": "",
        "ticket": None,
        "task": None,
        "tasks": [],
        "meeting": None,
        "reply": assistant_greeting(current_user, reply),
    }


def is_assistant_edit_request(value: str) -> bool:
    normalized = normalize_text(value)
    edit_words = [
        "edite", "editar", "altere", "alterar", "mude", "mudar", "atualize",
        "atualizar", "remarque", "remarcar", "reagende", "reagendar", "corrija", "corrigir",
    ]
    item_words = ["tarefa", "task", "reuniao", "meeting", "compromisso", "lembrete", "agenda"]
    return any(contains_phrase(normalized, word) for word in edit_words) and any(
        contains_phrase(normalized, word) for word in item_words
    )




def assistant_explicit_item_id(value: str, labels: list[str]) -> Optional[str]:
    label_pattern = "|".join(re.escape(normalize_text(label)) for label in labels)
    normalized = normalize_text(value)
    match = re.search(rf"(?:{label_pattern})\s*(?:id\s*)?#?([0-9a-f-]+)\b", normalized)
    if match:
        return match.group(1)
    match = re.search(r"#([0-9a-f-]+)\b", normalized)
    return match.group(1) if match else None


def assistant_title_change(value: str) -> Optional[str]:
    clean = " ".join(str(value or "").split())
    for pattern in [
        r"(?:novo\s+)?t.tulo\s+(?:para|como)\s+['\"]?(.+?)['\"]?(?=\s+(?:no\s+dia|dia|em|.s|as|com\s+o\s+link|link|prioridade|status|descri)\b|$)",
        r"(?:novo\s+)?nome\s+(?:para|como)\s+['\"]?(.+?)['\"]?(?=\s+(?:no\s+dia|dia|em|.s|as|com\s+o\s+link|link|prioridade|status|descri)\b|$)",
    ]:
        match = re.search(pattern, clean, flags=re.IGNORECASE)
        if match:
            return match.group(1).strip(" .'\"")[:180] or None
    return None


def assistant_best_item_match(rows, command: str, entity_label: str):
    normalized = normalize_text(command)
    ignored = {
        "edite", "editar", "altere", "alterar", "mude", "mudar", "atualize", "atualizar",
        "remarque", "remarcar", "reagende", "reagendar", "tarefa", "task", "reuniao",
        "meeting", "compromisso", "lembrete", "agenda", "para", "como", "com", "dia",
        "titulo", "nome", "horario", "data",
    }
    command_words = {word for word in normalized.split() if len(word) >= 3 and word not in ignored}
    scored = []
    for row in rows:
        title = str(getattr(row, "title", "") or "")
        normalized_title = normalize_text(title)
        title_words = {word for word in normalized_title.split() if len(word) >= 3}
        score = (100 if normalized_title and normalized_title in normalized else 0) + len(command_words & title_words)
        if score:
            scored.append((score, row))
    if not scored:
        raise HTTPException(status_code=400, detail=f"Nao encontrei {entity_label} no seu acesso. Informe o ID ou o titulo atual.")
    scored.sort(key=lambda item: item[0], reverse=True)
    if len(scored) > 1 and scored[0][0] == scored[1][0]:
        raise HTTPException(status_code=400, detail=f"Encontrei mais de um {entity_label}. Informe o ID para editar o item correto.")
    return scored[0][1]


def assistant_edited_schedule(value: str, current_start, current_end) -> tuple[Optional[datetime], Optional[datetime]]:
    schedule_value = re.sub(r"\b[0-9a-f]{8}-[0-9a-f-]{27,36}\b", " ", value, flags=re.IGNORECASE)
    normalized = normalize_text(schedule_value)
    has_date = bool(
        re.search(r"(?<!\d)\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?(?!\d)", normalized)
        or any(contains_phrase(normalized, marker) for marker in [
            "hoje", "amanha", "depois de amanha", "segunda", "terca", "quarta",
            "quinta", "sexta", "sabado", "domingo",
        ])
    )
    times = extract_assistant_time_matches(schedule_value)
    if not has_date and not times:
        return None, None
    start = meeting_datetime(current_start, "Inicio").astimezone(BOLT_DAILY_SUMMARY_TIMEZONE)
    end = meeting_datetime(current_end, "Termino").astimezone(BOLT_DAILY_SUMMARY_TIMEZONE)
    duration = end - start
    if has_date and times:
        return infer_assistant_schedule(schedule_value)
    if has_date:
        probe = f"{schedule_value} as {start.hour:02d}:{start.minute:02d}"
        updated_start, _ = infer_assistant_schedule(probe)
        return updated_start, updated_start + duration
    _, hour, minute = times[0]
    updated_local = start.replace(hour=hour, minute=minute, second=0, microsecond=0)
    return updated_local.astimezone(timezone.utc), (updated_local + duration).astimezone(timezone.utc)


def plan_assistant_edit(value: str, current_user: User, db, company_id: str) -> dict:
    normalized = normalize_text(value)
    meeting_request = any(contains_phrase(normalized, word) for word in [
        "reuniao", "meeting", "compromisso", "lembrete", "agenda",
    ])
    if meeting_request:
        explicit_id = assistant_explicit_item_id(value, ["reuniao", "meeting", "compromisso", "lembrete", "agenda"])
        rows = db.execute(text("""
            SELECT m.* FROM meetings m
            WHERE m.company_id=:company_id AND m.status='scheduled'
              AND (m.creator_user_id=:user_id OR :can_manage=TRUE)
            ORDER BY m.starts_at DESC
        """), {
            "company_id": company_id,
            "user_id": current_user.id,
            "can_manage": is_company_admin(db, current_user, company_id),
        }).all()
        target = next((row for row in rows if explicit_id and str(row.id).lower() == explicit_id.lower()), None)
        target = target or assistant_best_item_match(rows, value, "item da agenda")
        changes = {}
        new_title = assistant_title_change(value)
        if new_title:
            changes["title"] = new_title
        link_url = infer_assistant_link(value)
        if link_url:
            changes["link_url"] = link_url
        starts_at, ends_at = assistant_edited_schedule(value, target.starts_at, target.ends_at)
        if starts_at:
            changes.update({"starts_at": starts_at.isoformat(), "ends_at": ends_at.isoformat()})
        users = (
            db.query(User).join(CompanyUser, CompanyUser.user_id == User.id)
            .filter(User.is_active == True, CompanyUser.company_id == company_id, CompanyUser.status == ACTIVE_STATUS)
            .all()
        )
        assignees = infer_assignees(value, users)
        participant_markers = ["participante", "participantes", "inclua", "incluir", "convide", "convidar"]
        if any(contains_phrase(normalized, marker) for marker in participant_markers):
            if not assignees:
                raise HTTPException(status_code=400, detail="Nao encontrei os participantes citados na empresa.")
            changes["participant_ids"] = sorted({current_user.id, *[item.id for item in assignees]})
        if not changes:
            raise HTTPException(status_code=400, detail="Diga o que deseja alterar: titulo, data, horario, link ou participantes.")
        return {
            "intent": "meeting_edit", "assistant_name": MITTY_NAME,
            "target_id": str(target.id), "target_title": target.title,
            "subject": target.title, "title": changes.get("title", target.title), "changes": changes,
            "starts_at": changes.get("starts_at", meeting_datetime(target.starts_at, "Inicio").isoformat()),
            "ends_at": changes.get("ends_at", meeting_datetime(target.ends_at, "Termino").isoformat()),
            "assigned_to_names": [item.full_name for item in assignees],
        }

    explicit_id = assistant_explicit_item_id(value, ["tarefa", "task"])
    rows = db.query(TaskItem).filter(TaskItem.company_id == company_id).filter(
        or_(TaskItem.created_by_id == current_user.id, TaskItem.assigned_to_id == current_user.id)
    ).order_by(TaskItem.created_at.desc()).all()
    target = next((row for row in rows if explicit_id and str(row.id) == explicit_id), None)
    target = target or assistant_best_item_match(rows, value, "tarefa")
    changes = {}
    new_title = assistant_title_change(value)
    if new_title:
        changes["title"] = new_title[:200]
    status_aliases = {
        "concluida": "done", "finalizada": "done", "finalizado": "done",
        "em andamento": "doing", "fazendo": "doing", "revisao": "review", "backlog": "backlog",
    }
    for label, status in status_aliases.items():
        if contains_phrase(normalized, label):
            changes["status"] = status
            break
    if any(contains_phrase(normalized, word) for word in ["prioridade", "urgente", "alta", "baixa", "media"]):
        changes["priority"] = infer_priority(value, task_format=True)
    description_match = re.search(r"descri[c?][a?]o\s+(?:para|como)\s+(.+)$", value, flags=re.IGNORECASE)
    if description_match:
        changes["description"] = description_match.group(1).strip()
    if not changes:
        raise HTTPException(status_code=400, detail="Diga o que deseja alterar: titulo, descricao, prioridade ou status.")
    return {
        "intent": "task_edit", "assistant_name": MITTY_NAME,
        "target_id": target.id, "target_title": target.title,
        "subject": target.title, "title": changes.get("title", target.title), "changes": changes,
        "assigned_to_names": [assistant_address_name(current_user)],
    }
def assistant_clarification_payload(
    current_user: User,
    command: str,
    assistant_name: str,
    intent: str,
    message: str,
    missing_fields: list[str],
    *,
    continue_context: bool = True,
    plan: Optional[dict] = None,
) -> dict:
    return {
        "assistant_name": assistant_name,
        "address_name": assistant_address_name(current_user),
        "command": command,
        "intent": intent,
        "plan": plan or {},
        "executed": False,
        "needs_input": True,
        "missing_fields": missing_fields,
        "continue_context": continue_context,
        "requires_confirmation": False,
        "requires_urgency": "urgency" in missing_fields,
        "confirmation_summary": "",
        "ticket": None,
        "task": None,
        "tasks": [],
        "meeting": None,
        "reply": assistant_greeting(current_user, message),
    }


def assistant_subject_is_missing(value: str, intent: str) -> bool:
    normalized = normalize_text(value)
    action_words = {
        "ticket": ["ticket", "tiket", "ticekt", "chamado"],
        "task": ["tarefa", "task", "atividade"],
        "agenda_task": ["tarefa", "task", "agenda"],
        "meeting": ["reuniao", "meeting"],
        "reminder": ["compromisso", "lembrete", "evento", "agenda"],
    }.get(intent, [])
    removable = [
        "crie", "criar", "abra", "abrir", "adicione", "adicionar", "agende",
        "agendar", "marque", "marcar", "um", "uma", "o", "a", "por favor",
        *action_words,
    ]
    remainder = normalized
    for item in sorted(removable, key=len, reverse=True):
        remainder = re.sub(rf"(?<!\w){re.escape(normalize_text(item))}(?!\w)", " ", remainder)
    remainder = re.sub(r"\s+", " ", remainder).strip(" .,:;!?")
    return len(remainder) < 3

@app.post("/assistant/requests")
async def assistant_request(payload: dict, current_user: User = Depends(get_current_user)):
    message = str(payload.get("message") or "").strip()
    execute = bool(payload.get("execute", False))
    confirmed = bool(payload.get("confirmed", False))
    if not message:
        raise HTTPException(status_code=400, detail="Informe uma solicitação para a IA.")

    social_reply = assistant_social_message(message)
    if social_reply:
        assistant_username = str(payload.get("assistant_username") or "").strip().lower()
        assistant_name = (
            BOLT_NAME if assistant_username == BOLT_USERNAME
            else MITTY_NAME if assistant_username == MITTY_USERNAME
            else "Volt e Mitty"
        )
        return assistant_conversation_payload(current_user, message, assistant_name, social_reply)

    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        address_name = assistant_address_name(current_user)
        if is_assistant_edit_request(message):
            plan = plan_assistant_edit(message, current_user, db, company_id)
            intent = plan["intent"]
            changes = plan["changes"]
            change_labels = {
                "title": "titulo", "description": "descricao", "priority": "prioridade",
                "status": "status", "starts_at": "data/horario", "ends_at": "termino",
                "link_url": "link", "participant_ids": "participantes",
            }
            changed_text = ", ".join(change_labels.get(key, key) for key in changes)
            confirmation_summary = (
                f"Mitty alterara '{plan['target_title']}' (ID {plan['target_id']}): {changed_text}. "
                "A alteracao so sera aplicada depois da sua confirmacao."
            )
            response = {
                "assistant_name": MITTY_NAME,
                "address_name": address_name,
                "command": message,
                "intent": intent,
                "plan": plan,
                "executed": False,
                "needs_input": False,
                "missing_fields": [],
                "continue_context": True,
                "requires_confirmation": True,
                "requires_urgency": False,
                "confirmation_summary": confirmation_summary,
                "ticket": None,
                "task": None,
                "tasks": [],
                "meeting": None,
                "reply": assistant_greeting(current_user, confirmation_summary),
            }
            if not execute or not confirmed:
                return response

            if intent == "task_edit":
                task = db.query(TaskItem).filter(
                    TaskItem.id == int(plan["target_id"]),
                    TaskItem.company_id == company_id,
                ).first()
                if not task:
                    raise HTTPException(status_code=404, detail="Tarefa nao encontrada.")
                if task.created_by_id != current_user.id and task.assigned_to_id != current_user.id and not is_company_admin(db, current_user, company_id):
                    raise HTTPException(status_code=403, detail="Sem permissao para editar esta tarefa.")
                for field in ["title", "description", "priority", "status"]:
                    if field in changes:
                        setattr(task, field, changes[field])
                task.updated_at = datetime.utcnow()
                log_audit(db, current_user.id, "assistente_tarefa_editada", "task", task.id, company_id, {
                    "command": message, "changes": changes, "assistant": MITTY_NAME,
                })
                db.commit()
                db.refresh(task)
                response.update({
                    "executed": True,
                    "requires_confirmation": False,
                    "task": serialize_task(task, db),
                    "tasks": [serialize_task(task, db)],
                    "reply": assistant_greeting(current_user, f"editei a tarefa #{task.id} '{task.title}' com sucesso."),
                })
                return response

            meeting = db.execute(text("""
                SELECT * FROM meetings WHERE id=:id AND company_id=:company_id AND status='scheduled'
            """), {"id": plan["target_id"], "company_id": company_id}).first()
            if not meeting:
                raise HTTPException(status_code=404, detail="Item da agenda nao encontrado.")
            if meeting.creator_user_id != current_user.id and not is_company_admin(db, current_user, company_id):
                raise HTTPException(status_code=403, detail="Somente o criador ou admin da empresa pode editar este item.")
            updates = {}
            for field in ["title", "description", "link_url", "starts_at", "ends_at"]:
                if field in changes:
                    updates[field] = (
                        meeting_datetime(changes[field], "Data e horario")
                        if field in {"starts_at", "ends_at"}
                        else changes[field]
                    )
            starts_at = updates.get("starts_at", meeting_datetime(meeting.starts_at, "Inicio"))
            ends_at = updates.get("ends_at", meeting_datetime(meeting.ends_at, "Termino"))
            if ends_at <= starts_at:
                raise HTTPException(status_code=400, detail="O termino deve ser posterior ao inicio.")
            if "link_url" in updates:
                updates["link_url"] = normalize_meeting_link(updates["link_url"])
            assignments = [f"{field}=:{field}" for field in updates]
            params = {**updates, "id": plan["target_id"], "updated_at": datetime.now(timezone.utc)}
            assignments.extend([
                "updated_at=:updated_at", "reminder_day_sent=FALSE",
                "reminder_30m_sent=FALSE", "start_alert_sent=FALSE",
            ])
            db.execute(text(f"UPDATE meetings SET {', '.join(assignments)} WHERE id=:id"), params)
            if "participant_ids" in changes:
                participant_ids = {current_user.id, *[int(item) for item in changes["participant_ids"]]}
                db.execute(text("DELETE FROM meeting_participants WHERE meeting_id=:id"), {"id": plan["target_id"]})
                for user_id in participant_ids:
                    ensure_user_in_company(db, user_id, company_id)
                    db.execute(text("""
                        INSERT INTO meeting_participants (meeting_id, user_id, response_status)
                        VALUES (:meeting_id, :user_id, :status)
                    """), {
                        "meeting_id": plan["target_id"], "user_id": user_id,
                        "status": "accepted" if user_id == current_user.id else "pending",
                    })
            participant_ids = [row[0] for row in db.execute(text(
                "SELECT user_id FROM meeting_participants WHERE meeting_id=:id"
            ), {"id": plan["target_id"]}).all()]
            log_audit(db, current_user.id, "assistente_agenda_editada", "meeting", plan["target_id"], company_id, {
                "command": message, "changes": changes, "assistant": MITTY_NAME,
            })
            db.commit()
            row = db.execute(text("SELECT * FROM meetings WHERE id=:id"), {"id": plan["target_id"]}).first()
            serialized = serialize_meeting_row(db, row)
            local_start = meeting_datetime(serialized["starts_at"], "Inicio").astimezone(BOLT_DAILY_SUMMARY_TIMEZONE).strftime("%d/%m/%Y às %H:%M")
            response.update({
                "executed": True,
                "requires_confirmation": False,
                "meeting": serialized,
                "reply": assistant_greeting(current_user, f"atualizei '{serialized['title']}' para {local_start} e avisei os participantes."),
            })
            for user_id in participant_ids:
                await dispatch_meeting_reminder(
                    user_id, plan["target_id"], "Mitty - Agenda atualizada",
                    f"{serialized['title']} foi atualizado para {local_start}.",
                    meeting_datetime(serialized["starts_at"], "Inicio"), "updated",
                )
            wake_meeting_reminder_scheduler()
            return response

        if is_ticket_summary_request(message):
            ticket_summary = build_user_ticket_summary(db, current_user, company_id)
            return {
                "assistant_name": BOLT_NAME,
                "command": message,
                "intent": "ticket_summary",
                "plan": {},
                "executed": True,
                "requires_confirmation": False,
                "requires_urgency": False,
                "confirmation_summary": "",
                "ticket": None,
                "task": None,
                "tasks": [],
                "meeting": None,
                "ticket_summary": ticket_summary,
                "reply": assistant_greeting(current_user, build_ticket_summary_reply(ticket_summary)),
            }

        if is_mitty_summary_request(message):
            schedule_summary = build_mitty_schedule_summary(db, current_user, company_id, message)
            return {
                "assistant_name": MITTY_NAME,
                "command": message,
                "intent": "schedule_summary",
                "plan": {},
                "executed": True,
                "requires_confirmation": False,
                "requires_urgency": False,
                "confirmation_summary": "",
                "ticket": None,
                "task": None,
                "tasks": [],
                "meeting": None,
                "schedule_summary": schedule_summary,
                "reply": schedule_summary["reply"],
            }
        try:
            intent_hint = infer_intent(message)
        except HTTPException as exc:
            if "altera ou remove" in str(exc.detail):
                raise
            return assistant_clarification_payload(
                current_user,
                message,
                "IA",
                "conversation",
                "ainda não entendi qual ação você deseja. Diga se quer criar um ticket, uma task, um compromisso ou uma reunião.",
                ["action"],
            )

        assistant_hint = BOLT_NAME if intent_hint == "ticket" else MITTY_NAME
        if assistant_subject_is_missing(message, intent_hint):
            return assistant_clarification_payload(
                current_user,
                message,
                assistant_hint,
                intent_hint,
                "falta o assunto. Explique em uma frase o que precisa ser feito.",
                ["subject"],
            )

        if intent_hint in {"meeting", "reminder", "agenda_task"}:
            schedule_missing = assistant_schedule_requirements(message)
            if schedule_missing:
                missing_labels = []
                if "date" in schedule_missing:
                    missing_labels.append("a data, como amanhã ou 25/07")
                if "time" in schedule_missing:
                    missing_labels.append("o horário, como 14h, 14h30, 14:30 ou 1430")
                return assistant_clarification_payload(
                    current_user,
                    message,
                    MITTY_NAME,
                    intent_hint,
                    "está faltando " + " e ".join(missing_labels) + ".",
                    schedule_missing,
                )

        try:
            plan = plan_assistant_action(
                message,
                current_user,
                db,
                company_id,
                payload.get("urgency"),
            )
        except HTTPException as exc:
            if exc.status_code != 400 or "altera ou remove" in str(exc.detail):
                raise
            return assistant_clarification_payload(
                current_user,
                message,
                assistant_hint,
                intent_hint,
                str(exc.detail),
                ["details"],
            )
        intent = plan["intent"]
        assistant_name = plan["assistant_name"]
        ensure_creation_scope(db, current_user, company_id, plan["department"], intent)

        department_text = ", ".join(plan["departments"])
        assignee_text = ", ".join(plan["assigned_to_names"]) or current_user.full_name
        urgency_text = plan["ticket_priority"] or "a selecionar"

        if intent == "ticket":
            confirmation_summary = (
                f"Volt abrirá o ticket sobre '{plan['subject']}' para {assignee_text}. "
                f"Setor de referência: {department_text}. Urgência: {urgency_text}. "
                "Somente os usuários adicionados ao ticket receberão acesso e notificações."
            )
        elif intent == "task":
            confirmation_summary = (
                f"Mitty criará a tarefa pessoal sobre '{plan['subject']}' para {assignee_text} "
                "e notificará cada responsável. Confirme antes de executar."
            )
        else:
            starts_at = meeting_datetime(plan["starts_at"], "Inicio")
            ends_at = meeting_datetime(plan["ends_at"], "Termino")
            local_start = starts_at.astimezone(BOLT_DAILY_SUMMARY_TIMEZONE).strftime("%d/%m/%Y às %H:%M")
            local_end = ends_at.astimezone(BOLT_DAILY_SUMMARY_TIMEZONE).strftime("%d/%m/%Y às %H:%M")
            kind_label = {
                "meeting": "reunião",
                "agenda_task": "tarefa agendada",
                "reminder": "compromisso",
            }[intent]
            link_confirmation = f" Link: {plan['link_url']}." if plan.get("link_url") else ""
            external_confirmation = (
                f" Contato externo mencionado: {plan['external_counterpart']}."
                if plan.get("external_counterpart")
                else ""
            )
            confirmation_summary = (
                f"Mitty agendará a {kind_label} '{plan['subject']}' em {local_start}, com término em {local_end}. "
                f"Título: '{plan['title']}'. Participantes internos: {assignee_text}."
                f"{external_confirmation}{link_confirmation} Enviarei os lembretes aos participantes internos."
            )

        response = {
            "assistant_name": assistant_name,
            "address_name": address_name,
            "command": message,
            "intent": intent,
            "plan": plan,
            "executed": False,
            "needs_input": False,
            "missing_fields": [],
            "continue_context": True,
            "requires_confirmation": True,
            "requires_urgency": intent == "ticket" and not plan["ticket_priority"],
            "confirmation_summary": confirmation_summary,
            "ticket": None,
            "task": None,
            "tasks": [],
            "meeting": None,
            "reply": "",
        }

        if intent == "ticket" and not plan["ticket_priority"]:
            response.update({
                "needs_input": True,
                "missing_fields": ["urgency"],
                "requires_confirmation": False,
            })
            response["reply"] = assistant_greeting(
                current_user,
                f"falta o nível de urgência. Responda com: {', '.join(TICKET_URGENCY_LEVELS)}.",
            )
            return response

        if not execute or not confirmed:
            response["reply"] = assistant_greeting(
                current_user,
                confirmation_summary,
            )
            return response

        if intent == "ticket":
            ticket = Ticket(
                company_id=company_id,
                title=plan["title"],
                description=plan["description"],
                priority=plan["ticket_priority"],
                status="Aberto",
                department=plan["department"],
                channel="Assistente",
                created_by_id=current_user.id,
                assigned_to_id=None,
                attachment_file_id=None,
            )
            attachment_file_id = payload.get("attachment_file_id")
            if attachment_file_id:
                attachment = db.query(FileUpload).filter(
                    FileUpload.id == int(attachment_file_id),
                    FileUpload.company_id == company_id,
                ).first()
                if not attachment:
                    raise HTTPException(status_code=404, detail="Anexo do ticket nao encontrado.")
                ticket.attachment_file_id = attachment.id
            db.add(ticket)
            db.flush()
            sync_ticket_group_assignments(
                db,
                ticket,
                company_id,
                plan["assigned_to_ids"],
                plan["departments"],
            )
            log_audit(
                db,
                current_user.id,
                "assistente_ticket_criado",
                "ticket",
                ticket.id,
                company_id,
                {"command": message, "plan": plan, "assistant": BOLT_NAME},
            )
            bolt_notifications = create_bolt_ticket_messages(db, ticket, plan, current_user)
            db.commit()
            db.refresh(ticket)
            response["executed"] = True
            response["requires_confirmation"] = False
            response["requires_urgency"] = False
            response["ticket"] = serialize_ticket(ticket, db)
            response["reply"] = assistant_greeting(
                current_user,
                (
                    f"criei o ticket #{ticket.id} sobre {plan['subject']} para {assignee_text}. "
                    f"Setor de referência: {department_text}. Urgência: {plan['ticket_priority']}. "
                    "Os participantes explícitos foram notificados."
                ),
            )
            await dispatch_bolt_ticket_notifications(ticket, bolt_notifications)
            return response

        if intent == "task":
            task_assignee_ids = [current_user.id]
            created_tasks = []
            mitty_messages = []
            mitty = ensure_mitty_user(db, company_id)
            for assigned_user_id in task_assignee_ids:
                assigned_user = ensure_user_in_company(db, assigned_user_id, company_id)
                task = TaskItem(
                    company_id=company_id,
                    title=plan["title"],
                    description=plan["description"],
                    priority=plan["task_priority"],
                    category="Pessoal",
                    status="backlog",
                    due_date=plan["due_date"],
                    created_by_id=current_user.id,
                    assigned_to_id=assigned_user.id,
                )
                db.add(task)
                db.flush()
                log_audit(
                    db,
                    current_user.id,
                    "assistente_tarefa_criada",
                    "task",
                    task.id,
                    company_id,
                    {"command": message, "plan": plan, "assigned_to_id": assigned_user.id, "assistant": MITTY_NAME},
                )
                notice_text = assistant_greeting(
                    assigned_user,
                    f"uma nova tarefa foi criada para voce: {task.title}. Consulte a aba Tarefas.",
                )
                direct_message = Message(
                    company_id=company_id,
                    content=notice_text,
                    sender_id=mitty.id,
                    receiver_id=assigned_user.id,
                    message_type="task_notification",
                )
                db.add(direct_message)
                db.flush()
                mitty_messages.append((assigned_user.id, task, serialize_message(direct_message, db)))
                created_tasks.append(task)
            db.commit()
            for task in created_tasks:
                db.refresh(task)
            response["executed"] = True
            response["requires_confirmation"] = False
            serialized_tasks = [serialize_task(task, db) for task in created_tasks]
            response["task"] = serialized_tasks[0]
            response["tasks"] = serialized_tasks
            task_ids = ", ".join(f"#{task.id}" for task in created_tasks)
            response["reply"] = assistant_greeting(
                current_user,
                f"criei {len(created_tasks)} tarefa(s) pessoal(is) ({task_ids}) e notifiquei os responsaveis.",
            )
            for user_id, task, message_payload in mitty_messages:
                await manager.send_personal_message(
                    json.dumps({"type": "new_message", "message": message_payload}),
                    user_id,
                )
                await notify_task_user(
                    user_id,
                    "Mitty - Nova tarefa atribuida",
                    message_payload["content"],
                    task.id,
                )
            return response

        starts_at = meeting_datetime(plan["starts_at"], "Inicio")
        ends_at = meeting_datetime(plan["ends_at"], "Termino")
        meeting_id = str(uuid.uuid4())
        participant_ids = {current_user.id, *plan["assigned_to_ids"]}
        db.execute(text("""
            INSERT INTO meetings (
                id, company_id, creator_user_id, title, description, meeting_type,
                starts_at, ends_at, color, link_url
            )
            VALUES (
                :id, :company_id, :creator_user_id, :title, :description, :meeting_type,
                :starts_at, :ends_at, :color, :link_url
            )
        """), {
            "id": meeting_id,
            "company_id": company_id,
            "creator_user_id": current_user.id,
            "title": plan["title"],
            "description": plan["description"],
            "meeting_type": plan["meeting_type"],
            "starts_at": starts_at,
            "ends_at": ends_at,
            "color": plan["color"],
            "link_url": plan["link_url"],
        })
        for user_id in participant_ids:
            ensure_user_in_company(db, int(user_id), company_id)
            db.execute(text("""
                INSERT INTO meeting_participants (meeting_id, user_id, response_status)
                VALUES (:meeting_id, :user_id, :status)
            """), {
                "meeting_id": meeting_id,
                "user_id": user_id,
                "status": "accepted" if user_id == current_user.id else "pending",
            })

        mitty = ensure_mitty_user(db, company_id)
        local_start = starts_at.astimezone(BOLT_DAILY_SUMMARY_TIMEZONE).strftime("%d/%m/%Y às %H:%M")
        kind_label = "reuniao" if plan["meeting_type"] == "video" else (
            "tarefa" if plan["meeting_type"] == "task" else "compromisso"
        )
        kind_phrase = {
            "reuniao": "sua reunião",
            "tarefa": "sua tarefa",
            "compromisso": "seu compromisso",
        }.get(kind_label, "seu item")
        invitation_messages = []
        for user_id in participant_ids:
            participant = ensure_user_in_company(db, int(user_id), company_id)
            invitation_text = assistant_greeting(
                participant,
                f"{kind_phrase} sobre {plan['subject']} começa em {local_start}. Consulte a aba Agenda.",
            )
            if plan["link_url"]:
                invitation_text += f" Link: {plan['link_url']}"
            direct_message = Message(
                company_id=company_id,
                content=invitation_text,
                sender_id=mitty.id,
                receiver_id=user_id,
                message_type="meeting_invitation",
            )
            db.add(direct_message)
            db.flush()
            invitation_messages.append((user_id, invitation_text, serialize_message(direct_message, db)))

        log_audit(
            db,
            current_user.id,
            "assistente_meeting_criado",
            "meeting",
            meeting_id,
            company_id,
            {"command": message, "plan": plan, "assistant": MITTY_NAME},
        )
        db.commit()
        row = db.execute(text("SELECT * FROM meetings WHERE id=:id"), {"id": meeting_id}).first()
        response["meeting"] = serialize_meeting_row(db, row)
        response["executed"] = True
        response["requires_confirmation"] = False
        response["reply"] = assistant_greeting(
            current_user,
            f"agendei {kind_phrase} '{plan['subject']}' para {local_start} e notifiquei os participantes.",
        )
        for user_id, invitation_text, message_payload in invitation_messages:
            await manager.send_personal_message(
                json.dumps({"type": "new_message", "message": message_payload}),
                user_id,
            )
            await dispatch_meeting_reminder(
                user_id,
                meeting_id,
                f"Mitty - Novo {kind_label}",
                invitation_text,
                starts_at,
                "invite",
            )
        wake_meeting_reminder_scheduler()
        return response

@app.post("/assistant/chat")
async def assistant_chat(payload: dict, current_user: User = Depends(get_current_user)):
    message = str(payload.get("message") or "").strip()
    assistant_username = str(payload.get("assistant_username") or "").strip().lower()
    if not message:
        raise HTTPException(status_code=400, detail="Digite uma mensagem para o assistente.")
    if assistant_username not in {BOLT_USERNAME, MITTY_USERNAME}:
        raise HTTPException(status_code=400, detail="Assistente invalido.")

    request_payload = {
        "message": message,
        "company_id": payload.get("company_id"),
        "urgency": payload.get("urgency"),
        "execute": bool(payload.get("execute", False)),
        "confirmed": bool(payload.get("confirmed", False)),
        "attachment_file_id": payload.get("attachment_file_id"),
        "assistant_username": assistant_username,
    }
    # Valida o assistente escolhido em uma pre-analise que nunca executa a acao.
    # Assim, uma chamada direta manipulada nao executa a responsabilidade do outro bot.
    if request_payload["execute"] or request_payload["confirmed"]:
        preview_payload = {**request_payload, "execute": False, "confirmed": False}
        result = await assistant_request(preview_payload, current_user)
    else:
        result = await assistant_request(request_payload, current_user)
    intent = str(result.get("intent") or "conversation")
    allowed = (
        intent in {"ticket", "ticket_summary", "conversation"}
        if assistant_username == BOLT_USERNAME
        else intent in {"task", "meeting", "reminder", "agenda_task", "task_edit", "meeting_edit", "schedule_summary", "conversation"}
    )
    if not allowed:
        assistant_name = BOLT_NAME if assistant_username == BOLT_USERNAME else MITTY_NAME
        responsibility = (
            "Eu cuido de tickets. Para agenda, compromissos e tarefas, converse com a Mitty."
            if assistant_username == BOLT_USERNAME
            else "Eu cuido da sua agenda, compromissos e tarefas. Para tickets, converse com o Volt."
        )
        result = {
            "assistant_name": assistant_name,
            "command": message,
            "intent": "conversation",
            "plan": {},
            "executed": False,
            "needs_input": False,
            "requires_confirmation": False,
            "requires_urgency": False,
            "reply": assistant_greeting(current_user, responsibility),
        }
    elif request_payload["execute"] and request_payload["confirmed"]:
        result = await assistant_request(request_payload, current_user)

    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        assistant = (
            ensure_bolt_user(db, company_id)
            if assistant_username == BOLT_USERNAME
            else ensure_mitty_user(db, company_id)
        )
        user_message = Message(
            company_id=company_id,
            content=str(payload.get("display_message") or message).strip(),
            sender_id=current_user.id,
            receiver_id=assistant.id,
            message_type="assistant_chat",
        )
        db.add(user_message)
        db.flush()
        assistant_message = Message(
            company_id=company_id,
            content=str(result.get("reply") or "Estou pronto para ajudar."),
            sender_id=assistant.id,
            receiver_id=current_user.id,
            message_type="assistant_chat",
        )
        db.add(assistant_message)
        db.commit()
        db.refresh(user_message)
        db.refresh(assistant_message)
        user_payload = serialize_message(user_message, db)
        assistant_payload = serialize_message(assistant_message, db)

    for item in (user_payload, assistant_payload):
        await manager.send_personal_message(
            json.dumps({"type": "new_message", "message": item}),
            current_user.id,
        )
    return {**result, "chat_messages": [user_payload, assistant_payload]}

async def read_import_rows(file: UploadFile) -> list[dict]:
    content = await file.read()
    suffix = Path(file.filename or "").suffix.lower()

    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise HTTPException(status_code=400, detail="Importacao .xlsx requer a dependencia openpyxl no backend.")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(cell or "").strip().lower() for cell in rows[0]]
        result = []
        for values in rows[1:]:
            row = {headers[index]: str(value or "").strip() for index, value in enumerate(values) if index < len(headers)}
            if any(row.values()):
                result.append(row)
        return result

    text_payload = content.decode("utf-8-sig", errors="replace")
    sample = text_payload[:2048]
    delimiter = ";"
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=";,	,").delimiter
    except csv.Error:
        if "," in sample and ";" not in sample:
            delimiter = ","
        elif "\t" in sample:
            delimiter = "\t"
    reader = csv.DictReader(io.StringIO(text_payload), delimiter=delimiter)
    return [
        {str(key or "").strip().lower(): str(value or "").strip() for key, value in row.items()}
        for row in reader
        if any(str(value or "").strip() for value in row.values())
    ]


def preview_company_import_rows(rows: list[dict]) -> dict:
    preview = []
    for index, row in enumerate(rows, start=2):
        errors = []
        warnings = []
        name = row.get("nome_empresa", "").strip()
        status = normalize_status(row.get("status"))
        if not name:
            errors.append("nome_empresa obrigatorio")
        if status not in {ACTIVE_STATUS, INACTIVE_STATUS}:
            errors.append("status invalido")
        preview.append(
            {
                "row": index,
                "data": {
                    "nome_empresa": name,
                    "cnpj": row.get("cnpj", "").strip(),
                    "responsavel": row.get("responsavel", "").strip(),
                    "telefone_1": row.get("telefone_1", "").strip(),
                    "telefone_2": row.get("telefone_2", "").strip(),
                    "status": status,
                },
                "errors": errors,
                "warnings": warnings,
                "valid": not errors,
            }
        )
    return {
        "rows": preview,
        "valid_count": len([row for row in preview if row["valid"]]),
        "error_count": len([row for row in preview if row["errors"]]),
    }


def preview_user_import_rows(rows: list[dict], db) -> dict:
    preview = []
    for index, row in enumerate(rows, start=2):
        errors = []
        warnings = []
        email = row.get("email", "").strip().lower()
        company_id = row.get("id_empresa", "").strip()
        role = normalize_platform_role(row.get("nivel_usuario"))
        if not row.get("nome_usuario", "").strip():
            errors.append("nome_usuario obrigatorio")
        if not EMAIL_PATTERN.match(email):
            errors.append("email invalido")
        company = db.query(Company).filter(Company.id == company_id).first() if company_id else None
        if not company:
            errors.append("id_empresa nao encontrado")
        if role not in PLATFORM_ROLES:
            errors.append("nivel_usuario invalido")
        if not row.get("senha_primaria", "").strip():
            errors.append("senha_primaria obrigatoria")
        existing_user = db.query(User).filter(User.email == email).first() if email else None
        if existing_user:
            warnings.append("email existente: sera criado apenas o vinculo se necessario")
        preview.append(
            {
                "row": index,
                "data": {
                    "nome_usuario": row.get("nome_usuario", "").strip(),
                    "email": email,
                    "senha_primaria": row.get("senha_primaria", "").strip(),
                    "id_empresa": company_id,
                    "telefone": row.get("telefone", "").strip(),
                    "setor": row.get("setor", "").strip() or "Geral",
                    "nivel_usuario": role,
                    "status": normalize_status(row.get("status")),
                },
                "errors": errors,
                "warnings": warnings,
                "valid": not errors,
            }
        )
    return {
        "rows": preview,
        "valid_count": len([row for row in preview if row["valid"]]),
        "error_count": len([row for row in preview if row["errors"]]),
    }


def build_platform_mind_map(db) -> dict:
    companies = db.query(Company).order_by(Company.name.asc()).all()
    return {
        "id": "platform-root",
        "type": "platform",
        "label": "Admin Master",
        "children": [
            {
                "id": company.id,
                "type": "company",
                "label": company.name,
                "status": company.status,
                "children": [
                    {
                        "id": department.id,
                        "type": "department",
                        "company_id": company.id,
                        "label": department.name,
                        "status": department.status,
                        "children": [
                            {
                                "id": str(link.user_id),
                                "link_id": link.id,
                                "type": "user",
                                "company_id": company.id,
                                "department_id": department.id,
                                "department_name": department.name,
                                "label": user.full_name if user else f"Usuario {link.user_id}",
                                "email": user.email if user else None,
                                "role": link.role,
                                "status": link.status,
                            }
                            for link, user in (
                                db.query(CompanyUser, User)
                                .join(User, User.id == CompanyUser.user_id)
                                .filter(CompanyUser.company_id == company.id, CompanyUser.department_id == department.id)
                                .order_by(User.full_name.asc())
                                .all()
                            )
                        ],
                    }
                    for department in (
                        db.query(Department)
                        .filter(Department.company_id == company.id)
                        .order_by(Department.name.asc())
                        .all()
                    )
                ],
            }
            for company in companies
        ],
    }


MAX_CHAT_BACKUP_UNCOMPRESSED_BYTES = 128 * 1024 * 1024


def iter_chat_backup_messages(backup: ChatHistoryBackup):
    compressed_data = bytes(backup.compressed_data or b"")
    if backup.checksum_sha256:
        actual_checksum = hashlib.sha256(compressed_data).hexdigest()
        if actual_checksum != backup.checksum_sha256:
            raise ValueError("Checksum do backup nao confere.")
    with zipfile.ZipFile(io.BytesIO(compressed_data), "r") as archive:
        try:
            info = archive.getinfo("messages.json")
        except KeyError as exc:
            raise ValueError("Backup sem messages.json.") from exc
        if info.file_size > MAX_CHAT_BACKUP_UNCOMPRESSED_BYTES:
            raise ValueError("Backup excede o limite seguro para restauracao.")
        with archive.open(info, "r") as binary_stream:
            stream = io.TextIOWrapper(binary_stream, encoding="utf-8")
            decoder = json.JSONDecoder()
            buffer = ""
            position = 0
            started = False
            eof = False
            while True:
                if position:
                    buffer = buffer[position:]
                    position = 0
                chunk = stream.read(65536)
                if chunk:
                    buffer += chunk
                else:
                    eof = True
                while True:
                    while position < len(buffer) and buffer[position].isspace():
                        position += 1
                    if not started:
                        if position >= len(buffer):
                            break
                        if buffer[position] != "[":
                            raise ValueError("Formato de backup invalido.")
                        started = True
                        position += 1
                        continue
                    while position < len(buffer) and (buffer[position].isspace() or buffer[position] == ","):
                        position += 1
                    if position < len(buffer) and buffer[position] == "]":
                        return
                    if position >= len(buffer):
                        break
                    try:
                        item, end = decoder.raw_decode(buffer, position)
                    except json.JSONDecodeError:
                        if eof:
                            raise ValueError("JSON do backup esta incompleto.")
                        break
                    position = end
                    if isinstance(item, dict):
                        yield item
                if eof:
                    raise ValueError("Fim inesperado do backup.")


def chat_backup_message_is_relevant(message: dict, user_id: int, current_group_ids: set[int]) -> bool:
    if int(message.get("sender_id") or 0) == user_id or int(message.get("receiver_id") or 0) == user_id:
        return True
    group_id = int(message.get("group_id") or 0)
    if group_id:
        participant_ids = {int(item) for item in message.get("participant_user_ids", []) if str(item).isdigit()}
        return user_id in participant_ids or group_id in current_group_ids
    return message.get("receiver_id") in {None, ""}


def serialize_chat_backup(backup: ChatHistoryBackup, relevant_message_count: Optional[int] = None, error: Optional[str] = None) -> dict:
    return {
        "id": backup.id,
        "company_id": backup.company_id,
        "period_start": backup.period_start.isoformat() if backup.period_start else None,
        "period_end": backup.period_end.isoformat() if backup.period_end else None,
        "message_count": backup.message_count,
        "relevant_message_count": relevant_message_count,
        "original_size": backup.original_size,
        "compressed_size": backup.compressed_size,
        "format_version": backup.format_version or 1,
        "checksum_sha256": backup.checksum_sha256,
        "integrity": "error" if error else ("verified" if backup.checksum_sha256 else "legacy"),
        "error": error,
        "created_at": backup.created_at.isoformat() if backup.created_at else None,
    }


@app.get("/platform/chat-backups")
async def platform_list_chat_backups(
    company_id: str = Query(...),
    user_id: int = Query(...),
    current_user: User = Depends(get_current_user),
):
    ensure_admin(current_user)
    with SessionLocal() as db:
        company = db.query(Company).filter(Company.id == company_id).first()
        if not company:
            raise HTTPException(status_code=404, detail="Empresa nao encontrada.")
        target_user = ensure_user_in_company(db, user_id, company_id)
        current_group_ids = {
            row[0]
            for row in db.query(ChatGroupMember.group_id).join(ChatGroup, ChatGroup.id == ChatGroupMember.group_id).filter(
                ChatGroupMember.user_id == target_user.id,
                ChatGroup.company_id == company_id,
            ).all()
        }
        backups = db.query(ChatHistoryBackup).filter(
            ChatHistoryBackup.company_id == company_id
        ).order_by(ChatHistoryBackup.period_end.desc()).all()
        result = []
        for backup in backups:
            try:
                relevant_count = sum(
                    1 for message in iter_chat_backup_messages(backup)
                    if chat_backup_message_is_relevant(message, target_user.id, current_group_ids)
                )
                result.append(serialize_chat_backup(backup, relevant_count))
            except (ValueError, zipfile.BadZipFile, OSError) as exc:
                result.append(serialize_chat_backup(backup, error=str(exc)))
        return {
            "company": serialize_company(company),
            "user": serialize_user(target_user, db),
            "backups": result,
        }


@app.post("/platform/chat-backups/{backup_id}/restore")
async def platform_restore_chat_backup(
    backup_id: int,
    payload: dict,
    current_user: User = Depends(get_current_user),
):
    ensure_admin(current_user)
    company_id = str(payload.get("company_id") or "").strip()
    user_id = int(payload.get("user_id") or 0)
    if not company_id or not user_id:
        raise HTTPException(status_code=400, detail="Empresa e usuario sao obrigatorios.")
    with SessionLocal() as db:
        target_user = ensure_user_in_company(db, user_id, company_id)
        backup = db.query(ChatHistoryBackup).filter(
            ChatHistoryBackup.id == backup_id,
            ChatHistoryBackup.company_id == company_id,
        ).first()
        if not backup:
            raise HTTPException(status_code=404, detail="Backup nao encontrado para esta empresa.")
        current_group_ids = {
            row[0]
            for row in db.query(ChatGroupMember.group_id).join(ChatGroup, ChatGroup.id == ChatGroupMember.group_id).filter(
                ChatGroupMember.user_id == target_user.id,
                ChatGroup.company_id == company_id,
            ).all()
        }
        valid_user_ids = {row[0] for row in db.query(User.id).all()}
        active_group_ids = {
            row[0] for row in db.query(ChatGroup.id).filter(
                ChatGroup.company_id == company_id,
                ChatGroup.is_active == True,
            ).all()
        }
        existing_uids = {
            row[0] for row in db.query(Message.archive_uid).filter(
                Message.company_id == company_id,
                Message.archive_uid.isnot(None),
            ).all()
        }
        restored = 0
        skipped = 0
        restored_id_map = {}
        try:
            records = iter_chat_backup_messages(backup)
            for record in records:
                if not chat_backup_message_is_relevant(record, target_user.id, current_group_ids):
                    continue
                original_id = int(record.get("id") or 0)
                timestamp_text = str(record.get("timestamp") or "")
                archive_uid = str(record.get("archive_uid") or "").strip() or hashlib.sha256(
                    f"{company_id}:{original_id}:{timestamp_text}".encode("utf-8")
                ).hexdigest()
                if archive_uid in existing_uids:
                    skipped += 1
                    continue
                sender_id = int(record.get("sender_id") or 0)
                receiver_id = int(record.get("receiver_id") or 0) or None
                group_id = int(record.get("group_id") or 0) or None
                if sender_id not in valid_user_ids or (receiver_id and receiver_id not in valid_user_ids):
                    skipped += 1
                    continue
                if group_id and group_id not in active_group_ids:
                    skipped += 1
                    continue
                try:
                    timestamp = datetime.fromisoformat(timestamp_text.replace("Z", "+00:00"))
                except (TypeError, ValueError):
                    timestamp = datetime.now(timezone.utc)
                original_reply_id = int(record.get("reply_to_id") or 0)
                restored_message = Message(
                    company_id=company_id,
                    content=str(record.get("content") or ""),
                    sender_id=sender_id,
                    receiver_id=receiver_id,
                    group_id=group_id,
                    reply_to_id=restored_id_map.get(original_reply_id),
                    message_type=str(record.get("message_type") or "text")[:40],
                    file_path=None,
                    timestamp=timestamp,
                    is_read=bool(record.get("is_read")),
                    archive_uid=archive_uid,
                )
                db.add(restored_message)
                db.flush()
                restored_id_map[original_id] = restored_message.id
                existing_uids.add(archive_uid)
                restored += 1
        except (ValueError, zipfile.BadZipFile, OSError) as exc:
            db.rollback()
            raise HTTPException(status_code=422, detail=f"Backup invalido: {exc}") from exc
        target_user_id = target_user.id
        log_audit(db, current_user.id, "historico_chat_restaurado", "chat_history_backup", str(backup.id), company_id, {
            "target_user_id": target_user_id,
            "restored": restored,
            "skipped": skipped,
        })
        db.commit()
    await manager.send_personal_message(json.dumps({
        "type": "history_restored",
        "restored": restored,
        "backup_id": backup_id,
    }), target_user_id)
    return {"restored": restored, "skipped": skipped, "backup_id": backup_id, "user_id": target_user_id}


@app.get("/platform/overview")
async def platform_overview(current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    with SessionLocal() as db:
        companies = db.query(Company).order_by(Company.name.asc()).all()
        users = db.query(User).order_by(User.full_name.asc()).all()
        departments = db.query(Department).order_by(Department.name.asc()).all()
        department_links = db.query(DepartmentLink).order_by(DepartmentLink.created_at.desc()).all()
        links = db.query(CompanyUser).order_by(CompanyUser.created_at.desc()).all()
        logs = db.query(AuditLog).order_by(AuditLog.created_at.desc()).limit(80).all()
        return {
            "stats": {
                "companies": len(companies),
                "active_companies": len([company for company in companies if company.status == ACTIVE_STATUS]),
                "users": len(users),
                "active_users": len([user for user in users if user.is_active]),
                "links": len([link for link in links if link.status == ACTIVE_STATUS]),
                "departments": len(departments),
                "tickets": db.query(Ticket).count(),
                "messages": db.query(Message).count(),
            },
            "companies": [serialize_company(company) for company in companies],
            "users": [serialize_user(user, db) for user in users],
            "departments": [serialize_department(department) for department in departments],
            "department_links": [serialize_department_link(link, db) for link in department_links],
            "company_users": [serialize_company_user(link, db) for link in links],
            "audit_logs": [serialize_audit_log(log, db) for log in logs],
            "mind_map": build_platform_mind_map(db),
        }


@app.get("/platform/companies")
async def platform_list_companies(current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    with SessionLocal() as db:
        return [serialize_company(company) for company in db.query(Company).order_by(Company.name.asc()).all()]


@app.post("/platform/companies")
async def platform_create_company(payload: dict, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    name = str(payload.get("name") or payload.get("nome_empresa") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="Nome da empresa e obrigatorio.")
    with SessionLocal() as db:
        cnpj = str(payload.get("cnpj") or "").strip() or None
        if cnpj and db.query(Company).filter(Company.cnpj == cnpj).first():
            raise HTTPException(status_code=400, detail="CNPJ ja cadastrado.")
        company_id = str(uuid.uuid4())
        company = Company(
            id=company_id,
            tenant_global_id=company_id,
            name=name,
            cnpj=cnpj,
            responsible_name=str(payload.get("responsible_name") or payload.get("responsavel") or "").strip() or None,
            phone_primary=str(payload.get("phone_primary") or payload.get("telefone_1") or "").strip() or None,
            phone_secondary=str(payload.get("phone_secondary") or payload.get("telefone_2") or "").strip() or None,
            status=normalize_status(payload.get("status")),
        )
        db.add(company)
        db.flush()
        for department_name in DEFAULT_DEPARTMENTS:
            get_or_create_department(db, company.id, department_name)
        log_audit(db, current_user.id, "empresa_criada", "company", company.id, company.id, {"name": company.name})
        db.commit()
        db.refresh(company)
        return serialize_company(company)


@app.patch("/platform/companies/{company_id}")
async def platform_update_company(company_id: str, payload: dict, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    with SessionLocal() as db:
        company = db.query(Company).filter(Company.id == company_id).first()
        if not company:
            raise HTTPException(status_code=404, detail="Empresa nao encontrada.")
        if "name" in payload and not str(payload.get("name") or "").strip():
            raise HTTPException(status_code=400, detail="Nome da empresa e obrigatorio.")
        requested_cnpj = str(payload.get("cnpj") or "").strip() or None
        if "cnpj" in payload and requested_cnpj:
            duplicate = db.query(Company).filter(Company.cnpj == requested_cnpj, Company.id != company.id).first()
            if duplicate:
                raise HTTPException(status_code=400, detail="CNPJ ja cadastrado.")
        for field in ["name", "cnpj", "responsible_name", "phone_primary", "phone_secondary"]:
            if field in payload:
                setattr(company, field, str(payload.get(field) or "").strip() or None)
        if "status" in payload:
            company.status = normalize_status(payload.get("status"))
        company.updated_at = datetime.utcnow()
        log_audit(db, current_user.id, "empresa_editada", "company", company.id, company.id, {"name": company.name})
        db.commit()
        db.refresh(company)
        return serialize_company(company)


@app.delete("/platform/companies/{company_id}/history")
async def platform_clear_company_history(company_id: str, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)

    with SessionLocal() as db:
        company = db.query(Company).filter(Company.id == company_id).first()
        if not company:
            raise HTTPException(status_code=404, detail="Empresa nao encontrada.")
        cleared_at = datetime.utcnow()

        ticket_ids = [item[0] for item in db.query(Ticket.id).filter(Ticket.company_id == company_id).all()]
        ticket_message_ids = [
            item[0] for item in db.query(TicketMessage.id).filter(TicketMessage.company_id == company_id).all()
        ]
        message_ids = [item[0] for item in db.query(Message.id).filter(Message.company_id == company_id).all()]

        counts = {}
        counts["ticket_message_reactions"] = (
            db.query(TicketMessageReaction)
            .filter(TicketMessageReaction.ticket_message_id.in_(ticket_message_ids))
            .delete(synchronize_session=False)
            if ticket_message_ids else 0
        )
        counts["ticket_messages"] = db.query(TicketMessage).filter(
            TicketMessage.company_id == company_id
        ).delete(synchronize_session=False)
        counts["ticket_assignees"] = (
            db.query(TicketAssignee).filter(TicketAssignee.ticket_id.in_(ticket_ids)).delete(synchronize_session=False)
            if ticket_ids else 0
        )
        counts["ticket_departments"] = (
            db.query(TicketDepartment).filter(TicketDepartment.ticket_id.in_(ticket_ids)).delete(synchronize_session=False)
            if ticket_ids else 0
        )
        counts["tickets"] = db.query(Ticket).filter(Ticket.company_id == company_id).delete(synchronize_session=False)

        counts["message_reactions"] = (
            db.query(MessageReaction).filter(MessageReaction.message_id.in_(message_ids)).delete(synchronize_session=False)
            if message_ids else 0
        )
        counts["message_read_receipts"] = (
            db.query(MessageReadReceipt).filter(MessageReadReceipt.message_id.in_(message_ids)).delete(synchronize_session=False)
            if message_ids else 0
        )
        counts["daily_summaries"] = db.query(BoltDailyTicketSummary).filter(
            BoltDailyTicketSummary.company_id == company_id
        ).delete(synchronize_session=False)
        counts["messages"] = db.query(Message).filter(Message.company_id == company_id).delete(synchronize_session=False)
        counts["tasks"] = db.query(TaskItem).filter(TaskItem.company_id == company_id).delete(synchronize_session=False)

        counts["meeting_participants"] = int(db.execute(text("""
            SELECT COUNT(*)
            FROM meeting_participants
            WHERE meeting_id IN (SELECT id FROM meetings WHERE company_id = :company_id)
        """), {"company_id": company_id}).scalar() or 0)
        counts["meetings"] = int(db.execute(text(
            "SELECT COUNT(*) FROM meetings WHERE company_id = :company_id"
        ), {"company_id": company_id}).scalar() or 0)
        db.execute(text("""
            DELETE FROM meeting_participants
            WHERE meeting_id IN (SELECT id FROM meetings WHERE company_id = :company_id)
        """), {"company_id": company_id})
        db.execute(text("DELETE FROM meetings WHERE company_id = :company_id"), {"company_id": company_id})

        counts["files"] = db.query(FileUpload).filter(FileUpload.company_id == company_id).delete(synchronize_session=False)
        counts["backups"] = db.query(ChatHistoryBackup).filter(
            ChatHistoryBackup.company_id == company_id
        ).delete(synchronize_session=False)
        counts["attention_limits"] = db.query(ChatAttentionLimit).filter(
            ChatAttentionLimit.company_id == company_id
        ).delete(synchronize_session=False)

        deleted_total = sum(counts.values())
        company.history_cleared_at = cleared_at
        log_audit(db, current_user.id, "historico_empresa_apagado", "company_history", company_id, company_id, {
            "company_name": company.name,
            "deleted": counts,
            "deleted_total": deleted_total,
        })
        db.commit()
        company_name = company.name

    await manager.broadcast(json.dumps({
        "type": "company_history_cleared",
        "company_id": company_id,
        "cleared_at": ensure_utc_datetime(cleared_at).isoformat(),
    }), company_id=company_id)
    return {
        "message": "Historico operacional da empresa apagado permanentemente.",
        "company_id": company_id,
        "company_name": company_name,
        "cleared_at": ensure_utc_datetime(cleared_at).isoformat(),
        "deleted": counts,
        "deleted_total": deleted_total,
    }

@app.delete("/platform/companies/{company_id}")
async def platform_delete_company(company_id: str, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    if company_id == DEFAULT_COMPANY_ID:
        raise HTTPException(status_code=400, detail="A empresa padrao do sistema nao pode ser excluida.")

    with SessionLocal() as db:
        company = db.query(Company).filter(Company.id == company_id).first()
        if not company:
            raise HTTPException(status_code=404, detail="Empresa nao encontrada.")

        tenant_global_id = company.tenant_global_id or company.id
        company_user_ids = [
            item.user_id for item in db.query(CompanyUser).filter(CompanyUser.company_id == company_id).all()
        ]
        ticket_ids = [
            item[0] for item in db.query(Ticket.id).filter(Ticket.company_id == company_id).all()
        ]

        counts = {
            "ticket_messages": db.query(TicketMessage).filter(TicketMessage.company_id == company_id).delete(synchronize_session=False),
            "messages": db.query(Message).filter(Message.company_id == company_id).delete(synchronize_session=False),
            "tasks": db.query(TaskItem).filter(TaskItem.company_id == company_id).delete(synchronize_session=False),
            "chat_groups": db.query(ChatGroup).filter(ChatGroup.company_id == company_id).delete(synchronize_session=False),
            "backups": db.query(ChatHistoryBackup).filter(ChatHistoryBackup.company_id == company_id).delete(synchronize_session=False),
        }
        if ticket_ids:
            counts["ticket_assignees"] = db.query(TicketAssignee).filter(TicketAssignee.ticket_id.in_(ticket_ids)).delete(synchronize_session=False)
            counts["ticket_departments"] = db.query(TicketDepartment).filter(TicketDepartment.ticket_id.in_(ticket_ids)).delete(synchronize_session=False)
        else:
            counts["ticket_assignees"] = 0
            counts["ticket_departments"] = 0
        counts["tickets"] = db.query(Ticket).filter(Ticket.company_id == company_id).delete(synchronize_session=False)
        counts["files"] = db.query(FileUpload).filter(FileUpload.company_id == company_id).delete(synchronize_session=False)
        counts["department_links"] = db.query(DepartmentLink).filter(DepartmentLink.company_id == company_id).delete(synchronize_session=False)
        counts["company_users"] = db.query(CompanyUser).filter(CompanyUser.company_id == company_id).delete(synchronize_session=False)
        counts["audit_logs"] = db.query(AuditLog).filter(AuditLog.company_id == company_id).delete(synchronize_session=False)
        counts["departments"] = db.query(Department).filter(Department.company_id == company_id).delete(synchronize_session=False)
        db.delete(company)
        db.flush()

        deleted_users = 0
        for user_id in company_user_ids:
            user = db.query(User).filter(User.id == user_id).first()
            if not user or user.id == current_user.id or is_admin(user) or user.username in {BOLT_USERNAME, MITTY_USERNAME}:
                continue
            has_membership = db.query(CompanyUser.id).filter(CompanyUser.user_id == user_id).first()
            has_records = any([
                db.query(Message.id).filter(or_(Message.sender_id == user_id, Message.receiver_id == user_id)).first(),
                db.query(Ticket.id).filter(or_(Ticket.created_by_id == user_id, Ticket.assigned_to_id == user_id)).first(),
                db.query(TicketAssignee.id).filter(TicketAssignee.user_id == user_id).first(),
                db.query(TicketMessage.id).filter(TicketMessage.sender_id == user_id).first(),
                db.query(TaskItem.id).filter(or_(TaskItem.created_by_id == user_id, TaskItem.assigned_to_id == user_id)).first(),
                db.query(ChatGroup.id).filter(ChatGroup.created_by_id == user_id).first(),
                db.query(FileUpload.id).filter(FileUpload.uploaded_by == user_id).first(),
                db.query(DepartmentLink.id).filter(DepartmentLink.created_by_id == user_id).first(),
            ])
            if not has_membership and not has_records:
                db.query(AuditLog).filter(AuditLog.actor_user_id == user_id).update(
                    {AuditLog.actor_user_id: None}, synchronize_session=False
                )
                db.delete(user)
                deleted_users += 1

        db.commit()
        counts["users"] = deleted_users
        return {
            "message": "Empresa e dados do tenant excluidos permanentemente.",
            "company_id": company_id,
            "tenant_global_id": tenant_global_id,
            "deleted": counts,
        }


@app.get("/platform/departments")
async def platform_list_departments(company_id: Optional[str] = Query(default=None), current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    with SessionLocal() as db:
        query = db.query(Department)
        if company_id:
            query = query.filter(Department.company_id == company_id)
        return [serialize_department(item) for item in query.order_by(Department.name.asc()).all()]


@app.post("/platform/departments")
async def platform_create_department(payload: dict, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    company_id = str(payload.get("company_id") or "").strip()
    name = str(payload.get("name") or "").strip()
    if not company_id or not name:
        raise HTTPException(status_code=400, detail="Empresa e nome do setor sao obrigatorios.")
    with SessionLocal() as db:
        ensure_company_access(db, current_user, company_id)
        department = get_or_create_department(db, company_id, name, str(payload.get("description") or "").strip() or None)
        log_audit(db, current_user.id, "setor_criado", "department", department.id, company_id, {"name": department.name})
        db.commit()
        db.refresh(department)
        return serialize_department(department)


@app.get("/platform/department-links")
async def platform_list_department_links(current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    with SessionLocal() as db:
        links = db.query(DepartmentLink).order_by(DepartmentLink.created_at.desc()).all()
        return [serialize_department_link(link, db) for link in links]


@app.post("/platform/department-links")
async def platform_create_department_link(payload: dict, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    source_id = str(payload.get("source_department_id") or "").strip()
    target_id = str(payload.get("target_department_id") or "").strip()
    if not source_id or not target_id or source_id == target_id:
        raise HTTPException(status_code=400, detail="Selecione dois setores diferentes para criar o vinculo.")

    with SessionLocal() as db:
        source = db.query(Department).filter(Department.id == source_id).first()
        target = db.query(Department).filter(Department.id == target_id).first()
        if not source or not target:
            raise HTTPException(status_code=404, detail="Setor nao encontrado.")
        if source.company_id != target.company_id:
            raise HTTPException(status_code=400, detail="Os setores vinculados devem pertencer a mesma empresa.")
        existing = (
            db.query(DepartmentLink)
            .filter(
                DepartmentLink.company_id == source.company_id,
                or_(
                    (DepartmentLink.source_department_id == source_id) & (DepartmentLink.target_department_id == target_id),
                    (DepartmentLink.source_department_id == target_id) & (DepartmentLink.target_department_id == source_id),
                ),
            )
            .first()
        )
        if existing:
            raise HTTPException(status_code=400, detail="Estes setores ja estao vinculados.")

        link = DepartmentLink(
            id=str(uuid.uuid4()),
            company_id=source.company_id,
            source_department_id=source_id,
            target_department_id=target_id,
            label=str(payload.get("label") or "").strip() or None,
            created_by_id=current_user.id,
        )
        db.add(link)
        log_audit(db, current_user.id, "setores_vinculados", "department_link", link.id, source.company_id, {
            "source_department_id": source_id,
            "target_department_id": target_id,
        })
        db.commit()
        db.refresh(link)
        return serialize_department_link(link, db)


@app.delete("/platform/department-links/{link_id}")
async def platform_remove_department_link(link_id: str, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    with SessionLocal() as db:
        link = db.query(DepartmentLink).filter(DepartmentLink.id == link_id).first()
        if not link:
            raise HTTPException(status_code=404, detail="Vinculo entre setores nao encontrado.")
        log_audit(db, current_user.id, "vinculo_setores_removido", "department_link", link.id, link.company_id, {})
        db.delete(link)
        db.commit()
        return {"message": "Vinculo entre setores removido."}


@app.post("/platform/users")
async def platform_create_user(payload: dict, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    email = str(payload.get("email") or "").strip().lower()
    password = str(payload.get("password") or payload.get("senha_primaria") or "").strip()
    company_id = str(payload.get("company_id") or "").strip()
    if not EMAIL_PATTERN.match(email) or not password or not company_id:
        raise HTTPException(status_code=400, detail="Email, senha primaria e empresa sao obrigatorios.")
    role = normalize_platform_role(payload.get("role") or payload.get("nivel_usuario") or "user")
    if role == "master_admin":
        raise HTTPException(status_code=400, detail="Use is_platform_admin para criar Admin Master global.")
    with SessionLocal() as db:
        ensure_company_access(db, current_user, company_id)
        company = db.query(Company).filter(Company.id == company_id).first()
        if not company:
            raise HTTPException(status_code=404, detail="Empresa nao encontrada.")
        department = get_or_create_department(db, company_id, str(payload.get("department") or payload.get("setor") or "Geral"))
        full_name = str(payload.get("name") or payload.get("full_name") or payload.get("nome_usuario") or email).strip()
        user = db.query(User).filter(User.email == email).first()
        if not user:
            username = generate_unique_username(db, email, full_name, payload.get("username"))
            user = User(
                uuid=str(uuid.uuid4()),
                username=username,
                email=email,
                full_name=full_name,
                hashed_password=get_password_hash(password),
                phone=str(payload.get("phone") or payload.get("telefone") or "").strip() or None,
                phone_extension=str(payload.get("phone") or payload.get("telefone") or "").strip() or None,
                must_change_password=True,
                status=normalize_status(payload.get("status")),
                is_active=normalize_status(payload.get("status")) == ACTIVE_STATUS,
                access_level=legacy_access_for_role(role),
                department=department.name,
            )
            db.add(user)
            db.flush()
            log_audit(db, current_user.id, "usuario_criado", "user", user.id, company_id, {"email": user.email})
        link = (
            db.query(CompanyUser)
            .filter(CompanyUser.company_id == company_id, CompanyUser.user_id == user.id)
            .first()
        )
        if not link:
            link = CompanyUser(
                id=str(uuid.uuid4()),
                company_id=company_id,
                user_id=user.id,
                department_id=department.id,
                role=role,
                status=normalize_status(payload.get("status")),
            )
            db.add(link)
            log_audit(db, current_user.id, "vinculo_criado", "company_user", user.id, company_id, {"role": role})
        if payload_flag(payload.get("create_hub_login")) or payload_flag(payload.get("hub_enabled")):
            sync_hub_identity(company, user, role)
            log_audit(
                db, current_user.id, "login_hub_criado", "user", user.id, company_id,
                {
                    "email": user.email,
                    "tenant_global_id": company.tenant_global_id,
                    "role": role,
                },
            )
        db.commit()
        db.refresh(user)
        return serialize_user(user, db)


@app.patch("/platform/company-users/{link_id}")
async def platform_update_company_user(link_id: str, payload: dict, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    with SessionLocal() as db:
        link = db.query(CompanyUser).filter(CompanyUser.id == link_id).first()
        if not link:
            raise HTTPException(status_code=404, detail="Vinculo nao encontrado.")
        if "department_id" in payload:
            department = db.query(Department).filter(Department.id == str(payload.get("department_id")), Department.company_id == link.company_id).first()
            if not department:
                raise HTTPException(status_code=400, detail="Setor nao encontrado.")
            link.department_id = department.id
            log_audit(db, current_user.id, "setor_alterado", "company_user", link.id, link.company_id, {"department_id": department.id})
        if "role" in payload:
            link.role = normalize_platform_role(payload.get("role"))
            log_audit(db, current_user.id, "nivel_alterado", "company_user", link.id, link.company_id, {"role": link.role})
        if "status" in payload:
            link.status = normalize_status(payload.get("status"))
        db.commit()
        db.refresh(link)
        return serialize_company_user(link, db)


@app.delete("/platform/company-users/{link_id}")
async def platform_remove_company_user(link_id: str, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    with SessionLocal() as db:
        link = db.query(CompanyUser).filter(CompanyUser.id == link_id).first()
        if not link:
            raise HTTPException(status_code=404, detail="Vinculo nao encontrado.")
        link.status = INACTIVE_STATUS
        log_audit(db, current_user.id, "vinculo_removido", "company_user", link.id, link.company_id, {"user_id": link.user_id})
        db.commit()
        return {"message": "Vinculo removido."}


@app.post("/platform/users/{user_id}/reset-password")
async def platform_reset_user_password(user_id: int, payload: dict, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    new_password = str(payload.get("password") or payload.get("senha_primaria") or "").strip()
    if len(new_password) < 6:
        raise HTTPException(status_code=400, detail="Senha primaria deve ter pelo menos 6 caracteres.")
    with SessionLocal() as db:
        user = db.query(User).filter(User.id == user_id).first()
        if not user:
            raise HTTPException(status_code=404, detail="Usuario nao encontrado.")
        user.hashed_password = get_password_hash(new_password)
        user.auth_version = int(getattr(user, "auth_version", 0) or 0) + 1
        user.must_change_password = True
        user.updated_at = datetime.utcnow()
        company_id = payload.get("company_id") or (active_company_memberships(db, user)[0].company_id if active_company_memberships(db, user) else None)
        log_audit(db, current_user.id, "senha_resetada", "user", user.id, company_id, {"email": user.email})
        db.commit()
        return {"message": "Senha resetada.", "user": serialize_user(user, db)}


@app.post("/platform/import/companies/preview")
async def platform_import_companies_preview(file: UploadFile = File(...), current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    rows = await read_import_rows(file)
    return preview_company_import_rows(rows)


@app.post("/platform/import/companies/confirm")
async def platform_import_companies_confirm(payload: dict, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    preview = preview_company_import_rows(payload.get("rows") or [])
    if preview["error_count"]:
        return {**preview, "imported": 0, "message": "Corrija os erros antes de importar."}
    imported = 0
    with SessionLocal() as db:
        for row in preview["rows"]:
            data = row["data"]
            company = db.query(Company).filter(Company.cnpj == data["cnpj"]).first() if data["cnpj"] else None
            if not company:
                company_id = str(uuid.uuid4())
                company = Company(id=company_id, tenant_global_id=company_id, cnpj=data["cnpj"] or None)
                db.add(company)
                imported += 1
            company.name = data["nome_empresa"]
            company.responsible_name = data["responsavel"] or None
            company.phone_primary = data["telefone_1"] or None
            company.phone_secondary = data["telefone_2"] or None
            company.status = data["status"]
            company.updated_at = datetime.utcnow()
            for department_name in DEFAULT_DEPARTMENTS:
                get_or_create_department(db, company.id, department_name)
            log_audit(db, current_user.id, "empresa_importada", "company", company.id, company.id, {"name": company.name})
        db.commit()
    return {**preview, "imported": imported}


@app.post("/platform/import/users/preview")
async def platform_import_users_preview(file: UploadFile = File(...), current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    rows = await read_import_rows(file)
    with SessionLocal() as db:
        return preview_user_import_rows(rows, db)


@app.post("/platform/import/users/confirm")
async def platform_import_users_confirm(payload: dict, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    with SessionLocal() as db:
        preview = preview_user_import_rows(payload.get("rows") or [], db)
        if preview["error_count"]:
            return {**preview, "imported": 0, "linked": 0, "message": "Corrija os erros antes de importar."}
        imported = 0
        linked = 0
        for row in preview["rows"]:
            data = row["data"]
            company_id = data["id_empresa"]
            department = get_or_create_department(db, company_id, data["setor"])
            user = db.query(User).filter(User.email == data["email"]).first()
            if not user:
                username = generate_unique_username(db, data["email"], data["nome_usuario"])
                user = User(
                    uuid=str(uuid.uuid4()),
                    username=username,
                    email=data["email"],
                    full_name=data["nome_usuario"],
                    hashed_password=get_password_hash(data["senha_primaria"]),
                    phone=data["telefone"] or None,
                    phone_extension=data["telefone"] or None,
                    must_change_password=True,
                    status=data["status"],
                    is_active=data["status"] == ACTIVE_STATUS,
                    access_level=legacy_access_for_role(data["nivel_usuario"]),
                    department=department.name,
                )
                db.add(user)
                db.flush()
                imported += 1
                log_audit(db, current_user.id, "usuario_importado", "user", user.id, company_id, {"email": user.email})
            else:
                user.must_change_password = True
                user.updated_at = datetime.utcnow()
            link = (
                db.query(CompanyUser)
                .filter(CompanyUser.company_id == company_id, CompanyUser.user_id == user.id)
                .first()
            )
            if not link:
                db.add(
                    CompanyUser(
                        id=str(uuid.uuid4()),
                        company_id=company_id,
                        user_id=user.id,
                        department_id=department.id,
                        role=data["nivel_usuario"],
                        status=data["status"],
                    )
                )
                linked += 1
                log_audit(db, current_user.id, "vinculo_criado", "company_user", user.id, company_id, {"role": data["nivel_usuario"]})
        db.commit()
    return {**preview, "imported": imported, "linked": linked}


@app.get("/admin/overview")
async def admin_overview(current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    with SessionLocal() as db:
        users = db.query(User).order_by(User.full_name.asc()).all()
        tickets = db.query(Ticket).order_by(Ticket.created_at.desc()).limit(50).all()
        messages = db.query(Message).order_by(Message.timestamp.desc()).limit(50).all()
        departments = sorted(
            {item.department for item in users if item.department} | {department for department in DEFAULT_DEPARTMENTS}
        )

        return {
            "stats": {
                "users": len(users),
                "active_users": len([item for item in users if item.is_active]),
                "tickets": db.query(Ticket).count(),
                "open_tickets": db.query(Ticket).filter(Ticket.status != "Resolvido").count(),
                "messages": db.query(Message).count(),
                "departments": len(departments),
            },
            "departments": departments,
            "users": [serialize_user(item) for item in users],
            "tickets": [serialize_ticket(ticket, db) for ticket in tickets],
            "messages": serialize_messages(messages, db),
        }


@app.get("/coordinator/overview")
async def coordinator_overview(current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user)
        membership = get_company_membership(db, current_user, company_id)
        if not is_company_admin(db, current_user, company_id) and (not membership or membership.role != "coordinator"):
            raise HTTPException(status_code=403, detail="Acesso restrito a coordenadores.")
        department = user_department_name(db, current_user, company_id)
        users_query = (
            db.query(User)
            .join(CompanyUser, CompanyUser.user_id == User.id)
            .filter(CompanyUser.company_id == company_id, CompanyUser.status == ACTIVE_STATUS)
        )
        tickets_query = db.query(Ticket).filter(Ticket.company_id == company_id)
        # O conteudo dos chats nao e mais exposto como lista plana neste painel.
        # Administradores usam o Controle Interno, com selecao de usuario e auditoria.
        can_view_messages = False

        company_admin = is_company_admin(db, current_user, company_id)
        if not company_admin:
            membership = get_company_membership(db, current_user, company_id)
            if membership and membership.department_id:
                users_query = users_query.filter(CompanyUser.department_id == membership.department_id)
            explicit_ticket_ids = db.query(TicketAssignee.ticket_id).filter(TicketAssignee.user_id == current_user.id)
            tickets_query = tickets_query.filter(or_(
                Ticket.created_by_id == current_user.id,
                Ticket.assigned_to_id == current_user.id,
                Ticket.id.in_(explicit_ticket_ids),
            ))

        users = users_query.order_by(User.full_name.asc()).all()
        tickets = tickets_query.order_by(Ticket.created_at.desc()).limit(50).all()
        messages = []

        return {
            "department": department,
            "can_view_messages": can_view_messages,
            "stats": {
                "users": len(users),
                "tickets": len(tickets),
                "open_tickets": len([ticket for ticket in tickets if ticket.status != "Resolvido"]),
                "messages": len(messages),
            },
            "users": [serialize_user(item) for item in users],
            "tickets": [serialize_ticket(ticket, db) for ticket in tickets],
            "messages": serialize_messages(messages, db),
        }


@app.get("/chat/settings")
async def get_chat_settings(
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        company = db.query(Company).filter(Company.id == scoped_company_id).first()
        return {
            "company_id": scoped_company_id,
            "general_chat_enabled": False,
            "allow_user_message_editing": bool(company and company.allow_user_message_editing),
            "can_manage": is_company_admin(db, current_user, scoped_company_id),
        }


@app.put("/chat/settings")
async def update_chat_settings(payload: dict, current_user: User = Depends(get_current_user)):
    # Compatibilidade de rota para clientes antigos. O Chat Geral foi desativado
    # e nenhuma configuracao pode reativa-lo ou renomea-lo.
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        ensure_company_admin(db, current_user, company_id)
    raise HTTPException(status_code=410, detail="O Chat Geral foi desativado. Utilize apenas conversas diretas e grupos criados.")


@app.get("/groups/")
async def list_groups(company_id: Optional[str] = Query(default=None), current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        query = db.query(ChatGroup).filter(ChatGroup.company_id == scoped_company_id, ChatGroup.is_active == True)
        if not is_company_admin(db, current_user, scoped_company_id):
            member_group_ids = db.query(ChatGroupMember.group_id).filter(
                ChatGroupMember.user_id == current_user.id
            )
            query = query.filter(ChatGroup.id.in_(member_group_ids))
        groups = query.order_by(ChatGroup.name.asc()).all()
        payloads = serialize_groups(groups, db)
        group_ids = [group.id for group in groups]
        unread_by_group = {}
        if group_ids:
            unread_by_group = {
                group_id: count
                for group_id, count in (
                    db.query(Message.group_id, func.count(MessageMention.id))
                    .join(MessageMention, MessageMention.message_id == Message.id)
                    .filter(
                        MessageMention.user_id == current_user.id,
                        MessageMention.read_at.is_(None),
                        Message.company_id == scoped_company_id,
                        Message.group_id.in_(group_ids),
                    )
                    .group_by(Message.group_id)
                    .all()
                )
            }
        for item in payloads:
            item["unread_mention_count"] = unread_by_group.get(item["id"], 0)
        return payloads


@app.post("/groups/")
async def create_group(payload: dict, current_user: User = Depends(get_current_user)):
    name = str(payload.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="Nome do grupo e obrigatorio.")

    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        membership = get_company_membership(db, current_user, company_id)
        can_create_group = is_company_admin(db, current_user, company_id) or bool(
            membership and membership.role == "coordinator"
        )
        if not can_create_group:
            raise HTTPException(status_code=403, detail="Somente administradores e coordenadores podem criar grupos.")
        department = str(payload.get("department") or user_department_name(db, current_user, company_id) or "Geral").strip()
        if not is_company_admin(db, current_user, company_id) and not user_can_access_department(db, current_user, company_id, department):
            raise HTTPException(status_code=403, detail="Coordenador so pode criar grupos do proprio setor.")
        get_or_create_department(db, company_id, department)
        group = ChatGroup(
            company_id=company_id,
            name=name,
            description=str(payload.get("description") or "").strip() or None,
            department=department,
            image_data=str(payload.get("image_data") or "").strip() or None,
            created_by_id=current_user.id,
        )
        db.add(group)
        db.flush()
        member_ids = {int(item) for item in (payload.get("member_ids") or []) if str(item).isdigit()}
        member_ids.add(current_user.id)
        for member_id in member_ids:
            member = ensure_user_in_company(db, member_id, company_id)
            if not is_company_admin(db, current_user, company_id) and member.department != department:
                raise HTTPException(status_code=403, detail="Coordenador so pode incluir usuarios do proprio setor.")
            db.add(ChatGroupMember(group_id=group.id, user_id=member.id))
        db.commit()
        db.refresh(group)
        group_payload = serialize_group(group, db)
    event = json.dumps({"type": "group_upsert", "group": group_payload})
    for member_id in group_payload["member_ids"]:
        await manager.send_personal_message(event, member_id)
    return group_payload


@app.put("/groups/{group_id}")
async def update_group(group_id: int, payload: dict, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        group = db.query(ChatGroup).filter(ChatGroup.id == group_id, ChatGroup.company_id == company_id).first()
        if not group:
            raise HTTPException(status_code=404, detail="Grupo nao encontrado.")
        membership = get_company_membership(db, current_user, company_id)
        can_manage = is_company_admin(db, current_user, company_id) or bool(membership and membership.role == "coordinator")
        if not can_manage:
            raise HTTPException(status_code=403, detail="Somente administradores e coordenadores podem editar grupos.")
        if membership and membership.role == "coordinator" and not user_can_access_department(db, current_user, company_id, group.department):
            raise HTTPException(status_code=403, detail="Coordenador so pode editar grupos do proprio setor.")
        previous_member_ids = {
            row[0]
            for row in db.query(ChatGroupMember.user_id).filter(ChatGroupMember.group_id == group.id).all()
        }
        if "name" in payload:
            name = str(payload.get("name") or "").strip()
            if not name:
                raise HTTPException(status_code=400, detail="Nome do grupo e obrigatorio.")
            group.name = name[:120]
        if "description" in payload:
            group.description = str(payload.get("description") or "").strip() or None
        if "image_data" in payload:
            group.image_data = str(payload.get("image_data") or "").strip() or None
        if "member_ids" in payload:
            member_ids = {int(item) for item in payload.get("member_ids", []) if str(item).isdigit()}
            member_ids.add(current_user.id)
            members = [ensure_user_in_company(db, member_id, company_id) for member_id in member_ids]
            if membership and membership.role == "coordinator":
                for member in members:
                    if member.department != group.department:
                        raise HTTPException(status_code=403, detail="Coordenador so pode incluir usuarios do proprio setor.")
            db.query(ChatGroupMember).filter(ChatGroupMember.group_id == group.id).delete(synchronize_session=False)
            for member in members:
                db.add(ChatGroupMember(group_id=group.id, user_id=member.id))
        db.commit()
        db.refresh(group)
        group_payload = serialize_group(group, db)
    current_member_ids = set(group_payload["member_ids"])
    event = json.dumps({"type": "group_upsert", "group": group_payload})
    for member_id in current_member_ids:
        await manager.send_personal_message(event, member_id)
    removed_event = json.dumps({"type": "group_removed", "group_id": group_id})
    for member_id in previous_member_ids - current_member_ids:
        await manager.send_personal_message(removed_event, member_id)
    return group_payload


@app.delete("/groups/{group_id}")
async def delete_group(
    group_id: int,
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        group = db.query(ChatGroup).filter(
            ChatGroup.id == group_id,
            ChatGroup.company_id == scoped_company_id,
            ChatGroup.is_active == True,
        ).first()
        if not group:
            raise HTTPException(status_code=404, detail="Grupo nao encontrado.")
        membership = get_company_membership(db, current_user, scoped_company_id)
        can_manage = is_company_admin(db, current_user, scoped_company_id) or bool(
            membership and membership.role == "coordinator"
        )
        if not can_manage:
            raise HTTPException(status_code=403, detail="Somente administradores e coordenadores podem excluir grupos.")
        if membership and membership.role == "coordinator" and not user_can_access_department(
            db, current_user, scoped_company_id, group.department
        ):
            raise HTTPException(status_code=403, detail="Coordenador so pode excluir grupos do proprio setor.")
        member_ids = [
            row[0]
            for row in db.query(ChatGroupMember.user_id).filter(ChatGroupMember.group_id == group.id).all()
        ]
        group.is_active = False
        db.commit()
    event = json.dumps({"type": "group_removed", "group_id": group_id})
    for member_id in member_ids:
        await manager.send_personal_message(event, member_id)
    return {"message": "Grupo excluido com sucesso.", "group_id": group_id}


@app.post("/groups/{group_id}/read")
async def mark_group_messages_read(
    group_id: int,
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        group = ensure_chat_group_member(db, group_id, current_user.id, scoped_company_id)
        messages = db.query(Message).filter(
            Message.company_id == scoped_company_id,
            Message.group_id == group.id,
            Message.sender_id != current_user.id,
        ).all()
        existing_ids = {
            row[0] for row in db.query(MessageReadReceipt.message_id).filter(
                MessageReadReceipt.user_id == current_user.id,
                MessageReadReceipt.message_id.in_([message.id for message in messages] or [-1]),
            ).all()
        }
        message_ids = []
        for message in messages:
            if message.id not in existing_ids:
                db.add(MessageReadReceipt(message_id=message.id, user_id=current_user.id))
                message_ids.append(message.id)
        unread_mentions = (
            db.query(MessageMention)
            .join(Message, Message.id == MessageMention.message_id)
            .filter(
                MessageMention.user_id == current_user.id,
                MessageMention.read_at.is_(None),
                Message.company_id == scoped_company_id,
                Message.group_id == group.id,
            )
            .all()
        )
        for mention in unread_mentions:
            mention.read_at = datetime.utcnow()
        db.commit()
        member_ids = [row[0] for row in db.query(ChatGroupMember.user_id).filter(ChatGroupMember.group_id == group.id).all()]
        reader = db.query(User).filter(User.id == current_user.id).first()
        event = json.dumps({
            "type": "group_messages_read",
            "group_id": group.id,
            "message_ids": message_ids,
            "reader": {"id": reader.id, "name": display_user_name(reader), "full_name": reader.full_name, "profile_photo": reader.profile_photo},
        })
    if message_ids:
        for member_id in member_ids:
            await manager.send_personal_message(event, member_id)
    return {"message_ids": message_ids}

@app.delete("/messages/history/clear")
async def clear_chat_history(
    receiver_id: Optional[int] = Query(default=None),
    company_id: Optional[str] = Query(default=None),
    clear_archives: bool = Query(default=False),
    current_user: User = Depends(get_current_user),
):
    ensure_admin(current_user)
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        query = db.query(Message).filter(Message.company_id == scoped_company_id)
        if receiver_id:
            ensure_user_in_company(db, receiver_id, scoped_company_id)
            query = query.filter(
                Message.group_id.is_(None),
                or_(
                    (Message.sender_id == current_user.id) & (Message.receiver_id == receiver_id),
                    (Message.sender_id == receiver_id) & (Message.receiver_id == current_user.id),
                )
            )
        else:
            raise HTTPException(status_code=400, detail="Selecione uma conversa direta. O Chat Geral foi desativado.")
        deleted = query.delete(synchronize_session=False)
        if clear_archives and not receiver_id:
            db.query(ChatHistoryBackup).filter(ChatHistoryBackup.company_id == scoped_company_id).delete(synchronize_session=False)
        log_audit(db, current_user.id, "historico_chat_limpo", "message_history", None, scoped_company_id, {
            "receiver_id": receiver_id,
            "deleted": deleted,
            "clear_archives": clear_archives,
        })
        db.commit()
    await manager.broadcast(json.dumps({"type": "history_cleared", "receiver_id": receiver_id}), company_id=scoped_company_id)
    return {"deleted": deleted}


@app.post("/messages/read")
async def mark_messages_read(payload: dict, current_user: User = Depends(get_current_user)):
    sender_id = payload.get("sender_id")
    if not sender_id:
        raise HTTPException(status_code=400, detail="Remetente obrigatorio.")
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        sender = ensure_user_in_company(db, int(sender_id), company_id)
        sender_user_id = sender.id
        unread = db.query(Message).filter(
            Message.company_id == company_id,
            Message.sender_id == sender_user_id,
            Message.receiver_id == current_user.id,
            Message.is_read == False,
        ).all()
        ids = [message.id for message in unread]
        for message in unread:
            message.is_read = True
        db.commit()
    if ids:
        await manager.send_personal_message(json.dumps({"type": "message_read", "message_ids": ids}), sender_user_id)
    return {"message_ids": ids}

@app.patch("/messages/{message_id}")
async def edit_chat_message(
    message_id: int,
    payload: dict,
    current_user: User = Depends(get_current_user),
):
    content = str(payload.get("content") or "").strip()
    if not content:
        raise HTTPException(status_code=400, detail="A mensagem nao pode ficar vazia.")
    if len(content) > INLINE_CHAT_MESSAGE_MAX_LENGTH:
        raise HTTPException(
            status_code=400,
            detail=f"Edicoes sao limitadas a {INLINE_CHAT_MESSAGE_MAX_LENGTH} caracteres.",
        )

    with SessionLocal() as db:
        message = db.query(Message).filter(Message.id == message_id).first()
        if not message:
            raise HTTPException(status_code=404, detail="Mensagem nao encontrada.")
        company_id = ensure_company_access(db, current_user, message.company_id)
        company = db.query(Company).filter(Company.id == company_id).first()
        if not company or not company.allow_user_message_editing:
            raise HTTPException(status_code=403, detail="A edicao de mensagens esta desabilitada nesta empresa.")
        if message.sender_id != current_user.id:
            raise HTTPException(status_code=403, detail="Voce so pode editar mensagens enviadas por voce.")
        if message.message_type != "text":
            raise HTTPException(status_code=400, detail="Somente mensagens de texto podem ser editadas.")
        if content == str(message.content or "").strip():
            raise HTTPException(status_code=400, detail="A mensagem nao foi alterada.")
        message_timestamp = ensure_utc_datetime(message.timestamp) if message.timestamp else None
        if not message_timestamp or datetime.now(timezone.utc) - message_timestamp > timedelta(minutes=30):
            raise HTTPException(status_code=400, detail="O prazo de 30 minutos para editar esta mensagem terminou.")
        edit_count = db.query(func.count(MessageEdit.id)).filter(MessageEdit.message_id == message.id).scalar() or 0
        if edit_count >= 2:
            raise HTTPException(status_code=400, detail="Esta mensagem ja atingiu o limite de 2 edicoes.")

        group = None
        new_mention_ids = set()
        if message.group_id:
            group = ensure_chat_group_member(db, int(message.group_id), current_user.id, company_id)
            new_mention_ids = validated_group_mention_ids(
                db, group, payload.get("mention_user_ids"), company_id, current_user.id, content
            )
        db.add(MessageEdit(
            message_id=message.id,
            editor_user_id=current_user.id,
            previous_content=message.content,
        ))
        message.content = content
        added_mentions, removed_mentions = sync_message_mentions(db, message, new_mention_ids) if group else (set(), set())
        log_audit(db, current_user.id, "mensagem_chat_editada", "message", message.id, company_id, {
            "edit_number": int(edit_count) + 1,
        })
        db.commit()
        db.refresh(message)
        message_payload = serialize_message(message, db)
        receiver_id = message.receiver_id
        sender_id = message.sender_id
        group_id = message.group_id
        member_ids = []
        if group_id:
            member_ids = [
                row[0]
                for row in db.query(ChatGroupMember.user_id).filter(ChatGroupMember.group_id == group_id).all()
            ]

    event = json.dumps({
        "type": "message_updated",
        "message": message_payload,
        "newly_mentioned_user_ids": sorted(added_mentions),
        "removed_mention_user_ids": sorted(removed_mentions),
    })
    if group_id:
        for member_id in set(member_ids) | {sender_id}:
            await manager.send_personal_message(event, member_id)
    elif receiver_id:
        await manager.send_personal_message(event, sender_id)
        if receiver_id != sender_id:
            await manager.send_personal_message(event, receiver_id)
    else:
        await manager.broadcast(event, company_id=company_id)
    return message_payload


@app.post("/messages/{message_id}/reactions")
async def toggle_message_reaction(
    message_id: int,
    payload: dict,
    current_user: User = Depends(get_current_user),
):
    emoji = str(payload.get("emoji") or "").strip()
    if not emoji or len(emoji) > 32:
        raise HTTPException(status_code=400, detail="Informe um emoji valido.")

    with SessionLocal() as db:
        message = db.query(Message).filter(Message.id == message_id).first()
        if not message:
            raise HTTPException(status_code=404, detail="Mensagem nao encontrada.")
        company_id = ensure_company_access(db, current_user, message.company_id)
        if message.receiver_id and current_user.id not in {message.sender_id, message.receiver_id} and not is_admin(current_user):
            raise HTTPException(status_code=403, detail="Sem acesso a esta conversa.")

        existing = db.query(MessageReaction).filter(
            MessageReaction.message_id == message.id,
            MessageReaction.user_id == current_user.id,
            MessageReaction.emoji == emoji,
        ).first()
        if existing:
            db.delete(existing)
            action = "removed"
        else:
            db.add(MessageReaction(message_id=message.id, user_id=current_user.id, emoji=emoji))
            action = "added"
        db.commit()
        reactions = serialize_message_reactions(message.id, db)
        receiver_id = message.receiver_id
        sender_id = message.sender_id

    update = {
        "type": "reaction_update",
        "message_id": message_id,
        "reactions": reactions,
        "action": action,
        "user_name": current_user.full_name,
        "user_id": current_user.id,
        "emoji": emoji,
    }
    serialized = json.dumps(update)
    if receiver_id:
        await manager.send_personal_message(serialized, sender_id)
        if receiver_id != sender_id:
            await manager.send_personal_message(serialized, receiver_id)
    else:
        await manager.broadcast(serialized, company_id=company_id)
    return {"action": action, "reactions": reactions}


@app.get("/messages/history")
async def get_message_history_page(
    receiver_id: Optional[int] = Query(default=None),
    group_id: Optional[int] = Query(default=None),
    before_message_id: Optional[int] = Query(default=None, ge=1),
    limit: int = Query(default=50, ge=1, le=100),
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    if bool(receiver_id) == bool(group_id):
        raise HTTPException(status_code=400, detail="Selecione uma conversa direta ou um grupo.")
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        query = accessible_messages_query(db, current_user.id, scoped_company_id)
        if receiver_id:
            ensure_user_in_company(db, int(receiver_id), scoped_company_id)
            query = query.filter(
                Message.group_id.is_(None),
                or_(
                    and_(Message.sender_id == current_user.id, Message.receiver_id == int(receiver_id)),
                    and_(Message.sender_id == int(receiver_id), Message.receiver_id == current_user.id),
                ),
            )
        else:
            ensure_chat_group_member(db, int(group_id), current_user.id, scoped_company_id)
            query = query.filter(Message.group_id == int(group_id))
        return paginate_message_history(query, db, before_message_id=before_message_id, limit=limit)

@app.delete("/messages/{message_id}")
async def delete_chat_message(message_id: int, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    with SessionLocal() as db:
        message = db.query(Message).filter(Message.id == message_id).first()
        if not message:
            company_id = ensure_company_access(db, current_user)
            already_deleted = True
        else:
            company_id = message.company_id
            already_deleted = False
            db.delete(message)
            log_audit(db, current_user.id, "mensagem_chat_removida", "message", message_id, company_id, {})
            db.commit()
    await manager.broadcast(json.dumps({"type": "message_deleted", "message_id": message_id}), company_id=company_id)
    return {
        "message": "Mensagem já removida." if already_deleted else "Mensagem removida.",
        "already_deleted": already_deleted,
    }


@app.websocket("/messages/ws/{token}")
async def websocket_endpoint(
    websocket: WebSocket,
    token: str,
    client_id: Optional[str] = Query(default=None, max_length=80),
):
    user_data = None
    try:
        try:
            payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        except JWTError:
            # Complete the handshake before closing so an expired token is not
            # reported by ASGI as an unfinished handshake / HTTP 500.
            await websocket.accept()
            await websocket.close(code=1008, reason="Token expirado")
            return
        user_id = payload.get("sub")
        if not user_id:
            await websocket.accept()
            await websocket.close(code=1008, reason="Token sem usuario")
            return

        with SessionLocal() as db:
            try:
                user = db.query(User).filter(User.id == int(user_id)).first()
            except (TypeError, ValueError):
                user = None
            if not user:
                await websocket.accept()
                await websocket.close(code=1008, reason="Usuario da sessao nao encontrado")
                return
            if user.must_change_password:
                await websocket.accept()
                await websocket.close(code=1008, reason="Troca de senha obrigatoria")
                return
            try:
                company_id = ensure_company_access(db, user)
            except HTTPException as exc:
                await websocket.accept()
                await websocket.close(code=1008, reason=str(exc.detail)[:120])
                return
            user_data = {"id": user.id, "username": user.username, "full_name": user.full_name, "nickname": user.nickname, "company_id": company_id}

        await manager.connect(
            websocket,
            user_data["id"],
            user_data["company_id"],
            client_id=client_id,
        )
        await websocket.send_text(json.dumps({"type": "connection", "message": f"Conectado como {user_data['full_name']}!"}))

        with SessionLocal() as db:
            recent_messages = (
                accessible_messages_query(db, user_data["id"], user_data["company_id"])
                .order_by(Message.timestamp.desc())
                .limit(100)
                .all()
            )
            messages_data = serialize_messages(reversed(recent_messages), db)
            history_cleared_at = db.query(Company.history_cleared_at).filter(
                Company.id == user_data["company_id"]
            ).scalar()
        await websocket.send_text(json.dumps({
            "type": "message_history",
            "messages": messages_data,
            "company_id": user_data["company_id"],
            "history_cleared_at": ensure_utc_datetime(history_cleared_at).isoformat() if history_cleared_at else None,
        }))

        while True:
            data = await websocket.receive_text()
            try:
                message_data = json.loads(data)

                if message_data["type"] == "chat_message":
                    with SessionLocal() as db:
                        receiver_id = message_data.get("receiver_id")
                        group_id = message_data.get("group_id")
                        group = None
                        if receiver_id and group_id:
                            raise HTTPException(status_code=400, detail="Escolha conversa direta ou grupo.")
                        if not receiver_id and not group_id:
                            raise HTTPException(status_code=410, detail="O Chat Geral foi desativado. Selecione uma conversa direta ou grupo criado.")
                        if receiver_id:
                            receiver_id = int(receiver_id)
                            ensure_user_in_company(db, receiver_id, user_data["company_id"])
                        if group_id:
                            group_id = int(group_id)
                            group = ensure_chat_group_member(db, group_id, user_data["id"], user_data["company_id"])
                        reply_to_id = message_data.get("reply_to_id")
                        if reply_to_id:
                            reply_message = db.query(Message).filter(
                                Message.id == int(reply_to_id),
                                Message.company_id == user_data["company_id"],
                            ).first()
                            if not reply_message:
                                raise HTTPException(status_code=404, detail="Mensagem respondida nao encontrada.")
                        message_type = str(message_data.get("message_type") or "text")[:40]
                        file_path = message_data.get("file_path")
                        if message_type == "sticker":
                            reference = str(file_path or "")
                            if reference.startswith("default:"):
                                sticker_key = reference.split(":", 1)[1]
                                if sticker_key not in DEFAULT_STICKER_KEYS:
                                    await websocket.send_text(json.dumps({"type": "error", "detail": "Figurinha padrao invalida."}))
                                    continue
                            elif reference.startswith("sticker:"):
                                try:
                                    sticker_id = int(reference.split(":", 1)[1])
                                except (TypeError, ValueError) as exc:
                                    raise HTTPException(status_code=400, detail="Referencia de figurinha invalida.") from exc
                                sticker = db.query(ChatSticker).filter(
                                    ChatSticker.id == sticker_id,
                                    ChatSticker.company_id == user_data["company_id"],
                                ).first()
                                if not sticker or not sticker.is_active:
                                    raise HTTPException(status_code=404, detail="Figurinha nao encontrada nesta empresa.")
                                if sticker.owner_user_id != user_data["id"]:
                                    raise HTTPException(status_code=403, detail="Figurinha nao pertence ao usuario.")
                                sender = db.query(User).filter(User.id == user_data["id"]).first()
                                ensure_company_sticker_creation(db, sender, user_data["company_id"])
                            else:
                                raise HTTPException(status_code=400, detail="Referencia de figurinha obrigatoria.")
                        mention_user_ids = set()
                        if group:
                            mention_user_ids = validated_group_mention_ids(
                                db, group, message_data.get("mention_user_ids"), user_data["company_id"], user_data["id"], str(message_data.get("content") or "")
                            )
                        new_message = Message(
                            company_id=user_data["company_id"],
                            content=str(message_data.get("content") or "Figurinha"),
                            sender_id=user_data["id"],
                            receiver_id=receiver_id,
                            group_id=group_id,
                            reply_to_id=int(reply_to_id) if reply_to_id else None,
                            message_type=message_type,
                            file_path=file_path,
                        )
                        db.add(new_message)
                        db.flush()
                        if mention_user_ids:
                            sync_message_mentions(db, new_message, mention_user_ids)
                        db.commit()
                        db.refresh(new_message)
                        message_payload = serialize_message(new_message, db)
                        message_payload["client_id"] = message_data.get("client_id")
                        broadcast_message = {
                            "type": "new_message",
                            "message": message_payload,
                        }
    
                    if broadcast_message["message"].get("group_id"):
                        with SessionLocal() as db:
                            member_ids = {
                                row[0]
                                for row in db.query(ChatGroupMember.user_id).filter(
                                    ChatGroupMember.group_id == broadcast_message["message"]["group_id"]
                                ).all()
                            }
                        # O remetente precisa receber o payload canonico para substituir
                        # a mensagem otimista. Coordenadores e usuarios so recebem eventos
                        # dos grupos em que participam explicitamente.
                        member_ids.add(user_data["id"])
                        for member_id in member_ids:
                            await manager.send_personal_message(json.dumps(broadcast_message), member_id)
                    elif broadcast_message["message"]["receiver_id"]:
                        await manager.send_personal_message(json.dumps(broadcast_message), broadcast_message["message"]["receiver_id"])
                        await manager.send_personal_message(json.dumps(broadcast_message), user_data["id"])
                    else:
                        # Estado impossivel desde a desativacao do Chat Geral.
                        await websocket.send_text(json.dumps({"type": "error", "detail": "O Chat Geral foi desativado."}))
    
                elif message_data["type"] == "mark_read":
                    sender_id = int(message_data.get("sender_id") or 0)
                    if not sender_id:
                        raise HTTPException(status_code=400, detail="Remetente obrigatorio.")
                    with SessionLocal() as db:
                        sender = ensure_user_in_company(db, sender_id, user_data["company_id"])
                        sender_user_id = sender.id
                        unread = db.query(Message).filter(
                            Message.company_id == user_data["company_id"],
                            Message.sender_id == sender_user_id,
                            Message.receiver_id == user_data["id"],
                            Message.is_read == False,
                        ).all()
                        ids = [message.id for message in unread]
                        for message in unread:
                            message.is_read = True
                        db.commit()
                    if ids:
                        await manager.send_personal_message(json.dumps({"type": "message_read", "message_ids": ids}), sender_user_id)
    
                elif message_data["type"] == "switch_company":
                    with SessionLocal() as db:
                        active_user = db.query(User).filter(User.id == user_data["id"]).first()
                        if not active_user or not is_admin(active_user):
                            raise HTTPException(status_code=403, detail="Somente Admin Master pode alternar a empresa do chat.")
                        company_id = ensure_company_access(db, active_user, message_data.get("company_id"))
                        user_data["company_id"] = company_id
                        manager.connection_company_scope[id(websocket)] = company_id
                        recent_messages = (
                            accessible_messages_query(db, user_data["id"], company_id)
                            .order_by(Message.timestamp.desc())
                            .limit(100)
                            .all()
                        )
                        messages_data = serialize_messages(reversed(recent_messages), db)
                        history_cleared_at = db.query(Company.history_cleared_at).filter(
                            Company.id == company_id
                        ).scalar()
                    await websocket.send_text(json.dumps({
                        "type": "message_history",
                        "messages": messages_data,
                        "company_id": company_id,
                        "history_cleared_at": ensure_utc_datetime(history_cleared_at).isoformat() if history_cleared_at else None,
                    }))
                    await manager.send_online_users(websocket, company_id)
    
                elif message_data["type"] == "presence":
                    status = str(message_data.get("status") or "online").strip().lower()
                    if status not in {"online", "busy", "meeting"}:
                        status = "online"
                    manager.presence_status[user_data["id"]] = status
                    await manager.broadcast_user_status(user_data["id"], True, user_data["company_id"])
    
                elif message_data["type"] == "typing":
                    receiver_id = message_data.get("receiver_id")
                    group_id = message_data.get("group_id")
                    group_member_ids = []
                    if receiver_id and group_id:
                        raise HTTPException(status_code=400, detail="Indicador de digitacao deve pertencer a uma unica conversa.")
                    if receiver_id:
                        with SessionLocal() as db:
                            ensure_user_in_company(db, int(receiver_id), user_data["company_id"])
                        receiver_id = int(receiver_id)
                    if group_id:
                        group_id = int(group_id)
                        with SessionLocal() as db:
                            group = ensure_chat_group_member(db, group_id, user_data["id"], user_data["company_id"])
                            group_member_ids = [row[0] for row in db.query(ChatGroupMember.user_id).filter(ChatGroupMember.group_id == group.id).all()]
                    typing_message = {
                        "type": "typing",
                        "user_id": user_data["id"],
                        "username": user_data["username"],
                        "full_name": user_data["full_name"],
                        "nickname": user_data["nickname"],
                        "receiver_id": receiver_id,
                        "group_id": group_id,
                        "is_typing": bool(message_data.get("is_typing")),
                    }
                    if group_id:
                        for member_id in group_member_ids:
                            if member_id != user_data["id"]:
                                await manager.send_personal_message(json.dumps(typing_message), member_id)
                    elif receiver_id:
                        await manager.send_personal_message(json.dumps(typing_message), receiver_id)
                    else:
                        await manager.broadcast(json.dumps(typing_message), exclude_user=user_data["id"], company_id=user_data["company_id"])
    
                elif message_data["type"] == "attention_request":
                    receiver_id = int(message_data.get("receiver_id") or 0)
                    result = {
                        "type": "attention_result",
                        "receiver_id": receiver_id,
                        "success": False,
                    }
                    if not receiver_id or receiver_id == user_data["id"]:
                        result.update(code="invalid_receiver", message="Selecione outro usuario para chamar atencao.")
                        await websocket.send_text(json.dumps(result))
                        continue
    
                    now_utc = datetime.utcnow()
                    usage_date = datetime.now(timezone(timedelta(hours=-3))).date().isoformat()
                    with SessionLocal() as db:
                        receiver = ensure_user_in_company(db, receiver_id, user_data["company_id"])
                        limit = db.query(ChatAttentionLimit).filter(
                            ChatAttentionLimit.company_id == user_data["company_id"],
                            ChatAttentionLimit.sender_id == user_data["id"],
                            ChatAttentionLimit.receiver_id == receiver_id,
                        ).first()
                        if not limit:
                            limit = ChatAttentionLimit(
                                company_id=user_data["company_id"],
                                sender_id=user_data["id"],
                                receiver_id=receiver_id,
                                usage_date=usage_date,
                            )
                            db.add(limit)
                            db.flush()
    
                        if limit.blocked_until and now_utc < limit.blocked_until:
                            retry_after = max(1, int((limit.blocked_until - now_utc).total_seconds()))
                            result.update(
                                code="blocked",
                                message="Bot\u00e3o bloqueado por 24h por abuso de uso da fun\u00e7\u00e3o",
                                retry_after=retry_after,
                                blocked_until=limit.blocked_until.isoformat(),
                            )
                            db.commit()
                            await websocket.send_text(json.dumps(result))
                            continue
    
                        if limit.usage_date != usage_date:
                            limit.usage_date = usage_date
                            limit.daily_count = 0
                            limit.blocked_until = None
    
                        if limit.last_sent_at:
                            cooldown_remaining = 300 - int((now_utc - limit.last_sent_at).total_seconds())
                            if cooldown_remaining > 0:
                                result.update(
                                    code="cooldown",
                                    message="Aguarde 5 minutos para chamar a atencao deste usuario novamente.",
                                    retry_after=cooldown_remaining,
                                )
                                db.commit()
                                await websocket.send_text(json.dumps(result))
                                continue
    
                        if limit.daily_count >= 10:
                            limit.blocked_until = now_utc + timedelta(hours=24)
                            limit.updated_at = now_utc
                            db.commit()
                            result.update(
                                code="blocked",
                                message="Bot\u00e3o bloqueado por 24h por abuso de uso da fun\u00e7\u00e3o",
                                retry_after=24 * 60 * 60,
                                blocked_until=limit.blocked_until.isoformat(),
                            )
                            await websocket.send_text(json.dumps(result))
                            continue
    
                        limit.daily_count += 1
                        limit.last_sent_at = now_utc
                        limit.updated_at = now_utc
                        daily_count = limit.daily_count
                        receiver_name = receiver.full_name
                        db.commit()
    
                    request_id = uuid.uuid4().hex
                    await manager.send_personal_message(json.dumps({
                        "type": "attention_request",
                        "request_id": request_id,
                        "from_user": {
                            "id": user_data["id"],
                            "username": user_data["username"],
                            "full_name": user_data["full_name"],
                            "nickname": user_data["nickname"],
                        },
                    }), receiver_id)
                    await websocket.send_text(json.dumps({
                        "type": "attention_result",
                        "receiver_id": receiver_id,
                        "receiver_name": receiver_name,
                        "success": True,
                        "message": f"Atencao solicitada a {receiver_name}.",
                        "cooldown_seconds": 300,
                        "daily_count": daily_count,
                        "remaining_today": max(0, 10 - daily_count),
                    }))
                elif message_data["type"] == "voice_signal":
                    receiver_id = message_data.get("receiver_id")
                    signal_type = str(message_data.get("signal_type") or "").strip().lower()
                    if not receiver_id or signal_type not in {"offer", "answer", "ice", "cancel", "decline", "end", "invite", "join", "peer", "leave", "host_transfer", "room_event", "meeting_join", "meeting_presence"}:
                        raise HTTPException(status_code=400, detail="Sinal de chamada invalido.")
                    receiver_id = int(receiver_id)
                    if receiver_id == user_data["id"]:
                        raise HTTPException(status_code=400, detail="Nao e possivel ligar para si mesmo.")
                    # Sinalizacao WebRTC e efemera: somente valida o destinatario e retransmite pelo socket.
                    # Nenhum audio, historico ou registro de chamada e persistido no banco.
                    with SessionLocal() as db:
                        ensure_user_in_company(db, receiver_id, user_data["company_id"])
                        group_id = (message_data.get("signal") or {}).get("group_id")
                        if group_id:
                            ensure_chat_group_member(db, int(group_id), user_data["id"], user_data["company_id"])
                            ensure_chat_group_member(db, int(group_id), receiver_id, user_data["company_id"])
                        meeting_id = (message_data.get("signal") or {}).get("meeting_id")
                        if meeting_id:
                            allowed = db.execute(text("SELECT COUNT(*) FROM meeting_participants mp JOIN meetings m ON m.id=mp.meeting_id WHERE m.id=:meeting_id AND m.company_id=:company_id AND mp.user_id IN (:sender_id, :receiver_id)"), {"meeting_id": str(meeting_id), "company_id": user_data["company_id"], "sender_id": user_data["id"], "receiver_id": receiver_id}).scalar()
                            if allowed != 2:
                                raise HTTPException(status_code=403, detail="Participante sem acesso a esta reuniao.")
                    await manager.send_personal_message(json.dumps({
                        "type": "voice_signal",
                        "signal_type": signal_type,
                        "from_user": {
                            "id": user_data["id"],
                            "username": user_data["username"],
                            "full_name": user_data["full_name"],
                            "nickname": user_data["nickname"],
                        },
                        "signal": message_data.get("signal") or {},
                    }), receiver_id)
                elif message_data["type"] == "ping":
                    await websocket.send_text(json.dumps({"type": "pong"}))
                else:
                    await websocket.send_text(json.dumps({"type": "error", "detail": "Evento de WebSocket desconhecido."}))
            except HTTPException as exc:
                # Erros de permissao/validacao de uma acao nao devem derrubar
                # toda a conexao em tempo real. O cliente recebe o erro e
                # continua conectado para as demais conversas.
                await websocket.send_text(json.dumps({
                    "type": "error",
                    "detail": str(exc.detail),
                    "status_code": exc.status_code,
                }))
            except (json.JSONDecodeError, KeyError, TypeError, ValueError) as exc:
                await websocket.send_text(json.dumps({
                    "type": "error",
                    "detail": "Evento de WebSocket invalido.",
                }))

    except WebSocketDisconnect as exc:
        print(
            f"[VoltChat WS] desconectado user={user_data.get('id') if user_data else 'unknown'} "
            f"company={user_data.get('company_id') if user_data else 'unknown'} "
            f"client={client_id or 'legacy'} code={getattr(exc, 'code', None)}"
        )
    except Exception as exc:
        print(f"[VoltChat WS] erro na conexao: {exc}")
    finally:
        if user_data:
            still_online = manager.disconnect(user_data["id"], id(websocket))
            if not still_online:
                await manager.broadcast_user_status(user_data["id"], False, user_data["company_id"])


def purge_expired_files(db, batch_size: int = FILE_CLEANUP_BATCH_SIZE) -> int:
    """Remove um lote sem carregar blobs e sem deixar objetos R2 sem controle."""
    now = datetime.utcnow()
    expired_files = (
        db.query(FileUpload.id, FileUpload.file_path)
        .filter(FileUpload.expires_at.isnot(None), FileUpload.expires_at <= now)
        .order_by(FileUpload.expires_at.asc(), FileUpload.id.asc())
        .limit(max(1, batch_size))
        .all()
    )
    if not expired_files:
        return 0

    file_ids = []
    for row in expired_files:
        if is_r2_chat_file_path(row.file_path):
            if delete_r2_chat_file(row.file_path):
                file_ids.append(row.id)
            continue
        try:
            file_path = Path(row.file_path)
            if file_path.exists():
                file_path.unlink()
        except OSError:
            pass
        file_ids.append(row.id)

    if not file_ids:
        return 0
    db.query(Ticket).filter(Ticket.attachment_file_id.in_(file_ids)).update(
        {Ticket.attachment_file_id: None}, synchronize_session=False
    )
    db.query(TicketMessage).filter(TicketMessage.file_id.in_(file_ids)).update(
        {TicketMessage.file_id: None}, synchronize_session=False
    )
    db.query(FileUpload).filter(FileUpload.id.in_(file_ids)).delete(synchronize_session=False)
    db.commit()
    return len(file_ids)


def seconds_until_bolt_daily_summary(now: Optional[datetime] = None) -> int:
    current = ensure_utc_datetime(now or datetime.now(timezone.utc))
    target = current.replace(
        hour=BOLT_DAILY_SUMMARY_HOUR_UTC,
        minute=0,
        second=0,
        microsecond=0,
    )
    if target <= current:
        target += timedelta(days=1)
    return max(60, int((target - current).total_seconds()))


async def run_bolt_daily_ticket_summaries():
    prepared = await asyncio.to_thread(prepare_bolt_daily_ticket_summaries)
    await dispatch_bolt_daily_ticket_summaries(prepared)
    if prepared:
        print(f"[VoltChat] Volt: {len(prepared)} resumo(s) diário(s) enviado(s).")


async def periodic_bolt_daily_ticket_summaries():
    """Executa uma vez ao dia; a tabela de controle impede duplicidade entre instâncias."""
    current = datetime.now(timezone.utc)
    scheduled_today = current.replace(
        hour=BOLT_DAILY_SUMMARY_HOUR_UTC,
        minute=0,
        second=0,
        microsecond=0,
    )
    if current >= scheduled_today:
        try:
            await run_bolt_daily_ticket_summaries()
        except Exception as exc:
            print(f"[VoltChat] Falha no resumo diário do Volt: {exc}")

    while True:
        await asyncio.sleep(seconds_until_bolt_daily_summary())
        try:
            await run_bolt_daily_ticket_summaries()
        except Exception as exc:
            print(f"[VoltChat] Falha no resumo diário do Volt: {exc}")

def seconds_until_database_maintenance(now: Optional[datetime] = None) -> int:
    current = ensure_utc_datetime(now or datetime.now(timezone.utc))
    target = current.replace(
        hour=DATABASE_MAINTENANCE_HOUR_UTC,
        minute=DATABASE_MAINTENANCE_MINUTE_UTC,
        second=0,
        microsecond=0,
    )
    if target <= current:
        target += timedelta(days=1)
    return max(60, int((target - current).total_seconds()))


def run_expired_file_cleanup() -> int:
    with SessionLocal() as db:
        return purge_expired_files(db)


async def periodic_database_maintenance():
    """Executa uma unica janela leve por dia; backups continuam limitados a 30 dias."""
    await asyncio.sleep(seconds_until_database_maintenance())
    while True:
        try:
            await asyncio.to_thread(archive_due_chat_history)
        except Exception as exc:
            print(f"[VoltChat] Falha no backup mensal: {exc}")
        try:
            deleted_audits = await asyncio.to_thread(purge_expired_audit_logs)
            if deleted_audits:
                print(f"[VoltChat] Auditoria: {deleted_audits} registro(s) expirado(s) removido(s).")
        except Exception as exc:
            print(f"[VoltChat] Falha na limpeza da auditoria: {exc}")
        try:
            deleted_files = await asyncio.to_thread(run_expired_file_cleanup)
            if deleted_files:
                print(f"[VoltChat] Arquivos: {deleted_files} registro(s) expirado(s) removido(s).")
        except Exception as exc:
            print(f"[VoltChat] Falha na limpeza de arquivos: {exc}")
        engine.dispose()
        await asyncio.sleep(DATABASE_MAINTENANCE_INTERVAL_SECONDS)


def list_file_metadata_query(db, company_id: str):
    return (
        db.query(FileUpload)
        .options(defer(FileUpload.binary_data))
        .filter(
            FileUpload.company_id == company_id,
            or_(FileUpload.expires_at.is_(None), FileUpload.expires_at > datetime.utcnow()),
        )
    )


def serialize_file_upload(file_record: FileUpload, db=None, shared_at=None, shared_by_id=None) -> dict:
    uploader = db.query(User).filter(User.id == (shared_by_id or file_record.uploaded_by)).first() if db else None
    return {
        "id": file_record.id,
        "company_id": file_record.company_id,
        "filename": file_record.filename,
        "file_size": file_record.file_size,
        "content_type": file_record.content_type,
        "uploaded_by": file_record.uploaded_by,
        "uploaded_by_name": uploader.full_name if uploader else None,
        "upload_date": file_record.upload_date.isoformat() if file_record.upload_date else None,
        "expires_at": file_record.expires_at.isoformat() if file_record.expires_at else None,
        "shared_at": shared_at.isoformat() if shared_at else None,
    }


def decode_sticker_image(image_data: str) -> tuple[bytes, str]:
    match = re.fullmatch(r"data:(image/(?:webp|png|jpeg));base64,([A-Za-z0-9+/=\s]+)", str(image_data or "").strip())
    if not match:
        raise HTTPException(status_code=400, detail="Envie a figurinha em WebP, PNG ou JPEG.")
    try:
        content = base64.b64decode(re.sub(r"\s+", "", match.group(2)), validate=True)
    except (ValueError, TypeError) as exc:
        raise HTTPException(status_code=400, detail="Imagem de figurinha invalida.") from exc
    if not content or len(content) > 1024 * 1024:
        raise HTTPException(status_code=400, detail="A figurinha deve ter no maximo 1 MB.")
    content_type = match.group(1)
    signatures_valid = (
        (content_type == "image/webp" and content[:4] == b"RIFF" and content[8:12] == b"WEBP")
        or (content_type == "image/png" and content.startswith(b"\x89PNG\r\n\x1a\n"))
        or (content_type == "image/jpeg" and content.startswith(b"\xff\xd8\xff"))
    )
    if not signatures_valid:
        raise HTTPException(status_code=400, detail="O conteudo da figurinha nao corresponde ao formato informado.")
    return content, content_type


def serialize_chat_sticker(sticker: ChatSticker) -> dict:
    return {
        "id": sticker.id,
        "name": sticker.name,
        "content_type": sticker.content_type,
        "file_size": sticker.file_size,
        "created_at": sticker.created_at.isoformat() if sticker.created_at else None,
        "updated_at": sticker.updated_at.isoformat() if sticker.updated_at else None,
        "source_reference": sticker.source_reference,
        "source": f"sticker:{sticker.id}",
    }


@app.get("/stickers/config")
async def get_chat_sticker_config(
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        company = db.query(Company).filter(Company.id == scoped_company_id).first()
        return {
            "allow_user_creation": bool(company.allow_user_sticker_creation),
            "can_create": can_create_company_stickers(db, current_user, scoped_company_id),
        }


@app.get("/stickers")
async def list_chat_stickers(
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        if not can_create_company_stickers(db, current_user, scoped_company_id):
            return []
        stickers = db.query(ChatSticker).filter(
            ChatSticker.company_id == scoped_company_id,
            ChatSticker.owner_user_id == current_user.id,
            ChatSticker.is_active == True,
        ).order_by(ChatSticker.updated_at.desc(), ChatSticker.created_at.desc()).limit(100).all()
        return [serialize_chat_sticker(sticker) for sticker in stickers]


@app.post("/stickers")
async def create_chat_sticker(payload: dict, current_user: User = Depends(get_current_user)):
    name = str(payload.get("name") or "Minha figurinha").strip()[:80] or "Minha figurinha"
    content, content_type = decode_sticker_image(payload.get("image_data"))
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        ensure_company_sticker_creation(db, current_user, company_id)
        count = db.query(func.count(ChatSticker.id)).filter(
            ChatSticker.company_id == company_id,
            ChatSticker.owner_user_id == current_user.id,
            ChatSticker.is_active == True,
        ).scalar() or 0
        if count >= 100:
            raise HTTPException(status_code=400, detail="Limite de 100 figurinhas pessoais atingido.")
        sticker = ChatSticker(
            company_id=company_id,
            owner_user_id=current_user.id,
            name=name,
            image_data=content,
            content_type=content_type,
            file_size=len(content),
        )
        db.add(sticker)
        db.commit()
        db.refresh(sticker)
        return serialize_chat_sticker(sticker)


@app.post("/stickers/favorite")
async def favorite_chat_sticker(payload: dict, current_user: User = Depends(get_current_user)):
    try:
        message_id = int(payload.get("message_id"))
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="Mensagem de figurinha invalida.") from exc

    with SessionLocal() as db:
        message = db.query(Message).filter(Message.id == message_id).first()
        if not message or message.message_type != "sticker":
            raise HTTPException(status_code=404, detail="Figurinha nao encontrada na conversa.")
        company_id = ensure_company_access(db, current_user, message.company_id)
        if message.group_id:
            ensure_chat_group_member(db, int(message.group_id), current_user.id, company_id)
        elif message.receiver_id and current_user.id not in {message.sender_id, message.receiver_id}:
            raise HTTPException(status_code=403, detail="Sem acesso a esta conversa.")

        source_reference = str(message.file_path or "").strip()
        if not source_reference.startswith(("default:", "sticker:")):
            raise HTTPException(status_code=400, detail="Referencia de figurinha invalida.")

        if source_reference.startswith("sticker:"):
            ensure_company_sticker_creation(db, current_user, company_id)
            try:
                source_id = int(source_reference.split(":", 1)[1])
            except ValueError as exc:
                raise HTTPException(status_code=400, detail="Referencia de figurinha invalida.") from exc
            source_sticker = db.query(ChatSticker).filter(
                ChatSticker.id == source_id,
                ChatSticker.company_id == company_id,
            ).first()
            if not source_sticker:
                raise HTTPException(status_code=404, detail="Figurinha original nao encontrada.")
            if source_sticker.owner_user_id == current_user.id and source_sticker.is_active:
                return {**serialize_chat_sticker(source_sticker), "favorited": True, "already_saved": True}
            image_content = source_sticker.image_data
            content_type = source_sticker.content_type
        else:
            image_content, content_type = decode_sticker_image(payload.get("image_data"))

        existing = db.query(ChatSticker).filter(
            ChatSticker.company_id == company_id,
            ChatSticker.owner_user_id == current_user.id,
            ChatSticker.source_reference == source_reference,
            ChatSticker.is_active == True,
        ).first()
        if existing:
            return {**serialize_chat_sticker(existing), "favorited": True, "already_saved": True}

        count = db.query(func.count(ChatSticker.id)).filter(
            ChatSticker.company_id == company_id,
            ChatSticker.owner_user_id == current_user.id,
            ChatSticker.is_active == True,
        ).scalar() or 0
        if count >= 100:
            raise HTTPException(status_code=400, detail="Limite de 100 figurinhas pessoais atingido.")

        favorite = ChatSticker(
            company_id=company_id,
            owner_user_id=current_user.id,
            name=str(message.content or "Figurinha favorita").strip()[:80] or "Figurinha favorita",
            image_data=image_content,
            content_type=content_type,
            file_size=len(image_content),
            source_reference=source_reference,
        )
        db.add(favorite)
        try:
            db.commit()
        except IntegrityError:
            db.rollback()
            favorite = db.query(ChatSticker).filter(
                ChatSticker.company_id == company_id,
                ChatSticker.owner_user_id == current_user.id,
                ChatSticker.source_reference == source_reference,
                ChatSticker.is_active == True,
            ).first()
            if not favorite:
                raise
        db.refresh(favorite)
        return {**serialize_chat_sticker(favorite), "favorited": True, "already_saved": False}

@app.put("/stickers/{sticker_id}")
async def update_chat_sticker(sticker_id: int, payload: dict, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        sticker = db.query(ChatSticker).filter(ChatSticker.id == sticker_id, ChatSticker.owner_user_id == current_user.id).first()
        if not sticker or not sticker.is_active:
            raise HTTPException(status_code=404, detail="Figurinha nao encontrada.")
        ensure_company_access(db, current_user, sticker.company_id)
        ensure_company_sticker_creation(db, current_user, sticker.company_id)
        if "name" in payload:
            sticker.name = str(payload.get("name") or "Minha figurinha").strip()[:80] or "Minha figurinha"
        if payload.get("image_data"):
            content, content_type = decode_sticker_image(payload.get("image_data"))
            sticker.image_data = content
            sticker.content_type = content_type
            sticker.file_size = len(content)
        sticker.updated_at = datetime.utcnow()
        db.commit()
        db.refresh(sticker)
        return serialize_chat_sticker(sticker)


@app.delete("/stickers/{sticker_id}")
async def delete_chat_sticker(sticker_id: int, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        sticker = db.query(ChatSticker).filter(ChatSticker.id == sticker_id, ChatSticker.owner_user_id == current_user.id).first()
        if not sticker or not sticker.is_active:
            raise HTTPException(status_code=404, detail="Figurinha nao encontrada.")
        ensure_company_access(db, current_user, sticker.company_id)
        sticker.is_active = False
        sticker.updated_at = datetime.utcnow()
        db.commit()
        return {"success": True}


@app.get("/stickers/{sticker_id}/image")
async def get_chat_sticker_image(sticker_id: int, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        sticker = db.query(ChatSticker).filter(ChatSticker.id == sticker_id).first()
        if not sticker:
            raise HTTPException(status_code=404, detail="Figurinha nao encontrada.")
        ensure_company_access(db, current_user, sticker.company_id)
        return Response(
            content=sticker.image_data,
            media_type=sticker.content_type,
            headers={"Cache-Control": "private, max-age=86400", "Content-Disposition": "inline"},
        )

@app.post("/files/upload")
async def upload_file(
    file: UploadFile = File(...),
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    file_extension = Path(file.filename or "").suffix.lower()
    if file_extension not in ALLOWED_UPLOAD_EXTENSIONS:
        allowed = ", ".join(sorted(ALLOWED_UPLOAD_EXTENSIONS))
        raise HTTPException(status_code=400, detail=f"Formato nao permitido. Envie apenas: {allowed}.")

    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Arquivo muito grande. Limite de 10 MB.")

    unique_filename = f"{uuid.uuid4()}{file_extension}"
    content_type = file.content_type or "application/octet-stream"
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, company_id)
        if CHAT_FILE_STORAGE_MODE == "r2":
            file_path = upload_r2_chat_file(company_id, file.filename or unique_filename, content, content_type)
            binary_data = None
        else:
            upload_dir = Path(os.getenv("UPLOAD_DIR", "uploads"))
            upload_dir.mkdir(parents=True, exist_ok=True)
            local_path = upload_dir / unique_filename
            local_path.write_bytes(content)
            file_path = str(local_path)
            binary_data = content
        file_record = FileUpload(
            company_id=company_id,
            filename=file.filename or unique_filename,
            file_path=file_path,
            file_size=len(content),
            content_type=content_type,
            binary_data=binary_data,
            uploaded_by=current_user.id,
            expires_at=datetime.utcnow() + timedelta(days=15),
        )
        db.add(file_record)
        db.commit()
        db.refresh(file_record)
        return serialize_file_upload(file_record, db)


@app.get("/files/")
async def list_files(
    company_id: Optional[str] = Query(default=None),
    limit: int = Query(default=50, ge=1, le=100),
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        files = (
            list_file_metadata_query(db, scoped_company_id)
            .order_by(FileUpload.upload_date.desc())
            .limit(limit)
            .all()
        )
        return [serialize_file_upload(item, db) for item in files]


@app.get("/files/conversation")
async def list_conversation_files(
    user_id: Optional[int] = Query(default=None),
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        query = db.query(Message).filter(
            Message.company_id == scoped_company_id,
            Message.file_path.isnot(None),
        )
        if user_id:
            ensure_user_in_company(db, int(user_id), scoped_company_id)
            query = query.filter(
                or_(
                    (Message.sender_id == current_user.id) & (Message.receiver_id == int(user_id)),
                    (Message.sender_id == int(user_id)) & (Message.receiver_id == current_user.id),
                )
            )
        else:
            raise HTTPException(status_code=400, detail="Selecione uma conversa direta para consultar arquivos.")

        history = []
        seen_file_ids = set()
        for message in query.order_by(Message.timestamp.desc()).all():
            try:
                file_id = int(message.file_path)
            except (TypeError, ValueError):
                continue
            if file_id in seen_file_ids:
                continue
            file_record = db.query(FileUpload).filter(
                FileUpload.id == file_id,
                FileUpload.company_id == scoped_company_id,
                or_(FileUpload.expires_at.is_(None), FileUpload.expires_at > datetime.utcnow()),
            ).first()
            if file_record:
                seen_file_ids.add(file_id)
                history.append(serialize_file_upload(file_record, db, message.timestamp, message.sender_id))
        return history


@app.get("/files/download/{file_id}")
async def download_file(file_id: int, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        file_record = db.query(FileUpload).options(defer(FileUpload.binary_data)).filter(
            FileUpload.id == file_id,
            or_(FileUpload.expires_at.is_(None), FileUpload.expires_at > datetime.utcnow()),
        ).first()
        if not file_record:
            raise HTTPException(status_code=404, detail="Arquivo nao encontrado")
        ensure_company_access(db, current_user, file_record.company_id)
        file_path = file_record.file_path
        filename = file_record.filename
        content_type = file_record.content_type
        binary_data = None if is_r2_chat_file_path(file_path) else (bytes(file_record.binary_data or b"") if file_record.binary_data else None)

    if is_r2_chat_file_path(file_path):
        return RedirectResponse(presign_r2_chat_file(file_path), status_code=307)

    if not os.path.exists(file_path):
        if not binary_data:
            raise HTTPException(status_code=404, detail="Arquivo nao existe")
        return StreamingResponse(
            iter([binary_data]),
            media_type=content_type,
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Content-Length": str(len(binary_data)),
                "Cache-Control": "no-store",
            },
        )

    return FileResponse(path=file_path, filename=filename, media_type=content_type)



@app.get("/files/content/{file_id}")
async def stream_file_content(file_id: int, current_user: User = Depends(get_current_user)):
    """Entrega conteudo autenticado para midia embutida sem depender de CORS do R2.

    Arquivos R2 sao lidos em blocos; o Render nao carrega o objeto inteiro na memoria.
    Downloads tradicionais continuam usando URL assinada e seguem diretos ao R2.
    """
    with SessionLocal() as db:
        file_record = db.query(FileUpload).options(defer(FileUpload.binary_data)).filter(
            FileUpload.id == file_id,
            or_(FileUpload.expires_at.is_(None), FileUpload.expires_at > datetime.utcnow()),
        ).first()
        if not file_record:
            raise HTTPException(status_code=404, detail="Arquivo nao encontrado")
        ensure_company_access(db, current_user, file_record.company_id)
        file_path = file_record.file_path
        content_type = file_record.content_type or "application/octet-stream"
        file_size = int(file_record.file_size or 0)
        binary_data = None if is_r2_chat_file_path(file_path) else (bytes(file_record.binary_data or b"") if file_record.binary_data else None)

    if is_r2_chat_file_path(file_path):
        try:
            object_response = r2_client().get_object(
                Bucket=R2_BUCKET_NAME,
                Key=r2_chat_file_key(file_path),
            )
            body = object_response["Body"]
        except Exception as exc:
            print(f"[chat-files] Falha ao ler arquivo do R2: {exc}")
            raise HTTPException(status_code=404, detail="Arquivo nao existe") from exc

        def r2_chunks():
            try:
                while chunk := body.read(1024 * 1024):
                    yield chunk
            finally:
                body.close()

        headers = {"Cache-Control": "private, max-age=300"}
        if file_size > 0:
            headers["Content-Length"] = str(file_size)
        return StreamingResponse(r2_chunks(), media_type=content_type, headers=headers)

    if not os.path.exists(file_path):
        if not binary_data:
            raise HTTPException(status_code=404, detail="Arquivo nao existe")
        return StreamingResponse(
            iter([binary_data]),
            media_type=content_type,
            headers={"Content-Length": str(len(binary_data)), "Cache-Control": "private, max-age=300"},
        )

    return FileResponse(path=file_path, media_type=content_type, headers={"Cache-Control": "private, max-age=300"})

TENANT_USER_IMPORT_HEADERS = [
    "nome_completo",
    "usuario",
    "email",
    "ramal",
    "setor",
    "nivel_usuario",
]

TENANT_USER_DEFAULT_PASSWORD = "Alterar@123"


def preview_tenant_user_import_rows(rows: list[dict], db, company_id: str) -> dict:
    preview = []
    allowed_roles = {
        "user", "usuario", "usuário", "coordinator", "coordenador",
        "company_admin", "admin", "administrador",
    }
    for index, row in enumerate(rows, start=2):
        errors = []
        warnings = []
        name = str(row.get("nome_completo") or row.get("nome_usuario") or "").strip()
        username = str(row.get("usuario") or "").strip().lower()
        email = str(row.get("email") or "").strip().lower()
        department_name = str(row.get("setor") or "").strip()
        raw_role = str(row.get("nivel_usuario") or "user").strip().lower()
        role = normalize_platform_role(raw_role)

        if not name:
            errors.append("nome_completo obrigatorio")
        if not EMAIL_PATTERN.match(email):
            errors.append("email invalido")
        if not department_name:
            errors.append("setor obrigatorio")
        if raw_role not in allowed_roles or role == "master_admin":
            errors.append("nivel_usuario deve ser usuario, coordenador ou admin")

        existing_user = db.query(User).filter(User.email == email).first() if email else None
        if existing_user:
            existing_link = db.query(CompanyUser).filter(
                CompanyUser.company_id == company_id,
                CompanyUser.user_id == existing_user.id,
            ).first()
            warnings.append(
                "conta existente: dados e vinculo serao atualizados"
                if existing_link
                else "email ja existe em outro tenant: sera vinculado sem alterar a senha"
            )
        if username:
            username_owner = db.query(User).filter(User.username == username).first()
            if username_owner and (not existing_user or username_owner.id != existing_user.id):
                generated_username = generate_unique_username(db, email, name, username)
                warnings.append(f"usuario existente: sera criado como {generated_username}")
        if department_name and not db.query(Department).filter(
            Department.company_id == company_id,
            Department.name.ilike(department_name),
        ).first():
            warnings.append("setor ainda nao existe e sera criado")

        preview.append({
            "row": index,
            "data": {
                "nome_completo": name,
                "usuario": username,
                "email": email,
                "ramal": str(row.get("ramal") or "").strip(),
                "setor": department_name,
                "nivel_usuario": role,
            },
            "errors": errors,
            "warnings": warnings,
            "valid": not errors,
        })
    return {
        "rows": preview,
        "valid_count": len([row for row in preview if row["valid"]]),
        "error_count": len([row for row in preview if row["errors"]]),
    }


@app.get("/companies/available")
async def list_available_companies(current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        if is_admin(current_user):
            companies = db.query(Company).filter(Company.status == ACTIVE_STATUS).order_by(Company.name.asc()).all()
            return [{**serialize_company(company), "role": "master_admin"} for company in companies]

        memberships = active_company_memberships(db, current_user)
        result = []
        for membership in memberships:
            company = db.query(Company).filter(Company.id == membership.company_id).first()
            if company:
                result.append({**serialize_company(company), "role": membership.role})
        return sorted(result, key=lambda item: item["name"].lower())


def build_dashboard_ticket_breakdown(db, tickets: list[Ticket], current_user_id: int, perspective: str) -> dict:
    """Classifica tickets por estado e pela ultima mensagem, sem inventar indicadores."""
    latest_messages = {}
    ticket_ids = [ticket.id for ticket in tickets]
    if ticket_ids:
        for message in (
            db.query(TicketMessage)
            .filter(TicketMessage.ticket_id.in_(ticket_ids))
            .order_by(TicketMessage.created_at.desc(), TicketMessage.id.desc())
            .all()
        ):
            latest_messages.setdefault(message.ticket_id, message)

    result = {"total": len(tickets), "open": 0, "waiting_attendant": 0, "waiting_user": 0, "closed": 0}
    for ticket in tickets:
        if str(ticket.status or "").strip().lower() == "resolvido":
            result["closed"] += 1
            continue
        latest_message = latest_messages.get(ticket.id)
        if not latest_message:
            result["open"] += 1
            continue
        if perspective == "requester":
            waiting_key = "waiting_attendant" if latest_message.sender_id == current_user_id else "waiting_user"
        else:
            waiting_key = "waiting_user" if latest_message.sender_id == current_user_id else "waiting_attendant"
        result[waiting_key] += 1
    return result


def build_dashboard_recent_conversations(db, company_id: str, current_user_id: int) -> list[dict]:
    """Retorna apenas conversas diretas recentes. O Chat Geral foi desativado."""
    conversations = []
    seen_user_ids = set()
    direct_messages = (
        db.query(Message)
        .filter(
            Message.company_id == company_id,
            Message.receiver_id.isnot(None),
            or_(Message.sender_id == current_user_id, Message.receiver_id == current_user_id),
        )
        .order_by(Message.timestamp.desc(), Message.id.desc())
        .limit(120)
        .all()
    )
    for message in direct_messages:
        contact_id = message.receiver_id if message.sender_id == current_user_id else message.sender_id
        if not contact_id or contact_id in seen_user_ids:
            continue
        contact = db.query(User).filter(User.id == contact_id).first()
        if not contact:
            continue
        seen_user_ids.add(contact_id)
        conversations.append({
            "id": contact.id,
            "type": "direct",
            "name": display_user_name(contact),
            "full_name": contact.full_name,
            "username": contact.username,
            "profile_photo": contact.profile_photo,
            "status": contact.status,
            "last_message": str(message.content or "Anexo compartilhado").strip()[:120],
            "last_message_at": message.timestamp.isoformat() if message.timestamp else None,
        })
        if len(conversations) >= 24:
            break
    return conversations


def build_dashboard_birthdays(db, company_id: str) -> list[dict]:
    today = datetime.utcnow().date()
    today_prefix = today.strftime("%d-%m")
    users = (
        db.query(User)
        .join(CompanyUser, CompanyUser.user_id == User.id)
        .filter(
            CompanyUser.company_id == company_id,
            CompanyUser.status == ACTIVE_STATUS,
            User.is_active == True,
            User.birthday.isnot(None),
            User.birthday != "",
        )
        .order_by(User.full_name.asc())
        .all()
    )
    birthday_users = []
    for user in users:
        try:
            birthday = normalize_birthday(user.birthday)
        except ValueError:
            continue
        if not birthday or not birthday.startswith(today_prefix):
            continue
        parts = birthday.split("-")
        age = None
        if len(parts) == 3 and len(parts[2]) in {2, 4}:
            raw_year = int(parts[2])
            year = raw_year if len(parts[2]) == 4 else (2000 + raw_year if raw_year <= today.year % 100 else 1900 + raw_year)
            age = today.year - year
        birthday_users.append({
            "id": user.id,
            "name": display_user_name(user),
            "full_name": user.full_name,
            "profile_photo": user.profile_photo,
            "birthday": birthday,
            "age": age,
        })
    return birthday_users

@app.get("/dashboard/overview")
async def dashboard_overview(
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    """Retorna somente indicadores pertencentes ao tenant e ao escopo funcional do usuario."""
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        company = db.query(Company).filter(Company.id == scoped_company_id).first()
        membership = get_company_membership(db, current_user, scoped_company_id)
        company_admin = is_company_admin(db, current_user, scoped_company_id)
        coordinator = bool(membership and membership.role == "coordinator") and not company_admin
        department = user_department_name(db, current_user, scoped_company_id)

        ticket_query = db.query(Ticket).filter(Ticket.company_id == scoped_company_id)
        task_query = db.query(TaskItem).filter(TaskItem.company_id == scoped_company_id)

        if not company_admin:
            directly_assigned_ticket_ids = db.query(TicketAssignee.ticket_id).filter(
                TicketAssignee.user_id == current_user.id
            )
            ticket_scope = [
                Ticket.created_by_id == current_user.id,
                Ticket.assigned_to_id == current_user.id,
                Ticket.id.in_(directly_assigned_ticket_ids),
            ]
            task_query = task_query.filter(
                or_(TaskItem.assigned_to_id == current_user.id, TaskItem.created_by_id == current_user.id)
            )
            ticket_query = ticket_query.filter(or_(*ticket_scope))

        ticket_total = ticket_query.count()
        ticket_open = ticket_query.filter(Ticket.status != "Resolvido").count()
        ticket_in_progress = ticket_query.filter(Ticket.status == "Em andamento").count()
        task_total = task_query.count()
        task_active = task_query.filter(TaskItem.status != "done").count()
        task_in_progress = task_query.filter(TaskItem.status == "doing").count()
        recent_tickets = ticket_query.order_by(Ticket.created_at.desc()).limit(5).all()
        recent_tasks = task_query.order_by(TaskItem.created_at.desc()).limit(5).all()

        personal_request_tickets = (
            db.query(Ticket)
            .filter(Ticket.company_id == scoped_company_id, Ticket.created_by_id == current_user.id)
            .all()
        )
        direct_service_ticket_ids = [
            row[0]
            for row in db.query(TicketAssignee.ticket_id)
            .join(Ticket, Ticket.id == TicketAssignee.ticket_id)
            .filter(Ticket.company_id == scoped_company_id, TicketAssignee.user_id == current_user.id)
            .all()
        ]
        service_filters = [Ticket.assigned_to_id == current_user.id]
        if direct_service_ticket_ids:
            service_filters.append(Ticket.id.in_(direct_service_ticket_ids))
        personal_service_tickets = (
            db.query(Ticket)
            .filter(Ticket.company_id == scoped_company_id, or_(*service_filters))
            .all()
        )
        ticket_dashboards = {
            "my_requests": build_dashboard_ticket_breakdown(db, personal_request_tickets, current_user.id, "requester"),
            "my_services": build_dashboard_ticket_breakdown(db, personal_service_tickets, current_user.id, "attendant"),
        }

        # Coordenador tambem ve apenas tickets/tarefas em que participa explicitamente.
        scope = "company" if company_admin else "assigned" if coordinator else "user"
        role = "master_admin" if is_admin(current_user) else (membership.role if membership else "user")
        return {
            "company": serialize_company(company),
            "scope": {"type": scope, "role": role, "department": None},
            "stats": {
                "tickets": ticket_total,
                "open_tickets": ticket_open,
                "tickets_in_progress": ticket_in_progress,
                "tasks": task_total,
                "active_tasks": task_active,
                "tasks_in_progress": task_in_progress,
            },
            "recent_tickets": [serialize_ticket(ticket, db) for ticket in recent_tickets],
            "recent_tasks": [serialize_task(task, db) for task in recent_tasks],
            "recent_conversations": build_dashboard_recent_conversations(db, scoped_company_id, current_user.id),
            "ticket_dashboards": ticket_dashboards,
            "birthdays_today": build_dashboard_birthdays(db, scoped_company_id),
        }


@app.get("/company-admin/overview")
async def company_admin_overview(
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        ensure_company_admin(db, current_user, scoped_company_id)
        company = db.query(Company).filter(Company.id == scoped_company_id).first()
        departments = (
            db.query(Department)
            .filter(Department.company_id == scoped_company_id)
            .order_by(Department.status.asc(), Department.name.asc())
            .all()
        )
        membership_rows = (
            db.query(CompanyUser, User)
            .join(User, User.id == CompanyUser.user_id)
            .filter(CompanyUser.company_id == scoped_company_id)
            .order_by(User.full_name.asc())
            .all()
        )
        users = []
        for membership, user in membership_rows:
            user_payload = serialize_user(user, db)
            user_payload["membership"] = serialize_company_user(membership, db)
            users.append(user_payload)
        return {
            "company": serialize_company(company),
            "departments": [serialize_department(item) for item in departments],
            "users": users,
            "stats": {
                "departments": len(departments),
                "active_departments": len([item for item in departments if item.status == ACTIVE_STATUS]),
                "users": len(users),
                "active_users": len([item for item in users if item["membership"]["status"] == ACTIVE_STATUS]),
            },
        }


@app.patch("/company-admin/settings")
async def update_company_admin_settings(payload: dict, current_user: User = Depends(get_current_user)):
    changed = {}
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        ensure_company_admin(db, current_user, scoped_company_id)
        company = db.query(Company).filter(Company.id == scoped_company_id).first()
        if "allow_user_sticker_creation" in payload:
            company.allow_user_sticker_creation = payload_flag(payload.get("allow_user_sticker_creation"))
            changed["allow_user_sticker_creation"] = company.allow_user_sticker_creation
        if "allow_user_message_editing" in payload:
            company.allow_user_message_editing = payload_flag(payload.get("allow_user_message_editing"))
            changed["allow_user_message_editing"] = company.allow_user_message_editing
        if not changed:
            raise HTTPException(status_code=400, detail="Informe ao menos uma configuracao da empresa.")
        company.updated_at = datetime.utcnow()
        log_audit(
            db,
            current_user.id,
            "configuracoes_empresa_atualizadas",
            "company",
            company.id,
            company.id,
            changed,
        )
        db.commit()
        db.refresh(company)
        company_payload = serialize_company(company)
    await manager.broadcast(json.dumps({
        "type": "company_settings_updated",
        "company_id": scoped_company_id,
        **changed,
    }), company_id=scoped_company_id)
    return company_payload

def ensure_internal_control_target(db, company_id: str, user_id: int) -> User:
    membership = db.query(CompanyUser).filter(
        CompanyUser.company_id == company_id,
        CompanyUser.user_id == user_id,
    ).first()
    user = db.query(User).filter(User.id == user_id).first()
    if not membership or not user:
        raise HTTPException(status_code=404, detail="Usuario nao pertence a esta empresa.")
    return user


@app.get("/company-admin/internal-control/users/{user_id}/conversations")
async def company_admin_internal_control_conversations(
    user_id: int,
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        ensure_company_admin(db, current_user, scoped_company_id)
        target = ensure_internal_control_target(db, scoped_company_id, user_id)
        conversations = []

        direct_pairs = db.query(Message.sender_id, Message.receiver_id).filter(
            Message.company_id == scoped_company_id,
            Message.group_id.is_(None),
            Message.receiver_id.isnot(None),
            or_(Message.sender_id == target.id, Message.receiver_id == target.id),
        ).distinct().all()
        other_user_ids = {
            receiver_id if sender_id == target.id else sender_id
            for sender_id, receiver_id in direct_pairs
            if (receiver_id if sender_id == target.id else sender_id)
        }
        for other_id in other_user_ids:
            other_membership = db.query(CompanyUser).filter(
                CompanyUser.company_id == scoped_company_id,
                CompanyUser.user_id == other_id,
            ).first()
            other = db.query(User).filter(User.id == other_id).first()
            if not other_membership or not other:
                continue
            last_message = db.query(Message).filter(
                Message.company_id == scoped_company_id,
                Message.group_id.is_(None),
                or_(
                    and_(Message.sender_id == target.id, Message.receiver_id == other.id),
                    and_(Message.sender_id == other.id, Message.receiver_id == target.id),
                ),
            ).order_by(Message.timestamp.desc(), Message.id.desc()).first()
            conversations.append({
                "key": f"direct:{other.id}",
                "kind": "direct",
                "id": other.id,
                "name": display_user_name(other),
                "subtitle": other.full_name if display_user_name(other) != other.full_name else other.department,
                "profile_photo": other.profile_photo,
                "last_message": serialize_message(last_message, db) if last_message else None,
            })

        groups = db.query(ChatGroup).join(
            ChatGroupMember, ChatGroupMember.group_id == ChatGroup.id
        ).filter(
            ChatGroup.company_id == scoped_company_id,
            ChatGroup.is_active == True,
            ChatGroupMember.user_id == target.id,
        ).order_by(ChatGroup.name.asc()).all()
        for group in groups:
            last_message = db.query(Message).filter(
                Message.company_id == scoped_company_id,
                Message.group_id == group.id,
            ).order_by(Message.timestamp.desc(), Message.id.desc()).first()
            conversations.append({
                "key": f"group:{group.id}",
                "kind": "group",
                "id": group.id,
                "name": group.name,
                "subtitle": group.description or group.department or "Grupo",
                "profile_photo": group.image_data,
                "last_message": serialize_message(last_message, db) if last_message else None,
            })

        conversations.sort(
            key=lambda item: item["last_message"]["timestamp"] if item.get("last_message") else "",
            reverse=True,
        )
        log_audit(
            db,
            current_user.id,
            "controle_interno_usuario_acessado",
            "user",
            target.id,
            scoped_company_id,
            {"target_user_name": target.full_name, "conversation_count": len(conversations)},
        )
        db.commit()
        return {"user": serialize_user(target, db), "conversations": conversations}


@app.get("/company-admin/internal-control/users/{user_id}/messages")
async def company_admin_internal_control_messages(
    user_id: int,
    kind: str = Query(..., pattern="^(general|direct|group)$"),
    conversation_id: Optional[int] = Query(default=None),
    company_id: Optional[str] = Query(default=None),
    limit: int = Query(default=100, ge=1, le=200),
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        ensure_company_admin(db, current_user, scoped_company_id)
        target = ensure_internal_control_target(db, scoped_company_id, user_id)
        query = db.query(Message).filter(Message.company_id == scoped_company_id)

        if kind == "general":
            raise HTTPException(status_code=410, detail="O Chat Geral foi desativado.")
        elif kind == "direct":
            if not conversation_id:
                raise HTTPException(status_code=400, detail="Conversa direta invalida.")
            other = ensure_internal_control_target(db, scoped_company_id, conversation_id)
            query = query.filter(
                Message.group_id.is_(None),
                or_(
                    and_(Message.sender_id == target.id, Message.receiver_id == other.id),
                    and_(Message.sender_id == other.id, Message.receiver_id == target.id),
                ),
            )
            conversation_label = other.full_name
        else:
            if not conversation_id:
                raise HTTPException(status_code=400, detail="Grupo invalido.")
            group = ensure_chat_group_member(db, conversation_id, target.id, scoped_company_id)
            query = query.filter(Message.group_id == group.id)
            conversation_label = group.name

        messages = query.order_by(Message.timestamp.desc(), Message.id.desc()).limit(limit).all()
        log_audit(
            db,
            current_user.id,
            "controle_interno_conversa_visualizada",
            "user",
            target.id,
            scoped_company_id,
            {
                "kind": kind,
                "conversation_id": conversation_id,
                "conversation_label": conversation_label,
                "message_count": len(messages),
            },
        )
        db.commit()
        return {
            "read_only": True,
            "user": serialize_user(target, db),
            "messages": serialize_messages(reversed(messages), db),
        }

@app.post("/company-admin/departments")
async def company_admin_create_department(payload: dict, current_user: User = Depends(get_current_user)):
    name = str(payload.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="Nome do setor e obrigatorio.")
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        ensure_company_admin(db, current_user, company_id)
        existing = db.query(Department).filter(
            Department.company_id == company_id,
            Department.name.ilike(name),
        ).first()
        if existing:
            raise HTTPException(status_code=409, detail="Ja existe um setor com este nome.")
        department = Department(
            id=str(uuid.uuid4()),
            company_id=company_id,
            name=name,
            description=str(payload.get("description") or "").strip() or None,
            status=ACTIVE_STATUS,
        )
        db.add(department)
        log_audit(db, current_user.id, "setor_criado", "department", department.id, company_id, {"name": name})
        db.commit()
        db.refresh(department)
        return serialize_department(department)


@app.patch("/company-admin/departments/{department_id}")
async def company_admin_update_department(
    department_id: str,
    payload: dict,
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        department = db.query(Department).filter(Department.id == department_id).first()
        if not department:
            raise HTTPException(status_code=404, detail="Setor nao encontrado.")
        company_id = ensure_company_access(db, current_user, payload.get("company_id") or department.company_id)
        ensure_company_admin(db, current_user, company_id)
        if department.company_id != company_id:
            raise HTTPException(status_code=403, detail="Setor nao pertence a esta empresa.")

        old_name = department.name
        if "name" in payload:
            new_name = str(payload.get("name") or "").strip()
            if not new_name:
                raise HTTPException(status_code=400, detail="Nome do setor e obrigatorio.")
            duplicate = db.query(Department).filter(
                Department.company_id == company_id,
                Department.id != department.id,
                Department.name.ilike(new_name),
            ).first()
            if duplicate:
                raise HTTPException(status_code=409, detail="Ja existe um setor com este nome.")
            department.name = new_name
            if new_name != old_name:
                member_user_ids = [
                    item[0] for item in db.query(CompanyUser.user_id)
                    .filter(CompanyUser.company_id == company_id, CompanyUser.department_id == department.id)
                    .all()
                ]
                if member_user_ids:
                    db.query(User).filter(User.id.in_(member_user_ids)).update(
                        {User.department: new_name}, synchronize_session=False
                    )
                db.query(Ticket).filter(Ticket.company_id == company_id, Ticket.department == old_name).update(
                    {Ticket.department: new_name}, synchronize_session=False
                )
                db.query(TaskItem).filter(TaskItem.company_id == company_id, TaskItem.category == old_name).update(
                    {TaskItem.category: new_name}, synchronize_session=False
                )
                db.query(ChatGroup).filter(ChatGroup.company_id == company_id, ChatGroup.department == old_name).update(
                    {ChatGroup.department: new_name}, synchronize_session=False
                )

        if "description" in payload:
            department.description = str(payload.get("description") or "").strip() or None
        if "status" in payload:
            status = normalize_status(payload.get("status"))
            if status not in {ACTIVE_STATUS, INACTIVE_STATUS}:
                raise HTTPException(status_code=400, detail="Status de setor invalido.")
            if status == INACTIVE_STATUS:
                active_members = db.query(CompanyUser).filter(
                    CompanyUser.company_id == company_id,
                    CompanyUser.department_id == department.id,
                    CompanyUser.status == ACTIVE_STATUS,
                ).count()
                if active_members:
                    raise HTTPException(
                        status_code=409,
                        detail="Realocar os usuarios ativos antes de inativar o setor.",
                    )
            department.status = status

        log_audit(db, current_user.id, "setor_editado", "department", department.id, company_id, {
            "old_name": old_name,
            "name": department.name,
            "status": department.status,
        })
        db.commit()
        db.refresh(department)
        return serialize_department(department)


@app.patch("/company-admin/users/{user_id}/department")
async def company_admin_allocate_user(
    user_id: int,
    payload: dict,
    current_user: User = Depends(get_current_user),
):
    department_id = str(payload.get("department_id") or "").strip()
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        ensure_company_admin(db, current_user, company_id)
        membership = db.query(CompanyUser).filter(
            CompanyUser.company_id == company_id,
            CompanyUser.user_id == user_id,
        ).first()
        if not membership:
            raise HTTPException(status_code=404, detail="Usuario nao pertence a esta empresa.")
        department = db.query(Department).filter(
            Department.id == department_id,
            Department.company_id == company_id,
            Department.status == ACTIVE_STATUS,
        ).first()
        if not department:
            raise HTTPException(status_code=404, detail="Setor ativo nao encontrado.")
        membership.department_id = department.id
        user = db.query(User).filter(User.id == user_id).first()
        if user:
            user.department = department.name
            user.updated_at = datetime.utcnow()
        log_audit(db, current_user.id, "usuario_realocado", "company_user", membership.id, company_id, {
            "user_id": user_id,
            "department_id": department.id,
        })
        db.commit()
        db.refresh(membership)
        return serialize_company_user(membership, db)


@app.get("/company-admin/import/users/template")
async def company_admin_user_import_template(
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        ensure_company_admin(db, current_user, scoped_company_id)
        departments = [
            row.name for row in db.query(Department)
            .filter(Department.company_id == scoped_company_id, Department.status == ACTIVE_STATUS)
            .order_by(Department.name.asc())
            .all()
        ]

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Usuarios"
    sheet.append(TENANT_USER_IMPORT_HEADERS)
    sheet.append([
        "Maria da Silva",
        "maria.silva",
        "maria@empresa.com.br",
        "204",
        departments[0] if departments else "Operacao",
        "usuario",
    ])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2478F2")
    widths = [24, 20, 30, 12, 22, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    role_validation = DataValidation(
        type="list",
        formula1='"usuario,coordenador,admin"',
        allow_blank=False,
    )
    role_validation.error = "Escolha usuario, coordenador ou admin."
    role_validation.errorTitle = "Nivel de usuario invalido"
    sheet.add_data_validation(role_validation)
    role_validation.add("F2:F1001")
    guide = workbook.create_sheet("Orientacoes")
    guide.append(["Campo", "Orientacao"])
    guide_rows = [
        ("nome_completo", "Obrigatorio."),
        ("usuario", "Opcional. Se vazio, sera criado a partir do email."),
        ("email", "Obrigatorio e unico."),
        ("ramal", "Opcional."),
        ("setor", "Obrigatorio. Se ainda nao existir, sera criado apos confirmacao."),
        ("nivel_usuario", "Use usuario, coordenador ou admin."),
        ("Regras automaticas", "Status active, sem login no Hub e senha inicial Alterar@123."),
    ]
    for row in guide_rows:
        guide.append(row)
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 80

    payload = io.BytesIO()
    workbook.save(payload)
    payload.seek(0)
    return StreamingResponse(
        payload,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="modelo-importacao-usuarios.xlsx"'},
    )


@app.post("/company-admin/import/users/preview")
async def company_admin_user_import_preview(
    file: UploadFile = File(...),
    company_id: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
):
    rows = await read_import_rows(file)
    if len(rows) > 1000:
        raise HTTPException(status_code=400, detail="A planilha aceita no maximo 1000 usuarios por importacao.")
    with SessionLocal() as db:
        scoped_company_id = ensure_company_access(db, current_user, company_id)
        ensure_company_admin(db, current_user, scoped_company_id)
        return preview_tenant_user_import_rows(rows, db, scoped_company_id)


@app.post("/company-admin/import/users/confirm")
async def company_admin_user_import_confirm(payload: dict, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        ensure_company_admin(db, current_user, company_id)
        company = db.query(Company).filter(Company.id == company_id).first()
        preview = preview_tenant_user_import_rows(payload.get("rows") or [], db, company_id)
        if preview["error_count"]:
            return {**preview, "imported": 0, "linked": 0, "updated": 0, "message": "Corrija os erros antes de importar."}

        imported = 0
        linked = 0
        updated = 0
        for row in preview["rows"]:
            data = row["data"]
            department = get_or_create_department(db, company_id, data["setor"])
            user = db.query(User).filter(User.email == data["email"]).first()
            existing_link = None
            if user:
                existing_link = db.query(CompanyUser).filter(
                    CompanyUser.company_id == company_id,
                    CompanyUser.user_id == user.id,
                ).first()

            if not user:
                username = generate_unique_username(
                    db,
                    data["email"],
                    data["nome_completo"],
                    data["usuario"] or None,
                )
                user = User(
                    uuid=str(uuid.uuid4()),
                    username=username,
                    email=data["email"],
                    full_name=data["nome_completo"],
                    hashed_password=get_password_hash(TENANT_USER_DEFAULT_PASSWORD),
                    phone_extension=data["ramal"] or None,
                    must_change_password=True,
                    status=ACTIVE_STATUS,
                    is_active=True,
                    access_level=legacy_access_for_role(data["nivel_usuario"]),
                    department=department.name,
                )
                db.add(user)
                db.flush()
                imported += 1
            else:
                user.full_name = data["nome_completo"]
                user.phone_extension = data["ramal"] or None
                user.status = ACTIVE_STATUS
                user.is_active = True
                user.department = department.name
                user.access_level = legacy_access_for_role(data["nivel_usuario"])
                user.updated_at = datetime.utcnow()
                updated += 1

            link = existing_link
            if not link:
                link = CompanyUser(
                    id=str(uuid.uuid4()),
                    company_id=company_id,
                    user_id=user.id,
                    department_id=department.id,
                    role=data["nivel_usuario"],
                    status=ACTIVE_STATUS,
                )
                db.add(link)
                linked += 1
            else:
                link.department_id = department.id
                link.role = data["nivel_usuario"]
                link.status = ACTIVE_STATUS

            log_audit(db, current_user.id, "usuario_importado_tenant", "user", user.id, company_id, {
                "email": user.email,
                "department_id": department.id,
                "role": data["nivel_usuario"],
            })

        db.commit()
        return {
            **preview,
            "imported": imported,
            "linked": linked,
            "updated": updated,
            "message": "Importacao de usuarios concluida.",
        }

if __name__ == "__main__":
    create_default_users()
    uvicorn.run(app, host="0.0.0.0", port=int(os.getenv("PORT", "8001")))

# WhatsApp control center routes are intentionally metadata-only; Meta remains the message source of truth.
def whatsapp_row(row):
    if not row:
        return None
    data = dict(row._mapping)
    for key, value in list(data.items()):
        if isinstance(value, datetime):
            data[key] = value.isoformat()
    return data

@app.get("/whatsapp/overview")
async def whatsapp_overview(company_id: Optional[str] = None, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, company_id)
        params = {"company_id": company_id}
        settings = db.execute(text("SELECT company_id, phone_number_id, business_account_id, display_phone, status, bolt_mode, updated_at FROM whatsapp_settings WHERE company_id = :company_id"), params).first()
        tags = db.execute(text("SELECT id, name, color, created_at FROM whatsapp_tags WHERE company_id = :company_id ORDER BY name"), params).all()
        automations = db.execute(text("SELECT id, name, trigger_type, trigger_value, response_text, enabled, created_at, updated_at FROM whatsapp_automations WHERE company_id = :company_id ORDER BY created_at DESC"), params).all()
        curation = db.execute(text("SELECT id, external_contact_id, contact_name, incoming_excerpt, suggested_response, status, created_at FROM whatsapp_bolt_curation WHERE company_id = :company_id AND status = 'pending' ORDER BY created_at DESC LIMIT 100"), params).all()
        return {"company_id": company_id, "can_manage": is_company_admin(db, current_user, company_id), "settings": whatsapp_row(settings) or {"company_id": company_id, "status": "disconnected", "bolt_mode": "curation"}, "tags": [whatsapp_row(row) for row in tags], "automations": [whatsapp_row(row) for row in automations], "curation": [whatsapp_row(row) for row in curation]}

@app.put("/whatsapp/settings")
async def update_whatsapp_settings(payload: dict, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        ensure_company_admin(db, current_user, company_id)
        bolt_mode = str(payload.get("bolt_mode") or "curation").strip().lower()
        if bolt_mode not in {"off", "curation", "automatic"}:
            raise HTTPException(status_code=400, detail="Modo do Volt invalido.")
        db.execute(text("INSERT INTO whatsapp_settings (company_id, phone_number_id, business_account_id, display_phone, status, bolt_mode, updated_by_id, updated_at) VALUES (:company_id, :phone_number_id, :business_account_id, :display_phone, :status, :bolt_mode, :user_id, NOW()) ON CONFLICT (company_id) DO UPDATE SET phone_number_id=EXCLUDED.phone_number_id, business_account_id=EXCLUDED.business_account_id, display_phone=EXCLUDED.display_phone, status=EXCLUDED.status, bolt_mode=EXCLUDED.bolt_mode, updated_by_id=EXCLUDED.updated_by_id, updated_at=NOW()"), {"company_id": company_id, "phone_number_id": str(payload.get("phone_number_id") or "").strip() or None, "business_account_id": str(payload.get("business_account_id") or "").strip() or None, "display_phone": str(payload.get("display_phone") or "").strip() or None, "status": "connected" if payload.get("status") == "connected" else "disconnected", "bolt_mode": bolt_mode, "user_id": current_user.id})
        log_audit(db, current_user.id, "whatsapp_configurado", "whatsapp_settings", company_id, company_id, {"bolt_mode": bolt_mode})
        db.commit()
        return {"success": True}

@app.post("/whatsapp/tags")
async def create_whatsapp_tag(payload: dict, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        ensure_company_admin(db, current_user, company_id)
        name = str(payload.get("name") or "").strip()[:60]
        if not name:
            raise HTTPException(status_code=400, detail="Informe o nome da tag.")
        tag_id = str(uuid.uuid4())
        try:
            db.execute(text("INSERT INTO whatsapp_tags (id, company_id, name, color) VALUES (:id, :company_id, :name, :color)"), {"id": tag_id, "company_id": company_id, "name": name, "color": str(payload.get("color") or "#2563eb")[:20]})
            log_audit(db, current_user.id, "whatsapp_tag_criada", "whatsapp_tag", tag_id, company_id, {"name": name})
            db.commit()
        except IntegrityError as exc:
            db.rollback()
            raise HTTPException(status_code=409, detail="Ja existe uma tag com este nome.") from exc
        return {"id": tag_id, "name": name}

@app.post("/whatsapp/automations")
async def create_whatsapp_automation(payload: dict, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        ensure_company_admin(db, current_user, company_id)
        name = str(payload.get("name") or "").strip()[:120]
        response_text = str(payload.get("response_text") or "").strip()
        trigger_type = str(payload.get("trigger_type") or "keyword").strip().lower()
        if not name or not response_text:
            raise HTTPException(status_code=400, detail="Informe nome e mensagem da automacao.")
        if trigger_type not in {"keyword", "first_contact", "outside_hours"}:
            raise HTTPException(status_code=400, detail="Gatilho invalido.")
        automation_id = str(uuid.uuid4())
        db.execute(text("INSERT INTO whatsapp_automations (id, company_id, name, trigger_type, trigger_value, response_text, enabled, created_by_id) VALUES (:id, :company_id, :name, :trigger_type, :trigger_value, :response_text, :enabled, :user_id)"), {"id": automation_id, "company_id": company_id, "name": name, "trigger_type": trigger_type, "trigger_value": str(payload.get("trigger_value") or "").strip() or None, "response_text": response_text, "enabled": bool(payload.get("enabled", True)), "user_id": current_user.id})
        log_audit(db, current_user.id, "whatsapp_automacao_criada", "whatsapp_automation", automation_id, company_id, {"name": name})
        db.commit()
        return {"id": automation_id, "success": True}

@app.put("/whatsapp/curation/{curation_id}")
async def review_whatsapp_curation(curation_id: str, payload: dict, current_user: User = Depends(get_current_user)):
    ensure_admin(current_user)
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        ensure_company_admin(db, current_user, company_id)
        decision = str(payload.get("decision") or "").lower()
        if decision not in {"approved", "rejected"}:
            raise HTTPException(status_code=400, detail="Decisao invalida.")
        result = db.execute(text("UPDATE whatsapp_bolt_curation SET status=:status, reviewed_by_id=:user_id, reviewed_at=NOW(), suggested_response=COALESCE(:response_text, suggested_response) WHERE id=:id AND company_id=:company_id AND status='pending'"), {"status": decision, "user_id": current_user.id, "response_text": str(payload.get("response_text") or "").strip() or None, "id": curation_id, "company_id": company_id})
        if result.rowcount != 1:
            raise HTTPException(status_code=404, detail="Item de curadoria nao encontrado.")
        log_audit(db, current_user.id, "whatsapp_bolt_" + decision, "whatsapp_curation", curation_id, company_id)
        db.commit()
        return {"success": True, "status": decision}
# Meeting agenda: only metadata is persisted; WebRTC media stays ephemeral.
def meeting_datetime(value, field_name):
    try:
        parsed = datetime.fromisoformat(str(value or "").replace("Z", "+00:00"))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=f"{field_name} invalido.") from exc
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed


def serialize_meeting_row(db, row):
    data = dict(row._mapping)
    participants = db.execute(text("""
        SELECT u.id, u.full_name, u.nickname, u.username, u.profile_photo,
               mp.response_status, mp.joined_at, mp.left_at
        FROM meeting_participants mp JOIN users u ON u.id = mp.user_id
        WHERE mp.meeting_id = :meeting_id ORDER BY COALESCE(u.nickname, u.full_name, u.username)
    """), {"meeting_id": data["id"]}).all()
    data["participants"] = [whatsapp_row(item) for item in participants]
    for key in ("starts_at", "ends_at", "created_at", "updated_at"):
        if isinstance(data.get(key), datetime):
            data[key] = data[key].isoformat()
    return data


def normalize_meeting_link(value) -> Optional[str]:
    link_url = str(value or "").strip()
    if not link_url:
        return None
    if len(link_url) > 2048:
        raise HTTPException(status_code=400, detail="O link informado e muito longo.")
    parsed = urllib.parse.urlparse(link_url)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise HTTPException(status_code=400, detail="Informe um link valido iniciado por http:// ou https://.")
    return link_url


@app.get("/meetings")
async def list_meetings(start: Optional[str] = None, end: Optional[str] = None, company_id: Optional[str] = None, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, company_id)
        start_at = meeting_datetime(start, "Inicio") if start else datetime.now(timezone.utc) - timedelta(days=365)
        end_at = meeting_datetime(end, "Fim") if end else datetime.now(timezone.utc) + timedelta(days=730)
        rows = db.execute(text("""
            SELECT DISTINCT m.* FROM meetings m
            LEFT JOIN meeting_participants mp ON mp.meeting_id = m.id
            WHERE m.company_id = :company_id AND m.status = 'scheduled'
              AND m.starts_at < :end_at AND m.ends_at > :start_at
              AND (m.creator_user_id = :user_id OR mp.user_id = :user_id OR :can_manage = TRUE)
            ORDER BY m.starts_at
        """), {"company_id": company_id, "start_at": start_at, "end_at": end_at, "user_id": current_user.id, "can_manage": is_company_admin(db, current_user, company_id)}).all()
        return [serialize_meeting_row(db, row) for row in rows]


@app.post("/meetings")
async def create_meeting(payload: dict, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        title = str(payload.get("title") or "").strip()[:180]
        meeting_type = str(payload.get("meeting_type") or "reminder").strip().lower()
        link_url = normalize_meeting_link(payload.get("link_url"))
        starts_at = meeting_datetime(payload.get("starts_at"), "Data e horario inicial")
        ends_at = meeting_datetime(payload.get("ends_at"), "Data e horario final")
        if not title:
            raise HTTPException(status_code=400, detail="Informe o nome da reuniao.")
        if meeting_type not in {"reminder", "task", "video"}:
            raise HTTPException(status_code=400, detail="Tipo de reuniao invalido.")
        if ends_at <= starts_at:
            raise HTTPException(status_code=400, detail="O termino deve ser posterior ao inicio.")
        participant_ids = {current_user.id}
        for raw_id in payload.get("participant_ids") or []:
            participant = ensure_user_in_company(db, int(raw_id), company_id)
            participant_ids.add(participant.id)
        meeting_id = str(uuid.uuid4())
        db.execute(text("""
            INSERT INTO meetings (id, company_id, creator_user_id, title, description, meeting_type, starts_at, ends_at, color, link_url)
            VALUES (:id, :company_id, :creator_user_id, :title, :description, :meeting_type, :starts_at, :ends_at, :color, :link_url)
        """), {"id": meeting_id, "company_id": company_id, "creator_user_id": current_user.id, "title": title, "description": str(payload.get("description") or "").strip() or None, "meeting_type": meeting_type, "starts_at": starts_at, "ends_at": ends_at, "color": str(payload.get("color") or "#2563eb")[:20], "link_url": link_url})
        for user_id in participant_ids:
            db.execute(text("INSERT INTO meeting_participants (meeting_id, user_id, response_status) VALUES (:meeting_id, :user_id, :status)"), {"meeting_id": meeting_id, "user_id": user_id, "status": "accepted" if user_id == current_user.id else "pending"})
        invitation_messages = []
        mitty = ensure_mitty_user(db, company_id)
        local_start = starts_at.astimezone(BOLT_DAILY_SUMMARY_TIMEZONE).strftime("%d/%m/%Y às %H:%M")
        item_label = "reuniao" if meeting_type == "video" else ("tarefa" if meeting_type == "task" else "compromisso")
        invitation_text = f"voce foi incluido no {item_label} {title}, em {local_start}. Consulte os detalhes na aba Meeting."
        if link_url:
            invitation_text += f" Link: {link_url}"
        for user_id in participant_ids:
            if user_id == current_user.id:
                continue
            participant = ensure_user_in_company(db, user_id, company_id)
            personalized_text = assistant_greeting(participant, invitation_text)
            message = Message(company_id=company_id, content=personalized_text, sender_id=mitty.id, receiver_id=user_id, message_type="meeting_invitation")
            db.add(message)
            db.flush()
            invitation_messages.append((user_id, serialize_message(message, db)))
        log_audit(db, current_user.id, "meeting_criado", "meeting", meeting_id, company_id, {"type": meeting_type, "participants": sorted(participant_ids)})
        db.commit()
        row = db.execute(text("SELECT * FROM meetings WHERE id=:id"), {"id": meeting_id}).first()
        response = serialize_meeting_row(db, row)
        for user_id, message_payload in invitation_messages:
            await manager.send_personal_message(json.dumps({"type": "new_message", "message": message_payload}), user_id)
            await dispatch_meeting_reminder(user_id, meeting_id, "Mitty - Convite para Meeting", message_payload["content"], starts_at, "invite")
        wake_meeting_reminder_scheduler()
        return response




@app.patch("/meetings/{meeting_id}")
async def update_meeting(meeting_id: str, payload: dict, current_user: User = Depends(get_current_user)):
    participant_ids = []
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        meeting = db.execute(text("""
            SELECT * FROM meetings WHERE id=:id AND company_id=:company_id AND status='scheduled'
        """), {"id": meeting_id, "company_id": company_id}).first()
        if not meeting:
            raise HTTPException(status_code=404, detail="Item da agenda nao encontrado.")
        if meeting.creator_user_id != current_user.id and not is_company_admin(db, current_user, company_id):
            raise HTTPException(status_code=403, detail="Somente o criador ou admin da empresa pode editar este item.")

        updates = {}
        if "title" in payload:
            title = str(payload.get("title") or "").strip()[:180]
            if not title:
                raise HTTPException(status_code=400, detail="Informe o titulo do item.")
            updates["title"] = title
        if "description" in payload:
            updates["description"] = str(payload.get("description") or "").strip() or None
        if "meeting_type" in payload:
            meeting_type = str(payload.get("meeting_type") or "").strip().lower()
            if meeting_type not in {"reminder", "task", "video"}:
                raise HTTPException(status_code=400, detail="Tipo de item invalido.")
            updates["meeting_type"] = meeting_type
        if "color" in payload:
            updates["color"] = str(payload.get("color") or "#2563eb")[:20]
        if "link_url" in payload:
            updates["link_url"] = normalize_meeting_link(payload.get("link_url"))
        if "starts_at" in payload:
            updates["starts_at"] = meeting_datetime(payload.get("starts_at"), "Data e horario inicial")
        if "ends_at" in payload:
            updates["ends_at"] = meeting_datetime(payload.get("ends_at"), "Data e horario final")
        starts_at = updates.get("starts_at", meeting_datetime(meeting.starts_at, "Inicio"))
        ends_at = updates.get("ends_at", meeting_datetime(meeting.ends_at, "Termino"))
        if ends_at <= starts_at:
            raise HTTPException(status_code=400, detail="O termino deve ser posterior ao inicio.")

        assignments = [f"{field}=:{field}" for field in updates]
        params = {**updates, "id": meeting_id, "updated_at": datetime.now(timezone.utc)}
        assignments.extend([
            "updated_at=:updated_at", "reminder_day_sent=FALSE",
            "reminder_30m_sent=FALSE", "start_alert_sent=FALSE",
        ])
        db.execute(text(f"UPDATE meetings SET {', '.join(assignments)} WHERE id=:id"), params)

        if "participant_ids" in payload:
            requested_ids = {meeting.creator_user_id}
            for raw_id in payload.get("participant_ids") or []:
                requested_ids.add(ensure_user_in_company(db, int(raw_id), company_id).id)
            db.execute(text("DELETE FROM meeting_participants WHERE meeting_id=:id"), {"id": meeting_id})
            for user_id in requested_ids:
                db.execute(text("""
                    INSERT INTO meeting_participants (meeting_id, user_id, response_status)
                    VALUES (:meeting_id, :user_id, :status)
                """), {
                    "meeting_id": meeting_id, "user_id": user_id,
                    "status": "accepted" if user_id in {current_user.id, meeting.creator_user_id} else "pending",
                })

        participant_ids = [row[0] for row in db.execute(text(
            "SELECT user_id FROM meeting_participants WHERE meeting_id=:id"
        ), {"id": meeting_id}).all()]
        log_audit(db, current_user.id, "meeting_editado", "meeting", meeting_id, company_id, {
            "fields": sorted([key for key in payload if key != "company_id"]),
        })
        db.commit()
        row = db.execute(text("SELECT * FROM meetings WHERE id=:id"), {"id": meeting_id}).first()
        serialized = serialize_meeting_row(db, row)

    local_start = meeting_datetime(serialized["starts_at"], "Inicio").astimezone(BOLT_DAILY_SUMMARY_TIMEZONE).strftime("%d/%m/%Y às %H:%M")
    update_description = f"{serialized['title']} foi atualizado para {local_start}."
    if serialized.get("link_url"):
        update_description += f" Link da reunião: {serialized['link_url']}"
    for user_id in participant_ids:
        await dispatch_meeting_reminder(
            user_id, meeting_id, "Mitty - Agenda atualizada",
            update_description,
            meeting_datetime(serialized["starts_at"], "Inicio"), "updated",
        )
    wake_meeting_reminder_scheduler()
    return serialized
@app.put("/meetings/{meeting_id}/response")
async def respond_meeting(meeting_id: str, payload: dict, current_user: User = Depends(get_current_user)):
    response_status = str(payload.get("status") or "").lower()
    if response_status not in {"accepted", "declined", "joined", "left"}:
        raise HTTPException(status_code=400, detail="Resposta invalida.")
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, payload.get("company_id"))
        joined_at = datetime.now(timezone.utc) if response_status == "joined" else None
        left_at = datetime.now(timezone.utc) if response_status == "left" else None
        status_value = "accepted" if response_status in {"joined", "left"} else response_status
        result = db.execute(text("""
            UPDATE meeting_participants SET response_status=:status,
              joined_at=COALESCE(:joined_at, joined_at), left_at=COALESCE(:left_at, left_at)
            WHERE meeting_id=:meeting_id AND user_id=:user_id
              AND EXISTS (SELECT 1 FROM meetings WHERE id=:meeting_id AND company_id=:company_id)
        """), {"status": status_value, "joined_at": joined_at, "left_at": left_at, "meeting_id": meeting_id, "user_id": current_user.id, "company_id": company_id})
        if result.rowcount != 1:
            raise HTTPException(status_code=404, detail="Reuniao ou convite nao encontrado.")
        db.commit()
        return {"success": True, "status": response_status}


@app.delete("/meetings/{meeting_id}")
async def cancel_meeting(meeting_id: str, company_id: Optional[str] = None, current_user: User = Depends(get_current_user)):
    with SessionLocal() as db:
        company_id = ensure_company_access(db, current_user, company_id)
        row = db.execute(text("SELECT creator_user_id FROM meetings WHERE id=:id AND company_id=:company_id"), {"id": meeting_id, "company_id": company_id}).first()
        if not row:
            raise HTTPException(status_code=404, detail="Reuniao nao encontrada.")
        if row.creator_user_id != current_user.id and not is_company_admin(db, current_user, company_id):
            raise HTTPException(status_code=403, detail="Somente o criador ou admin da empresa pode cancelar.")
        participant_ids = [item[0] for item in db.execute(text(
            "SELECT user_id FROM meeting_participants WHERE meeting_id=:id"
        ), {"id": meeting_id}).all()]
        db.execute(text("DELETE FROM meeting_participants WHERE meeting_id=:id"), {"id": meeting_id})
        db.execute(text("DELETE FROM meetings WHERE id=:id"), {"id": meeting_id})
        log_audit(db, current_user.id, "meeting_excluido", "meeting", meeting_id, company_id)
        db.commit()
    event = json.dumps({"type": "meeting_deleted", "meeting_id": meeting_id})
    for participant_id in participant_ids:
        await manager.send_personal_message(event, participant_id)
    wake_meeting_reminder_scheduler()
    return {"success": True, "meeting_id": meeting_id}


async def dispatch_meeting_reminder(user_id, meeting_id, title, description, starts_at, kind):
    notification = {"id": f"meeting-{meeting_id}-{kind}-{user_id}", "type": "Meeting", "title": title, "description": description, "time": "Agora", "unread": True, "link": "/meetings", "meeting_id": meeting_id}
    await manager.send_personal_message(json.dumps({"type": "notification", "notification": notification}), user_id)


MEETING_REMINDER_IDLE_SECONDS = 24 * 60 * 60
meeting_reminder_wakeup = asyncio.Event()


def wake_meeting_reminder_scheduler():
    """Acorda o agendador quando a agenda muda, sem polling contínuo."""
    meeting_reminder_wakeup.set()


def next_pending_meeting_reminder_at(meetings, now: Optional[datetime] = None) -> Optional[datetime]:
    """Retorna somente o próximo marco de lembrete ainda pendente."""
    current = ensure_utc_datetime(now or datetime.now(timezone.utc))
    due_at = None
    for meeting in meetings:
        starts_at = ensure_utc_datetime(meeting["starts_at"])
        ends_at = ensure_utc_datetime(meeting["ends_at"])
        if ends_at <= current:
            continue
        milestones = (
            ("reminder_day_sent", starts_at - timedelta(days=1)),
            ("reminder_30m_sent", starts_at - timedelta(minutes=30)),
            ("start_alert_sent", starts_at),
        )
        for sent_field, candidate in milestones:
            if meeting.get(sent_field):
                continue
            if candidate <= current:
                return current
            if due_at is None or candidate < due_at:
                due_at = candidate
            break
    return due_at


def next_meeting_reminder_delay_seconds(now: Optional[datetime] = None) -> int:
    current = ensure_utc_datetime(now or datetime.now(timezone.utc))
    with SessionLocal() as db:
        rows = db.execute(text("""
            SELECT starts_at, ends_at, reminder_day_sent, reminder_30m_sent, start_alert_sent
            FROM meetings
            WHERE status='scheduled' AND ends_at > :now
        """), {"now": current}).mappings().all()
    due_at = next_pending_meeting_reminder_at(rows, current)
    if due_at is None:
        return MEETING_REMINDER_IDLE_SECONDS
    return max(1, int((due_at - current).total_seconds()))


async def run_meeting_reminders():
    now = datetime.now(timezone.utc)
    prepared = []
    with SessionLocal() as db:
        rows = db.execute(text("""
            SELECT * FROM meetings WHERE status='scheduled' AND ends_at > :now
              AND starts_at <= :day_limit
        """), {"now": now, "day_limit": now + timedelta(days=1)}).all()
        for row in rows:
            meeting = dict(row._mapping)
            seconds = (meeting["starts_at"] - now).total_seconds()
            reminder_kind = None
            flag = None
            if -60 <= seconds <= 60 and not meeting["start_alert_sent"]:
                reminder_kind, flag = "start", "start_alert_sent"
            elif 0 < seconds <= 1800 and not meeting["reminder_30m_sent"]:
                reminder_kind, flag = "30m", "reminder_30m_sent"
            elif 23 * 3600 <= seconds <= 24 * 3600 and not meeting["reminder_day_sent"]:
                reminder_kind, flag = "1d", "reminder_day_sent"
            if not reminder_kind:
                continue
            participant_ids = [item[0] for item in db.execute(text("SELECT user_id FROM meeting_participants WHERE meeting_id=:id AND response_status!='declined'"), {"id": meeting["id"]}).all()]
            when = meeting["starts_at"].astimezone(BOLT_DAILY_SUMMARY_TIMEZONE).strftime("%d/%m/%Y às %H:%M")
            item_label = "reuniao" if meeting["meeting_type"] == "video" else ("tarefa" if meeting["meeting_type"] == "task" else "compromisso")
            prefix = f"A {item_label} comeca agora" if reminder_kind == "start" else (f"Faltam 30 minutos para a {item_label}" if reminder_kind == "30m" else f"{item_label.capitalize()} amanha")
            description = f"{prefix}: {meeting['title']} - {when}."
            meeting_link = str(meeting.get("link_url") or "").strip()
            if meeting_link:
                description += f" Link da reunião: {meeting_link}"
            elif reminder_kind == "start" and meeting["meeting_type"] == "video":
                description += " Entre na chamada de vídeo pela aba Agenda."
            mitty = ensure_mitty_user(db, meeting["company_id"])
            for user_id in participant_ids:
                participant = ensure_user_in_company(db, user_id, meeting["company_id"])
                personalized_description = assistant_greeting(
                    participant,
                    description[0].lower() + description[1:] if description else description,
                )
                message = Message(company_id=meeting["company_id"], content=personalized_description, sender_id=mitty.id, receiver_id=user_id, message_type="meeting_notification")
                db.add(message)
                db.flush()
                notification_title = "Mitty - Reuniao" if meeting["meeting_type"] == "video" else ("Mitty - Tarefa" if meeting["meeting_type"] == "task" else "Mitty - Compromisso")
                prepared.append((user_id, meeting["id"], notification_title, personalized_description, meeting["starts_at"], reminder_kind, serialize_message(message, db)))
            db.execute(text(f"UPDATE meetings SET {flag}=TRUE WHERE id=:id"), {"id": meeting["id"]})
        db.commit()
    for user_id, meeting_id, title, description, starts_at, kind, message_payload in prepared:
        await manager.send_personal_message(json.dumps({"type": "new_message", "message": message_payload}), user_id)
        await dispatch_meeting_reminder(user_id, meeting_id, title, description, starts_at, kind)
    return len(prepared)


async def periodic_meeting_reminders():
    while True:
        try:
            await run_meeting_reminders()
        except Exception as exc:
            print(f"Falha nos lembretes de Meeting: {exc}")
        meeting_reminder_wakeup.clear()
        try:
            delay_seconds = await asyncio.to_thread(next_meeting_reminder_delay_seconds)
        except Exception as exc:
            print(f"Falha ao calcular proximo lembrete de Meeting: {exc}")
            delay_seconds = MEETING_REMINDER_IDLE_SECONDS
        try:
            await asyncio.wait_for(meeting_reminder_wakeup.wait(), timeout=delay_seconds)
        except asyncio.TimeoutError:
            pass
